from __future__ import annotations

from pathlib import Path
from unicodedata import normalize

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import Address, DocumentGenerationContext, Person
from sydel_doc_engine.generators.lot_02.regime_communautaire_common import (
    city_line,
    company_forme_sociale_complete,
    format_display_date,
    required_address,
    required_apport,
    required_company,
    required_regime_communautaire,
    required_text,
    street_line,
    validate_batch_enabled,
)
from sydel_doc_engine.rendering.docx_builder import (
    LETTER_WIDE_STYLE_PROFILE,
    add_paragraph,
    add_right_aligned_lines,
    add_spacer,
    add_subject_heading,
    keep_final_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.grammar import accord_terme_genre

OUTPUT_FILENAME = "lettre_renonciation_associe.docx"


class LettreRenonciationAssocieGenerator:
    """Generateur from-scratch de la lettre de renonciation du conjoint."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_batch_enabled(ctx)
        company = required_company(ctx.societe)
        apport = required_apport(ctx.apport)
        regime = required_regime_communautaire(ctx.regime_communautaire)
        if regime.renonciation is None:
            raise ValueError(
                "regime_communautaire.renonciation est obligatoire pour CODE-RC-001."
            )

        date_courrier = _date_courrier_avertissement(ctx)
        lieu_signature = required_text(
            regime.renonciation.lieu_signature,
            "regime_communautaire.renonciation.lieu_signature",
        )

        document = new_document(style_profile=LETTER_WIDE_STYLE_PROFILE)
        # R1 (Albane 2026-06-26) : c'est un courrier -> en haut a droite, le bloc
        # destinataire (nom + adresse) AU-DESSUS de la ligne « à <ville> », puis le
        # corps + l'objet descendus nettement plus bas.
        add_right_aligned_lines(
            document,
            _destinataire_block(ctx),
            space_after_pt=2,
        )
        add_spacer(document, space_after_pt=6)
        add_right_aligned_lines(
            document,
            [f"À {lieu_signature}"],
            space_after_pt=2,
        )
        # Descend le corps/objet (R1 : « que le texte avec objet soit bien plus bas »).
        add_spacer(document, space_after_pt=48)
        # m3 (Akainu ronde 2, 2026-07-12) : la lettre est signee par le CONJOINT qui renonce
        # -> « associé » s'accorde a SON genre (« associée » pour une conjointe), a l'objet
        # comme dans le corps (« devenir personnellement associé/associée »).
        conjoint_genre = _required_conjoint(ctx).genre
        add_subject_heading(
            document,
            (
                "Objet : Lettre de renonciation à revendiquer la qualité "
                f"d'{accord_terme_genre('associé', conjoint_genre)}"
            ),
            space_after_pt=12,
        )
        add_paragraph(document, _apporteur_appel(ctx))
        denomination = required_text(company.denomination, "societe.denomination")
        regime_matrimonial = _regime_matrimonial_display(
            required_text(
                regime.regime_matrimonial,
                "regime_communautaire.regime_matrimonial",
            )
        )
        add_paragraph(
            document,
            (
                f"Par courrier en date du {date_courrier}, tu m’as fait part du projet de "
                f"constitution de la société {denomination}, "
                f"{company_forme_sociale_complete(company)}, à laquelle tu souhaites t'associer "
                f"en apportant {required_text(apport.montant, 'apport.montant')} "
                f"({required_text(apport.montant_lettres, 'apport.montant_lettres')}) euros "
                f"dépendant de notre {regime_matrimonial}."
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        qualite_renoncee = required_text(
            regime.qualite_renoncee,
            "regime_communautaire.qualite_renoncee",
        )
        add_paragraph(
            document,
            (
                "Je te notifie, par la présente, mon intention de renoncer à la faculté de "
                f"devenir personnellement {accord_terme_genre(qualite_renoncee, conjoint_genre)} "
                "de cette société."
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        # M2 (Akainu SELARL ronde 4, 2026-07-12) : « mon conjoint » designe l'APPORTEUR
        # (personne_signataire, celui qui a fait l'apport) -> « ma conjointe » quand l'apporteur
        # est une femme. Distinct de l'accord « associé(e) » de l'objet (genre du CONJOINT
        # signataire). Coder l'INTENTION (regle 68), pas la seule tournure « associé ».
        apporteur_genre = ctx.personne_signataire.genre
        mon_conjoint = (
            "ma conjointe" if apporteur_genre == Gender.FEMININ else "mon conjoint"
        )
        add_paragraph(
            document,
            (
                "En tout état de cause, et conformément aux dispositions du Code civil, je "
                f"déclare donner mon consentement à l'apport effectué par {mon_conjoint}."
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        add_paragraph(document, "Fait pour servir et valoir ce que de droit.")
        # R2 (Albane 2026-06-26) : la mention « en N exemplaires » n'a pas de sens
        # pour un courrier -> retiree de la lettre de renonciation (uniquement).
        add_right_aligned_lines(document, [_conjoint_signature(ctx)], space_after_pt=0)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(document)
        document.save(output_path)
        return output_path


def _date_courrier_avertissement(ctx: DocumentGenerationContext) -> str:
    regime = required_regime_communautaire(ctx.regime_communautaire)
    if regime.date_courrier_avertissement is not None:
        return format_display_date(
            regime.date_courrier_avertissement,
            "regime_communautaire.date_courrier_avertissement",
        )
    if regime.avertissement is not None and regime.avertissement.date_signature is not None:
        return format_display_date(
            regime.avertissement.date_signature,
            "regime_communautaire.avertissement.date_signature",
        )
    raise ValueError(
        "regime_communautaire.date_courrier_avertissement ou "
        "regime_communautaire.avertissement.date_signature est obligatoire pour CODE-RC-001."
    )


def _regime_matrimonial_display(value: str) -> str:
    normalized = (
        normalize("NFKD", value).encode("ascii", "ignore").decode("ascii").lower()
    )
    if "communaute" in " ".join(normalized.split()):
        return "communauté"
    for prefix in ("sous le régime de ", "sous le regime de ", "régime de ", "regime de "):
        if value.lower().startswith(prefix):
            return value[len(prefix) :].strip()
    return value.strip()


def _apporteur_appel(ctx: DocumentGenerationContext) -> str:
    apporteur = ctx.personne_signataire
    civilite = required_text(apporteur.civilite, "apporteur.civilite_affichage")
    prenom = required_text(apporteur.prenom, "apporteur.prenom")
    nom = required_text(apporteur.nom, "apporteur.nom")
    return f"{civilite} {prenom} {nom},"


def _destinataire_block(ctx: DocumentGenerationContext) -> list[str]:
    """Bloc destinataire du courrier (R1) : nom + adresse de l'apporteur, place en
    haut a droite au-dessus de la ligne « à <ville> ».

    Le destinataire de la lettre de renonciation est l'associe apporteur
    (`personne_signataire`), a qui son conjoint adresse la renonciation."""
    apporteur = ctx.personne_signataire
    civilite = required_text(apporteur.civilite, "apporteur.civilite_affichage")
    prenom = required_text(apporteur.prenom, "apporteur.prenom")
    nom = required_text(apporteur.nom, "apporteur.nom")
    address = _apporteur_address(ctx)
    return [f"{civilite} {prenom} {nom}", street_line(address), city_line(address)]


def _apporteur_address(ctx: DocumentGenerationContext) -> Address:
    return required_address(
        ctx.personne_signataire.adresse_perso,
        "personne_signataire.adresse_perso",
    )


def _required_conjoint(ctx: DocumentGenerationContext) -> Person:
    conjoint = ctx.conjoint
    if conjoint is None:
        raise ValueError("conjoint est obligatoire pour CODE-RC-001.")
    return conjoint


def _conjoint_signature(ctx: DocumentGenerationContext) -> str:
    conjoint = _required_conjoint(ctx)
    prenom = required_text(conjoint.prenom, "conjoint.prenom")
    nom = required_text(conjoint.nom, "conjoint.nom")
    return f"{prenom} {nom}"
