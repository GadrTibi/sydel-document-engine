from __future__ import annotations

from datetime import date
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import Address, DocumentGenerationContext
from sydel_doc_engine.generators.lot_01.civilite import civilite_civile
from sydel_doc_engine.rendering.docx_builder import (
    add_framed_title,
    add_legal_reminder,
    add_paragraph,
    add_signature_block,
    add_spacer,
    keep_final_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.dates import format_date_fr
from sydel_doc_engine.utils.grammar import birth_label, filiation_label, subject_line

OUTPUT_FILENAME = "declaration_non_condamnation.docx"

DECLARATION_TEXT = (
    "Déclare sur l’honneur, conformément aux dispositions de l’article A.123-51 du Code de "
    "commerce, n’avoir fait l’objet d’aucune condamnation pénale ni de sanction civile ou "
    "administrative de nature à m’interdire – soit d’exercer une activité commerciale – soit de "
    "gérer, d’administrer ou de diriger une personne morale."
)

# Espace insecable (U+00A0) avant les deux-points : typographie francaise du modele source
# tokenise (« Rappel\xa0: »), perdue jusqu'ici par le generateur. Corrige globalement (tous types).
RAPPEL_TITLE_SUFFIX = chr(0x00A0) + ": Article L123-5 du code de commerce"
RAPPEL_PARAGRAPH_1 = (
    "Le fait de donner, de mauvaise foi, des indications inexactes ou incomplètes en vue d’une "
    "immatriculation, d’une radiation ou d’une mention complémentaire ou rectificative au registre "
    # Rafael 2026-07-09 (R5, groupement des milliers dès 4 chiffres, « partout ») : le
    # montant de l'amende légale (art. L123-5) est groupé « 4 500 euros » comme tout
    # montant ≥ 4 chiffres. Sens juridique inchangé (typographie française standard).
    "du commerce et des sociétés est puni d’une amende de 4 500 euros et d’un emprisonnement de "
    "six mois."
)
RAPPEL_PARAGRAPH_2 = (
    "Les dispositions des deuxième et troisième alinéas de l’article L.123-4 sont applicables "
    "dans les cas prévus au présent article."
)


class DeclarationNonCondamnationGenerator:
    """Générateur cible du DOC-001."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        person = ctx.personne_signataire
        address = person.adresse_perso
        if address is None:
            raise ValueError("personne_signataire.adresse_perso est obligatoire pour DOC-001.")

        # R3 (Albane 2026-07-07) : « Docteur » n'est pas une civilité — le slot
        # « Je soussigné __ » rend la civilité CIVILE (Monsieur/Madame, accordée
        # au genre du signataire), jamais le titre professionnel posé par le flux.
        civilite = civilite_civile(
            _required_text(person.civilite, "personne_signataire.civilite"),
            person.genre,
        )
        prenom = _required_text(person.prenom, "personne_signataire.prenom")
        nom = _required_text(person.nom, "personne_signataire.nom")
        date_naissance = _required_date(
            person.date_naissance,
            "personne_signataire.date_naissance",
        )
        ville_naissance = _required_text(
            person.ville_naissance,
            "personne_signataire.ville_naissance",
        )
        nationalite = _required_text(person.nationalite, "personne_signataire.nationalite")
        nom_pere = _required_text(person.nom_pere, "personne_signataire.nom_pere")
        nom_mere = _required_text(person.nom_mere, "personne_signataire.nom_mere")
        lieu_signature = _required_text(ctx.signature.lieu, "signature.lieu")
        adresse_perso = _compose_required_address(address)

        document = new_document()
        _add_title(document)
        # Aération sous le cadre du titre (retour Albane 2026-06-10).
        add_spacer(document, space_after_pt=10)
        _add_identity_block(
            document,
            subject=f"{subject_line(person.genre)} {civilite} {prenom} {nom}",
            # Retour Albane 2026-06-10 : « ne le {date} a {ville} ({departement}) »
            # — plus de point apres la ville, departement entre parentheses (si
            # renseigne).
            birth=(
                f"{birth_label(person.genre)} {date_naissance} "
                f"{_birth_city_prefix(person)} {ville_naissance}"
                f"{_birth_department_suffix(person)}"
            ),
            address=f"demeurant au {adresse_perso}",
            nationality=f"de nationalité {nationalite}",
            filiation_father=f"{filiation_label(person.genre)} {nom_pere}",
            filiation_mother=f"et de Madame {nom_mere}",
        )
        _add_paragraph(
            document,
            DECLARATION_TEXT,
            space_before=10,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        _add_signature_block(
            document,
            lieu_signature=lieu_signature,
            date_signature=format_date_fr(ctx.signature.date),
            image_path=ctx.signature.image_optionnelle,
        )
        # Mise en forme (Albane 2026-06-17, §7) : « descendre legerement » le
        # rappel legal en italique en menageant un espace avant le bloc.
        # add_legal_reminder est partage (autres docs) -> on aere ici seulement,
        # sans toucher le helper.
        add_spacer(document, space_after_pt=12)
        _add_legal_reminder(document)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(document)
        document.save(output_path)
        return output_path


def _required_text(value: str | None, field_name: str) -> str:
    # KAN-2 (Rafael 2026-07-13) : une donnée manquante NE bloque PAS la génération -> marqueur
    # visible « (À COMPLÉTER : … ) » sans crochets (comme required_text lot_03/04/05, R10), à
    # compléter à la main sur le DOCX, au lieu de lever.
    if value is None or not value.strip():
        return f"(À COMPLÉTER : {field_name})"
    return value.strip()


def _required_date(value: date | None, field_name: str) -> str:
    # KAN-2 : date manquante -> marqueur visible (non bloquant), à compléter à la main.
    if value is None:
        return f"(À COMPLÉTER : {field_name})"
    return format_date_fr(value)


def _compose_required_address(address: Address) -> str:
    # Numero de voie optionnel (champ fusionne « Numero et voie », retours
    # client 2026-06-11) : une adresse sans numero (lieu-dit) reste valide.
    num_voie = (address.num_voie or "").strip()
    voie = _required_text(address.voie, "personne_signataire.adresse_perso.voie")
    cp = _required_text(address.cp, "personne_signataire.adresse_perso.cp")
    ville = _required_text(address.ville, "personne_signataire.adresse_perso.ville")
    return f"{num_voie} {voie}, {cp} {ville}".strip()


def _birth_city_prefix(person) -> str:
    return "au" if person.ville_naissance_article_au else "\u00e0"


def _birth_department_suffix(person) -> str:
    departement = (getattr(person, "departement_naissance", None) or "").strip()
    return f" ({departement})" if departement else ""


def _add_title(document) -> None:
    add_framed_title(
        document,
        [
            "DECLARATION DE NON CONDAMNATION",
            "EN APPLICATION DE L’ARTICLE A.123-51 du Code de Commerce",
        ],
    )


def _add_identity_block(
    document,
    *,
    subject: str,
    birth: str,
    address: str,
    nationality: str,
    filiation_father: str,
    filiation_mother: str,
) -> None:
    # Bloc identite compact (retour Albane 2026-06-10 : retirer les interlignes
    # de « je soussigne » jusqu'a « de nationalite » et entre les noms des parents).
    # Mise en forme (Albane, retour « mise en forme » 2026-07) : AERER entre la
    # designation de l'associe/dirigeant (sujet -> naissance -> adresse ->
    # nationalite) et l'affiliation des parents (pere/mere). On garde la compacite
    # du bloc identite mais on menage un espace (6 pt) APRES la ligne « nationalite »,
    # juste AVANT la 1re ligne de filiation. Les lignes de filiation restent
    # compactes entre elles (pere -> mere, sa=0).
    _IDENTITY_STANDARD_SPACE_AFTER_PT = 6
    # Ordre FIGE du bloc identite. La ligne « nationalite » est la DERNIERE avant la
    # filiation : on la cible par son ROLE/POSITION (derniere ligne de designation),
    # pas par egalite de contenu (m3 : « line == nationality » etait fragile — une
    # valeur coincidente/vide aurait pu declencher l'espace au mauvais endroit).
    designation_lines = (
        (subject, True),
        (birth, False),
        (address, False),
        (nationality, False),
    )
    filiation_lines = (
        (filiation_father, False),
        (filiation_mother, False),
    )
    last_designation_index = len(designation_lines) - 1
    for index, (line, bold) in enumerate(designation_lines):
        # Espace uniquement APRES la DERNIERE ligne de designation (« de nationalite »),
        # juste avant la filiation ; le reste du bloc reste compact (sa=0).
        space_after = (
            _IDENTITY_STANDARD_SPACE_AFTER_PT if index == last_designation_index else 0
        )
        add_paragraph(document, line, bold=bold, space_after_pt=space_after)
    # Lignes de filiation compactes entre elles (pere -> mere, sa=0).
    for line, bold in filiation_lines:
        add_paragraph(document, line, bold=bold, space_after_pt=0)


def _add_paragraph(
    document,
    text: str,
    *,
    space_before: int = 0,
    bold: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> None:
    add_paragraph(
        document,
        text,
        space_before_pt=space_before,
        bold=bold,
        alignment=alignment,
    )


def _add_signature_block(
    document,
    *,
    lieu_signature: str,
    date_signature: str,
    image_path: Path | None,
) -> None:
    add_signature_block(
        document,
        [f"Fait à {lieu_signature}", f"Le {date_signature}"],
        image_path=image_path,
    )


def _add_legal_reminder(document) -> None:
    add_legal_reminder(
        document,
        title="Rappel",
        title_suffix=RAPPEL_TITLE_SUFFIX,
        paragraphs=[RAPPEL_PARAGRAPH_1, RAPPEL_PARAGRAPH_2],
    )
