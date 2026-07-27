from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from sydel_doc_engine.domain.models import (
    Address,
    CentreImpots,
    Company,
    DocumentGenerationContext,
    StatutsCivilsAssocie,
    StatutsCivilsContext,
)
from sydel_doc_engine.generators.lot_05.spfpl_libelles import libelle_metier
from sydel_doc_engine.rendering.docx_builder import (
    LETTER_WIDE_STYLE_PROFILE,
    add_framed_address_block,
    add_letter_place_date,
    add_paragraph,
    add_spacer,
    add_subject_heading,
    keep_final_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.dates import format_date_fr

OUTPUT_FILENAME = "lettre_option_is.docx"
DOCUMENT_CODE = "CODE-OPTION-IS-001"
# Micro holding (Albane 2026-06-29) : la lettre d'option IS fait partie du bundle de creation
# de la micro holding (societe civile) — meme modele que la SCI (« la societe civile … opte
# pour le regime de l'IS »).
SUPPORTED_STRUCTURES = {"SCI", "SCI IRIS", "MICRO_HOLDING"}
_EXPECTED_STATUTS_TYPE = {"SCI": "sci", "SCI IRIS": "sci_iris", "MICRO_HOLDING": "micro_holding"}


class LettreOptionIsGenerator:
    """Generateur from-scratch de la lettre d'option IS V1."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        _validate_context(ctx)
        company = _required_company(ctx.societe)
        tax_office = _required_tax_office(ctx.impots)
        statuts = _required_statuts_civils(ctx.statuts_civils)

        document = new_document(LETTER_WIDE_STYLE_PROFILE)
        _add_tax_office_block(document, tax_office)
        _add_place_date_and_subject(document, ctx.signature.lieu, ctx.signature.date)
        _add_body_intro(document, company)
        _add_identification_table(document, company, statuts)
        _add_body_close(document)
        # KAN-19 : nombre de gérants nommés (multi via dirigeants_nomines, sinon 1) -> accord.
        nb_gerants = len(ctx.dirigeants_nomines) if ctx.dirigeants_nomines else 1
        _add_signature(document, nb_gerants)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(document)
        document.save(output_path)
        return output_path


def _validate_context(ctx: DocumentGenerationContext) -> None:
    if ctx.structure not in SUPPORTED_STRUCTURES:
        supported = ", ".join(sorted(SUPPORTED_STRUCTURES))
        raise ValueError(f"dossier.structure doit etre dans [{supported}] pour {DOCUMENT_CODE}.")
    if ctx.dossier_options is None or not ctx.dossier_options.option_is:
        raise ValueError(f"dossier.options.option_is doit etre vrai pour {DOCUMENT_CODE}.")
    statuts = _required_statuts_civils(ctx.statuts_civils)
    expected_type = _EXPECTED_STATUTS_TYPE[ctx.structure]
    if statuts.type != expected_type:
        raise ValueError(f"statuts_civils.type doit etre {expected_type} pour {DOCUMENT_CODE}.")


def _required_company(company: Company | None) -> Company:
    if company is None:
        raise ValueError(f"societe est obligatoire pour {DOCUMENT_CODE}.")
    return company


def _required_tax_office(tax_office: CentreImpots | None) -> CentreImpots:
    if tax_office is None:
        raise ValueError(f"impots est obligatoire pour {DOCUMENT_CODE}.")
    return tax_office


def _required_statuts_civils(statuts: StatutsCivilsContext | None) -> StatutsCivilsContext:
    if statuts is None:
        raise ValueError(f"statuts_civils est obligatoire pour {DOCUMENT_CODE}.")
    if not statuts.associes:
        raise ValueError(f"statuts_civils.associes est obligatoire pour {DOCUMENT_CODE}.")
    return statuts


def _required_text(value: str | None, field_name: str) -> str:
    # KAN-2 (Rafael, rejete 2x) : « Tous les documents doivent pouvoir etre generes, meme si je ne
    # remplis AUCUN champ. » Une donnee manquante NE bloque PLUS -> marqueur metier visible
    # « (À COMPLÉTER : <libelle>) » (SANS crochet/point/underscore) au lieu de lever. Sortie
    # NOMINALE (valeur presente) BYTE-IDENTIQUE. Miroir de required_text (statuts_sel_exercice).
    if value is None or not value.strip():
        return f"(À COMPLÉTER : {libelle_metier(field_name)})"
    return value.strip()


def _required_int(value: int | None, field_name: str) -> int:
    # KAN-2 : ce helper ne sert QU'AUX controles de coherence (CALCUL) -> None-safe (-> 0), jamais
    # de crash. L'AFFICHAGE d'une quantite passe par `_quantite_parts` (marqueur si 0/None) pour ne
    # JAMAIS affirmer « 0 parts » dans un acte signable.
    if value is None:
        return 0
    return value


def _quantite_parts(value: int | None, libelle: str) -> str:
    """AFFICHAGE d'une quantite de parts dans le fil du texte (KAN-2).

    Une quantite NON RENSEIGNEE (None ou 0 = number_input jamais rempli) ne s'affirme JAMAIS
    (« 0 parts » est FAUX dans un acte signable) -> marqueur metier « (À COMPLÉTER : <libelle>) ».
    `libelle` = intitule metier deja humain (traverse `libelle_metier` inchange)."""
    if not value:
        return f"(À COMPLÉTER : {libelle_metier(libelle)})"
    return str(value)


# Retour Rafael R22-07 (2026-06-22) : le centre est TOUJOURS « Centre des Finances
# Publiques » (le service + l'adresse portent l'identification). On fige le libelle au
# lieu d'une variable saisie -> plus de champ « Centre » dans le formulaire.
CENTRE_FINANCES_PUBLIQUES = "Centre des Finances Publiques"

# R3 (Albane 2026-06-30) : la lettre d'option IS est emise PENDANT la constitution de la
# societe — elle n'est donc PAS encore immatriculee. Le SIREN ne doit JAMAIS afficher un
# numero (meme saisi par erreur), mais la constante « En cours d'immatriculation » (E
# majuscule, apostrophe COURBE U+2019, conforme au modele client
# docs/review/albane_micro_holding_2026-06-29/lettre_option_IS.docx). On retire donc
# l'exigence de non-vide sur company.siren pour ce document.
SIREN_EN_COURS_IMMATRICULATION = "En cours d’immatriculation"


# R2 (Albane 2026-06-30) : le bloc DESTINATAIRE (centre des impots) doit etre ENCADRE
# (enveloppe a fenetre) ET DESCENDU a la hauteur de la fenetre. Le calage a DROITE est
# conserve (« le côté me paraît bien »).
# DEFAUT DOCUMENTE (a valider Albane sur rendu) : la cote exacte n'est pas chiffree par
# Albane -> on vise ~4,5 cm du haut de page (fenetre standard FR). La marge haute du
# courrier est 2,5 cm (LETTER_WIDE_STYLE_PROFILE) ; il reste donc ~2,0 cm a descendre via
# un spacer pour amener le bloc a ~4,5 cm.
_FENETRE_DROP_TOP_CM = 2.0


def _add_tax_office_block(document: Any, tax_office: CentreImpots) -> None:
    add_framed_address_block(
        document,
        [
            _required_text(tax_office.service, "impots.service"),
            CENTRE_FINANCES_PUBLIQUES,
            _required_text(tax_office.adresse_ligne_1, "impots.adresse_ligne_1"),
            _required_text(tax_office.adresse_ligne_2, "impots.adresse_ligne_2"),
            (
                f"{_required_text(tax_office.cp, 'impots.cp')} "
                f"{_required_text(tax_office.ville, 'impots.ville')}"
            ),
        ],
        width_cm=7.5,
        drop_top_cm=_FENETRE_DROP_TOP_CM,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )
    add_spacer(document, space_after_pt=16)


def _add_place_date_and_subject(document: Any, lieu: str, signature_date: date) -> None:
    add_letter_place_date(
        document,
        f"Fait à {lieu}, le {format_date_fr(signature_date)}",
        space_after_pt=12,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )
    add_subject_heading(
        document,
        "Objet : Demande d'option pour le régime de l'impôt sur les sociétés",
        space_after_pt=12,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )


def _add_body_intro(document: Any, company: Company) -> None:
    # KAN-18 (Rafael 2026-07-15) : « Dans la lettre d'option IS, il faut retirer le nom de la
    # societe actuellement ecrit en gras, sous l'objet. Il ne doit pas apparaitre a cet
    # endroit. » -> le courrier ouvre directement sur « Madame, Monsieur, ».
    #
    # SUPERSEDE une demande INVERSE d'Albane (retour « mise en forme » 2026-07) : « le NOM DE
    # LA SOCIETE doit figurer EN GRAS dans la PREMIERE LIGNE du courrier » — c'est ce retour
    # qui avait ajoute ce paragraphe. Le retour le PLUS RECENT prime (regle 68) ; le supersede
    # est trace ici et signale a Rafael sur le ticket, jamais arbitre en silence.
    # La denomination reste presente dans la TABLE D'IDENTITE du courrier (elle n'y est pas
    # visee par le ticket, qui ne parle que de l'emplacement « sous l'objet »).
    add_paragraph(document, "Madame, Monsieur,", style_profile=LETTER_WIDE_STYLE_PROFILE)
    add_paragraph(
        document,
        (
            "Nous vous informons que la société civile dont vous trouverez la description "
            "ci-après opte pour le régime de l'Impôt sur les Sociétés, et souhaite que cette "
            "option produise ses effets à compter de l'exercice ouvert dès l'immatriculation."
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )
    add_paragraph(
        document,
        "Cette décision est prise en accord avec les associés participant.",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )
    add_paragraph(
        document,
        (
            "État d'identification de la société et liste des associés au 1er jour du premier "
            "exercice d'option :"
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )


def _add_identification_table(
    document: Any,
    company: Company,
    statuts: StatutsCivilsContext,
) -> None:
    _validate_capital_distribution(statuts)
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.columns[0].width = Cm(6)
    table.columns[1].width = Cm(10)

    capital = _required_text(
        company.capital_social or statuts.capital_social,
        "societe.capital_social",
    )
    denomination = _required_text(company.denomination, "societe.denomination")
    _add_table_row(table, "Dénomination", denomination)
    _add_table_row(table, "Adresse (siège ou principal établissement)", _company_address(company))
    # R3 (Albane 2026-06-30) : societe EN COURS DE CONSTITUTION -> jamais de numero SIREN,
    # toujours la constante (le numero saisi, s'il existe, est volontairement ignore ici).
    _add_table_row(table, "SIREN", SIREN_EN_COURS_IMMATRICULATION)
    # KAN-40 (Rafael 2026-07-27) : accord singulier/pluriel selon le nombre d'associes
    # (« l'associé » pour un seul, « les différents associés » a partir de deux).
    nb_associes = len(statuts.associes)
    # Elision correcte : « de l'associé » (1) / « des différents associés » (N) — jamais « de les ».
    associes_phrase = "de l'associé" if nb_associes == 1 else "des différents associés"
    label = (
        f"Nom, prénom et adresse {associes_phrase} de la société, "
        f"et répartition du capital de {capital} €"
    )
    # KAN-40 : le label ne se REPETE plus a chaque ligne associe (« fusionner les 2 cases en bas
    # a gauche pour eviter une repetition ») -> cellules gauche des lignes associes FUSIONNEES, le
    # label apparait une seule fois. Le NOM de l'associe est en GRAS (cellule valeur).
    left_cells: list[Any] = []
    for index, associe in enumerate(statuts.associes):
        row_cells = table.add_row().cells
        _fill_associe_value_cell(row_cells[1], _associe_table_text(associe, index))
        left_cells.append(row_cells[0])
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
    merged = left_cells[0]
    for cell in left_cells[1:]:
        merged = merged.merge(cell)
    merged.text = label
    merged.paragraphs[0].paragraph_format.space_after = Pt(2)

    add_spacer(document, space_after_pt=12)


def _fill_associe_value_cell(cell: Any, text: str) -> None:
    # KAN-40 : le NOM de l'associe (jusqu'a la 1re virgule : « Monsieur Jean Durand » / « La
    # société X ») en GRAS, le reste (adresse, qualite, parts) en normal.
    name, sep, rest = text.partition(", ")
    paragraph = cell.paragraphs[0]
    paragraph.add_run(name).bold = True
    if sep:
        paragraph.add_run(sep + rest)


def _add_table_row(table: Any, label: str, value: str) -> None:
    row_cells = table.add_row().cells
    row_cells[0].text = label
    row_cells[1].text = value
    for cell in row_cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)


def _validate_capital_distribution(statuts: StatutsCivilsContext) -> None:
    # KAN-2 : coherence de la repartition verifiee UNIQUEMENT quand les donnees sont renseignees.
    # A vide (total non saisi / parts absentes), aucun blocage : chaque quantite sortira en
    # marqueur cote AFFICHAGE. `_required_int` est None-safe (-> 0) ; on ne confronte que des
    # quantites reellement saisies (total > 0 et somme > 0).
    total = _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    associes_total = 0
    for index, associe in enumerate(statuts.associes):
        if associe.parts is None:
            continue
        associes_total += _required_int(
            associe.parts.nb,
            f"statuts_civils.associes[{index}].parts.nb",
        )
    if total and associes_total and associes_total != total:
        raise ValueError(
            "La repartition des parts doit correspondre a statuts_civils.nb_parts_total "
            f"pour {DOCUMENT_CODE}."
        )


def _company_address(company: Company) -> str:
    # KAN-2 : siege non renseigne -> marqueur metier, jamais de crash.
    if company.siege is None:
        return f"(À COMPLÉTER : {libelle_metier('societe.siege')})"
    return _address_display(company.siege, "societe.siege")


def _address_display(address: Address, field_name: str) -> str:
    if address.adresse_affichee:
        return address.adresse_affichee.strip()
    return (
        f"{_required_text(address.num_voie, f'{field_name}.num_voie')} "
        f"{_required_text(address.voie, f'{field_name}.voie')}, "
        f"{_required_text(address.cp, f'{field_name}.cp')} "
        f"{_required_text(address.ville, f'{field_name}.ville')}"
    )


def _parts_label(nb: int | None) -> str:
    # KAN-19 (Rafael) : accord singulier/pluriel selon le NOMBRE DE PARTS (« détenant 1 part »
    # / « détenant 2 parts »). Quantite absente/inconnue (marqueur) -> pluriel par defaut.
    return "part" if nb == 1 else "parts"


def _associe_table_text(associe: StatutsCivilsAssocie, index: int) -> str:
    field_name = f"statuts_civils.associes[{index}]"
    # KAN-2 : parts non renseignees -> marqueur (jamais « 0 parts » affirme), jamais de crash.
    nb_brut = associe.parts.nb if associe.parts is not None else None
    nb_parts = _quantite_parts(nb_brut, "nombre de parts détenues")
    parts_word = _parts_label(nb_brut)  # KAN-19 : accord singulier/pluriel
    if associe.type_personne == "personne_morale":
        return _associe_morale_table_text(associe, field_name, nb_parts, parts_word)
    return _associe_physique_table_text(associe, field_name, nb_parts, parts_word)


def _associe_physique_table_text(
    associe: StatutsCivilsAssocie,
    field_name: str,
    nb_parts: str,
    parts_word: str,
) -> str:
    address = associe.adresse_personnelle
    if associe.adresse_personnelle_affichee:
        address_display = associe.adresse_personnelle_affichee.strip()
    elif address is not None:
        address_display = _address_display(address, f"{field_name}.adresse_personnelle")
    else:
        # KAN-2 : adresse non renseignee -> marqueur metier, jamais de crash.
        address_display = f"(À COMPLÉTER : {libelle_metier(f'{field_name}.adresse_personnelle')})"
    qualite = _required_text(
        associe.parts.qualite_associe or associe.role_statutaire,
        f"{field_name}.parts.qualite_associe",
    )
    return (
        f"{_required_text(associe.civilite_affichage, f'{field_name}.civilite_affichage')} "
        f"{_required_text(associe.prenom, f'{field_name}.prenom')} "
        f"{_required_text(associe.nom, f'{field_name}.nom')}, demeurant au {address_display}, "
        f"{qualite}, détenant {nb_parts} {parts_word}."
    )


def _associe_morale_table_text(
    associe: StatutsCivilsAssocie,
    field_name: str,
    nb_parts: str,
    parts_word: str,
) -> str:
    # KAN-2 : siege non renseigne -> marqueur metier, jamais de crash.
    siege_display = (
        _address_display(associe.siege, f"{field_name}.siege")
        if associe.siege is not None
        else f"(À COMPLÉTER : {libelle_metier(f'{field_name}.siege')})"
    )
    return (
        f"La société {_required_text(associe.denomination, f'{field_name}.denomination')}, "
        f"ayant son siège social au {siege_display}, "
        f"détenant {nb_parts} {parts_word}."
    )


def _add_body_close(document: Any) -> None:
    add_paragraph(
        document,
        (
            "Nous vous remercions de bien vouloir nous confirmer que cette demande est validée, "
            "et vous prions de recevoir, Madame, Monsieur, l'assurance de notre parfaite "
            "considération."
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )


def _add_signature(document: Any, nb_gerants: int) -> None:
    # KAN-19 (Rafael 2026-07-15) : la signature reflète TOUS les gérants -> accord en NOMBRE
    # (« Les gérants » à partir de 2). « gérant » reste au MASCULIN même au pluriel/pour une
    # femme (KAN-23, @All : jamais de féminisation) -> jamais « gérantes ». Signalé à Rafael.
    add_spacer(document, space_after_pt=12)
    add_paragraph(
        document,
        "Les gérants" if nb_gerants > 1 else "Le gérant",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        style_profile=LETTER_WIDE_STYLE_PROFILE,
    )
