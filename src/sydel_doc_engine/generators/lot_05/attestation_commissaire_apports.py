from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import DocumentGenerationContext
from sydel_doc_engine.generators.lot_05.scm_cession_common import (
    mentions_conjoint,
    mentions_partenaire_pacse,
)
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    company_siege_display,
    format_display_date,
    person_address_display,
    person_short_identity,
    professional_entity_presentation,
    quantite_titres,
    required_apport_titres,
    required_apporteur,
    required_commissaire_aux_apports,
    required_societe_cible,
    required_societe_spfpl,
    required_text,
    validate_apport_context,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_hyphen_list_item,
    add_paragraph,
    keep_final_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.dates import format_date_fr

OUTPUT_FILENAME = "attestation_commissaire_apports.docx"


class AttestationCommissaireApportsGenerator:
    """Generateur from-scratch de la designation du commissaire aux apports."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_apport_context(ctx)
        apporteur = required_apporteur(ctx)
        societe_spfpl = required_societe_spfpl(ctx)
        societe_cible = required_societe_cible(ctx)
        apport_titres = required_apport_titres(ctx)
        commissaire = required_commissaire_aux_apports(ctx)
        apporteur_departement_naissance = required_text(
            apporteur.departement_naissance,
            "apporteur.departement_naissance",
        )
        # KAN-2 / M1 : profession de l'INDIVIDU apporteur = champ saisissable (marqueur si vide),
        # jamais l'attribut de TYPE profession_reglementee (« chirurgiens-dentistes » affirme a vide).
        apporteur_profession = required_text(
            apporteur.profession,
            "apporteur.profession",
        )
        cible_forme = required_text(societe_cible.forme_sociale, "societe_cible.forme_sociale")
        cible_name = required_text(societe_cible.denomination, "societe_cible.denomination")
        cible_numero_rcs = required_text(societe_cible.numero_rcs, "societe_cible.numero_rcs")

        docx = new_document()
        add_paragraph(docx, person_signature_header(apporteur))
        add_paragraph(docx, person_address_display(apporteur, "apporteur"))
        add_paragraph(
            docx,
            "Acte de désignation d'un commissaire aux apports",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            space_before_pt=10,
        )
        add_paragraph(
            docx,
            "Le soussigné, "
            f"{person_short_identity(apporteur, 'apporteur')}, "
            f"né le {format_display_date(apporteur.date_naissance, 'apporteur.date_naissance')} "
            f"à {required_text(apporteur.ville_naissance, 'apporteur.ville_naissance')} "
            f"({apporteur_departement_naissance}), "
            f"{apporteur_profession}, "
            f"de nationalité {required_text(apporteur.nationalite, 'apporteur.nationalite')}, "
            f"demeurant au {person_address_display(apporteur, 'apporteur')}, "
            f"{_apporteur_maritale(apporteur)}",
        )
        # R4 (Albane 2026-07-07, « orthographe irréprochable ») : le front pose la forme
        # abregee NON accentuee (« par actions simplifiee ») -> accent restaure a la sortie.
        spfpl_forme = required_text(
            societe_spfpl.forme_sociale, "societe_spfpl.forme_sociale"
        ).replace("simplifiee", "simplifiée")
        add_paragraph(
            docx,
            "seul futur associé de la société "
            f"{required_text(societe_spfpl.denomination, 'societe_spfpl.denomination')} "
            f"{spfpl_forme} "
            f"de {required_text(societe_spfpl.profession, 'societe_spfpl.profession')} "
            "en cours de formation,",
        )
        add_paragraph(docx, "a préalablement exposé et rappelé ce qui suit :")
        add_paragraph(
            docx,
            "Le soussigné a décidé de constituer une société de "
            f"{required_text(societe_spfpl.activite, 'societe_spfpl.activite')} "
            "moyennant l'apport suivant :",
        )
        add_hyphen_list_item(
            docx,
            f"{quantite_titres(apport_titres.nb_parts, 'nombre de parts apportées')} "
            f"parts de la {cible_forme} "
            f"dénommée \"{cible_name}\", "
            f"ayant son siège {company_siege_display(societe_cible, 'societe_cible')}, "
            "immatriculée au RCS de "
            f"{required_text(societe_cible.ville_rcs, 'societe_cible.ville_rcs')} "
            f"sous le numéro {cible_numero_rcs}.",
        )
        add_paragraph(docx, "Il a été convenu ce qui suit :")
        add_paragraph(
            docx,
            "Aux fins de réalisation de cet apport en nature à ladite société, "
            "le soussigné nomme :",
        )
        add_paragraph(
            docx,
            professional_entity_presentation(commissaire, "commissaire_aux_apports")
            + ", en qualité de commissaire aux apports.",
        )
        add_paragraph(
            docx,
            "À l'effet d'établir sous sa responsabilité un rapport sur la valeur "
            "dudit apport en nature, lequel sera annexé aux statuts de la société "
            "conformément à l'article L. 223-9 du Code de commerce.",
        )
        add_paragraph(docx, f"Fait à {ctx.signature.lieu}")
        add_paragraph(docx, f"Le {format_date_fr(ctx.signature.date)}")
        add_paragraph(docx, person_signature_header(apporteur), space_before_pt=12)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(docx)
        docx.save(output_path)
        return output_path


def person_signature_header(person) -> str:
    return (
        f"{required_text(person.prenom, 'apporteur.prenom')} "
        f"{required_text(person.nom, 'apporteur.nom')}"
    )


def _conjoint_nom(person) -> str:
    if person.conjoint is None:
        raise ValueError("apporteur.conjoint est obligatoire.")
    return required_text(person.conjoint.nom, "apporteur.conjoint.nom")


def _apporteur_maritale(apporteur) -> str:
    """Ligne matrimoniale de l'apporteur (comparution du commissaire aux apports).

    Akainu M1 round 2 (2026-07-02) : le menu matrimonial complet (R0702-02) ouvre le
    formulaire SPFPL au NON-MARIE, y compris en operation APPORT. Sans garde, un apporteur
    non marie rendait « <statut> avec (À COMPLÉTER : apporteur.conjoint.nom) » (conjoint
    fantome). On branche via le garde PARTAGE `mentions_conjoint` (R22-02), comme les actes de
    cession : marie -> « <statut> avec <nom conjoint> » BYTE-IDENTIQUE ; sinon -> statut seul.
    """
    situation = required_text(apporteur.situation_maritale, "apporteur.situation_maritale")
    if mentions_conjoint(apporteur.situation_maritale):
        return f"{situation} avec {_conjoint_nom(apporteur)}"
    # Albane 6.3/7.3 (RATIFIE 2026-07-06) : le PARTENAIRE PACSE s'affiche aussi (« <statut>
    # avec <nom> », meme wording que le marie ici — le modele ne porte pas de « sous le régime
    # de »). « Pas de mention sans nom » : si le partenaire n'a pas de nom -> statut seul.
    if mentions_partenaire_pacse(apporteur.situation_maritale):
        conjoint = apporteur.conjoint
        nom = (getattr(conjoint, "nom", None) or "").strip() if conjoint else ""
        if nom:
            return f"{situation} avec {nom}"
    return situation
