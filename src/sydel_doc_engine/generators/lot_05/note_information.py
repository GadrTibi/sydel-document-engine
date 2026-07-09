from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import DocumentGenerationContext
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    OPERATION_APPORT,
    OPERATION_CESSION,
    capital_after_lines,
    company_siege_display,
    montant_avec_euros,
    operation_party,
    person_signature,
    required_int,
    required_societe_cible,
    required_societe_spfpl,
    required_text,
    validate_note_context,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_hyphen_list_item,
    add_paragraph,
    add_spacer,
    new_document,
)

OUTPUT_FILENAME = "note_information.docx"

OPERATION_PHRASES = {
    OPERATION_CESSION: "d’acquérir",
    OPERATION_APPORT: "de recevoir en apport en nature",
}

OPERATION_NOMS = {
    OPERATION_CESSION: "ladite cession",
    OPERATION_APPORT: "ledit apport",
}


class NoteInformationGenerator:
    """Générateur from-scratch de la note d'information SPFPL."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        operation_type = validate_note_context(ctx)
        societe_spfpl = required_societe_spfpl(ctx)
        societe_cible = required_societe_cible(ctx)
        party = operation_party(ctx)
        nb_titres = _operation_nb_titres(ctx, operation_type)

        # Retour Albane 8.1 (2026-07) : « Reprendre la mise en forme selon le modèle d'origine
        # (lisible). » La note était générée À PLAT (10 paragraphes, aucune aération, titres et
        # signature non centrés, corps non justifié) alors que le MODÈLE source
        # (project/source_documents/lot_05/NOTE D'INFORMATION.docx) a 19 paragraphes avec des
        # LIGNES VIDES d'aération, des titres CENTRÉS/gras, un corps JUSTIFIÉ et un bloc signature
        # centré. On reproduit fidèlement cette FORME (le TEXTE et le fix 8.2 sur nb_titres restent
        # inchangés). Cartographie modèle -> code (indices modèle entre crochets) :
        #   [00] titre centré gras · [01] vide · [02]+[03] sous-titre + dénomination centrés gras
        #   [04][05] vides · [06] intro justifiée · [07] vide · [08] intro décomposition justifiée
        #   [09] vide · [10..] lignes tiret justifiées · [12][13] vides · [14] trait signature
        #   justifié gras · [15] nom signataire centré gras · [16] fonction centrée · [17][18] vides
        docx = new_document()
        add_paragraph(docx, "Note d’informations", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        add_spacer(docx)  # [01] ligne vide d'aération (modèle)
        add_paragraph(
            docx,
            "Constitution de la Société",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        add_paragraph(
            docx,
            required_text(societe_spfpl.denomination, "societe_spfpl.denomination"),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        add_spacer(docx)  # [04] ligne vide d'aération (modèle)
        add_spacer(docx)  # [05] ligne vide d'aération (modèle)
        add_paragraph(
            docx,
            (
                f"La {required_text(societe_spfpl.denomination, 'societe_spfpl.denomination')}, "
                "en cours de constitution, dont le siège est situé "
                f"{company_siege_display(societe_spfpl, 'societe_spfpl')}, au capital de "
                # M1 (Akainu doc-entier 2026-07-07) : unite obligatoire apres le montant ;
                # Rafael 2026-07-09 : ACCORDEE (« 1 euro » / « 60 000 euros »), jamais
                # « 1 euros » — montant_avec_euros idempotent.
                f"{montant_avec_euros(required_text(societe_spfpl.capital_social, 'societe_spfpl.capital_social'))}"  # noqa: E501
                ", "
                f"prévoit {OPERATION_PHRASES[operation_type]}, dès son immatriculation, "
                f"{nb_titres} parts de la "
                f"{required_text(societe_cible.denomination, 'societe_cible.denomination')}, "
                f"{required_text(societe_cible.forme_sociale, 'societe_cible.forme_sociale')} "
                f"de {_profession_reglementee(societe_cible)} "
                # M2 (idem) : unite « euros » sur le capital de la cible.
                f"au capital de {montant_avec_euros(_capital_social_cible(societe_cible))} "
                "divisé en "
                f"{required_int(societe_cible.nb_parts_total, 'societe_cible.nb_parts_total')} "
                "parts, dont le siège social est situé "
                f"{company_siege_display(societe_cible, 'societe_cible')}, immatriculée au "
                f"RCS de {required_text(societe_cible.ville_rcs, 'societe_cible.ville_rcs')} "
                f"sous le numéro {_numero_rcs_cible(societe_cible)}."
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        add_spacer(docx)  # [07] ligne vide d'aération (modèle)
        add_paragraph(
            docx,
            (
                f"Après {OPERATION_NOMS[operation_type]}, le capital de la "
                f"{required_text(societe_cible.denomination, 'societe_cible.denomination')} "
                "sera décomposé comme suit :"
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        add_spacer(docx)  # [09] ligne vide avant la liste (modèle)
        for line in capital_after_lines(ctx):
            add_hyphen_list_item(docx, line, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_spacer(docx)  # [12] ligne vide d'aération (modèle)
        add_spacer(docx)  # [13] ligne vide d'aération (modèle)
        # N2 (Rafael 2026-07-09) : le trait de signature etait JUSTIFY (rendu a GAUCHE) alors
        # que le nom du signataire est CENTRE -> desalignement. On CENTRE le trait pour qu'il
        # s'aligne exactement avec le nom du client (meme colonne centree, trait au-dessus du nom).
        add_paragraph(
            docx,
            "________________________",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            space_before_pt=12,
        )
        add_paragraph(
            docx,
            person_signature(party, _party_field_name(operation_type)),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        dirigeant = societe_spfpl.dirigeant
        fonction = required_text(
            dirigeant.fonction if dirigeant else None,
            "societe_spfpl.dirigeant.fonction",
        )
        add_paragraph(
            docx,
            (
                f"{fonction} de la "
                f"{required_text(societe_spfpl.denomination, 'societe_spfpl.denomination')}"
            ),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        add_spacer(docx)  # [17] ligne vide de fin (modèle)
        add_spacer(docx)  # [18] ligne vide de fin (modèle)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        docx.save(output_path)
        return output_path


def _operation_nb_titres(ctx: DocumentGenerationContext, operation_type: str) -> int:
    # Retour Albane 8.2 (2026-07) : en CESSION, la note d'information doit afficher le nombre de
    # PARTS CÉDÉES à la holding (`cession_parts.nb_parts`), pas le nombre d'actions de la SPFPL
    # (`operation_titres.nb_titres`). L'ordre de priorite etait inverse -> on prend cession_parts
    # EN PREMIER en cession. En apport, cession_parts n'existe pas : on garde operation_titres.
    if operation_type == OPERATION_CESSION:
        if ctx.cession_parts is not None and ctx.cession_parts.nb_parts is not None:
            return ctx.cession_parts.nb_parts
        if ctx.operation_titres is not None and ctx.operation_titres.nb_titres is not None:
            return ctx.operation_titres.nb_titres
    else:
        if ctx.operation_titres is not None and ctx.operation_titres.nb_titres is not None:
            return ctx.operation_titres.nb_titres
        if ctx.cession_parts is not None and ctx.cession_parts.nb_parts is not None:
            return ctx.cession_parts.nb_parts
    raise ValueError(
        "operation_titres.nb_titres ou cession_parts.nb_parts est obligatoire pour "
        "CODE-SPFPL-AGR-INFO-001."
    )


def _party_field_name(operation_type: str) -> str:
    if operation_type == OPERATION_CESSION:
        return "cedant"
    return "apporteur"


def _profession_reglementee(societe_cible) -> str:
    return required_text(
        societe_cible.profession_reglementee,
        "societe_cible.profession_reglementee",
    )


def _capital_social_cible(societe_cible) -> str:
    return required_text(
        societe_cible.capital_social,
        "societe_cible.capital_social",
    )


def _numero_rcs_cible(societe_cible) -> str:
    return required_text(
        societe_cible.numero_rcs,
        "societe_cible.numero_rcs",
    )
