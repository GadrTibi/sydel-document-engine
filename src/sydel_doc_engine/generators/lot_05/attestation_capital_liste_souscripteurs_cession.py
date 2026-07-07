from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import DocumentGenerationContext
from sydel_doc_engine.front_app.field_derivations import derive_gender_from_civilite
from sydel_doc_engine.generators.lot_01.civilite import est_titre_professionnel
from sydel_doc_engine.generators.lot_05.attestation_capital_liste_souscripteurs import (
    _ATTESTATION_GROUP_SPACER_PT,
    _adresse,
    _souscripteur_identite,
    _unique_souscripteur,
    _valeur_nominale_action,
)
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    company_siege_display,
    required_capital_souscription,
    required_int,
    required_societe_spfpl,
    required_text,
    validate_cession_context,
)
from sydel_doc_engine.rendering.docx_builder import add_paragraph, add_spacer, new_document
from sydel_doc_engine.utils.grammar import euro_word, subject_line

OUTPUT_FILENAME = "attestation_capital_liste_souscripteurs_cession.docx"


class AttestationCapitalListeSouscripteursCessionGenerator:
    """Generateur from-scratch de l'attestation capital / liste des souscripteurs SPFPL,
    VARIANTE CESSION.

    Retour Albane 11 (2026-07-06) : le bundle SPFPL par CESSION ne produisait AUCUNE
    attestation de capital (seul le bundle APPORT en avait une, DOC-042). Le modele source
    « Attestation sur le capital - cession - liste des souscripteurs.docx » existe : ce
    generateur en reproduit la FORME (Roboto 10, from-scratch comme l'apport) et le WORDING
    VERBATIM.

    Difference de fond avec l'apport (DOC-042) : ici le capital de la holding acquereuse est
    constitue EN NUMERAIRE (pas d'apport de titres en nature). Le corps reprend donc le modele
    cession :
    - « Capital social : X € en numeraire » (au lieu de « X euros ») ;
    - « ... entierement libere et depose dans les livres de la banque » ;
    - « ... a fait un apport de X euros en numeraire » (le montant numeraire = le capital) ;
    - « ... le VERSEMENT de la somme de X euros ... » (l'apport devient un versement) ;
    - le corps s'ouvre par « Je soussigné(e) ... » (accord de genre), absent de l'apport.

    Le montant en numeraire = societe_spfpl.capital_social (la holding cessionnaire est creee
    avec un capital integralement souscrit et libere en numeraire pour acquerir les parts).
    """

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_cession_context(ctx)
        societe_spfpl = required_societe_spfpl(ctx)
        capital = required_capital_souscription(ctx)
        souscripteur = _unique_souscripteur(capital.souscripteurs)
        president = capital.president or souscripteur
        spfpl_name = required_text(societe_spfpl.denomination, "societe_spfpl.denomination")
        spfpl_capital = required_text(societe_spfpl.capital_social, "societe_spfpl.capital_social")
        souscripteur_prenom = required_text(souscripteur.prenom, _souscripteur_field("prenom"))
        souscripteur_nom = required_text(souscripteur.nom, _souscripteur_field("nom"))
        # R3 (Albane 2026-07-07) : « Docteur » (titre) ne porte pas le genre — on
        # accorde alors sur le genre du signataire (= le président dans ce flux) ;
        # une civilité civile (Monsieur/Madame) continue de porter l'accord seule.
        president_civilite = required_text(
            president.civilite_affichage,
            "capital_souscription.president.civilite_affichage",
        )
        president_genre = (
            ctx.personne_signataire.genre
            if est_titre_professionnel(president_civilite)
            else derive_gender_from_civilite(president_civilite)
        )
        president_identite = _souscripteur_identite(
            president,
            "capital_souscription.president",
            genre=president_genre,
        )

        docx = new_document()
        # --- Bloc DESIGNATION de la societe (denomination / forme / activite / siege). ---
        add_paragraph(
            docx,
            spfpl_name,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        add_paragraph(
            docx,
            f"Société par actions simplifiée au capital de {spfpl_capital} euros",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        add_paragraph(
            docx,
            "Société de Participations Financières de Profession Libérale de "
            f"{required_text(societe_spfpl.profession, 'societe_spfpl.profession')}",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        add_paragraph(
            docx,
            f"Siège social : {company_siege_display(societe_spfpl, 'societe_spfpl')}",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        add_spacer(docx, space_after_pt=_ATTESTATION_GROUP_SPACER_PT)
        # --- Titre. ---
        add_paragraph(docx, "ATTESTATION", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        add_paragraph(
            docx,
            "Liste des souscripteurs",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        add_spacer(docx, space_after_pt=_ATTESTATION_GROUP_SPACER_PT)
        # --- Corps (verbatim modele cession). ---
        # [12] « Je soussigné(e) <president>, demeurant <adresse>, atteste que ... »
        # R3 (Albane 2026-07-07) : civilité CIVILE dans le slot « Je soussigné __ »
        # (jamais « Docteur », via _souscripteur_identite).
        add_paragraph(
            docx,
            f"{subject_line(president_genre)} "
            f"{president_identite}, "
            f"demeurant {_adresse(president, 'capital_souscription.president')}, "
            f"atteste que le capital de la société {spfpl_name} "
            "est réparti de la manière suivante\xa0:",
        )
        # [13] Capital social : X € en numeraire (nbsp avant « : » comme le modele).
        add_paragraph(docx, f"Capital social\xa0: {spfpl_capital} € en numéraire")
        # [14] Nombre d'actions: X actions d'un montant de <valeur> euros chacune
        # (verbatim modele : PAS d'espace avant « : »). R4 (Albane 2026-07-07,
        # explicite — supersede la fidélité modèle) : « de 100 d'euro chacune » ->
        # « de 100 euros chacune » ; accord singulier/pluriel via euro_word
        # (« 1 euro » / « 100 euros »), comme la variante SAS.
        valeur_nominale = _valeur_nominale_action(capital)
        add_paragraph(
            docx,
            "Nombre d’actions: "
            f"{required_int(capital.nb_actions_total, 'capital_souscription.nb_actions_total')} "
            f"actions d’un montant de {valeur_nominale} {euro_word(valeur_nominale)} chacune",
        )
        # [15] Repartition : X actions attribuees au Dr <prenom> <nom>, actionnaire unique.
        # Le modele cession ecrit « au Dr [civilite] [prenom] [nom] », mais [civilite] est
        # rempli par civilite_affichage = « Docteur » -> « au Dr Docteur » (doublon, meme
        # famille que le B1 « Docteur Docteur » deja corrige). On aligne sur l'apport DOC-042
        # (« au Dr [prenom] [nom] », sans civilite) -> sortie correcte « au Dr Camille Martin ».
        add_paragraph(
            docx,
            "Répartition\xa0: "
            f"{required_int(souscripteur.nb_actions, _souscripteur_field('nb_actions'))} "
            "actions attribuées au Dr "
            f"{souscripteur_prenom} {souscripteur_nom}, "
            "actionnaire unique",
        )
        # [16] Capital social de X € entierement libere et depose dans les livres de la banque
        # (espace final conserve verbatim du modele).
        add_paragraph(
            docx,
            f"Capital social\xa0de {spfpl_capital} € entièrement libéré et déposé "
            "dans les livres de la banque ",
        )
        # [17] Le Docteur <prenom> <nom> a fait un apport de <capital> euros en numeraire.
        add_paragraph(
            docx,
            f"Le Docteur {souscripteur_prenom} {souscripteur_nom} a fait un apport de "
            f"{spfpl_capital} euros en numéraire.",
            space_after_pt=_ATTESTATION_GROUP_SPACER_PT,
        )
        # [19] Le present etat ... le VERSEMENT de la somme de X euros ... certifie exact ...
        add_paragraph(
            docx,
            f"Le présent état qui constate la souscription d’actions de la société {spfpl_name}, "
            "ainsi que le versement de la somme de "
            f"{spfpl_capital} euros correspondant à la totalité du nominal desdites actions, est "
            "certifié exact, sincère et véritable par le Président, "
            f"{president_identite}",
        )
        # [20]-[23] Fait a / Le / signature. Le modele cession met « Fait a [ville_siege] » ;
        # ctx.signature.lieu est FORCE a la ville du siege cote front (SU3, spfpl_slice) ->
        # equivalent et identique a l'apport (DOC-042).
        add_paragraph(docx, f"Fait à {ctx.signature.lieu}")
        add_paragraph(docx, f"Le {ctx.signature.date.strftime('%d/%m/%Y')}")
        add_paragraph(docx, president_identite)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        docx.save(output_path)
        return output_path


def _souscripteur_field(field_name: str) -> str:
    return f"capital_souscription.souscripteurs[0].{field_name}"
