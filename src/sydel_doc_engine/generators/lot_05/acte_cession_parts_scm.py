# ruff: noqa: E501
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import DocumentGenerationContext
from sydel_doc_engine.generators.lot_04.statuts_sel_exercice_common import (
    situation_maritale_accentuee,
)
from sydel_doc_engine.generators.lot_05.scm_cession_common import (
    acte_signature_prestataire,
    add_body_paragraph,
    add_heading,
    address_display,
    associe_display,
    cedant_display,
    cessionnaire_forme,
    cessionnaire_representant_fonction,
    conjoint_display,
    format_display_date,
    mentions_conjoint,
    mentions_partenaire_pacse,
    partenaire_pacse_clause,
    required_text,
    save_clean_document,
    scm_cedee_address_for_acte,
    scm_cedee_forme,
    validate_acte_context,
)
from sydel_doc_engine.rendering.docx_builder import (
    DEFAULT_STYLE_PROFILE,
    add_framed_title,
    add_paragraph,
    add_signature_table,
    add_spacer,
    keep_final_signature_block_together,
    keep_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.departements import departement_nom
from sydel_doc_engine.utils.grammar import (
    accord_fonction,
    accord_participe_e,
    accord_terme_genre,
    elision_de,
    possessif_singulier,
)

OUTPUT_FILENAME = "acte_cession_parts_scm.docx"

# Mise en forme (Albane 2026-06-17, §13.1) : espace avant chaque grande section
# de l'acte (ORIGINE DE PROPRIETE, DECLARATIONS, CESSION, PRIX, FRAIS...) pour
# aerer un document juge trop serre. Local a l'acte : n'affecte ni le PV ni le
# courrier SCM (qui appellent add_heading sans space_before).
_SECTION_SPACE_BEFORE_PT = 12
# Albane 2026-06-26 §S10 : aerer APRES chaque titre de section (espace sous le heading,
# au-dela du space_after standard de 6 pt).
_SECTION_SPACE_AFTER_PT = 10


def _accord_euro(montant: str | None) -> str:
    """« euro » au singulier si le montant vaut exactement 1, « euros » sinon.

    M1 (Akainu 2026-06-26) : le prix unitaire par part peut valoir 1 (ex. 20 € / 20 parts),
    le verbatim Albane écrit « la part cédée vaut un (1) euro » au singulier — « 1 euros »
    etait une faute d'accord. N'affecte que l'unitaire ; le global garde « euros »."""
    v = (montant or "").strip().replace(" ", "").replace(",", ".")
    if "." in v:
        v = v.rstrip("0").rstrip(".")
    return "euro" if v == "1" else "euros"


def _section_heading(document, text: str) -> None:
    add_heading(
        document,
        text,
        space_before_pt=_SECTION_SPACE_BEFORE_PT,
        space_after_pt=_SECTION_SPACE_AFTER_PT,
    )


# Albane 2026-06-26 §S7 : les marqueurs de partie (« Soussigné de premiere part… »,
# « Soussignee de seconde part… », « Ci-apres denomme LA SOCIETE ») sont alignes a DROITE,
# en GRAS, avec un espace AVANT pour aerer/separer les blocs soussignes.
_PARTY_MARKER_SPACE_BEFORE_PT = 8


def _party_marker(document, text: str) -> None:
    add_paragraph(
        document,
        text,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        bold=True,
        space_before_pt=_PARTY_MARKER_SPACE_BEFORE_PT,
    )


# SC1 (Albane 2026-07-10) : espace entre chaque PARTIE (cedant | cessionnaire | LA SOCIETE)
# pour separer les blocs comparants, comme le modele.
_PARTY_SPACING_PT = 8
# SC1 : le bloc d'identite du cessionnaire est COMPACT (interligne superflu retire) — ses
# lignes se suivent serrees comme un bloc d'adresse, au lieu du space_after standard (6 pt).
_CESSIONNAIRE_LINE_SPACE_AFTER_PT = 2


class _Bullet(str):
    """Marqueur SC5 : element de liste d'article a rendre en PUCE (•)."""

    __slots__ = ()


def _cedant_genre(cedant) -> Gender:
    """Genre du cedant, derive de sa civilite d'affichage deja saisie (« Madame » -> feminin).

    SC3/SC4 (Albane 2026-07-10) : le cedant SCM ne porte pas de champ `genre` ; on lit sa
    civilite comme le PV AGE cession SCM (`_est_feminin`). Sert a accorder « il »->« elle »
    (SC3) et « soussigné »->« soussignée » (SC4) via `accord_terme_genre`."""
    civilite = (cedant.civilite_affichage or "").strip().casefold()
    feminines = ("madame", "mme", "mademoiselle", "mlle")
    return Gender.FEMININ if civilite.startswith(feminines) else Gender.MASCULIN


def _le_soussigne_premiere_part(genre: Gender) -> str:
    """« Le soussigné » -> « La soussignée » si le cedant est une femme (SC4).

    L'article ET le participe s'accordent : masculin -> « Le soussigné » (inchangé)."""
    article = "La" if genre == Gender.FEMININ else "Le"
    return f"{article} {accord_terme_genre('soussigné', genre)}"


def _add_party_line(document, segments, *, space_after_pt: int | None = None) -> None:
    """Paragraphe de corps JUSTIFIE compose de segments (texte, gras).

    SC1 : les NOMS des parties (et de la SCM) sont en GRAS, le reste de la ligne en normal."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(
        DEFAULT_STYLE_PROFILE.standard_space_after_pt if space_after_pt is None else space_after_pt
    )
    for text, bold in segments:
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = bool(bold)


def _add_bullet_item(document, text: str) -> None:
    """Element de liste d'article rendu en vraie PUCE (•) a retrait pendant (SC5)."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.paragraph_format.space_after = Pt(DEFAULT_STYLE_PROFILE.standard_space_after_pt)
    paragraph.add_run("• ")
    paragraph.add_run(text)


class ActeCessionPartsScmGenerator:
    """Generateur from-scratch de l'acte de cession de parts SCM V1."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        scm_cession = validate_acte_context(ctx)
        scm_cedee = scm_cession.scm_cedee
        cessionnaire = scm_cession.cessionnaire
        cedant = scm_cession.cedant
        parts_cedees = scm_cession.parts_cedees
        prix = scm_cession.prix
        if (
            scm_cedee is None
            or cessionnaire is None
            or cedant is None
            or parts_cedees is None
            or prix is None
            or cessionnaire.representant is None
        ):
            raise ValueError("scm_cession est incomplet pour l'acte de cession SCM.")

        cedant_name = cedant_display(cedant)
        # SC3/SC4 (Albane 2026-07-10) : genre du cedant pour accorder « il »->« elle » et
        # « soussigné »->« soussignée » ; masculin -> formes inchangees (byte-neutre).
        genre_cedant = _cedant_genre(cedant)
        # R22-02 : le conjoint n'est mentionne que si le cedant est marie (sinon « divorce
        # avec Madame X » fantome). Regle partagee mentions_conjoint (gold-aligned).
        # Albane 6.3/7.3 (RATIFIE 2026-07-06) : le PARTENAIRE PACSE s'affiche aussi, meme
        # wording « avec {partenaire} » (le PACS n'a pas de « sous le régime de … » ici).
        # « Pas de mention sans nom » : partenaire_pacse_clause renvoie "" si non renseigne.
        # m3 (Akainu 2026-07-12) : le statut matrimonial s'accorde au genre du cedant
        # (« marié » -> « mariée », « divorcé » -> « divorcée », « veuf » -> « veuve »),
        # PARITE avec les statuts (meme helper situation_maritale_accentuee). No-op au masculin ;
        # valeur hors menu renvoyee telle quelle (aucune invention).
        cedant_maritale = situation_maritale_accentuee(
            required_text(
                cedant.situation_maritale, "scm_cession.cedant.situation_maritale"
            ),
            feminine=genre_cedant == Gender.FEMININ,
        )
        if mentions_conjoint(cedant.situation_maritale):
            cedant_maritale_clause = f"{cedant_maritale} avec {conjoint_display(cedant)}"
        elif mentions_partenaire_pacse(cedant.situation_maritale):
            cedant_maritale_clause = (
                f"{cedant_maritale}{partenaire_pacse_clause(cedant.conjoint)}"
            )
        else:
            cedant_maritale_clause = cedant_maritale
        document = new_document()
        # Albane 2026-06-26 §S5/§S6 : aerer le cadre-titre (espace avant/apres, marges internes)
        # et ajouter une 3e ligne au cadre = la denomination de la SCM cedee (objet de la cession).
        scm_denomination = required_text(
            scm_cedee.denomination, "scm_cession.scm_cedee.denomination"
        )
        add_spacer(document, space_after_pt=12)
        add_framed_title(
            document,
            ["CESSION DES PARTS", "DE LA SOCIETE CIVILE DE MOYENS", scm_denomination],
            inner_spacing=True,
        )
        add_spacer(document, space_after_pt=12)

        add_body_paragraph(document, "Entre les soussignés :", bold=True)
        # SC1 : nom du cedant (partie) en GRAS ; le reste de sa ligne d'identite en normal.
        _add_party_line(
            document,
            [
                (cedant_name, True),
                (
                    f", {required_text(cedant.profession, 'scm_cession.cedant.profession')}, "
                    # Accord genre (Albane 2026-07-10, sweep intention) : « né/Inscrit » accordés
                    # au genre du cédant (« née »/« Inscrite » si femme) via accord_terme_genre.
                    f"{accord_terme_genre('né', genre_cedant)} le "
                    f"{format_display_date(cedant.date_naissance, 'scm_cession.cedant.date_naissance')} "
                    f"à {required_text(cedant.ville_naissance, 'scm_cession.cedant.ville_naissance')} "
                    f"({required_text(cedant.departement_naissance, 'scm_cession.cedant.departement_naissance')}), "
                    f"de nationalité {required_text(cedant.nationalite, 'scm_cession.cedant.nationalite')}, "
                    f"demeurant au {required_text(cedant.adresse_affichee, 'scm_cession.cedant.adresse_affichee')}, "
                    f"{cedant_maritale_clause}. "
                    f"{accord_terme_genre('Inscrit', genre_cedant)} au Tableau de l'ordre départemental des "
                    f"{_profession_ordre(ctx, cedant)} "
                    f"{elision_de(departement_nom(required_text(cedant.ordre.departemental if cedant.ordre else None, 'scm_cession.cedant.ordre.departemental')))} "
                    f"sous le numéro {required_text(cedant.ordre.numero if cedant.ordre else None, 'scm_cession.cedant.ordre.numero')} "
                    f"et sous le numéro RPPS {required_text(cedant.numero_rpps, 'scm_cession.cedant.numero_rpps')}.",
                    False,
                ),
            ],
        )
        # SC4 : « Soussigné » de première part accorde au genre du cedant (« Soussignée » si femme).
        _party_marker(
            document,
            f"{accord_terme_genre('Soussigné', genre_cedant)} de première part, ci-après dénommé « LE CÉDANT »,",
        )
        # SC1 : espace entre la partie CEDANT et la partie CESSIONNAIRE.
        add_spacer(document, space_after_pt=_PARTY_SPACING_PT)
        add_body_paragraph(document, "ET :", bold=True)
        # SC1 : denomination du cessionnaire (partie) en GRAS + bloc d'identite COMPACT.
        add_body_paragraph(
            document,
            required_text(cessionnaire.denomination, "scm_cession.cessionnaire.denomination"),
            bold=True,
            space_after_pt=_CESSIONNAIRE_LINE_SPACE_AFTER_PT,
        )
        add_body_paragraph(
            document,
            (
                f"{cessionnaire_forme(ctx, cessionnaire)} au capital de "
                f"{required_text(cessionnaire.capital_social, 'scm_cession.cessionnaire.capital_social')}"
                f"{' €' if ctx.structure == 'SELARL' else ''}"
            ),
            space_after_pt=_CESSIONNAIRE_LINE_SPACE_AFTER_PT,
        )
        add_body_paragraph(
            document,
            f"Ayant son siège au {address_display(cessionnaire.siege, 'scm_cession.cessionnaire.siege')}",
            space_after_pt=_CESSIONNAIRE_LINE_SPACE_AFTER_PT,
        )
        add_body_paragraph(
            document,
            f"En cours d'immatriculation au RCS de {required_text(cessionnaire.ville_rcs, 'scm_cession.cessionnaire.ville_rcs')}",
            space_after_pt=_CESSIONNAIRE_LINE_SPACE_AFTER_PT,
        )
        # SC4-adjacent (Albane 2026-07-10, accord genre) : le représentant du cessionnaire EST
        # le cedant (garde _validate_representant_matches_cedant) ; pour une femme, tout le
        # segment « Représentée par … » s'accorde — possessif « son »->« sa », fonction
        # « gérant »->« gérante » / « président »->« présidente », participe « domicilié »->
        # « domiciliée ». Masculin -> formes inchangées (byte-neutre). Sans quoi la règle de
        # conformité R15 (accord fonction, existante) flague le rendu féminin.
        _representant_fonction = accord_fonction(
            cessionnaire_representant_fonction(ctx, cessionnaire), genre_cedant
        )
        add_body_paragraph(
            document,
            (
                f"Représentée par {possessif_singulier(_representant_fonction, genre_cedant)} "
                f"{_representant_fonction}, "
                f"{cedant_name}, {accord_participe_e('domicilié', genre_cedant)} en cette qualité audit siège."
            ),
            space_after_pt=_CESSIONNAIRE_LINE_SPACE_AFTER_PT,
        )
        _party_marker(document, "Soussignée de seconde part, ci-après dénommé « LE CESSIONNAIRE »,")
        # SC1 : espace entre la partie CESSIONNAIRE et la partie LA SOCIETE + nom de la SCM en GRAS.
        add_spacer(document, space_after_pt=_PARTY_SPACING_PT)
        _add_party_line(
            document,
            [
                ("Ont procédé de la manière suivante à la cession des parts de la Société ", False),
                (required_text(scm_cedee.denomination, "scm_cession.scm_cedee.denomination"), True),
                (".", False),
            ],
        )
        _party_marker(document, "Ci-après dénommé « LA SOCIETE »,")

        _section_heading(document, "IL EST PREALABLEMENT EXPOSE CE QUI SUIT :")
        add_body_paragraph(
            document,
            (
                f"Par les présentes, {cedant_name} cède à la "
                f"{required_text(cessionnaire.denomination, 'scm_cession.cessionnaire.denomination')}, "
                f"{parts_cedees.nb} parts de la {required_text(scm_cedee.denomination, 'scm_cession.scm_cedee.denomination')}, telle que définie ci-après."
            ),
        )
        add_body_paragraph(
            document,
            (
                f"La Société {required_text(scm_cedee.denomination, 'scm_cession.scm_cedee.denomination')}, dont les parts cédées sont l'objet de la présente cession, est une "
                f"{scm_cedee_forme(ctx, scm_cedee)}, au capital de "
                f"{required_text(scm_cedee.capital_social, 'scm_cession.scm_cedee.capital_social')}"
                f"{' €' if ctx.structure == 'SELARL' else ''}, divisé en {scm_cedee.nb_parts_total} parts sociales, dont le siège est situé "
                f"{scm_cedee_address_for_acte(ctx, scm_cedee, cessionnaire)}, immatriculée au RCS de "
                f"{required_text(scm_cedee.ville_rcs, 'scm_cession.scm_cedee.ville_rcs')} sous le n° "
                f"{required_text(scm_cedee.numero_rcs, 'scm_cession.scm_cedee.numero_rcs')} et dont les cogérants sont "
                f"{_cogerants_display(scm_cession)}."
            ),
        )
        _add_origin_property(document, scm_cession, genre_cedant)
        _add_declarations_and_cession(document, scm_cession, cedant_name, genre_cedant)
        _add_price_and_payment(document, ctx, scm_cession, cedant_name)
        _add_source_tail(document, ctx, scm_cession, genre_cedant)
        # Aération (§13.1) : espace avant la zone de clôture / signature.
        add_spacer(document, space_after_pt=12)
        add_body_paragraph(document, f"Fait à {ctx.signature.lieu},")
        add_body_paragraph(
            document,
            f"En {required_text(scm_cession.nombre_exemplaires_lettres, 'scm_cession.nombre_exemplaires_lettres')} exemplaires originaux,",
        )
        add_body_paragraph(document, f"Le {scm_cession.date_acte_affichee or ''}")
        add_spacer(document, space_after_pt=18)
        # SC6 (Albane 2026-07-10) : 2 CADRES pour 2 signataires (une seule RANGEE de 2 cellules),
        # pas 4. Le cadre du HAUT (rangee de titres « Le cédant » / « Le cessionnaire ») est
        # SUPPRIME : le libelle est desormais integre DANS son propre cadre, au-dessus du nom et
        # de la zone manuscrite. Resultat : exactement 2 cadres cote a cote (§S12 : cote-a-cote
        # conserve). Albane laissait aussi le choix « ou pas de cadre » — on garde les cadres,
        # plus clairs, mais reduits a 2.
        cessionnaire_signataire = (
            "Représentée par "
            f"{required_text(cessionnaire.representant.civilite_courte, 'scm_cession.cessionnaire.representant.civilite_courte')} "
            f"{required_text(cessionnaire.representant.prenom, 'scm_cession.cessionnaire.representant.prenom')} "
            f"{required_text(cessionnaire.representant.nom, 'scm_cession.cessionnaire.representant.nom')}"
        )
        signature_table = add_signature_table(
            document,
            [
                [
                    f"Le cédant\n{cedant_name}",
                    f"Le cessionnaire\n{required_text(cessionnaire.denomination, 'scm_cession.cessionnaire.denomination')}\n{cessionnaire_signataire}",
                ],
            ],
            min_row_height_cm=3.0,
        )
        # KAN-36 : bloc signature final solidaire (une seule page). Les signataires sont dans une
        # TABLE : keep_final (paragraphe seul) ne la protege pas (il exclut le dernier paragraphe,
        # celui juste avant la table, et ne pose aucun cantSplit). On solidarise la TABLE (cantSplit
        # sur ses lignes + keepNext sur l'intro « Fait a … » / « Le … ») — convergence @All.
        keep_signature_block_together(document, signature_table)
        keep_final_signature_block_together(document)
        return save_clean_document(document, output_dir, OUTPUT_FILENAME)


def _profession_ordre(ctx: DocumentGenerationContext, cedant) -> str:
    if ctx.structure == "SELARL":
        return "chirurgiens-dentistes"
    return required_text(
        cedant.profession_reglementee_pluriel,
        "scm_cession.cedant.profession_reglementee_pluriel",
    )


def _cogerants_display(scm_cession) -> str:
    # SC2 (Albane 2026-07-10) : « aucun gérant saisi dans la SCM mais 3 noms inventés
    # apparaissent -> laisser VIERGE si non rempli ». L'ancien fallback fabriquait des
    # cogerants a partir des associes ([0], cedant, [2]) — des noms INVENTES (et une source
    # d'IndexError des que < 3 associes). On ne genere plus AUCUN nom par defaut : quand la
    # saisie ne fournit pas de cogerant, on rend le marqueur « a completer » (convention R10,
    # jamais un nom inventé).
    scm_cedee = scm_cession.scm_cedee
    cogerants = [c for c in (scm_cedee.cogerants if scm_cedee else []) if c and str(c).strip()]
    if cogerants:
        return ", ".join(cogerants)
    return required_text(None, "scm_cession.scm_cedee.cogerants")


def _add_origin_property(document, scm_cession, genre: Gender) -> None:
    _section_heading(document, "ORIGINE DE PROPRIETE")
    add_body_paragraph(
        document,
        "Aux termes des statuts le capital social de la SOCIETE est actuellement détenu comme suit :",
    )
    for index, associe in enumerate(scm_cession.associes_avant_cession, start=1):
        parts = associe.parts
        if parts is None:
            raise ValueError("scm_cession.associes_avant_cession.parts est obligatoire.")
        add_body_paragraph(
            document,
            f"{index}° {associe_display(associe, f'scm_cession.associes_avant_cession[{index - 1}]')}, représentant {parts.nb} parts sociales",
        )
    # SC3 (Albane 2026-07-10) : « le cédant déclare qu'il » -> « qu'elle » si le cedant est
    # une femme (accord genre via accord_terme_genre ; masculin -> « qu'il » inchangé).
    add_body_paragraph(
        document,
        f"{cedant_display(scm_cession.cedant)}, le CEDANT, déclare qu'{accord_terme_genre('il', genre)} est propriétaire des parts sociales pour les avoir souscrites à la constitution de la société.",
    )


def _add_declarations_and_cession(document, scm_cession, cedant_name: str, genre: Gender) -> None:
    _section_heading(document, "CECI EXPOSE, IL EST CONVENU CE QUI SUIT :")
    _section_heading(document, "DECLARATIONS")
    # SC5 (Albane 2026-07-10) : les items de la liste de declarations sont rendus en PUCES (•),
    # le chapeau « Le CEDANT déclare : » reste un paragraphe de corps.
    add_body_paragraph(document, "Le CEDANT déclare :")
    # SC4 « partout ailleurs » (Akainu 2026-07-12, M2) : les items de declaration referant au
    # CEDANT s'accordent au genre — pronom (« qu'il » -> « qu'elle ») ET adjectifs (« resident
    # francais » -> « residente francaise »). Codage par INTENTION via accord_terme_genre ;
    # no-op au masculin. Les items impersonnels (parts sociales...) restent inchanges.
    pronom = accord_terme_genre("il", genre)
    resident = accord_terme_genre("résident", genre)
    francais = accord_terme_genre("français", genre)
    for text in [
        f"qu'{pronom} dispose de la pleine capacité juridique d'aliéner ;",
        f"qu'{pronom} est {resident} {francais} ;",
        "que les parts sociales cédées sont libres de tout nantissement et de tout droit quelconque ;",
        "que les parts sociales cédées sont des biens propres.",
    ]:
        _add_bullet_item(document, text)
    _section_heading(document, "CESSION")
    # SC4 : « soussigné de première part » (le CEDANT) accorde au genre ; « soussignée de
    # deuxième part » designe la SOCIETE cessionnaire (feminine) et reste inchangé.
    add_body_paragraph(
        document,
        (
            f"Par les présentes, {cedant_name}, {accord_terme_genre('soussigné', genre)} de première part, cède et transporte sous les garanties ordinaires de fait ou de droit à la société "
            f"{required_text(scm_cession.cessionnaire.denomination, 'scm_cession.cessionnaire.denomination')}, soussignée de deuxième part qui accepte la pleine propriété de "
            f"{scm_cession.parts_cedees.nb} parts de la {required_text(scm_cession.scm_cedee.denomination, 'scm_cession.scm_cedee.denomination')}, numérotées de "
            f"{required_text(scm_cession.parts_cedees.plage, 'scm_cession.parts_cedees.plage')} inclus."
        ),
    )
    _section_heading(document, "PROPRIÉTÉ - JOUISSANCE")
    for text in [
        "Le cessionnaire sera propriétaire des parts cédées et en aura la jouissance à compter de ce jour.",
        "En conséquence, il aura seul droit à tous les dividendes qui seront mis en distribution sur ces parts après cette date.",
        "Le cessionnaire sera subrogé dans tous les droits et obligations attachés à l'actif cédé.",
    ]:
        add_body_paragraph(document, text)


def _add_price_and_payment(
    document,
    ctx: DocumentGenerationContext,
    scm_cession,
    cedant_name: str,
) -> None:
    prix = scm_cession.prix
    _section_heading(document, "PRIX")
    add_body_paragraph(
        document,
        (
            "La présente cession est consentie et acceptée moyennant le prix de "
            f"{required_text(prix.unitaire_lettres, 'scm_cession.prix.unitaire_lettres')} "
            f"({required_text(prix.unitaire, 'scm_cession.prix.unitaire')}) {_accord_euro(prix.unitaire)} par part cédée, soit le prix global de "
            f"{required_text(prix.global_lettres, 'scm_cession.prix.global_lettres')} "
            f"({required_text(prix.global_, 'scm_cession.prix.global')}) euros, payé comptant ce jour à {cedant_name} qui lui reconnaît et lui en donne bonne et valable quittance."
        ),
    )
    _section_heading(document, "PAIEMENT DU PRIX")
    add_body_paragraph(document, "Le prix est payé au moyen d'un prêt bancaire, établi par acte séparé, par virement.")
    credit = scm_cession.credit_vendeur
    if credit is not None and credit.actif:
        add_body_paragraph(document, _credit_vendeur_intro(ctx, credit))
        for text in [
            "Le Vendeur dispense l'Acquéreur de consentir une garantie sur le paiement du crédit-vendeur.",
            "Tout défaut de paiement, même partiel, de toute échéance mensuelle emportera l'exigibilité anticipée du solde du crédit-vendeur dû à cette date, en principal et intérêts, si bon semble au Vendeur, sans qu'il soit besoin d'aucune mise en demeure ou autre formalité.",
            _credit_vendeur_retard(ctx, credit),
            "A défaut de paiement par l'acquéreur, le Cédant pourra faire ordonner par la Justice, la cession des parts sociales, objet des présentes pour lui garantir le paiement du prix.",
        ]:
            add_body_paragraph(document, text)


def _credit_vendeur_intro(ctx: DocumentGenerationContext, credit) -> str:
    if ctx.structure == "SELAS":
        return (
            "Et pour partie d'un crédit-vendeur à hauteur de "
            f"{required_text(credit.montant, 'scm_cession.credit_vendeur.montant')} que les parties ont convenues de solder dans un délai maximum de "
            f"{required_text(credit.duree, 'scm_cession.credit_vendeur.duree')} à compter de la signature des présentes. Le montant annuel en principal du crédit-vendeur sera productif d'un intérêt annuel non capitalisé au taux de "
            f"{required_text(credit.taux, 'scm_cession.credit_vendeur.taux')}."
        )
    return (
        "Et pour partie d'un crédit-vendeur à hauteur de "
        f"{required_text(credit.montant, 'scm_cession.credit_vendeur.montant')} euros que les parties ont convenu de solder dans un délai maximum de "
        f"{required_text(credit.duree, 'scm_cession.credit_vendeur.duree')} ans à compter de la signature des présentes. Le montant annuel en principal du crédit-vendeur sera productif d'un intérêt annuel non capitalisé au taux de "
        f"{required_text(credit.taux, 'scm_cession.credit_vendeur.taux')} %."
    )


def _credit_vendeur_retard(ctx: DocumentGenerationContext, credit) -> str:
    if ctx.structure == "SELAS":
        return (
            "Au terme du délai de "
            f"{required_text(credit.duree, 'scm_cession.credit_vendeur.duree')}, les sommes restant dues porteront de plein droit et sans mise en demeure préalable, un intérêt de retard calculé sur la base du Taux d'intérêt légal publié par la Banque de France, majoré de "
            f"{required_text(credit.majoration_interet_retard, 'scm_cession.credit_vendeur.majoration_interet_retard')} à compter de la date d'échéance de ladite fraction et jusqu'à son paiement effectif. Les intérêts seront calculés au jour le jour et tout mois commencé sera dû en entier."
        )
    return (
        "Au terme du délai de "
        f"{required_text(credit.duree, 'scm_cession.credit_vendeur.duree')} ans, les sommes restant dues porteront de plein droit et sans mise en demeure préalable, un intérêt de retard calculé sur la base du Taux d'intérêt légal publié par la Banque de France, majoré de 3 points à compter de la date d'échéance de ladite fraction et jusqu'à son paiement effectif. Les intérêts seront calculés au jour le jour et tout mois commencé sera dû en entier."
    )


def _add_source_tail(document, ctx: DocumentGenerationContext, scm_cession, genre: Gender) -> None:
    # R9 (Albane 2026-07-07) : la clause de communication au Conseil de l'Ordre nomme le
    # departement de l'Ordre du CEDANT, en NOM avec la preposition correcte (« au Conseil
    # départemental de l'Ordre de Seine-et-Marne »), via la meme convention
    # `elision_de(departement_nom(…))` que la 12.4 SPFPL ratifiee.
    cedant = scm_cession.cedant
    ordre_departement = elision_de(
        departement_nom(
            required_text(
                cedant.ordre.departemental if cedant and cedant.ordre else None,
                "scm_cession.cedant.ordre.departemental",
            )
        )
    )
    sections = [
        (
            "DISPENSE DE GARANTIE D'ACTIF ET DE PASSIF",
            [
                "L'Acquéreur reconnaît avoir eu accès à tout renseignement, et renonce expressément et irrévocablement au bénéfice de toute garantie d'actif et de passif sur les parts cédées, ceci étant une condition substantielle à la conclusion des présentes sans laquelle le Vendeur n'aurait pas contracté.",
                "Les Parties déclarent avoir pris tout renseignement quant aux conséquences de la présente dispense de garantie d'actif et de passif et déclarent en faire leur affaire personnelle. En conséquence, les Parties donnent décharge pure et simple, entière et définitive et sans réserve au Rédacteur en ce qui concerne lesdites conséquences.",
                "Les Parties déclarent en outre faire leur affaire personnelle de l'intégralité des déclarations afférentes à la présente cession.",
            ],
        ),
        (
            "DÉCLARATIONS GÉNÉRALES",
            [
                # SC5 : items d'article en PUCES (•) ; les chapeaux « ... déclare(nt) : » restent
                # des paragraphes de corps. SC4 : « Le soussigné de première part » (le CEDANT)
                # accorde au genre -> « La soussignée » si femme ; le pluriel « Les soussignés »
                # (cedant + societe) reste au masculin generique (parties mixtes).
                "Les soussignés de première et seconde part déclarent, chacun en ce qui le concerne :",
                _Bullet("qu'ils ont la pleine capacité civile pour s'obliger dans le cadre des présentes et de leurs suites et, plus spécialement, qu'ils ne font pas présentement l'objet d'une procédure collective dans le cadre de la loi du 13 juillet 1967 ou de celle du 25 janvier 1985, ni ne sont susceptibles de l'être en raison de leurs professions et fonctions, ni ne sont en état de cessation de paiements ou déconfiture ;"),
                _Bullet("et qu'ils sont résidents français au sens de la réglementation des relations financières avec l'étranger."),
                f"{_le_soussigne_premiere_part(genre)} de première part déclare :",
                _Bullet("qu'il n'existe de son chef ou de celui des précédents propriétaires des parts cédées, aucune restriction d'ordre légal ou contractuel à la libre disposition de celles-ci, notamment par suite de promesses ou offres consenties à des tiers ou de saisies ;"),
                _Bullet("que les parts cédées sont libres de tout nantissement ou promesse de nantissement ;"),
                _Bullet("que la société dont les parts sont présentement cédées n'est pas en cessation de paiements, ni n'a fait l'objet d'une procédure de règlement amiable des entreprises en difficulté ou de redressement et liquidation judiciaire."),
            ],
        ),
        (
            "DÉCLARATION POUR L'ENREGISTREMENT",
            [
                "Pour la perception des droits d'enregistrement, le cédant atteste que les parts, objet de la présente cession, n'assurent pas la jouissance de droits immobiliers.",
                "Le cessionnaire s'engage à supporter tous les frais et droits d'enregistrement relatifs à la cession.",
            ],
        ),
        (
            "FORMALITÉS ET PUBLICITÉ",
            [
                "La présente cession sera signifiée à la société conformément aux dispositions de l'article 1690 du Code Civil. Toutefois, cette signification pourra être remplacée par le dépôt d'un original du présent acte au siège social contre remise par la gérance d'une attestation de ce dépôt.",
                "La gérance de la société se voit confier tous les pouvoirs en vue de remplir les formalités de publicité.",
            ],
        ),
        (
            "AFFIRMATION DE SINCERITE",
            [
                "Les parties affirment sous les peines édictées par l'article 1837 du Code Général des Impôts que le présent acte exprime l'intégralité du prix convenu.",
            ],
        ),
        (
            "COMMUNICATION DU PRESENT CONTRAT AU CONSEIL DE L'ORDRE",
            [
                # R9 (Albane 2026-07-07) : departement de l'Ordre du cedant, en nom.
                "Le présent contrat sera, sans délai, communiqué au Conseil départemental "
                f"de l'Ordre {ordre_departement} en vue de ses observations éventuelles.",
            ],
        ),
        (
            "FRAIS",
            [
                "Les frais, droits et honoraires des présentes et ceux qui en seront la conséquence, seront supportés par le cessionnaire, qui s'y oblige.",
            ],
        ),
        (
            "CONVENTION SUR LA PREUVE - SIGNATURE ELECTRONIQUE",
            [
                "Les Parties consentent expressément la faculté de procéder à la signature du présent acte par le système de signature électronique. Les Parties renoncent en conséquence expressément à signer et obtenir un quelconque acte original de ce dernier.",
                "Les Parties reconnaissent que le présent acte, tel que signé par voie électronique, constitue une preuve valable permettant d'apprécier les droits, les obligations et responsabilités des Parties et le consentement de leurs signataires.",
                f"Le présent acte est signé par chacune des Parties dans le cadre du processus de signature électronique via le service {acte_signature_prestataire(ctx, scm_cession)}.",
            ],
        ),
    ]
    for heading, paragraphs in sections:
        _section_heading(document, heading)
        for text in paragraphs:
            if isinstance(text, _Bullet):
                _add_bullet_item(document, text)
            else:
                add_body_paragraph(document, text)
