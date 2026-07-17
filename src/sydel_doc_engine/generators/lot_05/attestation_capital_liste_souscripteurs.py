from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import CapitalSouscripteur, DocumentGenerationContext
from sydel_doc_engine.generators.lot_01.civilite import civilite_civile
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    company_siege_display,
    elision_de,
    euro_word,
    montant_avec_euros,
    person_short_identity,
    quantite_titres,
    required_apport_titres,
    required_apporteur,
    required_capital_souscription,
    required_societe_cible,
    required_societe_spfpl,
    required_text,
    spfpl_forme_sociale_complete,
    validate_apport_context,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_paragraph,
    add_spacer,
    keep_final_signature_block_together,
    new_document,
)
from sydel_doc_engine.utils.dates import format_date_fr

OUTPUT_FILENAME = "attestation_capital_liste_souscripteurs.docx"

# Mise en forme (Albane, retour « mise en forme » 2026-07) : « ajouter des espaces »
# sur l'attestation capital / liste des souscripteurs. Deux aerations demandees :
# (a) espace APRES la phrase d'apport (« ... a fait un apport de ... euros ») et
#     APRES la ligne « Total des apports » ;
# (b) espace entre le TITRE (denomination), la DESIGNATION de la societe (forme /
#     activite / siege) et le CORPS du texte.
# On aere via des paragraphes-espaceurs (add_spacer) entre les GROUPES et via un
# space_after renforce sur les lignes visees, sans toucher au wording.
_ATTESTATION_GROUP_SPACER_PT = 10

# KAN-2 / B1 (Akainu 2026-07-15) : libellés métier des marqueurs de QUANTITÉ (constantes pour
# éviter l'apostrophe dans une f-string). Une quantité de titres non renseignée sort en
# « (À COMPLÉTER : <libellé>) » (`quantite_titres`), jamais « 0 actions » ni « 600 » inventé.
_LIBELLE_NB_ACTIONS_CAPITAL = "nombre d'actions composant le capital"
_LIBELLE_NB_ACTIONS_ATTRIBUEES = "nombre d'actions attribuées"


class AttestationCapitalListeSouscripteursGenerator:
    """Generateur from-scratch de l'attestation capital SPFPL V1."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_apport_context(ctx)
        societe_spfpl = required_societe_spfpl(ctx)
        societe_cible = required_societe_cible(ctx)
        apporteur = required_apporteur(ctx)
        apport_titres = required_apport_titres(ctx)
        capital = required_capital_souscription(ctx)
        souscripteur = _unique_souscripteur(capital.souscripteurs)
        president = capital.president or souscripteur
        spfpl_name = required_text(societe_spfpl.denomination, "societe_spfpl.denomination")
        spfpl_capital = required_text(societe_spfpl.capital_social, "societe_spfpl.capital_social")
        souscripteur_prenom = required_text(souscripteur.prenom, _souscripteur_field("prenom"))
        souscripteur_nom = required_text(souscripteur.nom, _souscripteur_field("nom"))
        apport_nature = required_text(
            capital.apports_nature_montant,
            "capital_souscription.apports_nature_montant",
        )
        apports_numeraire = required_text(
            capital.apports_numeraire_montant,
            "capital_souscription.apports_numeraire_montant",
        )

        docx = new_document()
        add_paragraph(
            docx,
            spfpl_name,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        add_paragraph(
            docx,
            f"Société par actions simplifiée au capital de {montant_avec_euros(spfpl_capital)}",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        # m1 (Akainu doc-entier 2026-07-09, propage depuis la variante cession) : designation
        # legale COMPLETE (convention P2) — profession au PLURIEL capitalisee + « par actions
        # simplifiée », comme l'acte et le titre des statuts (plus « … de chirurgien-dentiste »
        # au singulier minuscule, sans forme legale).
        add_paragraph(
            docx,
            spfpl_forme_sociale_complete(
                required_text(
                    apporteur.profession_reglementee_pluriel,
                    "apporteur.profession_reglementee_pluriel",
                )
            ),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        add_paragraph(
            docx,
            f"Siège social : {company_siege_display(societe_spfpl, 'societe_spfpl')}",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        # (b) Aeration entre la DESIGNATION de la societe (denomination / forme /
        # activite / siege) et le titre « ATTESTATION » + le corps.
        add_spacer(docx, space_after_pt=_ATTESTATION_GROUP_SPACER_PT)
        add_paragraph(docx, "ATTESTATION", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        add_paragraph(
            docx,
            "Liste des souscripteurs",
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
        )
        # (b) Aeration entre le bloc TITRE (« ATTESTATION » / « Liste des
        # souscripteurs ») et le CORPS du texte.
        add_spacer(docx, space_after_pt=_ATTESTATION_GROUP_SPACER_PT)
        # R3 durci (Rafael 2026-07-07, supersede A26-45/49) : « Docteur »/« Dr » ne sort
        # JAMAIS — tous les slots nommant une personne (désignation, répartition,
        # phrase d'apport, « par le Président, __ », signature) sont CIVILS
        # (Monsieur/Madame, genre du signataire).
        souscripteur_civilite = civilite_civile(
            required_text(
                souscripteur.civilite_affichage,
                _souscripteur_field("civilite_affichage"),
            ),
            ctx.personne_signataire.genre,
        )
        president_identite = _souscripteur_identite(
            president,
            "capital_souscription.president",
            genre=ctx.personne_signataire.genre,
        )
        add_paragraph(
            docx,
            f"{president_identite}, "
            f"demeurant au {_adresse(president, 'capital_souscription.president')}, "
            f"atteste que le capital de la société {spfpl_name} "
            "est réparti de la manière suivante :",
        )
        add_paragraph(docx, f"Capital social : {montant_avec_euros(spfpl_capital)}")
        add_paragraph(
            docx,
            # M1 (Akainu doc-entier 2026-07-07) : accord euro/euros via euro_word (le
            # « euros » fige rendait « 1 euros » pour une valeur nominale de 1) — aligne
            # sur le cousin cession DOC-051 (siloing regle 68 Q4 rattrape).
            "Nombre d’actions : "
            f"{quantite_titres(capital.nb_actions_total, _LIBELLE_NB_ACTIONS_CAPITAL)} "
            f"actions d’un montant {elision_de(str(_valeur_nominale_action(capital)))} "
            f"{euro_word(_valeur_nominale_action(capital))} chacune",
        )
        add_paragraph(
            docx,
            "Répartition : "
            f"{quantite_titres(souscripteur.nb_actions, _LIBELLE_NB_ACTIONS_ATTRIBUEES)} "
            f"actions attribuées à {souscripteur_civilite} "
            f"{souscripteur_prenom} {souscripteur_nom}, "
            "actionnaire unique",
        )
        add_paragraph(docx, "Apports en nature :", bold=True)
        add_paragraph(
            docx,
            f"{person_short_identity(apporteur, 'apporteur')} fait apport de "
            f"{quantite_titres(apport_titres.nb_parts, 'nombre de parts apportées')} parts de la "
            f"{required_text(societe_cible.forme_sociale, 'societe_cible.forme_sociale')} "
            f"dénommée {required_text(societe_cible.denomination, 'societe_cible.denomination')} "
            f"ayant son siège {company_siege_display(societe_cible, 'societe_cible')}, "
            "immatriculée au RCS de "
            f"{required_text(societe_cible.ville_rcs, 'societe_cible.ville_rcs')} "
            f"sous le numéro {required_text(societe_cible.numero_rcs, 'societe_cible.numero_rcs')} "
            f"pour une valeur de {montant_avec_euros(apport_nature)}.",
            # (a) Espace APRES la phrase d'apport (« ... fait apport de ... pour une
            # valeur de ... euros. »).
            space_after_pt=_ATTESTATION_GROUP_SPACER_PT,
        )
        add_paragraph(
            docx,
            f"Total des apports en nature {montant_avec_euros(apport_nature)}",
            # (a) Espace APRES la ligne « Total des apports ».
            space_after_pt=_ATTESTATION_GROUP_SPACER_PT,
        )
        add_paragraph(docx, f"Apports en numéraire : {apports_numeraire}")
        add_paragraph(
            docx,
            f"{souscripteur_civilite} {souscripteur_prenom} {souscripteur_nom} "
            "a fait la totalité des apports en nature.",
        )
        add_paragraph(
            docx,
            f"Le présent état qui constate la souscription d’actions de la société {spfpl_name}, "
            "ainsi que l’apport de la somme de "
            f"{montant_avec_euros(apport_nature)} correspondant à la totalité du nominal desdites actions, est "  # noqa: E501
            "certifié exact, sincère et véritable par le Président, "
            f"{president_identite}.",
        )
        add_paragraph(docx, f"Fait à {ctx.signature.lieu}")
        add_paragraph(docx, f"Le {format_date_fr(ctx.signature.date)}")
        # AT1 (Rafael 2026-07-09) : la ligne de SIGNATURE porte le nom SANS profession (la
        # profession reste dans « par le Président, … » juste au-dessus — une seule mention).
        add_paragraph(
            docx,
            _souscripteur_nom_civil(
                president,
                "capital_souscription.president",
                genre=ctx.personne_signataire.genre,
            ),
        )

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(docx)
        docx.save(output_path)
        return output_path


def _unique_souscripteur(souscripteurs: list[CapitalSouscripteur]) -> CapitalSouscripteur:
    if len(souscripteurs) != 1:
        raise ValueError(
            "capital_souscription.souscripteurs doit contenir exactement un "
            "souscripteur pour CODE-SPFPL-CORE-001."
        )
    return souscripteurs[0]


def _souscripteur_identite(
    souscripteur: CapitalSouscripteur,
    field_name: str,
    *,
    genre: Gender | None,
) -> str:
    # R3 (Albane 2026-07-07) : identité de SLOT de civilité — un titre professionnel
    # (« Docteur »/« Dr ») posé en civilite_affichage est substitué par la civilité
    # CIVILE accordée au genre ; une civilité déjà civile passe inchangée.
    civilite = civilite_civile(
        required_text(souscripteur.civilite_affichage, f"{field_name}.civilite_affichage"),
        genre,
    )
    return (
        f"{civilite} "
        f"{required_text(souscripteur.prenom, f'{field_name}.prenom')} "
        f"{required_text(souscripteur.nom, f'{field_name}.nom')} "
        f"{required_text(souscripteur.profession, f'{field_name}.profession')}"
    )


def _souscripteur_nom_civil(
    souscripteur: CapitalSouscripteur,
    field_name: str,
    *,
    genre: Gender | None,
) -> str:
    """AT1 (Rafael 2026-07-09) : identite de SIGNATURE = « <civilite civile> <prenom> <nom> »,
    SANS la profession. Le bloc signature affichait la profession DEUX FOIS pres du nom (« par le
    Président, <nom> <profession> » puis la ligne de signature « <nom> <profession> ») ; on retire
    la profession de la ligne de signature (le modele source la porte SANS profession) — une seule
    mention subsiste, dans la phrase « par le Président, … »."""
    civilite = civilite_civile(
        required_text(souscripteur.civilite_affichage, f"{field_name}.civilite_affichage"),
        genre,
    )
    return (
        f"{civilite} "
        f"{required_text(souscripteur.prenom, f'{field_name}.prenom')} "
        f"{required_text(souscripteur.nom, f'{field_name}.nom')}"
    )


def _adresse(souscripteur: CapitalSouscripteur, field_name: str) -> str:
    return required_text(
        souscripteur.adresse_personnelle_affichee,
        f"{field_name}.adresse_personnelle_affichee",
    )


def _souscripteur_field(field_name: str) -> str:
    return f"capital_souscription.souscripteurs[0].{field_name}"


def _valeur_nominale_action(capital) -> str:
    return required_text(
        capital.valeur_nominale_action,
        "capital_souscription.valeur_nominale_action",
    )
