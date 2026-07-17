from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import DocumentGenerationContext, SpfplPerson
from sydel_doc_engine.front_app.field_derivations import format_grouped_numeric_value
from sydel_doc_engine.generators.lot_01.civilite import civilite_civile
from sydel_doc_engine.generators.lot_05.scm_cession_common import (
    mentions_conjoint,
    mentions_partenaire_pacse,
    partenaire_pacse_clause,
)
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    company_siege_display,
    elision_de,
    euro_word,
    montant_lettres_avec_unite,
    quantite_titres,
    required_cedant,
    required_cession_parts,
    required_int,
    required_societe_cible,
    required_societe_spfpl,
    required_text,
    validate_cession_context,
)
from sydel_doc_engine.generators.lot_05.spfpl_libelles import libelle_metier
from sydel_doc_engine.rendering.docx_builder import ensure_demeurant_au, new_document_from_model
from sydel_doc_engine.utils.departements import departement_nom

OUTPUT_FILENAME = "acte_cession_parts_spfpl.docx"
_SOURCE_NAME = "Acte_cession_SPFPL_tiers_part_modele.docx"

# SP1/SP3/SP4 (Albane 2026-06-25) : ce generateur etait FROM-SCRATCH (texte code en
# dur, non accentue, ~15 sections, paraphrases vs le modele Albane qui en compte ~30 :
# GAP, RENONCIATION, SIGNIFICATION, AFFIRMATION DE SINCERITE...). « Tout revoir » = le
# rebatir en TOKEN-REPLACEMENT : on CHARGE le modele source (texte legal complet, accents
# et formulations d'Albane preserves par construction) et on remplace les placeholders.
# SP2 : le modele utilise [civilite_cedant] (M./Mme) en civilite et « Dr » abrege en
# repartition — JAMAIS « Le Docteur » en civilite ; la fidelite restaure donc la civilite
# civile demandee. Meme infra que les statuts civils (_source_path + remplacement).

# Lignes de repartition du capital : le modele porte 3 lignes figees
# (« - Dr [prenom] [nom] detenant [parts_personne_N] parts »). On les remplace par la
# repartition DYNAMIQUE (une ligne par associe reel via capital_before_lines).
_REPARTITION_MARKER = "[parts_personne_"

# R8 (Albane 2026-07-07) : « la répartition du capital apparaît deux fois, conserver un seul
# bloc ». Le modele porte DEUX recitals de repartition (EXPOSE PREALABLE « réparti à ce jour
# comme suit » + ORIGINE DE PROPRIETE « actuellement détenu comme suit ») ; on ne garde que
# celui de l'EXPOSE PREALABLE. La phrase d'intro du 2e bloc (ci-dessous) est retiree avec sa
# liste ; la phrase « [cedant] déclare qu'il est propriétaire… » qui suit est CONSERVEE.
# SUPERSEDE la fidelite « 2 recitals comme le modele » (le retour 07-07 prime).
_ORIGINE_REPARTITION_INTRO = "actuellement détenu comme suit"

# R2 (Albane 2026-07-07) : la SPFPL acquereuse n'est PAS immatriculee (le front pose
# numero_rcs="en cours") -> son identite porte « en cours de constitution » (formulation du
# corpus : note d'information SPFPL) et la ligne RCS devient « En cours d’immatriculation au
# RCS de <ville> » (formulation du modele source du contrat d'apport SPFPL), au lieu de
# « Immatriculée au RCS de <ville> sous le numéro en cours ». Un VRAI numero RCS (societe
# deja immatriculee) conserve la ligne du modele, remplie a l'identique. Predicat PARTAGE
# avec l'acte de cession d'ACTIONS (meme bloc identite, meme donnee front).
_RCS_EN_COURS_VALUES = frozenset({"", "en cours"})


def rcs_en_cours(numero_rcs: str | None) -> bool:
    """True si la societe n'est pas encore immatriculee (numero RCS absent ou « en cours »)."""
    return (numero_rcs or "").strip().lower() in _RCS_EN_COURS_VALUES


def _is_repartition_row(text: str, block_open: bool) -> bool:
    """Une ligne de repartition du capital : porte le placeholder fige
    « [parts_personne_N] », ou (bloc deja ouvert) une ligne « - … detenant … »
    residuelle des personnes 2/3 du modele a 3 lignes figees."""
    if _REPARTITION_MARKER in text:
        return True
    return block_open and bool(re.match(r"^\s*-\s", text)) and "detenant" in text.lower()


def _is_section_heading(paragraph) -> bool:
    """A5 : un titre de section de l'acte est un paragraphe NON vide dont chaque run porte le
    souligne (les titres du modele sont en gras+souligne ; les enonciations sont en maigre). Sert
    a fermer le mode « tirets » de DÉCLARATIONS DES PARTIES au titre de section suivant."""
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.underline for run in runs)


def _source_path() -> Path:
    path = Path("project/source_documents/lot_05") / _SOURCE_NAME
    if not path.exists():
        raise ValueError(f"modele source introuvable pour {OUTPUT_FILENAME}: {path}")
    return path


def _date_fr(value: object, field_name: str | None = None) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    text = str(value or "")
    # KAN-2 / M2 : une date manquante NE sort JAMAIS en BLANC quand un libelle est fourni ->
    # marqueur metier lisible « (À COMPLÉTER : date de naissance …) » (comme partout ailleurs).
    if not text.strip() and field_name is not None:
        return f"(À COMPLÉTER : {libelle_metier(field_name)})"
    return text


def _strip_euro(lettres: str) -> str:
    """Retire une unite « euro(s) » deja presente en fin de lettres (fixtures / robustesse)
    pour que l'accord `euro_word` sur le MONTANT ne double pas l'unite (« mille euros euros »).
    Ne touche pas les lettres seules (« mille » -> « mille »)."""
    return re.sub(r"\s+euros?$", "", lettres.strip(), flags=re.IGNORECASE)


# A4 (Rafael 2026-07-09) : dans l'acte de cession, tout NOMBRE ecrit EN LETTRES s'affiche en
# MAJUSCULES (« SOIXANTE MILLE », « CENT », « UN »), l'exemple client etant « SOIXANTE MILLE au
# lieu de soixante mille ». Seule la portion NOMBRE-EN-LETTRES passe en capitales : l'unite
# monetaire (« euro(s) », « centime(s) », « d'euro »), le connecteur « et » et les CHIFFRES entre
# parentheses (« (60 000) », « (1 000 €) ») restent inchanges (l'exemple exclut l'unite). PERIMETRE
# CONFIRME = cet acte uniquement ; une generalisation transverse (regle de conformite) est proposee
# separement, pas appliquee ici.
_MONETARY_UNIT_WORDS = frozenset(
    {
        "euro",
        "euros",
        "centime",
        "centimes",
        "d’euro",
        "d'euro",
        "d’euros",
        "d'euros",
        "et",
    }
)


def _upper_nombre_lettres(phrase: str) -> str:
    """Met en MAJUSCULES la portion NOMBRE-EN-LETTRES d'un fragment monetaire (A4).

    Les mots d'unite monetaire (« euro(s) », « centime(s) », « d'euro », « et ») restent en
    minuscules ; les chiffres et symboles ne portent pas de casse (« (60 000) », « € »
    inchanges). Idempotent (« SOIXANTE MILLE » -> « SOIXANTE MILLE »)."""

    # KAN-2 / B1 : une VALEUR absente sort en marqueur « (À COMPLÉTER : … ) » -> ne JAMAIS le
    # crier en MAJUSCULES (le libelle metier doit rester lisible ; l'A4 ne s'applique qu'aux
    # nombres reellement ecrits en lettres, jamais au marqueur d'un champ vide).
    if "À COMPLÉTER" in phrase:
        return phrase

    def _up(token: str) -> str:
        return token if token.lower() in _MONETARY_UNIT_WORDS else token.upper()

    return " ".join(_up(token) for token in phrase.split(" "))


def _capital_lettres_chiffres_euros(
    lettres: str | None, figure: str, field_name: str
) -> str:
    """Capital en « <lettres> (<chiffres groupes>) euro(s) » (Albane 2026-07-07).

    « soixante mille (60 000) euros » : lettres du champ `_lettres` du contexte (le front pose
    les mots NUS ; `_strip_euro` neutralise une unite deja presente dans une fixture), chiffres
    groupes par milliers via `format_grouped_numeric_value`, unite accordee au montant via
    `euro_word` (« euro » pour 1, « euros » sinon)."""
    lettres_nues = _strip_euro(required_text(lettres, field_name))
    return f"{lettres_nues} ({format_grouped_numeric_value(figure)}) {euro_word(figure)}"


def _person_address(person: SpfplPerson, field_name: str) -> str:
    if person.adresse_personnelle_affichee:
        return person.adresse_personnelle_affichee.strip()
    struct = person.adresse_personnelle
    if struct is not None and struct.adresse_affichee:
        return struct.adresse_affichee.strip()
    return required_text(
        person.adresse_personnelle_affichee,
        f"{field_name}.adresse_personnelle_affichee",
    )


def _repartition_lines(ctx: DocumentGenerationContext) -> list[str]:
    """Repartition du capital : « <civilite civile> <prenom> <nom> detenant N part(s) »,
    une ligne par associe reel.

    R3 « supprimer PARTOUT » (Rafael 2026-07-09) : le modele abregeait « Dr <prenom> <nom> » ;
    « Docteur »/« Dr » n'est jamais une civilite -> on rend la civilite CIVILE (Monsieur/Madame).
    `AssocieCible` ne porte pas de genre -> `civilite_civile(..., None)` = masculin par defaut
    (limite documentee : l'accord au feminin d'un associe cible n'est pas capturable ici). Une
    civilite deja civile (« Monsieur ») est renvoyee inchangee ; une civilite absente retombe sur
    « Monsieur » (defaut du menu SP2)."""
    associes = ctx.associes_cible or []
    if not associes:
        raise ValueError(f"associes_cible est obligatoire pour {OUTPUT_FILENAME}.")
    lines: list[str] = []
    for index, associe in enumerate(associes):
        field = f"associes_cible[{index}]"
        nb_parts = required_int(associe.nb_parts_avant, f"{field}.nb_parts_avant")
        # 12.5(a) (Albane 2026-07-06) : ne PAS emettre la ligne d'un associe/SPFPL
        # detenant 0 part (ex. l'acquereur pre-liste avant la cession) — une ligne
        # « detenant 0 parts » n'a pas de sens dans la repartition du capital.
        if nb_parts <= 0:
            continue
        label = "part" if nb_parts == 1 else "parts"
        if associe.type == "personne_morale":
            who = required_text(associe.denomination, f"{field}.denomination")
        else:
            prenom = required_text(associe.prenom, f"{field}.prenom")
            nom = required_text(associe.nom, f"{field}.nom")
            civilite = civilite_civile(associe.civilite_affichage or "", None) or "Monsieur"
            who = f"{civilite} {prenom} {nom}"
        lines.append(f"{who} détenant {nb_parts} {label}")
    return lines


class ActeCessionPartsSpfplGenerator:
    """Generateur token-replacement de l'acte de cession de parts SPFPL (fidele au
    modele source Albane, cf. SP1-SP4 2026-06-25)."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_cession_context(ctx)
        cedant = required_cedant(ctx)
        cession_parts = required_cession_parts(ctx)
        societe_spfpl = required_societe_spfpl(ctx)
        societe_cible = required_societe_cible(ctx)
        representant = societe_spfpl.representant
        if representant is None:
            raise ValueError("societe_spfpl.representant est obligatoire.")

        replacements = self._build_replacements(
            ctx, cedant, cession_parts, societe_spfpl, societe_cible, representant
        )
        repartition = _repartition_lines(ctx)

        # 12.1 (Albane 2026-07-06) : RESTAURATION DE LA MISE EN FORME (« quasi illisible »). Le
        # generateur reconstruisait le corps dans un `new_document()` a plat (tout JUSTIFY/None,
        # 0 paragraphe vide, aucun gras/centrage) et JETAIT les 145 paragraphes vides d'aeration
        # du modele -> forme perdue. On herite desormais la FORME du modele via
        # `new_document_from_model` (marges, page, styles, header/footer, police charte SYDEL) puis
        # on RECOPIE chaque paragraphe du modele — vides COMPRIS — en reproduisant son alignement,
        # son espacement et le gras/italique/souligne de CHAQUE run, avec token-replacement PAR RUN.
        # Aucun token du modele ne traverse une frontiere de run (verifie 2026-07-06) -> le
        # remplacement par run preserve la forme sans casser les fixes 12.2-12.7 ni l'elision n3.
        # A2 (Rafael 2026-07-09) : la designation du cedant (« Monsieur Prenom Nom, profession,
        # né le … ») etait ENTIEREMENT en gras -> gras superflu. On ne garde le gras que sur le
        # NOM + PRENOM (« Monsieur Prenom Nom »), le reste de la ligne repasse en maigre. Prefixe
        # reconstruit depuis les tokens simples deja resolus dans `replacements`.
        designation_prefix = (
            f"{replacements['[civilite_cedant]']} "
            f"{replacements['[prenom_cedant]']} "
            f"{replacements['[nom_cedant]']}"
        )

        source = Document(str(_source_path()))
        docx = new_document_from_model(_source_path())
        self._render_body(docx, source, replacements, repartition, designation_prefix)

        # SP3 (Akainu M1/M2) : PAS de bloc signature ajoute — le modele source porte DEJA sa
        # ligne de signature (« Dr <cedant> / La société <cessionnaire> / Représentée par … »),
        # rendue fidelement (accentuee) par le token-replacement ci-dessus. Un bloc ajoute la
        # dupliquait ET perdait les accents (« La societe »/« Representee »).
        self._assert_no_residual(docx)
        # Rafael/Albane 2026-07-09 : « demeurant [adresse] » -> « demeurant au [adresse] ».
        ensure_demeurant_au(docx)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        docx.save(output_path)
        return output_path

    def _render_body(  # noqa: C901 (dispatch fidele au modele : recital, A1, A2, A5)
        self,
        docx,
        source,
        replacements: dict[str, str],
        repartition: list[str],
        designation_prefix: str,
    ) -> None:
        """Recopie le corps du modele paragraphe par paragraphe (12.1) en appliquant les
        transformations fideles : UNE seule repartition dynamique (R8), saut de page avant l'acte
        (A1), gras limite au nom+prenom de la designation du cedant (A2), tirets sur les
        enonciations de « DÉCLARATIONS DES PARTIES » (A5)."""
        repartition_block_open = False
        repartition_emitted = False
        # A5 (Rafael 2026-07-09) : dans « DÉCLARATIONS DES PARTIES », chaque enonciation est
        # prefixee d'un tiret (« - … »). Le mode s'ouvre a la phrase d'intro et se ferme au titre
        # de section suivant.
        declarations_mode = False
        for paragraph in source.paragraphs:
            text = paragraph.text
            # R8 (Albane 2026-07-07) : la phrase d'intro du 2e recital (« Aux termes des statuts
            # le capital social de la SOCIETE est actuellement détenu comme suit : ») est retiree
            # avec sa liste — une seule repartition, celle de l'EXPOSE PREALABLE.
            if _ORIGINE_REPARTITION_INTRO in text:
                continue
            if _is_repartition_row(text, repartition_block_open):
                # Premiere ligne du PREMIER bloc de repartition -> emettre la liste dynamique
                # (fidele au modele : paragraphe JUSTIFY prefixe « -\t », une ligne par associe
                # reel). Lignes suivantes du meme bloc fige (personne_2/3) -> absorbees. Tout
                # bloc SUIVANT (ORIGINE DE PROPRIETE) -> entierement absorbe, sans re-emission
                # (R8 : une seule repartition dans l'acte).
                if not repartition_block_open and not repartition_emitted:
                    for line in repartition:
                        self._add_repartition_paragraph(docx, line)
                    repartition_emitted = True
                repartition_block_open = True
                continue
            repartition_block_open = False
            # A1 (Rafael 2026-07-09) : saut de page entre la 1re page de presentation et le debut
            # de l'acte (« ENTRE LES SOUSSIGNES : »).
            if text.strip().startswith("ENTRE LES SOUSSIGNES"):
                docx.add_page_break()
                self._copy_paragraph(docx, paragraph, replacements)
                continue
            # A2 : designation du cedant (nom+prenom en gras, reste maigre).
            if text.startswith("[civilite_cedant] [prenom_cedant] [nom_cedant],") and (
                "[profession_cedant]" in text
            ):
                self._copy_cedant_designation(
                    docx, paragraph, replacements, designation_prefix
                )
                continue
            # A5 : ouverture du mode « tirets » a la phrase d'intro de DÉCLARATIONS DES PARTIES.
            if "chacun en ce qui le concerne" in text:
                declarations_mode = True
                self._copy_paragraph(docx, paragraph, replacements)
                continue
            if declarations_mode:
                if _is_section_heading(paragraph):
                    declarations_mode = False  # fin du bloc : titre de section suivant.
                elif text.strip():
                    self._copy_paragraph(docx, paragraph, replacements, prefix="-\t")
                    continue
            self._copy_paragraph(docx, paragraph, replacements)

        # KAN-36 (Rafael 2026-07-16) : le bloc signature final (« Fait a … / Le … / En N
        # exemplaires / signataires cote a cote / Representee par … ») doit rester ENTIER sur une
        # seule page. keepNext sur chaque paragraphe de ce bloc sauf le dernier (le bloc est en fin
        # d'acte). On repere la DERNIERE occurrence de « Fait a » pour ne pas capter un homonyme.
        rendered = docx.paragraphs
        sig_start = next(
            (
                i
                for i in range(len(rendered) - 1, -1, -1)
                if rendered[i].text.strip().startswith(("Fait à", "Fait a"))
            ),
            None,
        )
        if sig_start is not None:
            for signature_paragraph in rendered[sig_start:-1]:
                signature_paragraph.paragraph_format.keep_with_next = True

    @staticmethod
    def _copy_paragraph_format(new_paragraph, source_paragraph) -> None:
        """Reproduit l'alignement, l'espacement (avant/apres, interligne) et le retrait du
        paragraphe source sur le paragraphe de sortie (12.1)."""
        new_paragraph.alignment = source_paragraph.alignment
        source_format = source_paragraph.paragraph_format
        new_format = new_paragraph.paragraph_format
        new_format.space_before = source_format.space_before
        new_format.space_after = source_format.space_after
        new_format.line_spacing = source_format.line_spacing
        # Retrait recopie DEFENSIVEMENT : certains paragraphes du modele portent une valeur de
        # retrait non entiere en twips (« 708.9999… ») que python-docx ne sait pas relire (leve
        # ValueError). Le retrait n'est pas structurant pour la lisibilite visee par 12.1
        # (alignement + aeration + gras/centrage le sont) -> on ignore un retrait illisible plutot
        # que d'echouer la generation.
        for attribute in ("left_indent", "first_line_indent"):
            try:
                setattr(new_format, attribute, getattr(source_format, attribute))
            except (ValueError, TypeError):
                continue

    @classmethod
    def _copy_paragraph(
        cls, docx, source_paragraph, replacements: dict[str, str], prefix: str = ""
    ) -> None:
        """Recopie un paragraphe du MODELE dans la sortie en preservant sa FORME (12.1).

        Reproduit l'alignement, l'espacement (avant/apres, interligne) et le retrait du
        paragraphe source, puis chaque run avec son gras / italique / souligne, en appliquant
        le token-replacement PAR RUN (aucun token ne traverse une frontiere de run -> la mise
        en forme est preservee sans casser les remplacements). Les paragraphes VIDES du modele
        (aeration) sont recopies tels quels (aucun `continue` : la forme d'aeration est le
        coeur du retour 12.1). Un `prefix` optionnel (ex. « -\\t » pour A5) est emis en maigre
        AVANT les runs du modele."""
        new_paragraph = docx.add_paragraph()
        cls._copy_paragraph_format(new_paragraph, source_paragraph)
        if prefix:
            new_paragraph.add_run(prefix)
        for source_run in source_paragraph.runs:
            new_run = new_paragraph.add_run(_replace(source_run.text, replacements))
            new_run.bold = source_run.bold
            new_run.italic = source_run.italic
            new_run.underline = source_run.underline

    @classmethod
    def _copy_cedant_designation(
        cls, docx, source_paragraph, replacements: dict[str, str], designation_prefix: str
    ) -> None:
        """A2 : recopie la designation du cedant avec le gras limite au NOM + PRENOM.

        Le paragraphe modele est ENTIEREMENT en gras ; on emet « <civilite> <prenom> <nom> » en
        gras puis le RESTE de la ligne (profession, naissance, adresse, situation maritale) en
        maigre. Si le texte rendu ne commence pas par le prefixe attendu (cas degrade), toute la
        ligne repasse en maigre (le gras superflu disparait dans tous les cas)."""
        new_paragraph = docx.add_paragraph()
        cls._copy_paragraph_format(new_paragraph, source_paragraph)
        full_text = _replace(source_paragraph.text, replacements)
        if full_text.startswith(designation_prefix):
            prefix_run = new_paragraph.add_run(designation_prefix)
            prefix_run.bold = True
            reste = full_text[len(designation_prefix):]
            if reste:
                new_paragraph.add_run(reste).bold = False
        else:
            new_paragraph.add_run(full_text).bold = False

    @staticmethod
    def _add_repartition_paragraph(docx, line: str) -> None:
        """Emet une ligne de repartition dynamique, FIDELE au modele : paragraphe JUSTIFY
        prefixe « -\\t » (le modele porte « -\\tDr <nom> détenant N parts » en JUSTIFY),
        au lieu de l'ancienne liste a retrait suspendu qui aplatissait la forme (12.1)."""
        paragraph = docx.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(f"-\t{line}")

    def _build_replacements(
        self, ctx, cedant, cession_parts, societe_spfpl, societe_cible, representant
    ) -> dict[str, str]:
        cible_capital = required_text(societe_cible.capital_social, "societe_cible.capital_social")
        ordre = cedant.ordre
        conjoint = cedant.conjoint
        # 12.4 + Albane 7.4/9.2 : departement de l'Ordre du cedant rendu par le NOM
        # (« Seine-et-Marne ») — calcule UNE fois, reutilise par l'identite, l'expose et la
        # clause de communication au Conseil de l'Ordre (R9, Albane 2026-07-07).
        ordre_departement = departement_nom(
            required_text(ordre.departement if ordre else None, "cedant.ordre.departement")
        )

        # 12.3 (Albane 2026-07-06) : le modele porte « d’[valeur_nominale_part_lettres] de valeur
        # nominale ». Le front pose la valeur nominale en lettres SANS unite (« cent » pour un
        # entier ; la FIGURE « 0,01 » pour un decimal depuis le containment 2026-07-06) -> il
        # manquait « euro(s) ». On compose via le helper partage `montant_lettres_avec_unite`,
        # SEUL point qui calcule la phrase monetaire DEPUIS LA FIGURE : ENTIER -> « cent euros » ;
        # DECIMAL (Albane 7.5) -> « un centime d'euro » (aucun double euro). Fixture portant deja
        # « euros » : le garde `"euro" in lettres` la laisse telle quelle (robustesse).
        valeur_nominale_lettres = required_text(
            societe_cible.valeur_nominale_part_lettres,
            "societe_cible.valeur_nominale_part_lettres",
        )
        if "euro" in valeur_nominale_lettres.lower():
            valeur_nominale_display = valeur_nominale_lettres.strip()
        else:
            valeur_nominale_display = montant_lettres_avec_unite(
                valeur_nominale_lettres,
                societe_cible.valeur_nominale_part,
            )
        # A4 : « cent euros » -> « CENT euros » (nombre en lettres en majuscules, unite intacte).
        valeur_nominale_display = _upper_nombre_lettres(valeur_nominale_display)

        # 12.2 (Albane 2026-07-06) : identite de l'acquereur (SPFPL en cours de constitution).
        # Le front injecte la forme ABREGEE (« par actions simplifiee », non accentuee) dans
        # `forme_sociale` -> la 2e ligne d'identite rendait « par actions simplifiee » au lieu de
        # la forme LEGALE COMPLETE. On rend la forme complete validee (identique au titre des
        # statuts SPFPL, cf. statuts_spfpl_templates) : « Société de Participations Financières de
        # Profession Libérale de <Profession-Plurielle> par actions simplifiée ». La profession
        # plurielle vient du cedant (donnee de dossier), titre-casee (« Chirurgiens-Dentistes »).
        profession_pluriel = required_text(
            cedant.profession_reglementee_pluriel,
            "cedant.profession_reglementee_pluriel",
        )
        forme_sociale_complete_acquereur = (
            "Société de Participations Financières de Profession Libérale de "
            f"{profession_pluriel.title()} par actions simplifiée"
        )
        # R2 (Albane 2026-07-07) : SPFPL acquereuse non immatriculee -> l'identite porte
        # « en cours de constitution » (apres la forme legale complete, comme la note
        # d'information du corpus) et la ligne RCS du modele (« Immatriculée au RCS de
        # <ville> sous le numéro en cours ») est remplacee par la formulation du corpus
        # « En cours d’immatriculation au RCS de <ville> » (modele source du contrat
        # d'apport SPFPL). Un vrai numero RCS conserve la ligne du modele a l'identique.
        ville_rcs_cessionnaire = required_text(
            societe_spfpl.ville_rcs, "societe_spfpl.ville_rcs"
        )
        if rcs_en_cours(societe_spfpl.numero_rcs):
            forme_sociale_complete_acquereur += " en cours de constitution"
            ligne_rcs_cessionnaire = (
                f"En cours d’immatriculation au RCS de {ville_rcs_cessionnaire}"
            )
        else:
            ligne_rcs_cessionnaire = (
                f"Immatriculée au RCS de {ville_rcs_cessionnaire} sous le numéro "
                f"{required_text(societe_spfpl.numero_rcs, 'societe_spfpl.numero_rcs')}"
            )

        # Capitaux en « lettres (chiffres groupes) euros » (Albane 2026-07-07) : le bloc
        # acquereur rendait « Au capital de 60000 » et l'expose « au capital social de 10000
        # divisé » -> on rend « soixante mille (60 000) euros » / « dix mille (10 000) euros ».
        # Lettres depuis les champs `_lettres` du contexte (poses par le front), chiffres
        # groupes via `format_grouped_numeric_value`, unite accordee via `euro_word`.
        # A4 : nombres en lettres en MAJUSCULES (« SOIXANTE MILLE (60 000) euros »).
        capital_cessionnaire_display = _upper_nombre_lettres(
            _capital_lettres_chiffres_euros(
                societe_spfpl.capital_social_lettres,
                required_text(societe_spfpl.capital_social, "societe_spfpl.capital_social"),
                "societe_spfpl.capital_social_lettres",
            )
        )
        capital_cedee_display = _upper_nombre_lettres(
            _capital_lettres_chiffres_euros(
                societe_cible.capital_social_lettres,
                cible_capital,
                "societe_cible.capital_social_lettres",
            )
        )

        # 12.6 (Albane 2026-07-06) / B1 (Akainu 2026-07-06) : prix — accord « euro(s) » sur le
        # MONTANT et ordre LETTRES puis CHIFFRES. Le front pose desormais les lettres SANS unite
        # figee (« un », « mille ») ; on ACCORDE l'unite au montant via `euro_word` (« un euro »
        # pour 1, « mille euros » pour >1) -> plus de « un euros ». `_strip_euro` neutralise une
        # unite deja presente dans une fixture (« mille euros ») pour eviter un double « euro ».
        # (a) PRIX UNITAIRE : modele « [prix_unitaire_part_lettres] ([prix_unitaire_part]) euro
        #     part cedee ». Fragment LISIBLE « <lettres> euro(s) (<chiffres> €) » -> la ligne rend
        #     « prix de un euro (1 €) part cedee » (unite accolee aux lettres, lien euro<->part
        #     preserve, plus de « (1 000) part cedee » sans unite).
        # (b) PRIX TOTAL : modele « [prix_cession] € ([prix_cession_lettres]) » (chiffres puis
        #     lettres). Albane veut LETTRES puis CHIFFRES -> « <lettres> euro(s) (<chiffres> €) ».
        prix_unitaire_chiffres = required_text(
            cession_parts.prix_unitaire, "cession_parts.prix_unitaire"
        )
        prix_unitaire_lettres = _strip_euro(
            required_text(
                cession_parts.prix_unitaire_lettres, "cession_parts.prix_unitaire_lettres"
            )
        )
        prix_unitaire_euro = euro_word(cession_parts.prix_unitaire)
        # A4 : « mille euros (1 000 €) » -> « MILLE euros (1 000 €) ».
        prix_unitaire_fragment = _upper_nombre_lettres(
            f"{prix_unitaire_lettres} {prix_unitaire_euro} ({prix_unitaire_chiffres} €)"
        )
        prix_cession_chiffres = required_text(
            cession_parts.prix_total, "cession_parts.prix_total"
        )
        prix_cession_lettres = _strip_euro(
            required_text(cession_parts.prix_total_lettres, "cession_parts.prix_total_lettres")
        )
        prix_cession_euro = euro_word(cession_parts.prix_total)
        # A4 : « soixante mille euros (60 000 €) » -> « SOIXANTE MILLE euros (60 000 €) ».
        prix_cession_fragment = _upper_nombre_lettres(
            f"{prix_cession_lettres} {prix_cession_euro} ({prix_cession_chiffres} €)"
        )
        # R0702-02 / Akainu M1 (2026-07-02) : le menu matrimonial complet ouvre le cas NON-MARIE.
        # Le modele porte « [situation_maritale_cedant] avec [conjoint...] » (« avec » LITTERAL) ->
        # un non-marie faisait fuiter « célibataire avec (À COMPLÉTER : …) » (conjoint fantome +
        # placeholder shippe). On branche la ligne comme l'acte d'ACTIONS (garde PARTAGE
        # mentions_conjoint, R22-02) : marie -> ligne complete byte-identique ; sinon statut seul.
        # La cle COMBINEE (plus longue) est traitee AVANT les tokens simples par `_replace`, donc
        # elle consomme tout le fragment « [situation_maritale_cedant] avec [conjoint...] » de P36.
        # Albane 6.3/7.3 (RATIFIE 2026-07-06) : le PARTENAIRE PACSE s'affiche aussi (« avec
        # {partenaire} », modele P36 sans regime). « Pas de mention sans nom » :
        # partenaire_pacse_clause -> "" si partenaire non renseigne (pacse nu).
        cedant_maritale = required_text(cedant.situation_maritale, "cedant.situation_maritale")
        conjoint_civilite = conjoint_prenom = conjoint_nom = ""
        if mentions_conjoint(cedant.situation_maritale):
            conjoint_civilite = required_text(
                conjoint.civilite_affichage if conjoint else None,
                "cedant.conjoint.civilite_affichage",
            )
            conjoint_prenom = required_text(
                conjoint.prenom if conjoint else None, "cedant.conjoint.prenom"
            )
            conjoint_nom = required_text(conjoint.nom if conjoint else None, "cedant.conjoint.nom")
            ligne_maritale_cedant = (
                f"{cedant_maritale} avec {conjoint_civilite} {conjoint_prenom} {conjoint_nom}"
            )
        elif mentions_partenaire_pacse(cedant.situation_maritale):
            ligne_maritale_cedant = (
                f"{cedant_maritale}{partenaire_pacse_clause(conjoint)}"
            )
        else:
            ligne_maritale_cedant = cedant_maritale
        # R3 « supprimer PARTOUT » (Rafael 2026-07-09) : la civilite du cedant est CIVILE
        # (Monsieur/Madame accorde au genre). Le token [civilite_cedant] la porte partout ; la
        # ligne de SIGNATURE du modele (« Dr [prenom_cedant] [nom_cedant] », P247) fige « Dr » en
        # LITERAL -> on la consomme via une cle COMBINEE (longest-first) qui rend « <civilite>
        # Prenom Nom », sans « Dr ».
        civilite_cedant = civilite_civile(
            required_text(cedant.civilite_affichage, "cedant.civilite_affichage"),
            cedant.genre,
        )
        repl = {
            # Cle COMBINEE (fragment matrimonial complet) : branche marie/non-marie, byte-identique
            # au modele pour un marie. Placee avant les tokens simples (longest-first, _replace).
            "[situation_maritale_cedant] avec [civilite_conjoint_cedant] "
            "[prenom_conjoint_cedant] [nom_conjoint_cedant]": ligne_maritale_cedant,
            # R3 : ligne signature « Dr [prenom_cedant] [nom_cedant] » -> civilite civile (P247).
            # Cle plus longue que « [civilite_cedant] … » -> traitee en premier par `_replace`.
            "Dr [prenom_cedant] [nom_cedant]": (
                f"{civilite_cedant} "
                f"{required_text(cedant.prenom, 'cedant.prenom')} "
                f"{required_text(cedant.nom, 'cedant.nom')}"
            ),
            # Cedant (personne physique)
            "[civilite_cedant]": civilite_cedant,
            "[prenom_cedant]": required_text(cedant.prenom, "cedant.prenom"),
            "[nom_cedant]": required_text(cedant.nom, "cedant.nom"),
            "[profession_cedant]": required_text(cedant.profession, "cedant.profession"),
            "[date_naissance_cedant]": _date_fr(cedant.date_naissance, "cedant.date_naissance"),
            "[ville_naissance_cedant]": required_text(
                cedant.ville_naissance, "cedant.ville_naissance"
            ),
            "[departement_naissance_cedant]": required_text(
                cedant.departement_naissance, "cedant.departement_naissance"
            ),
            "[nationalite_cedant]": required_text(cedant.nationalite, "cedant.nationalite"),
            "[adresse_cedant]": _person_address(cedant, "cedant"),
            # Fallback (P36 est deja consomme par la cle combinee ci-dessus ; ces tokens simples
            # ne restent que pour robustesse s'ils apparaissaient ailleurs). Valeurs conjoint = ""
            # pour un non-marie (jamais de « (À COMPLÉTER) » ni de conjoint fantome).
            "[situation_maritale_cedant]": cedant_maritale,
            "[numero_rpps_cedant]": required_text(
                ordre.numero_rpps if ordre else None, "cedant.ordre.numero_rpps"
            ),
            # 12.4 (Albane 2026-07-06) : le modele fige « du [departement] » devant le departement
            # de l'Ordre -> « des chirurgiens-dentistes du Paris » (agrammatical). On rend la
            # preposition correcte via `elision_de` (« de Paris », « d’Ardeche ») avec une cle
            # COMBINEE qui consomme le « du » fige du modele (traitee avant le token simple).
            # Albane 7.4/9.2 (2026-07-06) : le departement de l'Ordre s'affiche par le NOM
            # (« Seine-et-Marne »), plus par le numero (« 77 »). `departement_nom` enveloppe la
            # VALEUR ; `elision_de` reste la couche externe (12.4, preposition correcte).
            "du [ordre_departemental_cedant]": elision_de(ordre_departement),
            "du [departement_inscription_societe]": elision_de(ordre_departement),
            "[ordre_departemental_cedant]": ordre_departement,
            "[profession_reglementee]": required_text(
                cedant.profession_reglementee, "cedant.profession_reglementee"
            ),
            "[profession_reglementee_pluriel]": required_text(
                cedant.profession_reglementee_pluriel, "cedant.profession_reglementee_pluriel"
            ),
            "[civilite_conjoint_cedant]": conjoint_civilite,
            "[prenom_conjoint_cedant]": conjoint_prenom,
            "[nom_conjoint_cedant]": conjoint_nom,
            # Societe cedee (cible)
            "[denomination_societe_cedee]": required_text(
                societe_cible.denomination, "societe_cible.denomination"
            ),
            "[forme_sociale_complete]": required_text(
                societe_cible.forme_sociale_complete, "societe_cible.forme_sociale_complete"
            ),
            # Albane 2026-07-07 : capital de la societe cedee en « lettres (chiffres) euros »
            # (« au capital social de dix mille (10 000) euros divisé en 100 parts… »).
            "[capital_social_societe_cedee]": capital_cedee_display,
            "[nb_parts_total_societe_cedee]": quantite_titres(
                societe_cible.nb_parts_total, "nombre total de parts"
            ),
            "[ville_rcs_societe_cedee]": required_text(
                societe_cible.ville_rcs, "societe_cible.ville_rcs"
            ),
            "[numero_rcs_societe_cedee]": required_text(
                societe_cible.numero_rcs, "societe_cible.numero_rcs"
            ),
            "[adresse_siege]": company_siege_display(societe_cible, "societe_cible"),
            "[departement_inscription_societe]": ordre_departement,
            # R9 (Albane 2026-07-07) : la clause « COMMUNICATION DU PRESENT CONTRAT AU CONSEIL
            # DE L'ORDRE » nomme le departement de l'Ordre du CEDANT, en NOM avec la preposition
            # correcte (« au Conseil départemental de l’Ordre de Seine-et-Marne »), via la meme
            # convention `elision_de(departement_nom(…))` que la 12.4 ratifiee. Cle STATIQUE du
            # modele (sans crochets, un seul run) consommee avant les tokens (longest-first).
            "communiqué au Conseil départemental de l’Ordre en vue": (
                "communiqué au Conseil départemental de l’Ordre "
                f"{elision_de(ordre_departement)} en vue"
            ),
            # 12.3 (Albane 2026-07-06) : « d’[valeur…] de valeur nominale ». La valeur porte
            # desormais l'unite accordee (« cent euros » / « un euro », cf. valeur_nominale_display)
            # -> plus de « de cent de valeur nominale ». Elision : « cent euros » commence par une
            # consonne -> « DE cent euros » (pas « d’cent euros »). La cle combinee (plus longue)
            # est traitee en premier.
            "d’[valeur_nominale_part_lettres]": elision_de(valeur_nominale_display),
            "[valeur_nominale_part_lettres]": valeur_nominale_display,
            # Societe cessionnaire (SPFPL acquereur)
            "[denomination_societe_cessionnaire]": required_text(
                societe_spfpl.denomination, "societe_spfpl.denomination"
            ),
            # 12.2 (Albane 2026-07-06) : forme LEGALE COMPLETE de l'acquereur (SPFPL en cours de
            # constitution), pas l'abrege « par actions simplifiee » injecte par le front.
            # R2 (Albane 2026-07-07) : suffixee « en cours de constitution » si non immatriculee.
            "[forme_sociale_acquereur]": forme_sociale_complete_acquereur,
            # Albane 2026-07-07 : capital de l'acquereur en « lettres (chiffres) euros »
            # (« Au capital de soixante mille (60 000) euros »).
            "[capital_social_cessionnaire]": capital_cessionnaire_display,
            # R2 (Albane 2026-07-07) : la ligne RCS du modele est consommee EN BLOC (cle
            # combinee, traitee avant les tokens simples) -> « En cours d’immatriculation au
            # RCS de <ville> » pour la SPFPL en cours, plus jamais « sous le numéro en cours ».
            "Immatriculée au RCS de [ville_rcs_cessionnaire] sous le numéro "
            "[numero_rcs_acquereur]": ligne_rcs_cessionnaire,
            # Fallbacks (la cle combinee ci-dessus consomme la ligne du modele ; robustesse si
            # ces tokens apparaissaient isolement ailleurs).
            "[ville_rcs_cessionnaire]": ville_rcs_cessionnaire,
            "[numero_rcs_acquereur]": required_text(
                societe_spfpl.numero_rcs, "societe_spfpl.numero_rcs"
            ),
            "[adresse_siege_cessionnaire]": company_siege_display(societe_spfpl, "societe_spfpl"),
            "[civilite_acquereur_representant]": required_text(
                representant.civilite_affichage, "representant.civilite_affichage"
            ),
            "[civilite_representant_cessionnaire_courte]": required_text(
                representant.civilite_courte, "representant.civilite_courte"
            ),
            "[prenom_representant_cessionnaire]": required_text(
                representant.prenom, "representant.prenom"
            ),
            "[nom_representant_cessionnaire]": required_text(representant.nom, "representant.nom"),
            "[fonction_acquereur_representant]": required_text(
                representant.fonction, "representant.fonction"
            ),
            # Cession
            "[nb_parts_cedees]": quantite_titres(
                cession_parts.nb_parts, "nombre de parts cédées"
            ),
            # A4 : nombre de parts cedees en lettres en MAJUSCULES (« SOIXANTE (60) parts »).
            "[nb_parts_cedees_lettres]": _upper_nombre_lettres(
                required_text(
                    cession_parts.nb_parts_lettres, "cession_parts.nb_parts_lettres"
                )
            ),
            # 12.6 (Albane 2026-07-06) : cles COMBINEES (traitees avant les tokens simples) pour
            # (a) accorder « euro(s) » au prix UNITAIRE et (b) inverser l'ordre du prix TOTAL en
            # LETTRES puis CHIFFRES. La cle unitaire consomme le « euro » fige du modele (l'unite
            # accordee est deja dans `prix_unitaire_fragment`) ; la cle totale consomme « € (…) ».
            # n3 (Albane 2026-07-06) : ELISION euphonique « de un euro » -> « d’un euro ». Le
            # modele colle un « de » LITTERAL devant chaque prix (« le prix de … », « un prix
            # de … ») ; on etend la cle combinee pour consommer AUSSI ce « de » fige et rendre la
            # preposition via `elision_de` (« d’un euro » pour 1, « de mille euros » pour >1). La
            # cle la plus longue (avec « le prix »/« un prix ») est traitee avant les tokens.
            "le prix de [prix_unitaire_part_lettres] ([prix_unitaire_part]) euro": (
                f"le prix {elision_de(prix_unitaire_fragment)}"
            ),
            "un prix de [prix_cession] € ([prix_cession_lettres])": (
                f"un prix {elision_de(prix_cession_fragment)}"
            ),
            "[prix_unitaire_part]": required_text(
                cession_parts.prix_unitaire, "cession_parts.prix_unitaire"
            ),
            # A4 : fallbacks lettres en MAJUSCULES (les cles combinees ci-dessus consomment
            # deja les occurrences du modele ; ces tokens simples restent pour robustesse).
            "[prix_unitaire_part_lettres]": _upper_nombre_lettres(
                required_text(
                    cession_parts.prix_unitaire_lettres, "cession_parts.prix_unitaire_lettres"
                )
            ),
            "[prix_cession]": required_text(cession_parts.prix_total, "cession_parts.prix_total"),
            "[prix_cession_lettres]": _upper_nombre_lettres(
                required_text(
                    cession_parts.prix_total_lettres, "cession_parts.prix_total_lettres"
                )
            ),
            # Signature
            "[lieu_signature]": required_text(ctx.signature.lieu, "signature.lieu"),
            "[date_signature]": _date_fr(ctx.signature.date),
            # A4 : nombre d'exemplaires en lettres en MAJUSCULES (« En TROIS exemplaires »).
            "[nombre_exemplaires_lettres]": _upper_nombre_lettres(
                required_text(
                    cession_parts.nombre_exemplaires_lettres
                    or (ctx.document.nombre_exemplaires_lettres if ctx.document else None),
                    "cession_parts.nombre_exemplaires_lettres",
                )
            ),
            # 12.7 (Albane 2026-07-06) : phrase de paiement — « par le moyen … ou d’un virement »
            # (incorrect) -> wording valide d'Albane « Le prix est payé au moyen d’un chèque ou
            # virement. ». Texte FIXE du modele : on le remplace mot pour mot (apostrophe courbe
            # U+2019, byte-identique au modele). Aucun crochet -> hors garde anti-résidu.
            "Le prix est payé ce jour par le moyen d’un chèque ou d’un virement.": (
                "Le prix est payé au moyen d’un chèque ou virement."
            ),
        }
        return repl

    @staticmethod
    def _assert_no_residual(docx) -> None:
        full = "\n".join(p.text for p in docx.paragraphs)
        if "[" in full or "]" in full:
            residual = re.findall(r"\[[^\]]+\]", full)
            raise ValueError(f"placeholder source residuel dans {OUTPUT_FILENAME}: {residual}")


def _replace(text: str, replacements: dict[str, str]) -> str:
    out = text
    # Tokens les plus LONGS d'abord : une cle combinee « d’[valeur…] » (gestion de l'elision)
    # doit etre traitee AVANT le token simple « [valeur…] » qu'elle contient.
    for token in sorted(replacements, key=len, reverse=True):
        if token in out:
            out = out.replace(token, replacements[token])
    return out

