from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import DocumentGenerationContext
from sydel_doc_engine.generators.lot_05.spfpl_common import (
    capital_after_lines,
    company_siege_display,
    format_display_date,
    person_display,
    quantite_titres,
    required_cedant,
    required_cession_parts,
    required_int,
    required_societe_cible,
    required_societe_spfpl,
    required_text,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_centered_block,
    add_framed_title,
    add_hyphen_list_item,
    add_paragraph,
)


def add_societe_cible_header(docx, ctx: DocumentGenerationContext) -> None:
    societe_cible = required_societe_cible(ctx)
    siege = societe_cible.siege
    if siege is None:
        raise ValueError("societe_cible.siege est obligatoire pour CODE-SPFPL-AGR-INFO-001.")
    add_centered_block(
        docx,
        [
            # C1 (Rafael 2026-07-09) : uniformiser l'en-tete avec les autres docs — le TITRE
            # de la societe (denomination, 1re ligne) en GRAS (tuple (texte, bold, italic)).
            (required_text(societe_cible.denomination, "societe_cible.denomination"), True, False),
            required_text(societe_cible.forme_sociale, "societe_cible.forme_sociale"),
            f"Au capital de {_capital_social(societe_cible)} euros",
            (
                "Siège social : "
                f"{required_text(siege.num_voie, 'societe_cible.siege.num_voie')} "
                f"{required_text(siege.voie, 'societe_cible.siege.voie')}, "
                f"{required_text(siege.cp, 'societe_cible.siege.cp')} "
                f"{required_text(siege.ville, 'societe_cible.siege.ville')}"
            ),
            (
                "Immatriculée au RCS de "
                f"{required_text(societe_cible.ville_rcs, 'societe_cible.ville_rcs')} "
                "sous le n° "
                f"{required_text(societe_cible.numero_rcs, 'societe_cible.numero_rcs')}"
            ),
        ],
    )


def add_pv_title(docx, middle_line: str, ctx: DocumentGenerationContext) -> None:
    decision_date = ctx.decision.date if ctx.decision else None
    add_framed_title(
        docx,
        [
            "PROCÈS-VERBAL DE",
            middle_line,
            f"DU {format_display_date(decision_date, 'decision.date')}",
        ],
    )


def reunion_intro_lines(ctx: DocumentGenerationContext) -> tuple[str, str]:
    reunion = ctx.reunion
    if reunion is None:
        raise ValueError("reunion est obligatoire pour CODE-SPFPL-AGR-INFO-001.")
    annee = required_text(reunion.annee_lettres, "reunion.annee_lettres")
    # C2 (Rafael 2026-07-09) : la date en lettres apparaissait 2 fois — l'ANNEE etait ecrite
    # dans « L'an <annee>, » PUIS repetee dans « Le <jour mois> <annee>, ». Forme legale
    # canonique : l'annee une seule fois (« L'an <annee>, » / « Le <jour mois>, a <heure>, »).
    # On retire l'annee finale de la 2e ligne UNIQUEMENT si la date se termine par elle
    # (deterministe : si la date ne porte pas l'annee, aucune modif -> pas de doublon introduit).
    jour_mois = _strip_trailing_annee(
        required_text(reunion.date_lettres, "reunion.date_lettres"), annee
    )
    return (
        f"L'an {annee},",
        f"Le {jour_mois}, à {required_text(reunion.heure, 'reunion.heure')},",
    )


def _strip_trailing_annee(date_lettres: str, annee: str) -> str:
    """Retire l'annee finale (C2) de la date en lettres si elle y est repetee.

    « quatorze mai deux mille vingt-six » + annee « deux mille vingt-six » -> « quatorze mai ».
    Comparaison insensible a la casse ; si la date ne se termine PAS par l'annee, on la
    laisse intacte (on ne devine rien, on ne cree pas de doublon)."""
    date_clean = date_lettres.strip()
    annee_clean = annee.strip()
    if annee_clean and date_clean.lower().endswith(annee_clean.lower()):
        date_clean = date_clean[: len(date_clean) - len(annee_clean)].strip()
        date_clean = date_clean.rstrip(",").strip()
    return date_clean


def add_ordre_du_jour(docx, ctx: DocumentGenerationContext) -> None:
    # Akainu M1 (2026-06-25) : la denomination reelle du SPFPL beneficiaire etait ignoree
    # (« la SPFPL » hardcode) -> rendre la denomination du ctx.
    societe_spfpl = required_societe_spfpl(ctx)
    denomination = required_text(societe_spfpl.denomination, "societe_spfpl.denomination")
    # C3 (Rafael 2026-07-09) : les enonciations des decisions (ordre du jour) sont prefixees
    # d'un TIRET « - », comme le PV de nomination (coherence inter-PV). La phrase d'amorce
    # « Dès lors, il est décidé de ce qui suit : » n'est pas une enonciation -> pas de tiret.
    add_hyphen_list_item(docx, f"Agrément d'un nouvel associé, la {denomination} ;")
    add_hyphen_list_item(docx, "Modification corrélative des statuts ;")
    add_hyphen_list_item(docx, "Pouvoirs pour l'accomplissement des formalités.")
    add_paragraph(docx, "Dès lors, il est décidé de ce qui suit :")


def add_resolution_agrement(
    docx,
    ctx: DocumentGenerationContext,
    *,
    subject: str,
) -> None:
    societe_spfpl = required_societe_spfpl(ctx)
    societe_cible = required_societe_cible(ctx)
    cedant = required_cedant(ctx)
    cession_parts = required_cession_parts(ctx)
    # Retour Albane 13 (2026-07) : la mention « numérotées de [plage] inclus » doit être OMISE
    # entièrement quand la plage de parts n'est pas renseignée (sinon la phrase reste incomplète).
    # Plage renseignée -> phrase complète inchangée ; plage vide/None -> on saute la mention.
    plage_mention = _plage_mention(cession_parts)
    add_paragraph(docx, "PREMIÈRE RÉSOLUTION", bold=True, space_before_pt=10)
    add_paragraph(
        docx,
        (
            f"{subject} autorise la cession par {person_display(cedant, 'cedant')} de "
            f"{quantite_titres(cession_parts.nb_parts, 'nombre de parts cédées')} parts sociales "
            "qu'il détient de la "
            f"{required_text(societe_cible.denomination, 'societe_cible.denomination')}, à la "
            f"{required_text(societe_spfpl.denomination, 'societe_spfpl.denomination')}, "
            f"{plage_mention}à compter de ce jour."
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_paragraph(
        docx,
        (
            f"Par conséquent, {subject.lower()} agrée la société "
            f"{required_text(societe_spfpl.denomination, 'societe_spfpl.denomination')} "
            "en qualité de nouvelle associée à compter de ce jour."
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )


def add_article_7_bis(docx, ctx: DocumentGenerationContext, *, subject: str) -> None:
    societe_cible = required_societe_cible(ctx)
    # C4 (Rafael 2026-07-09) : le n° d'article du capital social de la SEL cible etait fige
    # EN DUR (« 7 bis ») aux 2 endroits (la phrase de modification ET l'en-tete du bloc). Il
    # est desormais saisissable au front (ctx.metadata["pv_article_capital_numero"]). Defaut
    # « 7 bis » conserve si non renseigne -> BYTE-NEUTRE sur les dossiers existants.
    article_numero = _article_capital_numero(ctx)
    add_paragraph(docx, "DEUXIÈME RÉSOLUTION", bold=True, space_before_pt=10)
    add_paragraph(
        docx,
        (
            f"En conséquence de la première résolution, {subject.lower()} décide, que "
            f"l'article {article_numero} des statuts sera modifié comme suit, à compter de "
            "ce jour :"
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_paragraph(docx, f"« Article {article_numero} - Capital social")
    add_paragraph(
        docx,
        (
            "Le capital social de la Société est fixé à "
            f"{required_text(societe_cible.capital_social, 'societe_cible.capital_social')} euros "
            f"({_capital_social_lettres(societe_cible)}) "
            "et est divisé en "
            f"{quantite_titres(societe_cible.nb_parts_total, 'nombre total de parts')} "
            "parts sociales d'un montant de "
            f"{_valeur_nominale_part(societe_cible)} "
            "euros chacune de nominal, entièrement libérées, attribuées aux Associés de "
            "la manière suivante :"
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    for line in capital_after_lines(ctx):
        add_hyphen_list_item(docx, line)
    add_paragraph(docx, "»")
    add_paragraph(docx, "Le reste de l'article est inchangé.")


def add_pouvoirs_resolution(docx, *, subject: str) -> None:
    add_paragraph(docx, "TROISIÈME RÉSOLUTION", bold=True, space_before_pt=10)
    add_paragraph(
        docx,
        (
            f"{subject} donne tous pouvoirs au porteur de copies ou d'extraits du présent "
            "procès-verbal pour remplir toutes formalités de droit."
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )


def add_societe_cible_context_sentence(docx, ctx: DocumentGenerationContext, text: str) -> None:
    societe_cible = required_societe_cible(ctx)
    add_paragraph(
        docx,
        text.format(
            denomination=required_text(societe_cible.denomination, "societe_cible.denomination"),
            capital=required_text(societe_cible.capital_social, "societe_cible.capital_social"),
            nb_parts=required_int(societe_cible.nb_parts_total, "societe_cible.nb_parts_total"),
            siege=company_siege_display(societe_cible, "societe_cible"),
        ),
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )


def _capital_social(societe_cible) -> str:
    return required_text(societe_cible.capital_social, "societe_cible.capital_social")


def _article_capital_numero(ctx: DocumentGenerationContext) -> str:
    """C4 : n° d'article du capital social de la SEL cible, saisissable au front via
    `ctx.metadata["pv_article_capital_numero"]`. Defaut « 7 bis » (byte-neutre) si le
    champ est absent ou vide."""
    numero = (ctx.metadata or {}).get("pv_article_capital_numero")
    if numero is None or not str(numero).strip():
        return "7 bis"
    return str(numero).strip()


def _plage_mention(cession_parts) -> str:
    # Retour Albane 13 : plage renseignee -> « numérotées de <plage> inclus » (mention complete,
    # espace final pour enchainer sur « à compter de ce jour. ») ; plage vide/None -> chaine vide
    # (la mention est entierement omise, on ne rend PAS le marqueur « (À COMPLÉTER ...) »).
    plage = cession_parts.plage_parts
    if plage is None or not plage.strip():
        return ""
    return f"numérotées de {plage.strip()} inclus "


def _capital_social_lettres(societe_cible) -> str:
    return required_text(
        societe_cible.capital_social_lettres,
        "societe_cible.capital_social_lettres",
    )


def _valeur_nominale_part(societe_cible) -> str:
    return required_text(
        societe_cible.valeur_nominale_part,
        "societe_cible.valeur_nominale_part",
    )
