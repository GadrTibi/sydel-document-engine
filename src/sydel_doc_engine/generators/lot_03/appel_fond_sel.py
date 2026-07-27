from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.models import (
    CessionAcquereur,
    CessionCabinet,
    CessionDestinataire,
    CessionFinancement,
    CessionVendeur,
    DocumentGenerationContext,
    Mandataire,
)
from sydel_doc_engine.generators.lot_01.civilite import civilite_civile
from sydel_doc_engine.generators.lot_03.bail_appel_common import (
    CABINET_DENTAIRE,
    CABINET_MEDICAL,
    cabinet_type,
    format_display_date,
    required_cession,
    required_text,
    validate_appel_fonds_context,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_centered_amount,
    add_header_logo,
    add_italic_instruction,
    add_paragraph,
    add_right_aligned_lines,
    add_subject_heading,
    keep_final_signature_block_together,
    new_document,
)

OUTPUT_FILENAME = "appel_fond_sel.docx"

# Libelle d'affichage du type de cabinet dans le corps de la lettre. Le type interne est
# normalise ("dentaire"/"medical") ; ici on rend la forme accentuee attendue dans le texte.
_CABINET_TYPE_LABELS = {
    CABINET_DENTAIRE: "dentaire",
    CABINET_MEDICAL: "médical",
}


class AppelFondSelGenerator:
    """Generateur from-scratch de l'appel de fonds SEL."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        validate_appel_fonds_context(ctx)
        cession = required_cession(ctx)
        cabinet_type_label = _CABINET_TYPE_LABELS[cabinet_type(ctx)]
        financement = _required_financement(cession.financement)
        destinataire = financement.destinataire
        cabinet = _required_cabinet(cession.cabinet)
        vendeur = _required_vendeur(cession.vendeur)
        acquereur = _required_acquereur(cession.acquereur)
        # §12.3 — l'appel de fonds est signe par le CONSEILLER SYDEL (le mandataire),
        # pas par le praticien/client. Le mandataire est deja capture cote front
        # (ctx.mandataire, defaut « Jordan ELBAZ »). On NE touche pas a
        # document.signataire (qui reste le praticien, consomme par les statuts).
        signataire = _required_mandataire(ctx.mandataire)

        docx = new_document()
        # Logo SYDEL en header, aligne a DROITE (retour UAT Rafael DOC-008).
        add_header_logo(docx, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        # Retours client 2026-06-11 (tickets 2.11 / 3.2) : les informations
        # bancaires ne bloquent plus la generation. Banque ou destinataire
        # absents -> lignes omises ; montant absent -> zone vide a completer.
        banque = financement.banque
        nom_banque = (banque.nom or "").strip() if banque is not None else ""
        destinataire_label = _destinataire_label(destinataire)
        cabinet_label = required_text(
            cabinet.denomination_ou_adresse_affichee,
            "cession.cabinet.denomination_ou_adresse_affichee",
        )
        vendeur_label = _vendeur_label(vendeur)
        acquereur_label = required_text(
            acquereur.denomination_societe,
            "cession.acquereur.denomination_societe",
        )

        # Bloc banque + lieu/date aligne a DROITE (retour UAT Rafael DOC-008).
        if nom_banque:
            add_paragraph(docx, nom_banque, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        add_paragraph(
            docx,
            (
                f"{required_text(ctx.signature.lieu, 'signature.lieu')}, le "
                f"{format_display_date(ctx.signature.date, 'signature.date')}"
            ),
            alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        if destinataire_label:
            add_italic_instruction(
                docx,
                f"A l’attention de {destinataire_label}",
            )
        add_subject_heading(docx, "Objet : demande de déblocage des fonds")
        add_paragraph(docx, "Cher Monsieur,")
        add_paragraph(
            docx,
            (
                "Nous vous remercions de bien vouloir procéder, ce jour, au déblocage des fonds "
                "d’un montant de :"
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        # Retour Albane 2026-06-17 (ticket lot 2, §12.2) : montant et symbole euros
        # sur la MEME ligne (un seul element centre, pas deux paragraphes).
        add_centered_amount(
            docx,
            [f"{(financement.montant_deblocage or '').strip()} €"],
        )
        add_paragraph(
            docx,
            (
                f"pour la cession du cabinet {cabinet_type_label} exploité au "
                f"{cabinet_label} de {vendeur_label} à la Société {acquereur_label}."
            ),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        add_paragraph(
            docx,
            "Nous vous prions d’agréer, Cher Monsieur, nos salutations distinguées.",
        )
        add_right_aligned_lines(
            docx,
            [
                (
                    f"{required_text(signataire.prenom, 'mandataire.prenom')} "
                    f"{required_text(signataire.nom, 'mandataire.nom')}"
                )
            ],
        )

        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(docx)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / OUTPUT_FILENAME
        docx.save(output_path)
        return output_path


# KAN-2 : un SOUS-OBJET absent NE bloque JAMAIS la generation -> instance vide (ses champs
# sortent en marqueurs « (À COMPLÉTER : … ) » via `required_text`), au lieu de lever. Le flux
# shell fournit toujours ces objets (pas de crash en prod) ; ce durcissement garantit le contrat
# « generable a champs vides » meme si un sous-objet manque.
def _required_financement(financement: CessionFinancement | None) -> CessionFinancement:
    return financement if financement is not None else CessionFinancement()


def _required_cabinet(cabinet: CessionCabinet | None) -> CessionCabinet:
    return cabinet if cabinet is not None else CessionCabinet()


def _required_vendeur(vendeur: CessionVendeur | None) -> CessionVendeur:
    return vendeur if vendeur is not None else CessionVendeur()


def _required_acquereur(acquereur: CessionAcquereur | None) -> CessionAcquereur:
    return acquereur if acquereur is not None else CessionAcquereur()


def _required_mandataire(mandataire: Mandataire | None) -> Mandataire:
    # §12.3 — signataire de l'appel de fonds = le conseiller SYDEL (mandataire).
    return mandataire if mandataire is not None else Mandataire()


def _destinataire_label(destinataire: CessionDestinataire | None) -> str:
    if destinataire is None:
        return ""
    parts = [
        (part or "").strip()
        for part in (destinataire.civilite_affichage, destinataire.prenom, destinataire.nom)
    ]
    return " ".join(part for part in parts if part)


def _vendeur_label(vendeur: CessionVendeur) -> str:
    # R3 « supprimer PARTOUT » (Rafael 2026-07-09) : le vendeur nomme dans « de <vendeur> »
    # (« de Docteur X » -> « de Monsieur/Madame X ») porte la civilite CIVILE, pas le titre
    # professionnel. `civilite_civile` accorde au genre du vendeur (defaut masculin si absent).
    civilite = required_text(
        vendeur.civilite_affichage, "cession.vendeur.civilite_affichage"
    )
    return (
        f"{civilite_civile(civilite, vendeur.genre)} "
        f"{required_text(vendeur.prenom, 'cession.vendeur.prenom')} "
        f"{required_text(vendeur.nom, 'cession.vendeur.nom')}"
    )
