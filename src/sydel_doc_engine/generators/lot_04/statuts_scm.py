from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    Company,
    DocumentGenerationContext,
    StatutsCivilsApport,
    StatutsCivilsAssocie,
    StatutsCivilsContext,
    StatutsCivilsParts,
    StatutsCivilsRepresentant,
)
from sydel_doc_engine.generators.lot_04.annexe_filter import is_creation_fee_annexe_line
from sydel_doc_engine.generators.lot_04.statuts_sel_exercice_common import (
    statuts_output_filename,
)
from sydel_doc_engine.generators.lot_05.spfpl_libelles import libelle_metier
from sydel_doc_engine.rendering.docx_builder import (
    add_paragraph,
    add_statuts_article_heading,
    add_statuts_body_paragraph,
    add_statuts_part_heading,
    add_statuts_signature_block,
    add_statuts_title_box,
    keep_final_signature_block_together,
    new_document_from_model,
)
from sydel_doc_engine.utils.dates import format_birthdate_fr
from sydel_doc_engine.utils.grammar import capitalize_first, euro_word, montant_avec_euros

DOCUMENT_CODE = "CODE-STATUTS-SCM-001"
MAX_ASSOCIES = 6
SOURCE_PATH = Path("project/source_documents/lot_04/Statuts SCM.docx")
OUTPUT_FILENAME = "statuts_scm.docx"

ASSOCIATE_SLICE = (25, 42)
APPORT_SLICE = (81, 90)
CAPITAL_SLICE = (95, 100)
SIGNATURE_SLICE = (262, 269)


class StatutsScmGenerator:
    """Generateur from-scratch des statuts SCM V1."""

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        data = _ResolvedStatutsScm.from_context(ctx)
        source_doc = Document(SOURCE_PATH)
        # R1 (Albane 2026-06-30) : on HERITE la forme du modele SCM (page custom, marges,
        # styles nommes, header avec le LOGO SYDEL deja present, footer vide) au lieu de
        # recopier le texte dans un new_document() au profil SYDEL qui ECRASAIT la forme.
        # Comme le modele porte deja son logo en header, on NE rappelle PAS add_header_logo
        # (sinon double logo) et on N'ecrase PAS le footer (vide dans le modele).
        output_doc = new_document_from_model(SOURCE_PATH)

        replacements = data.replacements()
        skip_until = -1
        for index, paragraph in enumerate(source_doc.paragraphs):
            if index < skip_until:
                continue
            if index == ASSOCIATE_SLICE[0]:
                _add_associate_block(output_doc, data)
                skip_until = ASSOCIATE_SLICE[1]
                continue
            if index == APPORT_SLICE[0]:
                _add_apport_block(output_doc, data)
                skip_until = APPORT_SLICE[1]
                continue
            if index == CAPITAL_SLICE[0]:
                _add_capital_block(output_doc, data)
                skip_until = CAPITAL_SLICE[1]
                continue
            if index == SIGNATURE_SLICE[0]:
                _add_signature_block(output_doc, data)
                skip_until = SIGNATURE_SLICE[1]
                continue

            text = paragraph.text.strip()
            if not text:
                continue
            rendered = _replace_placeholders(text, replacements)
            if is_creation_fee_annexe_line(rendered):  # O24-01 : annexe sans frais cabinet création
                continue
            _add_rendered_paragraph(output_doc, rendered)

        full_text = "\n".join(paragraph.text for paragraph in output_doc.paragraphs)
        if "[" in full_text or "]" in full_text:
            raise ValueError(f"placeholder source residuel dans le rendu {DOCUMENT_CODE}.")

        output_dir.mkdir(parents=True, exist_ok=True)
        # Retour Rafael 2026-07-07 : TOUS les statuts sont nommes
        # « Statuts <denomination>.docx » (helper partage, fallback historique si vide).
        output_path = output_dir / statuts_output_filename(data.denomination, OUTPUT_FILENAME)
        # KAN-36 : bloc signature final solidaire (une seule page).
        keep_final_signature_block_together(output_doc)
        output_doc.save(output_path)
        return output_path


class _ResolvedStatutsScm:
    def __init__(
        self,
        *,
        statuts: StatutsCivilsContext,
        denomination: str,
        denomination_courte: str,
        forme_sociale: str,
        siege_num_voie: str,
        siege_voie: str,
        siege_cp: str,
        siege_ville: str,
        signature_lieu: str,
        signature_date: str,
        associes: list[StatutsCivilsAssocie],
    ) -> None:
        self.statuts = statuts
        self.denomination = denomination
        self.denomination_courte = denomination_courte
        self.forme_sociale = forme_sociale
        self.siege_num_voie = siege_num_voie
        self.siege_voie = siege_voie
        self.siege_cp = siege_cp
        self.siege_ville = siege_ville
        self.signature_lieu = signature_lieu
        self.signature_date = signature_date
        self.associes = associes

    @classmethod
    def from_context(cls, ctx: DocumentGenerationContext) -> _ResolvedStatutsScm:
        # INVARIANT DE ROUTAGE conserve : le SELECTEUR de structure reste une garde fiable
        # (c'est le TYPE de dossier, pas un champ saisi) meme a formulaire vide.
        if ctx.structure != "SCM":
            raise ValueError(f"dossier.structure doit etre SCM pour {DOCUMENT_CODE}.")
        # KAN-2 : un OBJET requis absent NE bloque JAMAIS -> instance vide (ses champs sortent en
        # marqueurs « (À COMPLÉTER : … ) » via les helpers), au lieu de lever.
        societe = ctx.societe if ctx.societe is not None else Company()
        siege = societe.siege if societe.siege is not None else Address()
        statuts = ctx.statuts_civils if ctx.statuts_civils is not None else StatutsCivilsContext()
        # KAN-2 : le type n'est verifie que s'il est REELLEMENT saisi (sinon champ vide -> pas de
        # blocage). A vide, le routage a deja garanti la structure SCM ci-dessus.
        statuts_type = (statuts.type or "").strip().lower()
        if statuts_type and statuts_type != "scm":
            raise ValueError(f"statuts_civils.type doit etre scm pour {DOCUMENT_CODE}.")

        associes = list(statuts.associes)
        _validate_associes(associes)
        _validate_capital_totals(statuts, associes)
        _validate_statuts_fields(statuts)
        _validate_signatures(associes)

        return cls(
            statuts=statuts,
            denomination=_required_text(societe.denomination, "societe.denomination"),
            denomination_courte=_required_text(
                societe.denomination_courte,
                "societe.denomination_courte",
            ),
            forme_sociale=_required_text(
                statuts.forme_sociale or societe.forme_sociale,
                "statuts_civils.forme_sociale",
            ),
            siege_num_voie=_required_text(siege.num_voie, "societe.siege.num_voie"),
            siege_voie=_required_text(siege.voie, "societe.siege.voie"),
            siege_cp=_required_text(siege.cp, "societe.siege.cp"),
            siege_ville=_required_text(siege.ville, "societe.siege.ville"),
            signature_lieu=_required_text(ctx.signature.lieu, "signature.lieu"),
            signature_date=_format_display_date(ctx.signature.date, "signature.date"),
            associes=associes,
        )

    def replacements(self) -> dict[str, str]:
        depot = self.statuts.capital_depot
        capital_social = _required_text(
            self.statuts.capital_social,
            "statuts_civils.capital_social",
        )
        return {
            "[denomination_societe]": self.denomination,
            "[denomination_societe_courte]": self.denomination_courte,
            "[forme_sociale]": self.forme_sociale,
            # Rafael 2026-07-09 (R12/devise, en-tete) : le modele tokenise a perdu
            # l'espace (« [capital_social]euros ») -> cle COMBINEE traitee AVANT la
            # cle nue (ordre du dict), qui restaure « 1 000 euros ». Akainu M1 2026-07-09 :
            # accord euro/euros via montant_avec_euros (capital=1 -> « 1 euro », plus
            # jamais « 1 euros »), aligne sur le corps de l'art.6 et « Au capital social de ».
            "[capital_social]euros": montant_avec_euros(capital_social),
            "[capital_social]": capital_social,
            "[capital_lettres]": _required_text(
                self.statuts.capital_social_lettres,
                "statuts_civils.capital_social_lettres",
            ),
            "[nb_parts]": _quantite_parts(
                self.statuts.nb_parts_total, "statuts_civils.nb_parts_total"
            ),
            "[valeur_nominale_part]": _required_text(
                self.statuts.valeur_nominale_part,
                "statuts_civils.valeur_nominale_part",
            ),
            "[num_voie_siege]": self.siege_num_voie,
            "[voie_siege]": self.siege_voie,
            "[cp_siege]": self.siege_cp,
            "[ville_siege]": self.siege_ville,
            "[nom_banque]": _required_text(
                depot.banque_nom if depot else None,
                "statuts_civils.capital_depot.banque_nom",
            ),
            "[adresse_banque]": _required_text(
                depot.banque_adresse if depot else None,
                "statuts_civils.capital_depot.banque_adresse",
            ),
            "[lieu_signature]": self.signature_lieu,
        }


def _add_associate_block(document, data: _ResolvedStatutsScm) -> None:
    for index, associe in enumerate(data.associes):
        if index:
            add_paragraph(document, "ET", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        if _is_morale(associe):
            _add_morale_identity(document, associe)
        else:
            _add_physical_identity(document, associe)


def _add_apport_block(document, data: _ResolvedStatutsScm) -> None:
    # Rafael 2026-07-09 (SCM art. 6, devise automatique) : lettres + « euros »
    # (accord euro/euros au montant) et chiffres + « € », derives par le moteur —
    # l'utilisateur ne redige jamais l'unite.
    for associe in data.associes:
        apport = _required_apport(associe)
        montant = _required_text(apport.montant, "associes[].apport.montant")
        add_paragraph(
            document,
            f"{_apport_label(associe)} apporte à la Société la somme de "
            f"{_required_text(apport.montant_lettres, 'associes[].apport.montant_lettres')} "
            f"{euro_word(montant)}",
        )
        add_paragraph(
            document,
            f"ci- {montant} €.",
        )
    capital_lettres = _required_text(
        data.statuts.capital_social_lettres,
        "statuts_civils.capital_social_lettres",
    )
    capital_social = _required_text(
        data.statuts.capital_social,
        "statuts_civils.capital_social",
    )
    add_paragraph(
        document,
        f"Total des apports {capital_lettres} {euro_word(capital_social)} ({capital_social} €)",
    )
    depot = data.statuts.capital_depot
    banque_nom = _required_text(
        depot.banque_nom if depot else None,
        "statuts_civils.capital_depot.banque_nom",
    )
    banque_adresse = _required_text(
        depot.banque_adresse if depot else None,
        "statuts_civils.capital_depot.banque_adresse",
    )
    add_paragraph(
        document,
        "Cette somme a été intégralement déposée par les associés conformément à la loi, "
        "au crédit d’un compte ouvert au nom de la société en formation dans les livres "
        f"de la banque {banque_nom} {banque_adresse}.",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )


def _add_capital_block(document, data: _ResolvedStatutsScm) -> None:
    for associe in data.associes:
        parts = _required_parts(associe)
        # KAN-2 : quantite de parts a l'AFFICHAGE -> marqueur si vide, jamais « None »/« 0 parts ».
        add_paragraph(
            document,
            f"{_entity_label(associe)} {_quantite_parts(parts.nb, 'associes[].parts.nb')} parts",
        )
    add_paragraph(
        document,
        "Total du nombre de parts composant le capital social : "
        f"{_quantite_parts(data.statuts.nb_parts_total, 'statuts_civils.nb_parts_total')} parts",
    )


def _add_signature_block(document, data: _ResolvedStatutsScm) -> None:
    add_paragraph(document, f"Fait à {data.signature_lieu},")
    add_paragraph(document, f"Le {data.signature_date}")
    for associe in [a for a in data.associes if a.est_signataire]:
        add_statuts_signature_block(
            document,
            [_signature_label(associe)],
            mention_lines=["« Lu et approuvé »"],
        )
        add_paragraph(document, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=18)


def _add_morale_identity(document, associe: StatutsCivilsAssocie) -> None:
    capital_social = _required_text(associe.capital_social, "associes[].capital_social")
    add_paragraph(document, _required_text(associe.denomination, "associes[].denomination"))
    add_paragraph(
        document,
        f"{_required_text(associe.forme_juridique, 'associes[].forme_juridique')} de "
        f"{_required_text(associe.profession, 'associes[].profession')}",
    )
    # Rafael 2026-07-09 (transverse devise) : unite DERIVEE si montant nu (idempotent).
    add_paragraph(
        document,
        f"Au capital social de {montant_avec_euros(capital_social)}",
    )
    add_paragraph(
        document,
        f"Siège social : {_address_display(associe.siege, 'associes[].siege')}",
    )
    add_paragraph(
        document,
        f"{_required_text(associe.numero_rcs, 'associes[].numero_rcs')} R.C.S. de "
        f"{_required_text(associe.ville_rcs, 'associes[].ville_rcs')}",
    )
    representant = _required_representant(associe)
    add_paragraph(
        document,
        "Représentée par "
        f"{_required_text(representant.civilite_affichage, 'associes[].representant.civilite')} "
        f"{_required_text(representant.prenom, 'associes[].representant.prenom')} "
        f"{_required_text(representant.nom, 'associes[].representant.nom')}, "
        "en sa qualité de "
        f"{_required_text(representant.fonction, 'associes[].representant.fonction')}",
    )


def _add_physical_identity(document, associe: StatutsCivilsAssocie) -> None:
    gender = associe.genre or Gender.MASCULIN
    born = "Née" if gender == Gender.FEMININ else "Né"
    add_paragraph(document, _signature_label(associe))
    # Rafael 2026-07-09 (R12) : chaque element du bloc liste commence par une MAJUSCULE
    # (« Chirurgien-dentiste de profession », « Célibataire »), comme les lignes voisines
    # (7.2 Albane, patron SPFPL/SELAS).
    add_paragraph(
        document,
        f"{capitalize_first(_required_text(associe.profession, 'associes[].profession'))}"
        " de profession",
    )
    add_paragraph(
        document,
        f"{born} le {_format_birthdate(associe.date_naissance, 'associes[].date_naissance')} "
        f"à {_required_text(associe.ville_naissance, 'associes[].ville_naissance')}",
    )
    add_paragraph(document, f"Demeurant au {_person_address(associe)}")
    add_paragraph(
        document,
        f"De nationalité {_required_text(associe.nationalite, 'associes[].nationalite')}",
    )
    add_paragraph(
        document,
        capitalize_first(
            _required_text(associe.situation_maritale, "associes[].situation_maritale")
        ),
    )


def _validate_associes(associes: list[StatutsCivilsAssocie]) -> None:
    if not associes:
        raise ValueError(f"au moins un associe est obligatoire pour {DOCUMENT_CODE}.")
    if len(associes) > MAX_ASSOCIES:
        raise ValueError(f"les statuts SCM sont limites a 6 associes pour {DOCUMENT_CODE}.")
    for associe in associes:
        _required_apport(associe)
        _required_parts(associe)
        _required_text(associe.profession, "associes[].profession")
        if _is_morale(associe):
            _required_text(associe.denomination, "associes[].denomination")
            _required_text(associe.forme_juridique, "associes[].forme_juridique")
            _required_text(associe.capital_social, "associes[].capital_social")
            _address_display(associe.siege, "associes[].siege")
            _required_text(associe.numero_rcs, "associes[].numero_rcs")
            _required_text(associe.ville_rcs, "associes[].ville_rcs")
            _required_representant(associe)
        else:
            _required_text(associe.civilite_affichage, "associes[].civilite_affichage")
            _required_text(associe.prenom or associe.prenoms, "associes[].prenom")
            _required_text(associe.nom, "associes[].nom")
            _format_display_date(associe.date_naissance, "associes[].date_naissance")
            _required_text(associe.ville_naissance, "associes[].ville_naissance")
            _person_address(associe)
            _required_text(associe.nationalite, "associes[].nationalite")
            _required_text(associe.situation_maritale, "associes[].situation_maritale")


def _validate_capital_totals(
    statuts: StatutsCivilsContext,
    associes: list[StatutsCivilsAssocie],
) -> None:
    # KAN-2 : les coherences ne se controlent que sur des valeurs REELLES. A formulaire vide
    # (aucun total saisi, montants non numeriques), il n'y a rien a comparer -> on ne bloque
    # jamais (le document se genere, les zones vides sortent en marqueurs). Les controles
    # complets restent exerces des que TOUTES les valeurs concernees sont renseignees.
    if (
        statuts.nb_parts_total
        and associes
        and all(a.parts is not None and a.parts.nb for a in associes)
    ):
        total_parts = sum((a.parts.nb or 0) for a in associes)
        if total_parts != statuts.nb_parts_total:
            raise ValueError(
                "la somme des parts doit correspondre a statuts_civils.nb_parts_total "
                f"pour {DOCUMENT_CODE}."
            )

    if (
        _is_numeric_amount(statuts.capital_social)
        and associes
        and all(a.apport is not None and _is_numeric_amount(a.apport.montant) for a in associes)
    ):
        total_apports = sum(_amount_to_int(a.apport.montant) for a in associes)
        if total_apports != _amount_to_int(statuts.capital_social):
            raise ValueError(
                "la somme des apports doit correspondre a statuts_civils.capital_social "
                f"pour {DOCUMENT_CODE}."
            )


def _validate_statuts_fields(statuts: StatutsCivilsContext) -> None:
    _required_text(statuts.forme_sociale, "statuts_civils.forme_sociale")
    _required_text(statuts.capital_social, "statuts_civils.capital_social")
    _required_text(statuts.capital_social_lettres, "statuts_civils.capital_social_lettres")
    _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    _required_text(statuts.valeur_nominale_part, "statuts_civils.valeur_nominale_part")
    depot = statuts.capital_depot
    _required_text(depot.banque_nom if depot else None, "statuts_civils.capital_depot.banque_nom")
    _required_text(
        depot.banque_adresse if depot else None,
        "statuts_civils.capital_depot.banque_adresse",
    )


def _validate_signatures(associes: list[StatutsCivilsAssocie]) -> None:
    signataires = [associe for associe in associes if associe.est_signataire]
    if len(signataires) != len(associes):
        raise ValueError(
            f"les signatures SCM doivent correspondre aux associes pour {DOCUMENT_CODE}."
        )
    for associe in signataires:
        _signature_label(associe)


def _required_apport(associe: StatutsCivilsAssocie) -> StatutsCivilsApport:
    # KAN-2 : apport absent -> instance vide (montant/lettres sortent en marqueurs au site
    # d'affichage), jamais lever.
    return associe.apport if associe.apport is not None else StatutsCivilsApport()


def _required_parts(associe: StatutsCivilsAssocie) -> StatutsCivilsParts:
    # KAN-2 : parts absentes -> instance vide (quantite rendue par _quantite_parts), jamais lever.
    return associe.parts if associe.parts is not None else StatutsCivilsParts()


def _required_representant(associe: StatutsCivilsAssocie) -> StatutsCivilsRepresentant:
    # KAN-2 : representant absent -> instance vide (ses champs sortent en marqueurs), jamais lever.
    return (
        associe.representant
        if associe.representant is not None
        else StatutsCivilsRepresentant()
    )


def _signature_label(associe: StatutsCivilsAssocie) -> str:
    if _is_morale(associe):
        denomination = _required_text(associe.denomination, "associes[].denomination")
        representant = _required_representant(associe)
        civilite = _required_text(
            representant.civilite_affichage,
            "associes[].representant.civilite",
        )
        return (
            f"{denomination}, représentée par "
            f"{civilite} "
            f"{_required_text(representant.prenom, 'associes[].representant.prenom')} "
            f"{_required_text(representant.nom, 'associes[].representant.nom')}"
        )
    prenoms = associe.prenoms or associe.prenom
    return (
        f"{_required_text(associe.civilite_affichage, 'associes[].civilite_affichage')} "
        f"{_required_text(prenoms, 'associes[].prenoms')} "
        f"{_required_text(associe.nom, 'associes[].nom')}"
    )


def _entity_label(associe: StatutsCivilsAssocie) -> str:
    """Designation de l'associe dans les blocs apports / parts.

    Le modele source (paras 81 et 95) n'y porte QUE la denomination de la personne
    morale (sans representant) et l'identite de la personne physique. Le wording
    "representee par ..." n'existe, dans le modele, qu'en comparution (para 30) et
    en signature. On l'aligne donc strictement sur la source.
    """
    if _is_morale(associe):
        return _required_text(associe.denomination, "associes[].denomination")
    prenoms = associe.prenoms or associe.prenom
    return (
        f"{_required_text(associe.civilite_affichage, 'associes[].civilite_affichage')} "
        f"{_required_text(prenoms, 'associes[].prenoms')} "
        f"{_required_text(associe.nom, 'associes[].nom')}"
    )


def _apport_label(associe: StatutsCivilsAssocie) -> str:
    """Sujet de la ligne d'apport. Source para 81 : "La [denomination] apporte ..."
    pour une personne morale ; para 84 : identite nue pour une personne physique.
    """
    if _is_morale(associe):
        return f"La {_entity_label(associe)}"
    return _entity_label(associe)


def _is_morale(associe: StatutsCivilsAssocie) -> bool:
    return associe.type_personne == "personne_morale"


def _person_address(associe: StatutsCivilsAssocie) -> str:
    if associe.adresse_personnelle_affichee:
        return associe.adresse_personnelle_affichee.strip()
    return _address_display(associe.adresse_personnelle, "associes[].adresse_personnelle")


def _address_display(address: Address | None, field_name: str) -> str:
    # KAN-2 : adresse absente -> marqueur metier, jamais lever.
    if address is None:
        return _required_text(None, field_name)
    if address.adresse_affichee:
        return address.adresse_affichee.strip()
    return (
        f"{_required_text(address.num_voie, f'{field_name}.num_voie')} "
        f"{_required_text(address.voie, f'{field_name}.voie')} "
        f"{_required_text(address.cp, f'{field_name}.cp')} "
        f"{_required_text(address.ville, f'{field_name}.ville')}"
    )


def _format_display_date(value: date | str | None, field_name: str) -> str:
    # KAN-2 : date absente -> marqueur metier, jamais lever.
    if value is None:
        return _required_text(None, field_name)
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return _required_text(value, field_name)


def _format_birthdate(value: date | str | None, field_name: str) -> str:
    """Date de NAISSANCE des statuts SCM en « JJ mois AAAA » (Albane 2026-07-09, B1).

    Une valeur numerique saisie (« 10/03/1975 », ISO ou objet date) est reformatee en
    francais lettre ; une date deja lettree reste inchangee. La date de SIGNATURE conserve
    son format court (``_format_display_date``).
    """
    # KAN-2 : date de naissance absente -> marqueur metier, jamais lever.
    if value is None or (isinstance(value, str) and not value.strip()):
        return _required_text(None, field_name)
    return format_birthdate_fr(value)


def _required_text(value: str | None, field_name: str) -> str:
    # KAN-2 (R10) : une donnee manquante NE bloque PAS -> marqueur METIER lisible SANS crochets
    # (« (À COMPLÉTER : … ) », pour ne pas declencher le garde-fou anti-placeholder source) au lieu
    # de lever. Sortie NOMINALE (valeur presente) inchangee : byte-identique au modele.
    if value is None or not str(value).strip():
        return f"(À COMPLÉTER : {libelle_metier(field_name)})"
    return str(value).strip()


def _required_int(value: int | None, field_name: str) -> int:
    # KAN-2 : point de CALCUL / coherence -> None-safe (0), jamais lever. NE PAS s'en servir pour
    # AFFICHER une quantite de parts (« 0 parts » dans un acte signable serait FAUX) : passer par
    # _quantite_parts a l'affichage.
    return value if value is not None else 0


def _quantite_parts(value: int | None, field_name: str) -> str:
    # KAN-2 : AFFICHAGE d'une quantite de parts. Non renseignee (None ou < 1) -> marqueur METIER,
    # jamais « 0 parts »/« None parts ». Une quantite reelle (>= 1) est rendue inchangee -> sortie
    # NOMINALE byte-identique.
    if value is None or value < 1:
        return f"(À COMPLÉTER : {libelle_metier(field_name)})"
    return str(value)


def _is_numeric_amount(value: str | None) -> bool:
    # KAN-2 : True seulement si `value` porte un montant NUMERIQUE reel (ni vide, ni marqueur
    # « (À COMPLÉTER : … ) »). Sert a n'exercer les coherences capital que sur du saisi.
    if value is None or not str(value).strip():
        return False
    return _normalize_amount(str(value)).isdigit()


def _normalize_amount(text: str) -> str:
    return (
        text.replace(" ", "")
        .replace("\xa0", "")
        .replace("euros", "")
        .replace("euro", "")
        .replace("EUR", "")
        .replace("€", "")
        .strip()
    )


def _amount_to_int(value: str | None) -> int:
    text = _required_text(value, "montant")
    normalized = (
        text.replace(" ", "")
        .replace("\u00a0", "")
        .replace("euros", "")
        .replace("euro", "")
        .replace("EUR", "")
        .replace("€", "")
        .strip()
    )
    if not normalized.isdigit():
        raise ValueError(f"montant numerique attendu pour {DOCUMENT_CODE}: {value}")
    return int(normalized)


def _add_rendered_paragraph(document, text: str) -> None:
    if text == "STATUTS":
        add_statuts_title_box(document, text)
    elif text.startswith("TITRE "):
        add_statuts_part_heading(document, text)
    elif text.startswith("Article "):
        add_statuts_article_heading(document, text, left_indent_cm=0.25)
    else:
        add_statuts_body_paragraph(document, text)


def _replace_placeholders(text: str, replacements: dict[str, str]) -> str:
    rendered = text
    for placeholder, value in replacements.items():
        rendered = rendered.replace(placeholder, value)
    return rendered
