from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    Company,
    DocumentGenerationContext,
    ExerciceSocial,
    Person,
    StatutsSasuHoldingContext,
)
from sydel_doc_engine.generators.lot_04.statuts_sel_exercice_common import (
    statuts_output_filename,
)
from sydel_doc_engine.generators.lot_05.spfpl_libelles import libelle_metier
from sydel_doc_engine.rendering.docx_builder import keep_final_signature_block_together
from sydel_doc_engine.rendering.docx_template_fill import fill_docx_template
from sydel_doc_engine.utils.months import FRENCH_MONTHS

DOCUMENT_CODE = "CODE-STATUTS-SASU-HOLDING-001"
OUTPUT_FILENAME = "statuts_sasu_holding.docx"
EXPECTED_STRUCTURE = "SASU_HOLDING"

# Modele source tokenise (modele officiel Albane 2026-06-29, SASU Holding generaliste).
# Resolu independamment du cwd : parents[4] = racine du repo.
_SOURCE_MODEL = (
    Path(__file__).resolve().parents[4]
    / "project"
    / "source_documents"
    / "lot_04"
    / "statuts SASU Holding.docx"
)


class StatutsSasuHoldingGenerator:
    """Generateur des statuts SASU Holding (SAS unipersonnelle, holding patrimoniale).

    Token-replacement pur sur le modele officiel Albane (l'acte est unipersonnel : un seul
    associe = le president, aucun bloc repetitif). Le genre de l'associe unique accorde les
    tournures figees du modele (« Le soussigne »/« La soussignee », « il/elle a decide »,
    « ne/nee le »).
    """

    def generate(self, ctx: DocumentGenerationContext, output_dir: Path) -> Path:
        if ctx.structure != EXPECTED_STRUCTURE:
            raise ValueError(
                f"dossier.structure doit etre {EXPECTED_STRUCTURE} pour {DOCUMENT_CODE}."
            )
        # KAN-2 (Rafael 2026-07-14) : « Tous les documents doivent pouvoir être générés, MÊME SI
        # je ne remplis AUCUN champ. » -> aucun `raise` sur donnée manquante. Un objet absent
        # tombe sur une instance VIDE dont chaque champ ressortira en « (À COMPLÉTER : … ) »
        # (`_required_text`), jamais un crash. La sortie NOMINALE (champs remplis) reste
        # byte-identique : ces replis ne s'activent que quand l'objet est None.
        statuts = ctx.statuts_sasu_holding or StatutsSasuHoldingContext()
        societe = ctx.societe or Company()
        associe = ctx.personne_signataire or _empty_person()
        siege = societe.siege or Address()
        adresse_perso = associe.adresse_perso or Address()

        replacements = _build_replacements(statuts, societe, associe, ctx, siege, adresse_perso)
        gender_pairs = [(associe.genre, _SOUSSIGNE_PAIRS)]

        output_dir.mkdir(parents=True, exist_ok=True)
        # Retour Rafael 2026-07-07 : TOUS les statuts sont nommes
        # « Statuts <denomination>.docx » (helper partage, fallback historique si vide).
        output_path = output_dir / statuts_output_filename(
            societe.denomination, OUTPUT_FILENAME
        )
        fill_docx_template(
            _SOURCE_MODEL, replacements, output_path, gender_pairs=gender_pairs
        )
        # Rafael 2026-07-09 : (1) retirer TOUT surlignage herite du modele source
        # (runs surlignes d'edition, ex. « MLG ») ; (2) l'annexe demarre TOUJOURS en
        # debut de nouvelle page (saut de page avant le titre « ANNEXE »).
        _postprocess_sasu_statuts(output_path)
        return output_path


def _postprocess_sasu_statuts(output_path: Path) -> None:
    """Post-traitement de forme des statuts SASU Holding (Rafael 2026-07-09).

    - Retire tout surlignage (`w:highlight`) des runs : le modele source Albane porte
      des residus de surlignage d'edition (run « MLG ») que le token-replacement
      preserve — aucun surlignage ne doit sortir dans un acte livre.
    - Le titre « ANNEXE » demarre en debut de NOUVELLE page : `page_break_before` pose
      sur son paragraphe (robuste a la pagination, contrairement a un saut manuel).
    Le TEXTE juridique n'est pas touche (forme uniquement).
    """
    document = Document(str(output_path))
    for paragraph in _iter_paragraphs_with_tables(document):
        for run in paragraph.runs:
            rpr = run._element.find(qn("w:rPr"))  # noqa: SLF001 - acces XML python-docx
            if rpr is None:
                continue
            for highlight in rpr.findall(qn("w:highlight")):
                rpr.remove(highlight)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper().startswith("ANNEXE"):
            paragraph.paragraph_format.page_break_before = True
            break
    # KAN-36 : bloc signature final solidaire (une seule page).
    keep_final_signature_block_together(document)
    document.save(str(output_path))


def _iter_paragraphs_with_tables(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


# Accords de genre de l'associe unique (modele fige au masculin). Pour une associee unique
# personne physique feminin -> accorder les tournures de comparution.
_SOUSSIGNE_PAIRS: list[tuple[str, str]] = [
    ("Le soussigné", "La soussignée"),
    ("qu’il a décidé", "qu’elle a décidé"),
    ("né le", "née le"),
    ("l’associé unique pourra", "l’associée unique pourra"),
]


def _build_replacements(
    statuts: StatutsSasuHoldingContext,
    societe: Company,
    associe: Person,
    ctx: DocumentGenerationContext,
    siege: Address,
    adresse_perso: Address,
) -> dict[str, str]:
    # KAN-2 : exercice absent -> instance vide (ses champs ressortent en marqueur), jamais None.
    exercice = ctx.exercice_social or ExerciceSocial()
    capital = _required_text(statuts.capital_social or societe.capital, "capital social")
    return {
        "[denomination_societe]": _required_text(
            societe.denomination, "dénomination de la société"
        ),
        "[forme_sociale]": _required_text(statuts.forme_sociale, "forme sociale"),
        "[capital_social]": capital,
        "[capital_lettres]": _required_text(
            statuts.capital_social_lettres, "capital social en toutes lettres"
        ),
        "[num_voie_siege]": _text(siege.num_voie),
        "[voie_siege]": _required_text(siege.voie, "voie du siège"),
        "[ville_siege]": _required_text(siege.ville, "ville du siège"),
        "[cp_siege]": _required_text(siege.cp, "code postal du siège"),
        "[civilite]": _required_text(associe.civilite, "civilité de l'associé"),
        "[prenom]": _required_text(associe.prenom, "prénom de l'associé"),
        "[nom]": _required_text(associe.nom, "nom de l'associé"),
        "[date_naissance]": _french_date(associe.date_naissance, "date de naissance de l'associé"),
        "[ville_naissance]": _required_text(
            associe.ville_naissance, "ville de naissance de l'associé"
        ),
        "[nationalite]": _required_text(associe.nationalite, "nationalité de l'associé"),
        "[num_voie_perso]": _text(adresse_perso.num_voie),
        "[voie_perso]": _required_text(adresse_perso.voie, "voie de l'adresse personnelle"),
        "[cp_perso]": _required_text(adresse_perso.cp, "code postal personnel"),
        "[ville_perso]": _required_text(adresse_perso.ville, "ville personnelle"),
        "[nb_actions]": _display_int(statuts.nb_actions, "nombre d'actions"),
        "[nom_banque]": _required_text(statuts.nom_banque, "nom de la banque"),
        "[debut_exercice]": _required_text(exercice.debut, "début de l'exercice social"),
        "[fin_exercice]": _required_text(exercice.fin, "fin de l'exercice social"),
        "[date_cloture_exercice_1]": _required_text(
            exercice.date_cloture_premier_exercice, "date de clôture du premier exercice"
        ),
        "[lieu_signature]": _required_text(ctx.signature.lieu, "lieu de signature"),
        "[date_signature]": _french_date(ctx.signature.date, "date de signature"),
        "[qualite_associe]": _required_text(statuts.qualite_associe, "qualité de l'associé"),
        "[fonction_dirigeant]": _required_text(statuts.fonction_dirigeant, "fonction du dirigeant"),
    }


# KAN-2 (Rafael 2026-07-14) — « Tous les documents doivent pouvoir être générés, MÊME SI je ne
# remplis AUCUN champ. » Un champ absent NE lève plus : il ressort en marqueur métier visible
# « (À COMPLÉTER : <libellé> ) » (via `libelle_metier`, aligné sur le pattern SPFPL /
# `statuts_sel_exercice_common.required_text`). Le libellé passé est déjà métier -> `libelle_metier`
# le laisse traverser et garantit qu'il reste SANS point / underscore / crochet / chiffre.
def _marqueur(field_name: str) -> str:
    return f"(À COMPLÉTER : {libelle_metier(field_name)})"


def _french_date(value: date | str | None, field_name: str) -> str:
    if value is None:
        return _marqueur(field_name)
    if isinstance(value, date):
        jour = "1er" if value.day == 1 else str(value.day)  # convention francaise du 1er
        return f"{jour} {FRENCH_MONTHS[value.month]} {value.year}"
    return _required_text(value, field_name)


def _required_text(value: str | None, field_name: str) -> str:
    if value is None or not str(value).strip():
        return _marqueur(field_name)
    return str(value).strip()


def _display_int(value: int | None, field_name: str) -> str:
    # AFFICHAGE d'une quantité : non renseignée (None) OU nulle (0, valeur d'un `number_input`
    # vide) -> marqueur, JAMAIS un « 0 » affirmé (un acte signable qui dit « 0 actions » est faux,
    # Akainu B1). Renseignée -> valeur brute, byte-identique au gold.
    if value is None or value < 1:
        return _marqueur(field_name)
    return str(value)


def _empty_person() -> Person:
    # KAN-2 : associé absent -> instance vide (genre masculin neutre, identité vide) dont chaque
    # champ ressort en marqueur ; `Person` exige genre/civilite/prenom/nom, on les fournit vides.
    return Person(genre=Gender.MASCULIN, civilite="", prenom="", nom="")


def _text(value: str | None) -> str:
    return str(value).strip() if value is not None else ""


__all__ = ["StatutsSasuHoldingGenerator", "DOCUMENT_CODE", "EXPECTED_STRUCTURE"]
