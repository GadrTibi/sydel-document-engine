from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    DocumentGenerationContext,
    StatutsCivilsAssocie,
    StatutsCivilsContext,
)
from sydel_doc_engine.generators.lot_04.annexe_filter import is_creation_fee_annexe_line
from sydel_doc_engine.generators.lot_04.statuts_sel_exercice_common import (
    statuts_output_filename,
)
from sydel_doc_engine.rendering.docx_builder import (
    add_paragraph,
    add_statuts_article_heading,
    add_statuts_body_paragraph,
    add_statuts_hanging_list_item,
    add_statuts_matrix_table,
    add_statuts_part_heading,
    add_statuts_signature_block,
    add_statuts_signature_grid,
    add_statuts_title_box,
    new_document,
)
from sydel_doc_engine.utils.grammar import euro_word

DOCUMENT_CODE = "CODE-STATUTS-CIVILS-CORE-001"
MAX_ASSOCIES = 6

# Espace insecable (U+00A0) : ponctuation fine francaise du modele Albane micro holding
# (« Siege social<NBSP>: », « comme suit<NBSP>: »). Defini via chr() pour rester ASCII-safe.
_NBSP = chr(0x00A0)

# Marqueur editorial interne du modele source SCI (fin Art. 31, source para 547) :
# "... jusqu'au [date]. A RETIRER SI LA SOCIETE EST A L'IR". C'est une INSTRUCTION INTERNE de
# redaction, pas du texte juridique destine au client -> on la retire de la sortie tout en gardant
# la clause qu'elle annote. Ancree sur "A RETIRER" jusqu'a la fin du paragraphe (tolerant a
# l'apostrophe droite/typographique et a la casse). NB : un eventuel conditionnement IR/IS de la
# clause elle-meme est une decision METIER, non tranchee ici
# (cf. _PASSE2_VERIFICATION_REPORT sect.3).
_EDITORIAL_MARKER_RE = re.compile(r"\s*A RETIRER SI LA SOCIETE EST A L.IR\s*$", re.IGNORECASE)


@dataclass(frozen=True)
class StatutsCivilTemplate:
    source_name_contains: tuple[str, ...]
    output_filename: str
    expected_structure: str
    expected_type: str
    associate_slice: tuple[int, int]
    apport_slice: tuple[int, int]
    capital_slice: tuple[int, int]
    signature_slice: tuple[int, int] | None = None
    append_signatures_after: int | None = None
    statuts_title_box_before: int | None = None


SCS_TEMPLATE = StatutsCivilTemplate(
    source_name_contains=("SCS",),
    output_filename="statuts_scs.docx",
    expected_structure="SCS",
    expected_type="scs",
    associate_slice=(14, 18),
    apport_slice=(43, 58),
    capital_slice=(63, 76),
    append_signatures_after=253,
    # FORME : le titre encadre "STATUTS" du modele source vit dans une TABLE (cellule T0), entre
    # le bloc d'en-tete (paras 1-5) et "LES SOUSSIGNES" (para 12). Le moteur n'itere que les
    # paragraphes source -> il ne le voit pas. On le restaure en l'injectant avant le para 12.
    statuts_title_box_before=12,
)

SCI_TEMPLATE = StatutsCivilTemplate(
    source_name_contains=("SCI",),
    output_filename="statuts_sci.docx",
    expected_structure="SCI",
    expected_type="sci",
    associate_slice=(25, 44),
    apport_slice=(97, 111),
    capital_slice=(120, 131),
    signature_slice=(612, 623),
)

# Micro holding (Albane 2026-06-29) : VRAI modele Albane « societe civile de portefeuille a
# capital variable » (26 articles, distinct du SCI). Modele source tokenise depuis le DOCX
# fourni par Albane (Statuts_Micro_holding.docx). Slices derivees du modele Albane :
#   - statuts_title_box_before=9 : le titre « STATUTS » vit dans une TABLE (entre l'en-tete et
#     « LES SOUSSIGNES ») -> non vu par l'iteration des paragraphes, on l'injecte avant le P9.
#   - associate (17,33) : comparution SPFPL (morale) + praticien (physique), wording MH propre.
#   - apport (86,96) : art. 6 (« X apporte la somme de … / Ci … euros » + total + depot).
#   - capital (100,114) : art. 7 (variable min/max/effectif + division + repartition par associe).
#   - signature (461,465) : « Fait a [lieu] / Le [date] » + signataires cote a cote.
# L'objet social (art. 2) est desormais VERBATIM dans le modele Albane (plus de token/variante).
MICRO_HOLDING_TEMPLATE = StatutsCivilTemplate(
    source_name_contains=("MICRO", "HOLDING"),
    output_filename="statuts_micro_holding.docx",
    expected_structure="MICRO_HOLDING",
    expected_type="micro_holding",
    associate_slice=(17, 33),
    apport_slice=(86, 96),
    capital_slice=(100, 114),
    signature_slice=(461, 465),
    statuts_title_box_before=9,
)

SCI_IRIS_TEMPLATE = StatutsCivilTemplate(
    source_name_contains=("SCI", "IRIS"),
    output_filename="statuts_sci_iris.docx",
    expected_structure="SCI IRIS",
    expected_type="sci_iris",
    associate_slice=(24, 42),
    apport_slice=(95, 109),
    capital_slice=(120, 131),
    # Source para 626 = "A [lieu], le [date]" est AVANT l'ancien debut de slice (629) -> il etait
    # rendu par le chemin source PUIS re-rendu par _add_signature_block => date en double. On etend
    # le debut a 626 pour que la ligne date source soit englobee (skip) et rendue une seule fois par
    # le bloc signature. Cf. _PASSE2_VERIFICATION_REPORT.md sect.4.
    signature_slice=(626, 636),
)


def generate_statuts_civil_docx(  # noqa: C901
    ctx: DocumentGenerationContext,
    output_dir: Path,
    template: StatutsCivilTemplate,
) -> Path:
    data = _ResolvedStatutsCivil.from_context(ctx, template)
    source = _source_path(template)
    source_doc = Document(source)
    output_doc = new_document()
    output_doc.sections[0].footer.paragraphs[0].text = f"{data.denomination} - Statuts constitutifs"

    replacements = data.common_replacements()
    skip_until = -1
    for index, paragraph in enumerate(source_doc.paragraphs):
        if index < skip_until:
            continue
        if (
            template.statuts_title_box_before is not None
            and index == template.statuts_title_box_before
        ):
            add_statuts_title_box(output_doc, "STATUTS")
        if index == template.associate_slice[0]:
            _add_associate_block(output_doc, data)
            skip_until = template.associate_slice[1]
            continue
        if index == template.apport_slice[0]:
            _add_apport_block(output_doc, data)
            skip_until = template.apport_slice[1]
            continue
        if index == template.capital_slice[0]:
            _add_capital_block(output_doc, data)
            skip_until = template.capital_slice[1]
            continue
        if template.signature_slice is not None and index == template.signature_slice[0]:
            _add_signature_block(output_doc, data)
            skip_until = template.signature_slice[1]
            continue

        text = paragraph.text.strip()
        if not text:
            continue
        rendered = _replace_placeholders(text, replacements)
        rendered = _strip_editorial_marker(rendered)
        if not rendered:
            continue
        if is_creation_fee_annexe_line(rendered):  # O24-01 : annexe sans frais cabinet création
            continue
        _add_rendered_paragraph(output_doc, rendered, paragraph)
        if template.expected_type == "sci_iris" and index == 561:
            _add_resultat_groupes_block(output_doc, data)
        if (
            template.append_signatures_after is not None
            and index == template.append_signatures_after
        ):
            _add_signature_block(output_doc, data)

    full_text = "\n".join(paragraph.text for paragraph in output_doc.paragraphs)
    if "[" in full_text or "]" in full_text:
        raise ValueError(f"placeholder source residuel dans le rendu {DOCUMENT_CODE}.")

    output_dir.mkdir(parents=True, exist_ok=True)
    # SCS6 (Albane 2026-06-25) : pour la SCS, le fichier statuts porte la denomination
    # (« Statuts <denomination>.docx ») au lieu du nom fixe « statuts_scs.docx » — meme
    # convention que la SELARL, deja ratifiee (2026-06-10, helper statuts_output_filename).
    # Le « _ » du verbatim « Statuts_{denomination_sociale} » est sa notation placeholder
    # (cf. le snake_case « denomination_sociale ») -> on aligne sur la convention espace
    # ratifiee. Les autres civiles (SCI / SCI IRIS / SCM) gardent leur nom fixe.
    filename = (
        statuts_output_filename(data.denomination, template.output_filename)
        if template.expected_structure == "SCS"
        else template.output_filename
    )
    output_path = output_dir / filename
    output_doc.save(output_path)
    return output_path


class _ResolvedStatutsCivil:
    def __init__(
        self,
        *,
        template: StatutsCivilTemplate,
        statuts: StatutsCivilsContext,
        denomination: str,
        forme_sociale: str,
        adresse_siege: str,
        siege_num_voie: str,
        siege_voie: str,
        siege_cp: str,
        siege_ville: str,
        ville_rcs: str,
        signature_lieu: str,
        signature_date: str,
        associes: list[StatutsCivilsAssocie],
        objet_social: str | None = None,
    ) -> None:
        self.template = template
        self.statuts = statuts
        self.objet_social = objet_social
        self.denomination = denomination
        self.forme_sociale = forme_sociale
        self.adresse_siege = adresse_siege
        self.siege_num_voie = siege_num_voie
        self.siege_voie = siege_voie
        self.siege_cp = siege_cp
        self.siege_ville = siege_ville
        self.ville_rcs = ville_rcs
        self.signature_lieu = signature_lieu
        self.signature_date = signature_date
        self.associes = associes

    @classmethod
    def from_context(
        cls,
        ctx: DocumentGenerationContext,
        template: StatutsCivilTemplate,
    ) -> _ResolvedStatutsCivil:
        if ctx.structure != template.expected_structure:
            raise ValueError(
                f"dossier.structure doit etre {template.expected_structure} pour {DOCUMENT_CODE}."
            )
        if ctx.statuts_civils is None:
            raise ValueError(f"statuts_civils est obligatoire pour {DOCUMENT_CODE}.")
        statuts_type = _required_text(ctx.statuts_civils.type, "statuts_civils.type").lower()
        if statuts_type != template.expected_type:
            raise ValueError(
                f"statuts_civils.type doit etre {template.expected_type} pour {DOCUMENT_CODE}."
            )
        if ctx.societe is None:
            raise ValueError(f"societe est obligatoire pour {DOCUMENT_CODE}.")
        if ctx.societe.siege is None:
            raise ValueError(f"societe.siege est obligatoire pour {DOCUMENT_CODE}.")
        associes = list(ctx.statuts_civils.associes)
        _validate_associes(associes, template)
        _validate_capital_totals(ctx.statuts_civils, associes)
        if template.expected_type == "scs":
            _validate_scs(ctx.statuts_civils, associes)
        if template.expected_type == "sci":
            _validate_sci(associes)
        if template.expected_type == "sci_iris":
            _validate_sci_iris(ctx.statuts_civils, associes)
        if template.expected_type == "micro_holding":
            _validate_micro_holding(ctx.statuts_civils)
        _validate_template_fields(ctx.statuts_civils, template)

        return cls(
            template=template,
            statuts=ctx.statuts_civils,
            denomination=_required_text(ctx.societe.denomination, "societe.denomination"),
            forme_sociale=_required_text(
                ctx.statuts_civils.forme_sociale or ctx.societe.forme_sociale,
                "statuts_civils.forme_sociale",
            ),
            adresse_siege=_address_display(ctx.societe.siege, "societe.siege"),
            siege_num_voie=_required_text(ctx.societe.siege.num_voie, "societe.siege.num_voie"),
            siege_voie=_required_text(ctx.societe.siege.voie, "societe.siege.voie"),
            siege_cp=_required_text(ctx.societe.siege.cp, "societe.siege.cp"),
            siege_ville=_required_text(ctx.societe.siege.ville, "societe.siege.ville"),
            ville_rcs=_required_text(ctx.societe.ville_rcs, "societe.ville_rcs"),
            signature_lieu=ctx.signature.lieu,
            signature_date=_format_display_date(ctx.signature.date, "signature.date"),
            associes=associes,
            objet_social=ctx.statuts_civils.objet_social,
        )

    def common_replacements(self) -> dict[str, str]:
        statuts = self.statuts
        depot = statuts.capital_depot
        return {
            "[denomination_societe]": self.denomination,
            "[forme_sociale]": self.forme_sociale,
            "[mention_capital_variable]": _text_or_empty(statuts.mention_capital_variable),
            "[capital_social]": _required_text(
                statuts.capital_social,
                "statuts_civils.capital_social",
            ),
            "[capital_lettres]": _required_text(
                statuts.capital_social_lettres,
                "statuts_civils.capital_social_lettres",
            ),
            "[capital_autorise]": _text_or_empty(statuts.capital_autorise),
            "[capital_autorise_lettres]": _text_or_empty(statuts.capital_autorise_lettres),
            "[capital_social_maximal]": _text_or_empty(statuts.capital_maximal),
            "[capital_social_maximal_lettres]": _text_or_empty(statuts.capital_maximal_lettres),
            "[nb_parts]": str(
                _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
            ),
            "[nb_parts_total]": str(
                _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
            ),
            "[nb_parts_lettres]": _required_text(
                statuts.nb_parts_total_lettres,
                "statuts_civils.nb_parts_total_lettres",
            ),
            "[nb_parts_total_lettres]": _required_text(
                statuts.nb_parts_total_lettres,
                "statuts_civils.nb_parts_total_lettres",
            ),
            "[valeur_nominale_part]": _required_text(
                statuts.valeur_nominale_part,
                "statuts_civils.valeur_nominale_part",
            ),
            "[valeur_nominale_part_lettres]": _text_or_empty(statuts.valeur_nominale_part_lettres),
            "[plage_parts_total]": _text_or_empty(statuts.plage_parts_totale),
            "[parts_debut]": str(_first_part_number(self.associes)),
            "[parts_fin]": str(_last_part_number(self.associes)),
            "[adresse_siege]": self.adresse_siege,
            "[num_voie_siege]": self.siege_num_voie,
            "[voie_siege]": self.siege_voie,
            "[cp_siege]": self.siege_cp,
            "[ville_siege]": self.siege_ville,
            "[ville_rcs]": self.ville_rcs,
            "[duree_societe]": _text_or_empty(statuts.duree_societe),
            "[nom_banque]": _required_text(
                depot.banque_nom if depot else None,
                "statuts_civils.capital_depot.banque_nom",
            ),
            "[adresse_banque]": _required_text(
                depot.banque_adresse if depot else None,
                "statuts_civils.capital_depot.banque_adresse",
            ),
            "[date_cloture_exercice_1]": _required_text(
                statuts.date_cloture_premier_exercice,
                "statuts_civils.date_cloture_premier_exercice",
            ),
            "[lieu_signature]": _required_text(
                self.signature_lieu,
                "signature.lieu",
            ),
            "[date_signature]": self.signature_date,
            "[nombre_exemplaires_lettres]": _text_or_empty(statuts.nombre_exemplaires_lettres),
            "[denomination_cabinet_mandataire]": _text_or_empty(
                statuts.denomination_cabinet_mandataire
            ),
        }


def _bold_paragraph(paragraph) -> None:
    """Met tous les runs d'un paragraphe en gras (lignes d'identite de comparution)."""
    for run in paragraph.runs:
        run.bold = True


def _add_associate_block(document, data: _ResolvedStatutsCivil) -> None:
    micro_holding = data.template.expected_type == "micro_holding"
    for associe in data.associes:
        if _is_morale(associe):
            if micro_holding:
                _add_morale_identity_micro_holding(document, associe)
            else:
                _add_morale_identity(document, associe)
        elif micro_holding:
            _add_physical_identity_micro_holding(document, associe)
        else:
            _add_physical_identity(document, associe)


def _mh_morale_denomination(associe: StatutsCivilsAssocie) -> str:
    return _required_text(associe.denomination, "associes[].denomination")


def _mh_short_label(associe: StatutsCivilsAssocie) -> str:
    """Libelle court micro holding (apport / capital / signature).

    Personne morale -> « La <denomination> » (article feminin, denomination seule, SANS
    « representee par … » contrairement au libelle generique). Personne physique ->
    « <civilite> <prenoms> <nom> » (sans la mention « epouse … » reservee a la comparution).
    """
    if _is_morale(associe):
        return f"La {_mh_morale_denomination(associe)}"
    prenoms = associe.prenoms or associe.prenom
    return (
        f"{_required_text(associe.civilite_affichage, 'associes[].civilite_affichage')} "
        f"{_required_text(prenoms, 'associes[].prenoms')} "
        f"{_required_text(associe.nom, 'associes[].nom')}"
    )


def _add_morale_identity_micro_holding(document, associe: StatutsCivilsAssocie) -> None:
    # Comparution morale micro holding (modele Albane P017-P022) : 6 lignes distinctes.
    _bold_paragraph(add_paragraph(document, f"- {_mh_morale_denomination(associe)}"))
    add_paragraph(document, _required_text(associe.forme_juridique, "associes[].forme_juridique"))
    capital_morale = _required_text(associe.capital_social, "associes[].capital_social")
    add_paragraph(document, f"Au capital de {capital_morale} euros")
    add_paragraph(
        document, f"Siège social : {_address_display(associe.siege, 'associes[].siege')}"
    )
    add_paragraph(
        document,
        f"Immatriculée au RCS de {_required_text(associe.ville_rcs, 'associes[].ville_rcs')} "
        f"sous le numéro {_required_text(associe.numero_rcs, 'associes[].numero_rcs')}",
    )
    if associe.representant is None:
        raise ValueError(
            f"associes[].representant est obligatoire pour une personne morale {DOCUMENT_CODE}."
        )
    add_paragraph(
        document,
        f"Représentée par son "
        f"{_required_text(associe.representant.fonction, 'associes[].representant.fonction')}, "
        f"{_required_text(associe.representant.civilite_affichage, 'representant.civilite')} "
        f"{_required_text(associe.representant.prenom, 'associes[].representant.prenom')} "
        f"{_required_text(associe.representant.nom, 'associes[].representant.nom')}",
    )


def _add_physical_identity_micro_holding(document, associe: StatutsCivilsAssocie) -> None:
    # Comparution physique micro holding (modele Albane P025-P029). Libelle d'identite avec
    # virgule finale ; date de naissance verbatim (« 1er decembre 1978 »).
    gender = associe.genre or Gender.MASCULIN
    born = "Née" if gender == Gender.FEMININ else "Né"
    _bold_paragraph(add_paragraph(document, f"{_signature_label(associe)},"))
    add_paragraph(
        document,
        f"{born} le {_format_display_date(associe.date_naissance, 'associes[].date_naissance')} "
        f"à {_required_text(associe.ville_naissance, 'associes[].ville_naissance')} "
        f"({_required_text(associe.departement_naissance, 'associes[].departement_naissance')})",
    )
    add_paragraph(
        document,
        f"De nationalité {_required_text(associe.nationalite, 'associes[].nationalite')}",
    )
    add_paragraph(
        document,
        _required_text(associe.situation_maritale, "associes[].situation_maritale"),
    )
    add_paragraph(document, f"Demeurant {_person_address(associe)}")


def _add_apport_block(document, data: _ResolvedStatutsCivil) -> None:
    if data.template.expected_type == "scs":
        _add_apport_block_scs(document, data)
        return
    if data.template.expected_type == "micro_holding":
        _add_apport_block_micro_holding(document, data)
        return
    for associe in data.associes:
        _add_apport_line(document, associe, expected_type=data.template.expected_type)
    capital_social = _required_text(data.statuts.capital_social, "statuts_civils.capital_social")
    # SCI / SCI IRIS : total fidele "SOIT AU TOTAL [capital] euros" (source SCI para 110,
    # SCI IRIS para 108). La clause de depot est deja presente dans le modele source (SCI para 111,
    # SCI IRIS para 110) et n'est PAS couverte par la slice apport -> elle est rendue fidelement,
    # accents compris, par le chemin source standard. Cf. _CIVILS_FIX_SPEC_V1.md.
    add_paragraph(
        document,
        f"SOIT AU TOTAL {capital_social} euros",
    )


def _add_apport_block_scs(document, data: _ResolvedStatutsCivil) -> None:
    # SCS (source paras 41-57 ; slice (43,58)).
    # L'en-tete "Le capital social est constitue par les apports en numeraires suivants : /
    # Associes commandites :" est le paragraphe source 41, HORS slice -> il est deja rendu
    # fidelement (accents compris) par le chemin source standard. Le bloc ne le reduplique PAS.
    for associe in _associes_by_role(data.associes, "commandite"):
        _add_apport_line(document, associe, expected_type="scs")
    total_commandites = _required_text(
        data.statuts.total_apports_commandites,
        "statuts_civils.total_apports_commandites",
    )
    # Source para 49 : "Le montant total verse par le commandite est de \t\t\t  [total]."
    add_paragraph(
        document,
        f"Le montant total versé par le commandité est de \t\t\t  {total_commandites}.",
    )
    # Source para 51 : "Associé commanditaire\xa0:" (singulier, accent, NBSP avant deux-points).
    add_paragraph(document, "Associé commanditaire :")
    commanditaires = _associes_by_role(data.associes, "commanditaire")
    for associe in commanditaires:
        _add_apport_line(document, associe, expected_type="scs", commanditaire=True)
    # Source para 56 : "Le montant total verse par le commanditaire est de \t\t\t   [montant]."
    total_commanditaires = _format_amount_total(commanditaires, commanditaire=True)
    add_paragraph(
        document,
        f"Le montant total versé par le commanditaire est de \t\t\t   {total_commanditaires}.",
    )
    capital_social = _required_text(data.statuts.capital_social, "statuts_civils.capital_social")
    # Source para 57 : "Total des apports en numeraires\xa0: \t\t\t\t\t  [capital]" (NBSP avant
    # les deux-points) puis depot SCS.
    add_paragraph(
        document,
        f"Total des apports en numéraires : \t\t\t\t\t  {capital_social}",
    )
    capital_lettres = _required_text(
        data.statuts.capital_social_lettres,
        "statuts_civils.capital_social_lettres",
    )
    depot = data.statuts.capital_depot
    banque_nom = _required_text(
        depot.banque_nom if depot else None,
        "statuts_civils.capital_depot.banque_nom",
    )
    banque_adresse = _required_text(
        depot.banque_adresse if depot else None,
        "statuts_civils.capital_depot.banque_adresse",
    )
    add_paragraph(
        document,
        f"Cette somme de {capital_lettres} ({capital_social}) a été intégralement versée dès avant "
        "ce jour à un compte ouvert au nom de la Société en formation, à la Banque "
        f"{banque_nom}, {banque_adresse}.",
    )


def _add_capital_block(document, data: _ResolvedStatutsCivil) -> None:
    if data.template.expected_type == "scs":
        _add_capital_block_scs(document, data)
        return
    if data.template.expected_type == "micro_holding":
        _add_capital_block_micro_holding(document, data)
        return
    for associe in data.associes:
        parts = _required_parts(associe)
        add_paragraph(document, _signature_label(associe))
        if data.template.expected_type == "sci_iris":
            add_paragraph(
                document,
                "A concurrence de "
                f"{_required_text(parts.nb_lettres, 'associes[].parts.nb_lettres')} parts, "
                f"ci\t{parts.nb} parts Numérotées de "
                f"{_required_int(parts.debut, 'associes[].parts.debut')} à "
                f"{_required_int(parts.fin, 'associes[].parts.fin')}.",
            )
        else:
            # SCI plain : le modele source (Modele statuts SCI.docx, para 120-121) rend
            # "[label]" puis "A concurrence de [lettres] parts, ci<TAB>[nb] parts " (sans
            # numerotation). L'ancien rendu "- [label], / Proprietaire de [lettres] parts
            # sociales [nb] parts sociales" etait un wording INVENTE (croise depuis la SCS),
            # absent du modele SCI -> remplace par le wording source verifie.
            add_paragraph(
                document,
                "A concurrence de "
                f"{_required_text(parts.nb_lettres, 'associes[].parts.nb_lettres')} parts, "
                f"ci\t{parts.nb} parts ",
            )
    add_paragraph(
        document,
        "SOIT AU TOTAL "
        f"{_required_int(data.statuts.nb_parts_total, 'statuts_civils.nb_parts_total')} parts",
    )


def _add_apport_block_micro_holding(document, data: _ResolvedStatutsCivil) -> None:
    # Art. 6 APPORTS (modele Albane P086-P095). Une ligne « <label> apporte la somme de
    # <lettres> » + « Ci\t<montant> euros » par associe, puis total + clause de depot.
    # NB FIDELITE : le modele source utilise des points de conduite (« Ci…… 1010 euros »)
    # comme remplissage visuel d'alignement ; on les rend par une TABULATION (convention du
    # moteur, cf. SCI « ci\t<montant> euros ») -> seul ecart cosmetique, sans valeur juridique.
    for associe in data.associes:
        apport = _required_apport(associe)
        add_paragraph(
            document,
            f"{_mh_short_label(associe)} apporte la somme de "
            f"{_required_text(apport.montant_lettres, 'associes[].apport.montant_lettres')} euros",
        )
        add_paragraph(
            document,
            f"\tCi\t{_required_text(apport.montant, 'associes[].apport.montant')} euros",
        )
    capital_social = _required_text(data.statuts.capital_social, "statuts_civils.capital_social")
    add_paragraph(document, f"Total des apports : \t{capital_social} euros")
    depot = data.statuts.capital_depot
    banque_nom = _required_text(
        depot.banque_nom if depot else None, "statuts_civils.capital_depot.banque_nom"
    )
    banque_adresse = _required_text(
        depot.banque_adresse if depot else None, "statuts_civils.capital_depot.banque_adresse"
    )
    add_paragraph(
        document,
        f"Cette somme de {capital_social} € a été déposée par les associés conformément à la "
        "loi, au crédit d’un compte ouvert au nom de la société en formation auprès de la "
        f"banque {banque_nom}, {banque_adresse}.",
    )


def _add_capital_block_micro_holding(document, data: _ResolvedStatutsCivil) -> None:
    # Art. 7 CAPITAL SOCIAL variable (modele Albane P100-P113). Lettres en MAJUSCULES (verbatim
    # modele). Repartition des parts par associe (« <label> \t<nb> parts »).
    statuts = data.statuts
    capital = _required_text(statuts.capital_social, "statuts_civils.capital_social")
    capital_lettres = _required_text(
        statuts.capital_social_lettres, "statuts_civils.capital_social_lettres"
    ).upper()
    capital_max = _required_text(statuts.capital_maximal, "statuts_civils.capital_maximal")
    capital_max_lettres = _required_text(
        statuts.capital_maximal_lettres, "statuts_civils.capital_maximal_lettres"
    ).upper()
    nb_parts = _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    vnp = _required_text(statuts.valeur_nominale_part, "statuts_civils.valeur_nominale_part")
    vnp_lettres = _required_text(
        statuts.valeur_nominale_part_lettres, "statuts_civils.valeur_nominale_part_lettres"
    ).upper()
    add_paragraph(document, "Le capital social est variable.")
    add_paragraph(document, f"Le capital social minimal est fixé à {capital_lettres} ({capital}€).")
    add_paragraph(
        document,
        f"Le capital social maximal est fixé à {capital_max_lettres} EUROS ({capital_max}€).",
    )
    add_paragraph(
        document,
        f"Le capital social effectif est fixé à {capital_lettres} ({capital} €) euros à la "
        "constitution de la Société.",
    )
    add_paragraph(
        document,
        f"Le capital social est divisé en {nb_parts} parts de {vnp}€ "
        f"({vnp_lettres} {euro_word(vnp).upper()}) chacune.",
    )
    add_paragraph(document, f"Elles sont réparties entre les associés comme suit{_NBSP}:")
    for associe in data.associes:
        parts = _required_parts(associe)
        add_paragraph(document, f"{_mh_short_label(associe)} \t{parts.nb} parts")
    add_paragraph(document, f"Composant le capital social effectif\t{nb_parts} parts")


def _add_capital_block_scs(document, data: _ResolvedStatutsCivil) -> None:
    # SCS (source paras 63-75 ; slice (63,76)).
    # Preambule source para 63 (capital effectif + division + numerotation + attribution) :
    # actuellement supprime par l'ancien rendu -> reintroduit fidelement (accents compris).
    capital_social = _required_text(data.statuts.capital_social, "statuts_civils.capital_social")
    capital_lettres = _required_text(
        data.statuts.capital_social_lettres,
        "statuts_civils.capital_social_lettres",
    )
    nb_parts_total = _required_int(data.statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    nb_parts_total_lettres = _required_text(
        data.statuts.nb_parts_total_lettres,
        "statuts_civils.nb_parts_total_lettres",
    )
    valeur_nominale_part = _required_text(
        data.statuts.valeur_nominale_part,
        "statuts_civils.valeur_nominale_part",
    )
    valeur_nominale_part_lettres = _required_text(
        data.statuts.valeur_nominale_part_lettres,
        "statuts_civils.valeur_nominale_part_lettres",
    )
    plage_parts_total = _required_text(
        data.statuts.plage_parts_totale,
        "statuts_civils.plage_parts_totale",
    )
    add_paragraph(
        document,
        f"Le capital social effectif est fixé à {capital_lettres}({capital_social}) euros. "
        f"Il est divisé en {nb_parts_total_lettres} ({nb_parts_total}) parts sociales de "
        f"{valeur_nominale_part_lettres} ({valeur_nominale_part}) "
        f"{euro_word(valeur_nominale_part)} chacune de valeur nominale, "
        f"numérotées de {plage_parts_total}, lesquelles sont attribuées aux associés comme suit :",
    )
    for associe in data.associes:
        parts = _required_parts(associe)
        # Source paras 63 (suite) / 67 / 71 : "- [label], [qualite],".
        qualite = f", {parts.qualite_associe}" if parts.qualite_associe else ""
        add_paragraph(document, f"- {_signature_label(associe)}{qualite},")
        # Source paras 64 / 68 / 72 : "Proprietaire de [lettres] parts sociales<TAB>[nb] parts
        # sociales " (TAB entre lettres et nombre, espace final).
        add_paragraph(
            document,
            "Propriétaire de "
            f"{_required_text(parts.nb_lettres, 'associes[].parts.nb_lettres')} parts sociales\t"
            f"{parts.nb} parts sociales ",
        )
        # Source paras 65 / 69 / 73 : "Numerotees de [plage]".
        if parts.plage_affichee:
            add_paragraph(document, f"Numérotées de {parts.plage_affichee}")
    # Source para 75 : "Total des parts sociales\xa0composant le capital\xa0:\t\t\t[nb] parts
    # sociales" (NBSP apres "sociales" et avant les deux-points).
    add_paragraph(
        document,
        f"Total des parts sociales composant le capital :\t\t\t{nb_parts_total} parts "
        "sociales",
    )


def _mh_signature_label(associe: StatutsCivilsAssocie) -> str:
    if _is_morale(associe):
        return _mh_morale_denomination(associe)
    prenoms = associe.prenoms or associe.prenom
    return (
        f"{_required_text(associe.civilite_affichage, 'associes[].civilite_affichage')} "
        f"{_required_text(prenoms, 'associes[].prenoms')} "
        f"{_required_text(associe.nom, 'associes[].nom')}"
    )


def _add_signature_block(document, data: _ResolvedStatutsCivil) -> None:
    if data.template.expected_type == "micro_holding":
        # Modele Albane P461-P464 : « Fait a <lieu> » / « Le <date> » sur deux lignes, puis les
        # signataires cote a cote (physiques d'abord, morales ensuite) separes par des tabulations.
        add_paragraph(document, f"Fait à {data.signature_lieu}")
        add_paragraph(document, f"Le {data.signature_date}")
        physiques = [a for a in data.associes if a.est_signataire and not _is_morale(a)]
        morales = [a for a in data.associes if a.est_signataire and _is_morale(a)]
        signers = [_mh_signature_label(a) for a in (physiques + morales)]
        if signers:
            add_paragraph(document, "\t\t\t\t\t".join(signers))
        return
    if data.template.signature_slice is not None:
        add_paragraph(document, f"A {data.signature_lieu}, le {data.signature_date}")
    signers = [_signature_label(a) for a in data.associes if a.est_signataire]
    if data.template.expected_type == "scs":
        add_statuts_signature_grid(document, signers, mention="Lu et approuvé")
        return
    for signer in signers:
        add_statuts_signature_block(
            document,
            [signer],
            bold=True,
            underline=True,
        )


def _add_resultat_groupes_block(document, data: _ResolvedStatutsCivil) -> None:
    rows = []
    for group in data.statuts.resultat_groupes_parts:
        parts_debut = _required_int(
            group.parts_debut,
            "statuts_civils.resultat_groupes_parts[].parts_debut",
        )
        parts_fin = _required_int(
            group.parts_fin,
            "statuts_civils.resultat_groupes_parts[].parts_fin",
        )
        quote_part = _required_text(
            group.quote_part_resultat_exceptionnel,
            "statuts_civils.resultat_groupes_parts[].quote_part_resultat_exceptionnel",
        )
        rows.append((f"Parts numérotées de {parts_debut} à {parts_fin}", quote_part))
    if data.statuts.resultat_quote_part_exceptionnel_total:
        rows.append(("Total", data.statuts.resultat_quote_part_exceptionnel_total))
    add_statuts_matrix_table(
        document,
        ("Groupe de parts", "Quote-part du résultat exceptionnel"),
        rows,
    )


def _add_apport_line(
    document,
    associe: StatutsCivilsAssocie,
    *,
    expected_type: str,
    commanditaire: bool = False,
) -> None:
    apport = associe.apport
    if apport is None:
        raise ValueError(f"associes[].apport est obligatoire pour {DOCUMENT_CODE}.")
    montant = (apport.montant_commanditaire or apport.montant) if commanditaire else apport.montant
    montant_lettres = (
        (apport.montant_commanditaire_lettres or apport.montant_lettres)
        if commanditaire
        else apport.montant_lettres
    )
    montant = _required_text(montant, "associes[].apport.montant")
    montant_lettres = _required_text(montant_lettres, "associes[].apport.montant_lettres")
    if expected_type in {"sci", "sci_iris"}:
        # Modele source SCI (para 97-99) : "[label]" / "La somme de [lettres] euros," /
        # "ci<TAB>[montant] euros". L'ancien format "- [label] apporte, / la somme de
        # [lettres], [montant]" etait le format SCS, croise par erreur sur la SCI
        # (ni "euros", ni "ci", "apporte" invente).
        add_paragraph(document, _signature_label(associe))
        add_paragraph(document, f"La somme de {montant_lettres} euros,")
        add_paragraph(document, f"ci\t{montant} euros")
    else:
        # Format SCS source para 43-44 : "- [label] apporte," puis
        # "la somme de [lettres], <TAB>[montant]" (virgule + espace + TAB).
        add_paragraph(document, f"- {_signature_label(associe)} apporte,")
        add_paragraph(document, f"la somme de {montant_lettres}, \t{montant}")


def _add_physical_identity(document, associe: StatutsCivilsAssocie) -> None:
    gender = associe.genre or Gender.MASCULIN
    born = "Née" if gender == Gender.FEMININ else "Né"
    # R22-06 : la ligne d'identite du comparant est en gras dans la source (comparution).
    _bold_paragraph(add_paragraph(document, _signature_label(associe)))
    add_paragraph(
        document,
        f"{born} le {_format_display_date(associe.date_naissance, 'associes[].date_naissance')} "
        f"à {_required_text(associe.ville_naissance, 'associes[].ville_naissance')} "
        f"({_required_text(associe.departement_naissance, 'associes[].departement_naissance')})",
    )
    add_paragraph(
        document,
        f"De nationalité {_required_text(associe.nationalite, 'associes[].nationalite')}",
    )
    add_paragraph(
        document,
        _required_text(associe.situation_maritale, "associes[].situation_maritale"),
    )
    add_paragraph(document, f"Demeurant {_person_address(associe)}")


def _add_morale_identity(document, associe: StatutsCivilsAssocie) -> None:
    # R22-06 : la ligne d'identite du comparant est en gras dans la source (comparution).
    _bold_paragraph(add_paragraph(document, _signature_label(associe)))
    add_paragraph(
        document,
        f"{_required_text(associe.forme_juridique, 'associes[].forme_juridique')} "
        f"au capital de {_required_text(associe.capital_social, 'associes[].capital_social')}, "
        f"ayant son siege {_address_display(associe.siege, 'associes[].siege')}, "
        f"immatriculee au RCS de {_required_text(associe.ville_rcs, 'associes[].ville_rcs')} "
        f"sous le numero {_required_text(associe.numero_rcs, 'associes[].numero_rcs')}.",
    )
    if associe.representant is None:
        raise ValueError(
            f"associes[].representant est obligatoire pour une personne morale {DOCUMENT_CODE}."
        )
    add_paragraph(
        document,
        "Representee par "
        f"{_required_text(associe.representant.civilite_affichage, 'representant.civilite')} "
        f"{_required_text(associe.representant.prenom, 'associes[].representant.prenom')} "
        f"{_required_text(associe.representant.nom, 'associes[].representant.nom')}, "
        f"{_required_text(associe.representant.fonction, 'associes[].representant.fonction')}.",
    )


def _validate_associes(
    associes: list[StatutsCivilsAssocie],
    template: StatutsCivilTemplate,
) -> None:
    if not associes:
        raise ValueError(f"au moins un associe est obligatoire pour {DOCUMENT_CODE}.")
    if len(associes) > MAX_ASSOCIES:
        raise ValueError(f"les statuts civils sont limites a 6 associes pour {DOCUMENT_CODE}.")
    for associe in associes:
        _required_parts(associe)
        # SCI standard + associe personne morale = AUTORISE (ratifie Rafael 2026-06-08) :
        # rendu via _add_morale_identity, comme SCM / SCI IRIS. (Ancien garde « hors
        # source observee V1 » leve : Rafael confirme le schema SCI -> holding -> SPFPL.)


def _validate_capital_totals(
    statuts: StatutsCivilsContext,
    associes: list[StatutsCivilsAssocie],
) -> None:
    total_parts = sum(_required_int(_required_parts(a).nb, "associes[].parts.nb") for a in associes)
    expected_parts = _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    if total_parts != expected_parts:
        raise ValueError(
            "la somme des parts doit correspondre a statuts_civils.nb_parts_total "
            f"pour {DOCUMENT_CODE}."
        )
    total_apports = sum(_amount_to_int(_required_apport(a).montant) for a in associes)
    expected_capital = _amount_to_int(statuts.capital_social)
    if total_apports != expected_capital:
        raise ValueError(
            "la somme des apports doit correspondre a statuts_civils.capital_social "
            f"pour {DOCUMENT_CODE}."
        )


def _validate_scs(statuts: StatutsCivilsContext, associes: list[StatutsCivilsAssocie]) -> None:
    if not _associes_by_role(associes, "commandite"):
        raise ValueError(f"au moins un associe commandite est obligatoire pour {DOCUMENT_CODE}.")
    if not _associes_by_role(associes, "commanditaire"):
        raise ValueError(f"au moins un associe commanditaire est obligatoire pour {DOCUMENT_CODE}.")
    _required_text(statuts.total_apports_commandites, "statuts_civils.total_apports_commandites")
    _required_text(statuts.capital_maximal, "statuts_civils.capital_maximal")
    _required_text(statuts.capital_maximal_lettres, "statuts_civils.capital_maximal_lettres")


def _validate_sci(associes: list[StatutsCivilsAssocie]) -> None:
    # SCI standard + associe personne morale = AUTORISE (ratifie Rafael 2026-06-08 :
    # une SCI classique peut avoir une autre societe comme associee, schema frequent
    # SCI -> micro-holding -> SPFPL). Aucune contrainte specifique ici : l'identite
    # morale est rendue par _add_morale_identity (qui exige forme/capital/siege/RCS/
    # representant et leve une erreur claire si un champ manque), exactement comme pour
    # SCM et SCI IRIS. La fidelite exacte du wording reste a confirmer par Rafael/Albane.
    _ = associes
    return None


def _validate_micro_holding(statuts: StatutsCivilsContext) -> None:
    # Micro holding (VRAI modele Albane 2026-06-29) : l'objet social (art. 2 « societe civile de
    # portefeuille ») est desormais VERBATIM dans le modele source -> plus de token/variante a
    # valider. Le capital maximal (= 10x le minimum) et ses lettres sont requis (art. 7) :
    _required_text(statuts.capital_maximal, "statuts_civils.capital_maximal")
    _required_text(statuts.capital_maximal_lettres, "statuts_civils.capital_maximal_lettres")
    _required_text(
        statuts.valeur_nominale_part_lettres, "statuts_civils.valeur_nominale_part_lettres"
    )


def _validate_sci_iris(
    statuts: StatutsCivilsContext,
    associes: list[StatutsCivilsAssocie],
) -> None:
    if not any(_is_morale(associe) for associe in associes):
        raise ValueError(
            f"SCI IRIS requiert l'associe personne morale source en V1 pour {DOCUMENT_CODE}."
        )
    if not statuts.resultat_groupes_parts:
        raise ValueError(
            f"statuts_civils.resultat_groupes_parts est obligatoire pour {DOCUMENT_CODE}."
        )
    for group in statuts.resultat_groupes_parts:
        _required_int(group.parts_debut, "statuts_civils.resultat_groupes_parts[].parts_debut")
        _required_int(group.parts_fin, "statuts_civils.resultat_groupes_parts[].parts_fin")
        _required_text(
            group.quote_part_resultat_exceptionnel,
            "statuts_civils.resultat_groupes_parts[].quote_part_resultat_exceptionnel",
        )


def _validate_template_fields(
    statuts: StatutsCivilsContext,
    template: StatutsCivilTemplate,
) -> None:
    _required_text(statuts.capital_social, "statuts_civils.capital_social")
    _required_text(statuts.capital_social_lettres, "statuts_civils.capital_social_lettres")
    _required_int(statuts.nb_parts_total, "statuts_civils.nb_parts_total")
    _required_text(statuts.nb_parts_total_lettres, "statuts_civils.nb_parts_total_lettres")
    _required_text(statuts.valeur_nominale_part, "statuts_civils.valeur_nominale_part")
    _required_text(
        statuts.date_cloture_premier_exercice, "statuts_civils.date_cloture_premier_exercice"
    )
    if template.expected_type in {"sci", "sci_iris", "micro_holding"}:
        _required_text(statuts.mention_capital_variable, "statuts_civils.mention_capital_variable")
        _required_text(statuts.capital_autorise, "statuts_civils.capital_autorise")
        _required_text(statuts.capital_autorise_lettres, "statuts_civils.capital_autorise_lettres")
    if template.expected_type == "scs":
        _required_text(
            statuts.valeur_nominale_part_lettres, "statuts_civils.valeur_nominale_part_lettres"
        )
        _required_text(statuts.plage_parts_totale, "statuts_civils.plage_parts_totale")
        _required_text(statuts.duree_societe, "statuts_civils.duree_societe")
        _required_text(
            statuts.nombre_exemplaires_lettres, "statuts_civils.nombre_exemplaires_lettres"
        )
        _required_text(
            statuts.denomination_cabinet_mandataire,
            "statuts_civils.denomination_cabinet_mandataire",
        )


def _source_path(template: StatutsCivilTemplate) -> Path:
    lot_dir = Path("project/source_documents/lot_04")
    candidates = []
    for path in lot_dir.glob("*.docx"):
        name = path.name.upper()
        if all(token.upper() in name for token in template.source_name_contains):
            if template.expected_type == "sci" and "IRIS" in name:
                continue
            candidates.append(path)
    if len(candidates) != 1:
        raise ValueError(f"source DOCX introuvable ou ambigue pour {DOCUMENT_CODE}: {candidates}")
    return candidates[0]


def _add_rendered_paragraph(document, text: str, source_paragraph=None) -> None:
    if text == "STATUTS":
        add_statuts_title_box(document, text)
    elif text.startswith("TITRE "):
        add_statuts_part_heading(document, text)
    elif text.startswith("ARTICLE ") or text.startswith("Article "):
        add_statuts_article_heading(document, text, left_indent_cm=0.25)
    elif text.startswith("- "):
        add_statuts_hanging_list_item(document, text[2:])
    else:
        _add_source_styled_body(document, text, source_paragraph)


def _add_source_styled_body(document, text: str, source_paragraph) -> None:
    """Rend un paragraphe de corps en PRESERVANT la mise en forme de la source.

    R22-06 (Rafael 2026-06-22, « toute la première page ») : l'en-tête (dénomination /
    forme / capital / siège) est CENTRÉ dans la source, « LES SOUSSIGNES » est en gras
    souligné, les intitulés sont en gras — le moteur les aplatissait en justifié Normal.
    On recopie l'alignement + le gras + le souligné de la source au lieu de les perdre.
    Le corps des articles (justifié, non gras) reste inchangé.
    """
    alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bold = False
    underline = False
    title_like = False
    if source_paragraph is not None:
        style_name = (
            source_paragraph.style.name.lower() if source_paragraph.style else ""
        )
        title_like = "title" in style_name
        src_alignment = source_paragraph.alignment
        if src_alignment is not None:
            alignment = src_alignment
        elif title_like:
            # Titre source (dénomination) : centré via le style, pas via l'alignement.
            alignment = WD_ALIGN_PARAGRAPH.CENTER
        text_runs = [run for run in source_paragraph.runs if run.text.strip()]
        # Gras/souligné de la source recopiés sur la ligne (intitulés, « LES SOUSSIGNES »,
        # noms en comparution…). Le corps des articles source n'a aucun run en gras, donc
        # pas de sur-gras du corps.
        bold = any(bool(run.bold) for run in text_runs)
        underline = any(bool(run.underline) for run in text_runs)
    paragraph = add_statuts_body_paragraph(document, text, alignment=alignment)
    if title_like:
        bold = True
    if bold or underline:
        for run in paragraph.runs:
            if bold:
                run.bold = True
            if underline:
                run.underline = True


def _replace_placeholders(text: str, replacements: dict[str, str]) -> str:
    rendered = text
    for placeholder, value in replacements.items():
        rendered = rendered.replace(placeholder, value)
    return rendered


def _strip_editorial_marker(text: str) -> str:
    """Retire le marqueur editorial interne SCI sans toucher a la clause annotee."""
    return _EDITORIAL_MARKER_RE.sub("", text).rstrip()


def _required_text(value: str | None, field_name: str) -> str:
    if value is None or not str(value).strip():
        raise ValueError(f"{field_name} est obligatoire pour {DOCUMENT_CODE}.")
    return str(value).strip()


def _text_or_empty(value: str | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _required_int(value: int | None, field_name: str) -> int:
    if value is None:
        raise ValueError(f"{field_name} est obligatoire pour {DOCUMENT_CODE}.")
    return value


def _required_parts(associe: StatutsCivilsAssocie):
    if associe.parts is None:
        raise ValueError(f"associes[].parts est obligatoire pour {DOCUMENT_CODE}.")
    _required_int(associe.parts.nb, "associes[].parts.nb")
    _required_text(associe.parts.nb_lettres, "associes[].parts.nb_lettres")
    return associe.parts


def _required_apport(associe: StatutsCivilsAssocie):
    if associe.apport is None:
        raise ValueError(f"associes[].apport est obligatoire pour {DOCUMENT_CODE}.")
    _required_text(associe.apport.montant, "associes[].apport.montant")
    _required_text(associe.apport.montant_lettres, "associes[].apport.montant_lettres")
    return associe.apport


def _format_display_date(value: date | str | None, field_name: str) -> str:
    if value is None:
        raise ValueError(f"{field_name} est obligatoire pour {DOCUMENT_CODE}.")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return _required_text(value, field_name)


def _address_display(address: Address | None, field_name: str) -> str:
    if address is None:
        raise ValueError(f"{field_name} est obligatoire pour {DOCUMENT_CODE}.")
    if address.adresse_affichee:
        return address.adresse_affichee.strip()
    return (
        f"{_required_text(address.num_voie, f'{field_name}.num_voie')} "
        f"{_required_text(address.voie, f'{field_name}.voie')} - "
        f"{_required_text(address.cp, f'{field_name}.cp')} "
        f"{_required_text(address.ville, f'{field_name}.ville')}"
    )


def _person_address(associe: StatutsCivilsAssocie) -> str:
    if associe.adresse_personnelle_affichee:
        return associe.adresse_personnelle_affichee.strip()
    return _address_display(associe.adresse_personnelle, "associes[].adresse_personnelle")


def _signature_label(associe: StatutsCivilsAssocie) -> str:
    if _is_morale(associe):
        denomination = _required_text(associe.denomination, "associes[].denomination")
        if associe.representant is None:
            return denomination
        return (
            f"{denomination}, representee par "
            f"{_required_text(associe.representant.civilite_affichage, 'representant.civilite')} "
            f"{_required_text(associe.representant.prenom, 'associes[].representant.prenom')} "
            f"{_required_text(associe.representant.nom, 'associes[].representant.nom')}"
        )
    prenoms = associe.prenoms or associe.prenom
    return (
        f"{_required_text(associe.civilite_affichage, 'associes[].civilite_affichage')} "
        f"{_required_text(prenoms, 'associes[].prenoms')} "
        f"{_required_text(associe.nom, 'associes[].nom')}"
    )


def _is_morale(associe: StatutsCivilsAssocie) -> bool:
    return associe.type_personne == "personne_morale"


def _associes_by_role(
    associes: list[StatutsCivilsAssocie],
    role: str,
) -> list[StatutsCivilsAssocie]:
    return [a for a in associes if (a.role_statutaire or "").lower() == role]


def _first_part_number(associes: list[StatutsCivilsAssocie]) -> int:
    values = [a.parts.debut for a in associes if a.parts and a.parts.debut is not None]
    return min(values) if values else 1


def _last_part_number(associes: list[StatutsCivilsAssocie]) -> int:
    values = [a.parts.fin for a in associes if a.parts and a.parts.fin is not None]
    if values:
        return max(values)
    return sum(_required_int(_required_parts(a).nb, "associes[].parts.nb") for a in associes)


def _format_amount_total(
    associes: list[StatutsCivilsAssocie],
    *,
    commanditaire: bool = False,
) -> str:
    # Source para 56 : "Le montant total verse par le commanditaire est de ... [montant]".
    # Pour un commanditaire unique, c'est son montant ; pour plusieurs, leur somme (montant total).
    total = 0
    for associe in associes:
        apport = _required_apport(associe)
        if commanditaire:
            montant = apport.montant_commanditaire or apport.montant
        else:
            montant = apport.montant
        total += _amount_to_int(montant)
    return str(total)


def _amount_to_int(value: str | None) -> int:
    text = _required_text(value, "montant")
    normalized = (
        text.replace(" ", "")
        .replace("\u00a0", "")
        .replace(".", "")  # separateur de milliers \u00ab 1.020 \u00bb (montants entiers en euros)
        .replace("euros", "")
        .replace("euro", "")
        .replace("EUR", "")
        .replace("€", "")
        .strip()
    )
    if not normalized.isdigit():
        raise ValueError(f"montant numerique attendu pour {DOCUMENT_CODE}: {value}")
    return int(normalized)
