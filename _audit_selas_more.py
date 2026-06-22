import sys, os; sys.path.insert(0, os.path.abspath("src"))
import re, tempfile
from pathlib import Path
from docx import Document
import _audit_selas_coherence as base
from sydel_doc_engine.front_app import selas_multi_slice as slc

def docx_text(path):
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.extend(p.text for p in c.paragraphs if p.text)
    return "\n".join(parts)

out = Path(tempfile.mkdtemp(prefix="selas_m_", dir="."))
dossier = slc.generate_dossier(base.payload, out)
texts = {Path(p).name: docx_text(p) for p in dossier.docx_paths}

def show(name, pat, win=80):
    norm = " ".join(texts[name].split())
    print(f"\n--- {name} :: /{pat}/ ---")
    for m in re.finditer(pat, norm, re.IGNORECASE):
        s=max(0,m.start()-win); e=min(len(norm),m.end()+win)
        print("  ..."+norm[s:e]+"...")

# Marie Lefevre address as shown in statuts (associe block)
show("statuts_selas_multi.docx", r"Marie Lefevre[^.]{0,140}")
# Marie address in avertissement (conjoint foyer)
show("lettre_avertissement_conjoint.docx", r"(Victor Hugo|Republique|R[ée]publique)")
# how forme prepended -> source of SELAS SELAS
