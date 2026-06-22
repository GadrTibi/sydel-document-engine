import sys, os; sys.path.insert(0, os.path.abspath("src"))
import re, tempfile
from pathlib import Path
from docx import Document
import _audit_selas_coherence as base
from sydel_doc_engine.front_app import selas_multi_slice as slc

def docx_text(path):
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.extend(p.text for p in c.paragraphs if p.text)
    return "\n".join(parts)

out = Path(tempfile.mkdtemp(prefix="selas_ctx_", dir="."))
dossier = slc.generate_dossier(base.payload, out)
texts = {Path(p).name: docx_text(p) for p in dossier.docx_paths}

def show(name, pattern, win=70):
    txt = texts[name]
    norm = " ".join(txt.split())
    print(f"\n--- {name} :: /{pattern}/ ---")
    for m in re.finditer(pattern, norm, re.IGNORECASE):
        s = max(0, m.start()-win); e = min(len(norm), m.end()+win)
        print("  ..."+norm[s:e]+"...")

# 1) doubled SELAS in avertissement
show("lettre_avertissement_conjoint.docx", r"SELAS SELAS")
show("lettre_avertissement_conjoint.docx", r"d[eé]nomm[ée]e?[^.]{0,80}")
# 2) DNC president address vs siege
show("declaration_non_condamnation.docx", r"demeurant[^.]{0,90}")
show("declaration_non_condamnation.docx", r"33000[^.]{0,40}")
# 3) statuts: how denomination appears (the 'S' / 'T' variants)
show("statuts_selas_multi.docx", r"DES TILLEULS [ST]\b", 40)
# 4) capital wording in renonciation/avertissement
show("lettre_renonciation_associe.docx", r"capital[^.]{0,80}")
show("lettre_avertissement_conjoint.docx", r"capital[^.]{0,80}")
# 5) president named in lettres? (should be associe2 Marie Lefevre as renoncant)
show("lettre_renonciation_associe.docx", r"(Marie|Claire|Lefevre|Durand)")
# 6) siege in pv vs statuts
show("pv_nomination_gerant.docx", r"si[èe]ge[^.]{0,80}")
show("statuts_selas_multi.docx", r"si[èe]ge social[^.]{0,80}")
