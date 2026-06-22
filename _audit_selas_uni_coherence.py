import sys, os
sys.path.insert(0, os.path.abspath("src"))

import re
from datetime import date
from pathlib import Path

from docx import Document

from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni


def payload():
    return {
        "denomination": "SELAS MARTIN",
        "capital_social": "1000",
        "nb_actions_total": 1000,
        "duree": "99 ans",
        "ville_rcs": "Paris",
        "lieu_exercice_adresse": "12 avenue de la Republique, 75011 Paris",
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "banque_nom": "BANQUE EXEMPLE",
        "banque_adresse": "1 boulevard Haussmann, 75009 Paris",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 decembre",
        "exercice_cloture": "31 decembre 2026",
        "civilite": "Monsieur",
        "prenom": "Camille",
        "nom": "Martin",
        "date_naissance": date(1980, 1, 2),
        "ville_naissance": "Paris",
        "departement_naissance": "75",
        "nationalite": "francaise",
        "titre_affichage": "Docteur",
        "adresse_num": "5",
        "adresse_voie": "5 rue Royale",
        "adresse_cp": "75008",
        "adresse_ville": "Paris",
        "nom_pere": "Pierre Martin",
        "nom_mere": "Anne Martin",
        "situation_maritale": "marie",
        "regime_matrimonial": "communaute legale",
        "conjoint_civilite": "Madame",
        "conjoint_prenom": "Alice",
        "conjoint_nom": "Martin",
        "ordre_conseil": "Ordre des medecins",
        "departement_ordre": "Paris",
        "numero_ordre": "12345",
        "numero_rpps": "10000000001",
        "ordre_ville": "Paris",
        "ordre_cp": "75008",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 14),
        "decision_date": date(2026, 5, 14),
        # active conditional to get DOC-005/006 too
        "regime_communautaire": True,
        "ordre_president_feminin": False,
        "valeur_nominale_action": "1",
        "mandataire_prenom": "Jean",
        "mandataire_nom": "Conseiller",
    }


def docx_full_text(path):
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def search(text, patterns):
    """Return first match per pattern label."""
    out = {}
    for label, pat in patterns.items():
        m = re.search(pat, text, re.IGNORECASE)
        out[label] = m.group(0).strip() if m else None
    return out


def main():
    out_dir = Path("artifacts/_audit_selas_uni")
    out_dir.mkdir(parents=True, exist_ok=True)
    plan = uni.build_selas_uni_medecin_plan(payload())
    print("can_generate:", plan.can_generate, "| codes:", plan.document_codes)
    gen = uni.generate_dossier(payload(), out_dir)

    texts = {}
    for p in gen.docx_paths:
        texts[p.name] = docx_full_text(p)
        print("GENERATED:", p.name, "len", len(texts[p.name]))

    # Shared fields to probe, with regex per field.
    field_patterns = {
        "denomination": r"SELAS MARTIN",
        "capital_chiffres": r"1\s?000\s*(?:€|euros?)",
        "nb_actions": r"1\s?000\s+actions",
        "valeur_nominale": r"(?:UN|1)\s*(?:€|euros?)\s+(?:de nominal|chacune|par action)?",
        "siege_full": r"10[, ]+rue de la Paix[^\n,]*75002\s*Paris",
        "siege_ville_cp": r"75002\s*Paris",
        "rcs_ville": r"R\.?C\.?S\.?\s*Paris|RCS de Paris|registre du commerce[^\n]*Paris",
        "associe_nom": r"Camille\s+Martin|Martin\s+Camille|M(?:onsieur|\.)?\s+Camille\s+Martin",
        "conjoint": r"Alice\s+Martin",
        "signature_date": r"14\s+mai\s+2026|14/05/2026",
        "cloture_date": r"31\s+d[ée]cembre\s+2026",
    }

    print("\n=== FIELD EXTRACTION PER DOC ===")
    results = {}
    for name, txt in texts.items():
        results[name] = search(txt, field_patterns)
    # print table
    fields = list(field_patterns.keys())
    for f in fields:
        print(f"\n--- {f} ---")
        for name in texts:
            print(f"  {name:42s} : {results[name][f]!r}")

    # Dump raw text files for manual probing
    for name, txt in texts.items():
        (out_dir / (name + ".txt")).write_text(txt, encoding="utf-8")
    print("\nDumped txt to", out_dir)


if __name__ == "__main__":
    main()
