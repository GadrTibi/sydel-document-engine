"""KAN-2 — Garde de conformité : générer le dossier SPFPL SANS AUCUN champ rempli.

Rafael (rejeté 2×) : « Tous les documents doivent pouvoir être générés, même si je ne remplis
AUCUN champ. » Un champ vide sort en MARQUEUR « (À COMPLÉTER : <libellé métier>) », JAMAIS une
valeur FAUSSE/inventée NI un BLANC silencieux (« Fait à », « Le » nus).

Ce garde génère les dossiers cession ET apport avec un payload VIDE FIDÈLE (les SÉLECTEURS
`operation`/`operation_spfpl_type`/`is_apport` restent — ce sont le TYPE de doc, pas des champs)
et asserte, sur TOUT le texte de TOUS les docx — **paragraphes ET cellules de tableau** —, le
CONCEPT (règle 68 §3 : on code l'intention, pas la liste des tournures déjà signalées) :

  * S  : aucun ANCRAGE SIGNATURE NU — pas de « Fait à » / « Le » suivi de rien (lieu/date
    manquants sortent en marqueur, jamais un blanc silencieux) ;
  * B1 : aucune QUANTITÉ DE TITRES affirmée — ni chiffrée (« 600 actions », « (0) parts »),
    ni en lettres (« six cents »), ni « zéro (0) » — elle sort en marqueur ;
  * B2 : aucun STATUT MATRIMONIAL affirmé (célibataire / marié·e / divorcé·e / veuf·ve /
    pacsé·e) — le statut manquant sort en marqueur ;
  * B4 : aucune CIVILITÉ inventée pour un signataire (« Monsieur/Madame ») — seuls les cas
    LÉGITIMES ancrés (destinataire de l'Ordre, mandataire SYDEL, parents définitionnels) ;
  * M1 : aucune PROFESSION D'INDIVIDU affirmée — le boilerplate LÉGAL des statuts et l'Ordre
    sont légitimes, mais le TYPE ne fuit pas sur une personne dont l'identité est un marqueur ;
  * M2 : aucun MARQUEUR TECHNIQUE — pas de point, underscore, chiffre d'index ni casse de token
    dans « (À COMPLÉTER : … ) » ; que des libellés métier lisibles.

Il attrape MÉCANIQUEMENT toute réintroduction d'un défaut de la classe (règle 68 §3) — la suite
nominale ne les voyait pas, et coder les seuls littéraux déjà vus laissait passer les nouvelles
formulations (leçon 2026-07-09) et les blancs silencieux (Akainu 4e passe 2026-07-17).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest
from docx import Document

_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_ROOT / "tests" / "unit"))

import test_multi_type_front as mtf  # noqa: E402

from sydel_doc_engine.front_app import spfpl_slice  # noqa: E402

# Monsieur/Madame LÉGITIMES sur formulaire vide (pas des civilités inventées d'un signataire) :
# destinataire de l'Ordre, mandataire SYDEL hardcodé, civilités DÉFINITIONNELLES des parents.
# ANCRÉS (règle 68 §3 / m2) : « fils/fille de Monsieur », jamais un « de Monsieur » nu qui
# masquerait un « représentée de Monsieur <nom inventé> ».
_CIVILITE_LEGITIMES = (
    "Monsieur le Président",
    "Madame la Présidente",
    "Monsieur Jordan",
    "fils de Monsieur",
    "fille de Monsieur",
    "et de Madame",
)

# Statuts matrimoniaux — AUCUN ne doit être affirmé sur un formulaire vide (le champ
# situation_maritale manquant sort en marqueur). Concept, pas le seul « célibataire ».
_STATUTS_MATRIMONIAUX = re.compile(
    r"\b(célibataire|marié|mariée|divorcé|divorcée|veuf|veuve|pacsé|pacsée)\b",
    re.IGNORECASE,
)


def _empty_payload(structure: str) -> dict:
    full = mtf._spfpl_payload(structure)
    payload = {
        key: ("" if isinstance(value, str) else (value if isinstance(value, bool) else None))
        for key, value in full.items()
    }
    payload["structure"] = structure
    for selector in ("operation", "operation_spfpl_type", "is_apport"):
        if selector in full:
            payload[selector] = full[selector]
    payload["cession_data"] = {}
    payload["apport_data"] = {}
    return payload


def _text_units(structure: str, tmp_path: Path) -> list[tuple[str, str]]:
    """Toutes les unités de texte (paragraphes + cellules de tableau) de tous les docx.

    Le garde balaie les DEUX : une valeur régressée dans une cellule serait invisible si l'on
    ne lisait que `Document(...).paragraphs`.
    """
    out = tmp_path / structure.replace(" ", "_")
    out.mkdir(parents=True, exist_ok=True)
    generated = spfpl_slice.generate_dossier(_empty_payload(structure), out)
    assert generated.docx_paths, f"aucun document généré pour {structure} vide"
    pairs: list[tuple[str, str]] = []
    for path in generated.docx_paths:
        doc = Document(str(path))
        for p in doc.paragraphs:
            pairs.append((path.name, p.text))
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        pairs.append((path.name, p.text))
    return pairs


@pytest.mark.parametrize("structure", ["SPFPL cession", "SPFPL apport"])
def test_spfpl_dossier_generates_from_empty_form(structure: str, tmp_path: Path) -> None:
    pairs = _text_units(structure, tmp_path)
    text = "\n".join(u for _, u in pairs)

    # S : ancrage signature NU (blanc silencieux) — étiquette « Fait à » / « Le » sans valeur.
    # « Fait à » porte TOUJOURS un lieu inline -> un « Fait à » nu est un blanc, partout.
    # « Le » nu est TOLÉRÉ pour le SEUL statuts CESSION (fidélité modèle ratifiée Rafael
    # 2026-07-09, test ...without_signature_date : bloc signature « Le » suivi du signataire,
    # date volontairement absente). Le statuts APPORT, lui, rend « Le (À COMPLÉTER : date de
    # signature) » (template statuts_spfpl_templates:786) -> il reste SOUS contrôle, comme
    # contrat / acte / attestations. Exemption cadrée sur le doc ratifié, pas « statuts » large
    # (Akainu 5e passe m2 : sinon un « Le » nu régressé dans le statuts apport passerait vert).
    dangling = [
        f"{doc}: {u.strip()!r}"
        for doc, u in pairs
        if u.strip() in ("Fait à", "Fait à,")
        or (u.strip() in ("Le", "le") and "statuts_spfpl_cession" not in doc.lower())
    ]
    assert not dangling, (
        f"{structure} vide : ANCRAGE SIGNATURE NU (blanc silencieux, doit être un marqueur) : "
        f"{dangling}"
    )

    # B1 : quantité de titres affirmée (chiffre, lettres, ou « (0) »/« zéro (0) »). On code le
    # CONCEPT « titres » (actions / parts / titres — le corpus emploie les trois : « Titres
    # Cédés », « valeur d'un titre apporté »), pas les seules tournures vues (Akainu 5e passe m1).
    quantite_chiffree = re.findall(r"\b\d[\d\s]*\s+(?:actions|parts|titres)\b", text)
    assert not quantite_chiffree, (
        f"{structure} vide : quantité de titres CHIFFRÉE affirmée (doit être un marqueur) : "
        f"{quantite_chiffree[:3]}"
    )
    assert "six cents" not in text.lower(), (
        f"{structure} vide : nombre d'actions en LETTRES affirmé « six cents » (marqueur attendu)"
    )
    assert (
        "(0) actions" not in text
        and "(0) parts" not in text
        and "(0) titres" not in text
        and "zéro (0)" not in text
    ), f"{structure} vide : quantité de titres « (0) » affirmée (doit être un marqueur)"

    # B2 : statut matrimonial affirmé (concept complet, pas seulement « célibataire »).
    matrimonial = _STATUTS_MATRIMONIAUX.findall(text)
    assert not matrimonial, (
        f"{structure} vide : statut matrimonial affirmé {sorted(set(matrimonial))} "
        f"(le champ manquant doit sortir en marqueur, aucun statut par défaut)"
    )

    # B4 : civilité « Monsieur/Madame » inventée pour un signataire (hors cas légitimes ancrés).
    residual = text
    for legit in _CIVILITE_LEGITIMES:
        residual = residual.replace(legit, "")
    assert "Monsieur" not in residual and "Madame" not in residual, (
        f"{structure} vide : civilité « Monsieur/Madame » inventée pour un signataire"
    )

    # M1 : profession d'INDIVIDU affirmée (descripteur d'identité), hors boilerplate légal.
    individu_profession = re.findall(
        r"(?<!la profession de )[Cc]hirurgiens?-[Dd]entistes? de profession", text
    ) + re.findall(r"\), [Cc]hirurgiens?-[Dd]entistes?, de nationalité", text)
    assert not individu_profession, (
        f"{structure} vide : profession d'individu « chirurgien-dentiste » affirmée "
        f"(le TYPE ne doit pas fuir sur une personne dont l'identité est un marqueur) : "
        f"{individu_profession[:3]}"
    )

    # M2 : marqueur technique (point / underscore / chiffre d'index / casse de token).
    marqueurs = re.findall(r"\(À COMPLÉTER : [^)]*\)", text)
    techniques = [
        m
        for m in marqueurs
        if ("." in m or "_" in m or re.search(r"\d", m) or re.search(r"[a-z][A-Z]{2}", m))
    ]
    assert not techniques, (
        f"{structure} vide : marqueur(s) TECHNIQUE(S) (doit être un libellé métier lisible) : "
        f"{sorted(set(techniques))[:5]}"
    )
