"""Corpus de conformité transverse — génération RÉELLE de tous les bundles.

Génère le bundle COMPLET de chaque type d'entreprise via les slices front réels
(``generate_dossier`` des slices), en RÉUTILISANT les payloads des tests front
existants (``test_multi_type_front`` / ``test_clean_front_app`` /
``test_sasu_holding_plan`` / ``test_statuts_micro_holding``) — jamais de payload
réinventé. Le corpus est un dict ``{type_key: {nom_doc: texte}}``.

La suite de conformité (``test_conformite_transverse.py``) applique ensuite les
règles transversales client (R1..R9) à CHAQUE document de CHAQUE type. Un retour
client codifié = une assertion permanente : il ne peut plus revenir silencieusement
sur aucun type.

NB chemins : ce module insère ``src`` et ``tests/unit`` dans ``sys.path`` pour
importer les payloads des tests unitaires par nom de module (même mécanique que la
collecte pytest en mode ``prepend``).
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_ROOT / "src"), str(_ROOT / "tests" / "unit")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from docx import Document  # noqa: E402


def docx_text(path: Path) -> str:
    """Texte intégral d'un DOCX : paragraphes + tables + en-têtes + pieds de page."""
    document = Document(path)
    texts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs if p.text)
    for section in document.sections:
        texts.extend(p.text for p in section.header.paragraphs if p.text)
        texts.extend(p.text for p in section.footer.paragraphs if p.text)
    return "\n".join(texts)


def _bundle(generated) -> dict[str, str]:
    return {p.name: docx_text(p) for p in generated.docx_paths}


# Ordre stable des types du corpus (sert à la paramétrisation pytest — PAS de
# génération à la collecte : la liste est statique, le corpus est bâti en fixture).
CORPUS_KEYS: tuple[str, ...] = (
    "selarl",
    "selarl_regime",
    "selarl_cession_medical",
    "selarl_cession_dentaire",
    "selarl_cession_scm",
    "selas_multi",
    "selas_multi_phys",
    "selas_uni_medecin",
    "selas_uni_dentiste",
    "spfpl_cession",
    # Variante valeur nominale = 1 € (surface R6 : « actions de un euro »), bundle
    # restreint aux statuts pour ne pas dupliquer les autres cellules du type.
    "spfpl_cession_vn1",
    "spfpl_apport",
    "sas",
    "sasu_holding",
    "sci",
    "sci_iris",
    "scm",
    "scs",
    "micro_holding",
)

# R11 (Albane, Direction Juridique, + Rafael 2026-07-09) : nombre de GÉRANTS PERSONNES
# PHYSIQUES du payload de chaque type — le bundle doit porter EXACTEMENT une DNC par
# gérant physique. SUPERSEDE le retour Rafael du matin « 1 DNC par associé » (verrou R11
# initial) : la DNC ne se génère QUE pour les gérants, pas pour les associés non-gérants
# (Rafael a confirmé « Albane a raison »). Tenu ALIGNÉ sur les payloads assemblés par
# ``build_corpus`` (tout changement de gouvernance/roster doit mettre cette table à jour ;
# le test R11 échoue sinon, c'est son rôle).
#   - civils multi (sci/scm/scs/micro_holding) : 1 gérant (gerant_index, 1er physique).
#   - selas_multi_phys : 1 dirigeant (seul le président Durand est coché dirigeant).
#   - unipersonnels : l'associé unique EST gérant → 1.
EXPECTED_DNC_GERANTS: dict[str, int] = {
    "selarl": 1,  # unipersonnelle — le praticien est gérant
    "selarl_regime": 1,
    "selarl_cession_medical": 1,
    "selarl_cession_dentaire": 1,
    "selarl_cession_scm": 1,
    "selas_multi": 1,  # 1 PP (président/dirigeant) + 1 PM (pas de DNC PM)
    "selas_multi_phys": 1,  # Durand + Martin, seul Durand (président) est dirigeant
    "selas_uni_medecin": 1,
    "selas_uni_dentiste": 1,
    "spfpl_cession": 1,  # associé unique (multi bloqué par le moteur)
    "spfpl_cession_vn1": 0,  # bundle FILTRÉ aux statuts seuls (aucune DNC attendue)
    "spfpl_apport": 1,
    "sas": 1,  # actionnaire unique = président
    "sasu_holding": 1,
    "sci": 1,  # Durand + Martin, 1 seul gérant (Durand)
    "sci_iris": 1,  # 1 PM + 1 PP (le PP est gérant)
    "scm": 1,  # 2 associés, 1 gérant
    "scs": 1,  # commandité (gérant) + commanditaire (pas de DNC)
    "micro_holding": 1,  # Durand + Martin, 1 gérant (gérante Durand)
}


def _normalise_civil(payload: dict) -> dict:
    """Aligne le payload civil des tests sur ce que l'UI RÉELLE envoie.

    Les payloads unitaires portent des raccourcis non représentatifs qui parasiteraient
    la conformité (R4) sans révéler de défaut GÉNÉRATEUR :
    - ``forme_sociale`` saisi « societe civile » (non accentué) alors que le formulaire
      la DÉRIVE (§18.1, ``civil_forme_sociale`` → « société civile ») → on retire la clé
      pour exercer la dérivation réelle ;
    - ``nationalite``/``situation_maritale`` des associés : l'UI réelle passe par des
      dérivations accentuées (``render_nationalite_selectbox`` → « française » ;
      ``_situation_display_civil`` → « célibataire ») — jamais la forme nue non accentuée.
    """
    payload.pop("forme_sociale", None)
    for associe in payload["associes"]:
        if associe.type_personne == "personne_physique":
            associe.nationalite = "française"
            associe.situation_maritale = "célibataire"
    return payload


def build_corpus(base_dir: Path) -> dict[str, dict[str, str]]:  # noqa: C901 - assemblage séquentiel de 19 bundles, pas de logique
    """Génère une fois le bundle de chaque type et retourne {type: {doc: texte}}."""
    # Imports différés : payloads des tests front existants (réutilisés tels quels).
    import test_clean_front_app as cfa
    import test_multi_type_front as mtf
    import test_sasu_holding_plan as sasu_tests
    import test_statuts_micro_holding as micro_tests

    from sydel_doc_engine.front_app import (
        civil_statuts_slice as css,
    )
    from sydel_doc_engine.front_app import (
        sas_slice,
        sasu_holding_slice,
        selas_multi_slice,
        spfpl_slice,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_dentiste_slice as sd,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_medecin_slice as sm,
    )
    from sydel_doc_engine.front_app.selarl_slice import (
        PROFESSION_MEDECIN,
        generate_selarl_dossier,
    )
    from sydel_doc_engine.scenarios.selarl import build_selarl_scenario

    corpus: dict[str, dict[str, str]] = {}

    # --- SELARL (chemin historique dédié) --------------------------------------
    # ``siege_voie`` du payload de test vaut « avenue du Siege » (non accentué) : pur
    # artefact d'adresse de test qui parasiterait R4 — on nomme une voie neutre.
    selarl_overrides = {"siege_voie": "avenue de Breteuil"}
    corpus["selarl"] = _bundle(
        generate_selarl_dossier(
            cfa._valid_selarl_input(PROFESSION_MEDECIN, **selarl_overrides),
            base_dir / "selarl",
        )
    )
    # Régime communautaire : ajoute lettre de renonciation + avertissement conjoint.
    corpus["selarl_regime"] = _bundle(
        generate_selarl_dossier(
            cfa._valid_selarl_input(
                PROFESSION_MEDECIN, regime_communautaire=True, **selarl_overrides
            ),
            base_dir / "selarl_regime",
        )
    )
    # Cessions (scénarios figés du moteur) : actes + bail + appel de fonds + docs SCM.
    # Normalisation d'artefacts de fixture (couche SOUS l'UI réelle) :
    # - ``siege_voie`` du scénario vaut « avenue du Siege » (adresse de test non
    #   accentuée) → voie neutre ;
    # - le cédant SCM du scénario porte « francaise » / « marie » nus alors que le
    #   chemin UI réel dérive des libellés accentués (``_scm_cedant_situation_
    #   maritale_display``, O24-11) — le scénario, lui, court-circuite la dérivation.
    def _selarl_scenario_ui(key: str):
        import dataclasses

        data = dataclasses.replace(build_selarl_scenario(key), siege_voie="avenue de Breteuil")
        if data.scm_cession_context is not None:
            data.scm_cession_context.cedant.nationalite = "française"
            data.scm_cession_context.cedant.situation_maritale = "marié"
        return data

    corpus["selarl_cession_medical"] = _bundle(
        generate_selarl_dossier(
            _selarl_scenario_ui("selarl_medecin_cession_cabinet_medical"),
            base_dir / "selarl_cession_medical",
        )
    )
    corpus["selarl_cession_dentaire"] = _bundle(
        generate_selarl_dossier(
            _selarl_scenario_ui("selarl_dentiste_cession_cabinet_dentaire"),
            base_dir / "selarl_cession_dentaire",
        )
    )
    corpus["selarl_cession_scm"] = _bundle(
        generate_selarl_dossier(
            _selarl_scenario_ui("selarl_dentiste_cession_scm"),
            base_dir / "selarl_cession_scm",
        )
    )

    # --- SELAS pluripersonnelle -------------------------------------------------
    # Variante mixte (1 physique + 1 personne morale) : chemin identité morale.
    corpus["selas_multi"] = _bundle(
        selas_multi_slice.generate_dossier(mtf._selas_payload(), base_dir / "selas_multi")
    )
    # Variante entièrement physique : ajoute l'attestation souscripteurs (DOC-045).
    corpus["selas_multi_phys"] = _bundle(
        selas_multi_slice.generate_dossier(
            mtf._selas_payload_n(
                [mtf._selas_phys("Claire", "Durand", 60), mtf._selas_phys("Paul", "Martin", 40)]
            ),
            base_dir / "selas_multi_phys",
        )
    )

    # --- SELAS unipersonnelles (médecin / dentiste) ------------------------------
    corpus["selas_uni_medecin"] = _bundle(
        sm.generate_dossier(mtf._selas_uni_payload(), base_dir / "selas_uni_medecin")
    )
    corpus["selas_uni_dentiste"] = _bundle(
        sd.generate_dossier(mtf._selas_uni_payload(), base_dir / "selas_uni_dentiste")
    )

    # --- SPFPL dentistes (cession + apport) --------------------------------------
    def _spfpl_payload_ui(structure: str) -> dict:
        # L'UI réelle sert la nationalité par déroulant accentué (« française ») et la
        # forme complète de la cible est une saisie dont la forme représentative est
        # accentuée — le payload de test porte des raccourcis non accentués qui
        # parasiteraient R4 sans révéler de défaut générateur.
        payload = mtf._spfpl_payload(structure)
        payload["nationalite"] = "française"
        payload["cession_data"] = {
            **payload["cession_data"],
            "cible_forme_complete": "société d'exercice libéral à responsabilité limitée",
        }
        return payload

    corpus["spfpl_cession"] = _bundle(
        spfpl_slice.generate_dossier(_spfpl_payload_ui("SPFPL cession"), base_dir / "spfpl_cession")
    )
    # Rafael 2026-07-09 : actionnaire MARIE sous communaute legale -> DOC-005/006
    # (renonciation + avertissement conjoint) entrent dans le corpus. La double unite
    # « soixante mille euros (60 000) euros » de la lettre avait echappe a R7 parce que
    # le bundle scanne ne CONTENAIT pas la lettre (payload sans regime communautaire).
    # Valeurs accentuees = ce que l'UI reelle derive (situation_display /
    # regime_matrimonial_from_status).
    spfpl_apport_payload = _spfpl_payload_ui("SPFPL apport")
    spfpl_apport_payload.update(
        {
            "situation_maritale": "marié",
            "regime_matrimonial": "la communauté légale",
            "regime_communautaire": True,
        }
    )
    corpus["spfpl_apport"] = _bundle(
        spfpl_slice.generate_dossier(spfpl_apport_payload, base_dir / "spfpl_apport")
    )
    # Variante valeur nominale = 1 € : surface du retour élision (« actions de un
    # euro » attendu « d'un euro », R6). Bundle restreint aux STATUTS : les autres
    # documents dupliqueraient les cellules du type spfpl_cession sans rien prouver.
    vn1_payload = _spfpl_payload_ui("SPFPL cession")
    vn1_payload.update(
        {"capital_social": "600", "nb_actions_total": 600, "valeur_nominale_action": "1"}
    )
    corpus["spfpl_cession_vn1"] = {
        name: text
        for name, text in _bundle(
            spfpl_slice.generate_dossier(vn1_payload, base_dir / "spfpl_cession_vn1")
        ).items()
        # R10 (Rafael 2026-07-07) : le fichier statuts porte la denomination
        # (« Statuts SPFPL MARTIN.docx ») — filtre par prefixe, robuste au payload.
        if name.startswith("Statuts ")
    }

    # --- SAS (SPFPL médecins forme SAS) + SASU holding ---------------------------
    # Même normalisation nationalité (déroulant accentué dans l'UI réelle).
    sas_payload = dict(mtf._sas_payload())
    sas_payload["nationalite"] = "française"
    corpus["sas"] = _bundle(sas_slice.generate_dossier(sas_payload, base_dir / "sas"))
    corpus["sasu_holding"] = _bundle(
        sasu_holding_slice.generate_dossier(sasu_tests._payload(), base_dir / "sasu_holding")
    )

    # --- Types civils (socle partagé) ---------------------------------------------
    sci_payload = mtf._civil_base(
        "SCI",
        "sci",
        [
            mtf._pp("Jean", "Durand", 40, 1, 40, 400),
            mtf._pp("Alice", "Martin", 60, 41, 100, 600),
        ],
    )
    # Option IS activée : la lettre d'option IS (DOC-022) entre dans le bundle scanné.
    sci_payload.update(mtf._OPTION_IS_INPUTS)
    corpus["sci"] = _bundle(css.generate_dossier(_normalise_civil(sci_payload), base_dir / "sci"))

    corpus["sci_iris"] = _bundle(
        css.generate_dossier(
            _normalise_civil(
                mtf._civil_base(
                    "SCI IRIS",
                    "sci_iris",
                    [mtf._pm(40, 1, 40, 400), mtf._pp("Alice", "Martin", 60, 41, 100, 600)],
                )
            ),
            base_dir / "sci_iris",
        )
    )

    # SCM à 2 associés physiques + inter-SEL actif : bundle maximal (statuts + demande
    # ordre + pacte + liste dépenses + contrat frais communs + règlement intérieur).
    scm_payload = mtf._civil_base(
        "SCM",
        "scm",
        [
            mtf._pp("Jean", "Durand", 50, 1, 50, 500),
            mtf._pp("Alice", "Martin", 50, 51, 100, 500),
        ],
    )
    scm_payload.update(mtf._SCM_INTER_SEL_INPUTS)
    corpus["scm"] = _bundle(css.generate_dossier(_normalise_civil(scm_payload), base_dir / "scm"))

    corpus["scs"] = _bundle(
        css.generate_dossier(
            _normalise_civil(
                mtf._civil_base(
                    "SCS",
                    "scs",
                    [
                        mtf._pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
                        mtf._pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
                    ],
                )
            ),
            base_dir / "scs",
        )
    )

    # Le payload micro holding des tests statuts ne porte pas les saisies du tronc
    # commun (DNC / procuration) : on les complète comme le fait le formulaire réel
    # (mêmes champs que ``_civil_base`` de test_multi_type_front).
    micro_payload = micro_tests._payload()
    micro_payload.update(
        {
            "signataire_nom_pere": "Pierre Durand",
            "signataire_nom_mere": "Anne Durand",
            "signataire_adresse_num": "1",
            "signataire_adresse_voie": "rue Exemple",
            "signataire_adresse_cp": "75000",
            "signataire_adresse_ville": "Paris",
            "signataire_fonction": "gerant",
        }
    )
    corpus["micro_holding"] = _bundle(
        css.generate_dossier(micro_payload, base_dir / "micro_holding")
    )

    missing = set(CORPUS_KEYS) - set(corpus)
    if missing:  # défense : la liste statique et le bâtisseur doivent rester alignés
        raise RuntimeError(f"types du corpus non générés : {sorted(missing)}")
    return corpus


# ---------------------------------------------------------------------------
# Corpus « montant unitaire = 1 » (R13, accord euro/euros — Rafael 2026-07-09)
# ---------------------------------------------------------------------------
#
# But : régénérer chaque type avec un CAPITAL / APPORT de 1 euro pour surfacer TOUT
# « 1 euros » (accord singulier faux) — hardcodé en Python OU figé dans un modèle
# source. Un spot silo-é (« 1 euros » resté sur un seul type/document) échoue ICI, sur
# TOUS les types concernés à la fois (gate registre de propagation, règle 68 Q4).
#
# Périmètre : les types dont le capital/apport se force PROPREMENT à 1 (1 part de 1 €,
# ou apport de 1 € par associé) sans casser un validateur de roster ni un champ
# « lettres » figé. Les scénarios de CESSION SELARL (figés) et la micro holding (capital
# variable, jamais 1 €) ne sont PAS régénérés à 1 € : leurs générateurs sont partagés et
# déjà couverts par les types ci-dessous (statuts_civils_common, autorisation, PV…).
CORPUS_CAP1_KEYS: tuple[str, ...] = (
    "selarl",
    "selarl_regime",
    "selas_multi",
    "selas_multi_phys",
    "selas_uni_medecin",
    "selas_uni_dentiste",
    "spfpl_cession",
    "spfpl_apport",
    "sas",
    "sasu_holding",
    "sci",
    "sci_iris",
    "scm",
    "scs",
)


def _apport_1_euro(associe):
    """Copie pydantic d'un associé dont l'apport est forcé à 1 euro (1 action/part)."""
    from sydel_doc_engine.domain.models import StatutsCivilsApport

    update: dict[str, object] = {
        "apport": StatutsCivilsApport(montant="1", montant_lettres="un")
    }
    if getattr(associe, "nb_actions", None) is not None:
        update["nb_actions"] = 1
    return associe.model_copy(update=update)


def build_corpus_cap1(base_dir: Path) -> dict[str, dict[str, str]]:  # noqa: C901 - assemblage séquentiel, pas de logique
    """Génère le bundle de chaque type de ``CORPUS_CAP1_KEYS`` avec un montant = 1 €."""
    import test_clean_front_app as cfa
    import test_multi_type_front as mtf
    import test_sasu_holding_plan as sasu_tests

    from sydel_doc_engine.front_app import (
        civil_statuts_slice as css,
    )
    from sydel_doc_engine.front_app import (
        sas_slice,
        sasu_holding_slice,
        selas_multi_slice,
        spfpl_slice,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_dentiste_slice as sd,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_medecin_slice as sm,
    )
    from sydel_doc_engine.front_app.selarl_slice import (
        PROFESSION_MEDECIN,
        generate_selarl_dossier,
    )

    corpus: dict[str, dict[str, str]] = {}

    # --- SELARL création : 1 part de 1 euro (capital = 1). --------------------
    selarl_cap1 = {
        "siege_voie": "avenue de Breteuil",
        "capital_social": "1",
        "nb_parts_total": 1,
        "valeur_nominale_part": "1",
    }
    corpus["selarl"] = _bundle(
        generate_selarl_dossier(
            cfa._valid_selarl_input(PROFESSION_MEDECIN, **selarl_cap1),
            base_dir / "selarl",
        )
    )
    corpus["selarl_regime"] = _bundle(
        generate_selarl_dossier(
            cfa._valid_selarl_input(PROFESSION_MEDECIN, regime_communautaire=True, **selarl_cap1),
            base_dir / "selarl_regime",
        )
    )

    # --- SELAS pluripersonnelle : apport de 1 euro par associé (capital = 2). ---
    _cap2 = {"capital_social": "2", "nb_actions_total": 2, "valeur_nominale_action": "1"}
    selas_mixte = mtf._selas_payload()
    selas_mixte["associes"] = [_apport_1_euro(a) for a in selas_mixte["associes"]]
    selas_mixte.update(_cap2)
    corpus["selas_multi"] = _bundle(
        selas_multi_slice.generate_dossier(selas_mixte, base_dir / "selas_multi")
    )
    selas_phys = mtf._selas_payload_n(
        [_apport_1_euro(mtf._selas_phys("Claire", "Durand", 1)),
         _apport_1_euro(mtf._selas_phys("Paul", "Martin", 1))]
    )
    selas_phys.update(_cap2)
    corpus["selas_multi_phys"] = _bundle(
        selas_multi_slice.generate_dossier(selas_phys, base_dir / "selas_multi_phys")
    )

    # --- SELAS unipersonnelles : capital = 1 euro (1 action de 1 €). -----------
    def _selas_uni_cap1(slice_mod, key: str):
        payload = mtf._selas_uni_payload()
        payload.update(
            {"capital_social": "1", "nb_actions_total": 1, "valeur_nominale_action": "1"}
        )
        return slice_mod.generate_dossier(payload, base_dir / key)

    corpus["selas_uni_medecin"] = _bundle(_selas_uni_cap1(sm, "selas_uni_medecin"))
    corpus["selas_uni_dentiste"] = _bundle(_selas_uni_cap1(sd, "selas_uni_dentiste"))

    # --- SPFPL cession / apport : capital = 1 euro (1 action de 1 €). ----------
    def _spfpl_cap1(structure: str, key: str, extra: dict | None = None):
        payload = mtf._spfpl_payload(structure)
        payload["nationalite"] = "française"
        payload["cession_data"] = {
            **payload["cession_data"],
            "cible_forme_complete": "société d'exercice libéral à responsabilité limitée",
        }
        payload.update(
            {
                "capital_social": "1",
                "nb_actions_total": 1,
                "valeur_nominale_action": "1",
                "apport_montant": "1",
                "apport_nb_parts": 1,
                "apport_valeur_globale": "1",
            }
        )
        if extra:
            payload.update(extra)
        return spfpl_slice.generate_dossier(payload, base_dir / key)

    corpus["spfpl_cession"] = _bundle(_spfpl_cap1("SPFPL cession", "spfpl_cession"))
    corpus["spfpl_apport"] = _bundle(
        _spfpl_cap1(
            "SPFPL apport",
            "spfpl_apport",
            {
                "situation_maritale": "marié",
                "regime_matrimonial": "la communauté légale",
                "regime_communautaire": True,
            },
        )
    )

    # --- SAS : capital = 1 euro (numéraire 1 €, aucun apport en nature). -------
    sas_payload = dict(mtf._sas_payload())
    sas_payload["nationalite"] = "française"
    sas_payload.update(
        {
            "capital_social": "1",
            "nb_actions_total": 1,
            "valeur_nominale_action": "1",
            "apports_nature_montant": "0",
            "apports_numeraire_montant": "1",
            "apport_nb_parts": 1,
        }
    )
    corpus["sas"] = _bundle(sas_slice.generate_dossier(sas_payload, base_dir / "sas"))

    # --- SASU holding : capital = 1 euro (1 action). --------------------------
    sasu_payload = dict(sasu_tests._payload())
    sasu_payload.update({"capital_social": "1", "nb_actions": 1})
    corpus["sasu_holding"] = _bundle(
        sasu_holding_slice.generate_dossier(sasu_payload, base_dir / "sasu_holding")
    )

    # --- Types civils : 1 part de 1 euro par associé. -------------------------
    def _civil_cap1(structure: str, stype: str, associes, extra: dict | None = None):
        payload = mtf._civil_base(structure, stype, associes)
        payload.update(
            {
                "capital_social": str(len(associes)),
                "nb_parts_total": len(associes),
                "valeur_nominale_part": "1",
            }
        )
        if extra:
            payload.update(extra)
        return _bundle(css.generate_dossier(_normalise_civil(payload), base_dir / stype))

    corpus["sci"] = _civil_cap1("SCI", "sci", [mtf._pp("Jean", "Durand", 1, 1, 1, "1")])
    corpus["sci_iris"] = _civil_cap1(
        "SCI IRIS",
        "sci_iris",
        [mtf._pm(1, 1, 1, "1"), mtf._pp("Alice", "Martin", 1, 2, 2, "1")],
    )
    corpus["scm"] = _civil_cap1(
        "SCM",
        "scm",
        [mtf._pp("Jean", "Durand", 1, 1, 1, "1"), mtf._pp("Alice", "Martin", 1, 2, 2, "1")],
    )
    corpus["scs"] = _civil_cap1(
        "SCS",
        "scs",
        [
            mtf._pp("Jean", "Durand", 1, 1, 1, "1", role="commandite"),
            mtf._pp("Alice", "Martin", 1, 2, 2, "1", role="commanditaire"),
        ],
    )

    missing = set(CORPUS_CAP1_KEYS) - set(corpus)
    if missing:
        raise RuntimeError(f"types du corpus cap1 non générés : {sorted(missing)}")
    return corpus
