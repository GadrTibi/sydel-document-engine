"""M4 (KAN-36) — Garde de conformité : le bloc signature reste solidaire (une seule page).

Convergence @All (2026-07-16) : la garde est le RECENSEMENT complet des documents signés,
pas un échantillon de cas reproduits. Elle génère un dossier RÉEL de CHAQUE structure
atteignable au front (via les slices `generate_dossier`, payloads = ceux des tests unitaires —
aucun payload réinventé) :

  * SELARL uni cession médicale, SELARL MULTI-associés (≥3 membres) ;
  * SELAS multi (physiques), SELAS uni médecin, SELAS uni dentiste ;
  * SPFPL cession, SPFPL apport (contrat d'apport) ;
  * SAS (SPFPL médecins forme SAS), SASU holding ;
  * SCM, SCS (bloc signature en grille).

Pour chaque document qui porte un INDICATEUR de signature — ancre « Fait à … » / « A/À …,
le … » / clôture de PROCÈS-VERBAL / mentions « Bon pour » / « Lu et approuvé » / clôtures de
LETTRE (« prie d'agréer », « veuillez agréer », « salutations distinguées »…) — la garde
vérifie la COHÉSION RÉELLE du bloc signature (pas juste « un keepNext quelconque ») :

  * chemin PARAGRAPHE : du début du bloc (dernière ancre, sinon les ~6 derniers paragraphes
    non vides) jusqu'à la fin du bloc (borné au 1er saut de page / à l'ANNEXE), TOUS les
    paragraphes sauf le dernier portent keepNext (généralise le cas B1 à TOUS les docs) ;
  * chemin TABLE : une grille de signature dont TOUTES les lignes portent cantSplit ET dont
    le paragraphe d'intro immédiatement précédent porte keepNext.

Un document signé qui ne satisfait NI l'un NI l'autre => échec listant le(s) document(s).
``KNOWN_SIGNATURE_GAPS`` reste VIDE : plus aucun générateur signé n'est toléré non protégé.

Exécutable en CI sans Word ni rendu PDF : on inspecte les propriétés de pagination
(keepNext / cantSplit) directement dans le DOCX (python-docx / OOXML).
"""

from __future__ import annotations

import dataclasses
import re
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# Indicateurs texte de « document signé » (paragraphes OU cellules de table). Couvre les
# ancres d'ouverture, les mentions de signature ET les clôtures de LETTRE (sinon un courrier
# signé — ex. courrier SDE cession SCM — passe inaperçu faute de « Fait à »).
SIGNATURE_MARKERS: tuple[str, ...] = (
    "Fait à",
    "Fait a",
    "signé après lecture",
    "dressé le présent procès-verbal",
    "Bon pour",
    "Lu et approuvé",
    # Clôtures de LETTRE (courriers signés sans ancre « Fait à »).
    "prie d'agréer",
    "prie d’agréer",
    "prions d'agréer",
    "prions d’agréer",
    "veuillez agréer",
    "salutations distinguées",
    "sentiments dévoués",
)

# Convergence @All 2026-07-16 : plus AUCUNE exclusion. Tous les générateurs signés révélés par
# la garde durcie ont été protégés (statuts SAS, courrier SDE cession SCM, SELARL multi-associés,
# contrat d'apport SPFPL, grille de signature SCS). Toute régression rouvrira la garde.
KNOWN_SIGNATURE_GAPS: frozenset[str] = frozenset()


def _full_text(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def _is_signed(path: Path) -> bool:
    # Un doc est signé si (a) un paragraphe est une ANCRE de signature (même intention que
    # ``_is_signature_anchor`` — inclut « A <lieu>, le <date> » des statuts SCI/SCI IRIS, que les
    # marqueurs-substrings ratent) OU (b) le texte porte un marqueur de signature/clôture.
    document = Document(str(path))
    if any(_is_signature_anchor(p.text.strip()) for p in document.paragraphs):
        return True
    return any(marker in _full_text(path) for marker in SIGNATURE_MARKERS)


# --------------------------------------------------------------------------- #
# Détecteur de cohésion KAN-36 (chemin paragraphe OU chemin table)
# --------------------------------------------------------------------------- #


def _is_signature_anchor(stripped: str) -> bool:
    """DÉBUT de bloc signature (même logique que ``keep_final_signature_block_together``) :
    toute clôture « Fait … », l'en-tête d'acte « A/À <lieu>, le <date> », la clôture de PV."""
    return (
        (stripped.startswith("Fait ") and "générateur" not in stripped)
        # « A <lieu>, le <date> » : une DATE (chiffre) suit « le » — sinon on false-matche de la
        # prose de corps (« A l'expiration du delai …, le conjoint … »). Même intention que le
        # helper ``keep_final_signature_block_together`` (règle 68 : l'intention, pas la tournure).
        or (
            stripped.startswith(("A ", "À "))
            and re.search(r",\s*le\s+\d", stripped) is not None
        )
        or "signé après lecture" in stripped
        or "dressé le présent procès-verbal" in stripped
        or "il a été dressé le présent" in stripped
    )


def _ordered_blocks(document) -> list[tuple[str, object]]:
    """Blocs du corps DANS L'ORDRE du document (paragraphes ET tables interleavés).

    ``document.paragraphs`` / ``document.tables`` perdent l'ordre relatif p<->tbl : on marche
    les enfants du corps pour situer une grille de signature vs son intro, et pour borner le
    bloc à l'ANNEXE (qui peut être un paragraphe OU une table titre encadrée)."""
    blocks: list[tuple[str, object]] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("P", Paragraph(child, document)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("T", Table(child, document)))
    return blocks


def _is_annexe(kind: str, obj: object) -> bool:
    """Frontière « ANNEXE » (fin du bloc signature) — paragraphe OU cadre-titre en table."""
    if kind == "P":
        return obj.text.strip().upper().startswith("ANNEXE")  # type: ignore[union-attr]
    try:
        return obj.rows[0].cells[0].text.strip().upper().startswith("ANNEXE")  # type: ignore[union-attr]
    except (IndexError, AttributeError):
        return False


def _has_page_break(paragraph) -> bool:
    if paragraph.paragraph_format.page_break_before:
        return True
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))  # noqa: SLF001


def _all_rows_cantsplit(table) -> bool:
    rows = table.rows
    return bool(rows) and all(
        row._tr.xpath("./w:trPr/w:cantSplit") for row in rows  # noqa: SLF001
    )


def _table_block_protected(blocks: list[tuple[str, object]]) -> bool:
    """Chemin TABLE : une grille dont TOUTES les lignes portent cantSplit ET dont le paragraphe
    d'intro immédiatement précédent porte keepNext (« Fait à … » / « A …, le … » solidaire)."""
    for index, (kind, obj) in enumerate(blocks):
        if kind != "T" or not _all_rows_cantsplit(obj):
            continue
        prev = index - 1
        while prev >= 0 and blocks[prev][0] != "P":
            prev -= 1
        if prev >= 0 and blocks[prev][1].paragraph_format.keep_with_next:  # type: ignore[union-attr]
            return True
    return False


def _signature_block_start(blocks: list[tuple[str, object]]) -> int | None:
    """Index (dans ``blocks``) du DÉBUT du bloc signature : dernière ancre, sinon les ~6
    derniers paragraphes non vides (fallback lettres sans ancre)."""
    para_positions = [i for i, (kind, _) in enumerate(blocks) if kind == "P"]
    # Une ancre n'est une SIGNATURE que dans la partie BASSE du doc (un « Fait à … » / « A …, le
    # <date> » en HAUT = en-tête de date d'une lettre, pas la signature). Même seuil que le helper.
    anchor_floor = len(blocks) * 0.4
    for i in reversed(para_positions):
        if i < anchor_floor:
            break
        if _is_signature_anchor(blocks[i][1].text.strip()):  # type: ignore[union-attr]
            return i
    non_empty = [i for i in para_positions if blocks[i][1].text.strip()]  # type: ignore[union-attr]
    if len(non_empty) < 2:
        return None
    return non_empty[max(0, len(non_empty) - 6)]


def _signature_tables_protected(
    blocks: list[tuple[str, object]], start: int, term: int, table_indices: list[int]
) -> bool:
    """Toutes les tables du bloc [start:term) sont des grilles de signataires PROTÉGÉES : cantSplit
    sur toutes leurs lignes ET intro (paragraphe précédent) keepNext, ET les paragraphes d'intro
    avant la 1re table forment une chaîne keepNext. Retourne False si une table n'est pas protégée
    (elle peut être une table de DONNÉES et non de signature — l'appelant retombe alors sur le
    chemin paragraphe)."""
    for i in table_indices:
        if not _all_rows_cantsplit(blocks[i][1]):
            return False
        prev = i - 1
        while prev >= start and blocks[prev][0] != "P":
            prev -= 1
        if prev < start or not blocks[prev][1].paragraph_format.keep_with_next:  # type: ignore[union-attr]
            return False
    first_table = table_indices[0]
    intro = [
        blocks[i][1]
        for i in range(start, first_table)
        if blocks[i][0] == "P" and blocks[i][1].text.strip()  # type: ignore[union-attr]
    ]
    return all(p.paragraph_format.keep_with_next for p in intro)


def _has_kan36_protection(path: Path) -> bool:  # noqa: C901
    """Protection KAN-36 RÉELLE ? Cohésion du bloc signature, pas « un keepNext quelconque ».

    On délimite le bloc signature [début, fin) sur les blocs ORDONNÉS du corps (paragraphes ET
    tables). Fin = 1er saut de page OU ANNEXE (paragraphe ou cadre-titre en table) après le début.
    Deux formes de bloc :

    * bloc contenant une TABLE de signataires : CHAQUE table du bloc porte cantSplit sur TOUTES
      ses lignes ET est solidarisée à son intro (paragraphe précédent keepNext) ; en outre les
      paragraphes d'intro avant la 1re table forment une chaîne keepNext. (Un keep_final seul —
      paragraphe — NE protège PAS une table : il exclut le dernier paragraphe, justement celui
      juste avant la table, et ne pose aucun cantSplit — vécu acte cession parts / PV AGE SCM.)
    * bloc purement paragraphes : TOUS les paragraphes sauf le dernier portent keepNext
      (généralise le cas B1 à tous les documents signés)."""
    document = Document(str(path))
    blocks = _ordered_blocks(document)
    start = _signature_block_start(blocks)
    if start is None:
        return _table_block_protected(blocks)

    term = len(blocks)
    for i in range(start + 1, len(blocks)):
        kind, obj = blocks[i]
        if (kind == "P" and _has_page_break(obj)) or _is_annexe(kind, obj):
            term = i
            break

    # Chemin TABLE : si le bloc contient une grille de signataires PROTÉGÉE (cantSplit sur toutes
    # ses lignes + intro keepNext), le doc est protégé. Si la table du bloc n'est PAS une table de
    # signature protégée (ex. table de DONNÉES « identification de la société » au milieu d'une
    # lettre d'option IS), on NE conclut PAS à l'absence de protection : on RETOMBE sur le chemin
    # paragraphe (la vraie signature est alors un paragraphe « Le gérant » en fin, chaîné keepNext).
    table_indices = [i for i in range(start, term) if blocks[i][0] == "T"]
    if table_indices and _signature_tables_protected(blocks, start, term, table_indices):
        return True

    # Bloc PUREMENT paragraphes : on reproduit EXACTEMENT la délimitation du helper en ESPACE
    # PARAGRAPHES (document.paragraphs), pas en espace-blocs — sinon les tables situées AILLEURS
    # dans le doc décalent le comptage du fallback « 6 derniers non vides » (le garde tombait un
    # paragraphe trop haut sur les lettres, ex. lettre d'option IS). Ancre (partie basse) sinon
    # fallback, bornée au 1er saut de page / ANNEXE. Même logique que keep_final_signature_block.
    paras = document.paragraphs
    p_start: int | None = None
    anchor_floor = len(paras) * 0.4
    for i in range(len(paras) - 1, -1, -1):
        if i < anchor_floor:
            break
        if _is_signature_anchor(paras[i].text.strip()):
            p_start = i
            break
    if p_start is None:
        non_empty = [i for i, p in enumerate(paras) if p.text.strip()]
        if len(non_empty) < 2:
            return False
        p_start = non_empty[max(0, len(non_empty) - 6)]
    p_term = len(paras)
    for i in range(p_start + 1, len(paras)):
        if _has_page_break(paras[i]) or paras[i].text.strip().upper().startswith("ANNEXE"):
            p_term = i
            break
    block_paras = [p for p in paras[p_start:p_term] if p.text.strip()]
    return len(block_paras) >= 2 and all(
        p.paragraph_format.keep_with_next for p in block_paras[:-1]
    )


# --------------------------------------------------------------------------- #
# Recensement : un dossier RÉEL de CHAQUE structure atteignable au front
# --------------------------------------------------------------------------- #


def _selarl_multi_input():
    """SELARL multi-associés à 3 membres (praticien 60 parts + 2 membres 25 + 15 = 100).

    Réutilise le payload valide des tests front (``_membre_bernard`` + ``_valid_selarl_input``,
    ``test_clean_front_app``) — parts/apports ajustés pour sommer au capital (règle du moteur)."""
    import test_clean_front_app as cfa

    from sydel_doc_engine.domain.models import StatutsCivilsApport, StatutsCivilsParts
    from sydel_doc_engine.front_app.selarl_slice import PROFESSION_MEDECIN

    membre_1 = cfa._membre_bernard().model_copy(
        update={
            "apport": StatutsCivilsApport(montant="250", montant_lettres="deux cent cinquante"),
            "parts": StatutsCivilsParts(nb=25, nb_lettres="vingt-cinq"),
        }
    )
    membre_2 = cfa._membre_bernard().model_copy(
        update={
            "prenom": "Marc",
            "nom": "Petit",
            "numero_ordre": "ORD-777",
            "numero_rpps": "20000000003",
            "adresse_personnelle_affichee": "3 rue Neuve, 69002 Lyon",
            "apport": StatutsCivilsApport(montant="150", montant_lettres="cent cinquante"),
            "parts": StatutsCivilsParts(nb=15, nb_lettres="quinze"),
        }
    )
    return cfa._valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        praticien_apport="600",
        membres_additionnels=(membre_1, membre_2),
    )


def _generate_bundles(base: Path) -> dict[str, list[Path]]:  # noqa: C901
    """Bundle RÉEL par structure (payloads = ceux des tests unitaires, cf. _conformance_corpus)."""
    import test_multi_type_front as mtf
    import test_sasu_holding_plan as sasu_tests
    import test_statuts_micro_holding as micro_tests
    from _conformance_corpus import _normalise_civil

    from sydel_doc_engine.front_app import (
        civil_statuts_slice as css,
    )
    from sydel_doc_engine.front_app import (
        sas_slice,
        sasu_holding_slice,
        selas_multi_slice,
        spfpl_slice,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_dentiste_slice as sd,
    )
    from sydel_doc_engine.front_app import (
        selas_uni_medecin_slice as sm,
    )
    from sydel_doc_engine.front_app.selarl_slice import generate_selarl_dossier
    from sydel_doc_engine.scenarios.selarl import build_selarl_scenario

    bundles: dict[str, list[Path]] = {}

    # SELAS pluripersonnelle physique : statuts SELAS multi + attestation souscripteurs +
    # PV nomination + procuration + DNC + demande inscription + autorisation.
    bundles["selas_multi"] = list(
        selas_multi_slice.generate_dossier(
            mtf._selas_payload_n(
                [mtf._selas_phys("Claire", "Durand", 60), mtf._selas_phys("Paul", "Martin", 40)]
            ),
            base / "selas_multi",
        ).docx_paths
    )

    # SELAS unipersonnelles (médecin / dentiste) : statuts SEL en mode mono + satellites signés.
    bundles["selas_uni_medecin"] = list(
        sm.generate_dossier(mtf._selas_uni_payload(), base / "selas_uni_medecin").docx_paths
    )
    bundles["selas_uni_dentiste"] = list(
        sd.generate_dossier(mtf._selas_uni_payload(), base / "selas_uni_dentiste").docx_paths
    )

    # SPFPL cession : note d'information + PV agrément + acte cession parts + attestation + statuts.
    spfpl_payload = mtf._spfpl_payload("SPFPL cession")
    spfpl_payload["nationalite"] = "française"
    spfpl_payload["cession_data"] = {
        **spfpl_payload["cession_data"],
        "cible_forme_complete": "société d'exercice libéral à responsabilité limitée",
    }
    bundles["spfpl_cession"] = list(
        spfpl_slice.generate_dossier(spfpl_payload, base / "spfpl_cession").docx_paths
    )

    # SPFPL apport : contrat d'apport SPFPL (DOC-041) + attestations + statuts (associé MARIÉ
    # communauté légale -> lettres renonciation/avertissement conjoint dans le bundle).
    spfpl_apport_payload = mtf._spfpl_payload("SPFPL apport")
    spfpl_apport_payload["nationalite"] = "française"
    spfpl_apport_payload["cession_data"] = {
        **spfpl_apport_payload["cession_data"],
        "cible_forme_complete": "société d'exercice libéral à responsabilité limitée",
    }
    spfpl_apport_payload.update(
        {
            "situation_maritale": "marié",
            "regime_matrimonial": "la communauté légale",
            "regime_communautaire": True,
        }
    )
    bundles["spfpl_apport"] = list(
        spfpl_slice.generate_dossier(spfpl_apport_payload, base / "spfpl_apport").docx_paths
    )

    # SAS (SPFPL médecins forme SAS) : statuts SAS (« Fait à » + président + « Bon pour
    # acceptation ») + attestation souscripteurs SAS + PV rémunération + procuration + DNC.
    sas_payload = dict(mtf._sas_payload())
    sas_payload["nationalite"] = "française"
    bundles["sas"] = list(sas_slice.generate_dossier(sas_payload, base / "sas").docx_paths)

    # SASU holding : statuts SASU + liste souscripteurs + PV rémunération + autorisation + DNC.
    bundles["sasu_holding"] = list(
        sasu_holding_slice.generate_dossier(sasu_tests._payload(), base / "sasu_holding").docx_paths
    )

    # SCM : statuts SCM + autorisation + demande inscription + pacte associés (table) + satellites.
    bundles["scm"] = list(
        css.generate_dossier(
            _normalise_civil(
                mtf._civil_base(
                    "SCM",
                    "scm",
                    [
                        mtf._pp("Jean", "Durand", 50, 1, 50, 500),
                        mtf._pp("Alice", "Martin", 50, 51, 100, 500),
                    ],
                )
            ),
            base / "scm",
        ).docx_paths
    )

    # SCS : liste des souscripteurs SCS + statuts SCS (bloc signature en GRILLE cantSplit).
    bundles["scs"] = list(
        css.generate_dossier(
            _normalise_civil(
                mtf._civil_base(
                    "SCS",
                    "scs",
                    [
                        mtf._pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
                        mtf._pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
                    ],
                )
            ),
            base / "scs",
        ).docx_paths
    )

    # SCI : statuts SCI (bloc signature « A <lieu>, le <date> » + signataires côte à côte) + option
    # IS + tronc commun (DNC / procuration / autorisation).
    sci_payload = mtf._civil_base(
        "SCI",
        "sci",
        [
            mtf._pp("Jean", "Durand", 40, 1, 40, 400),
            mtf._pp("Alice", "Martin", 60, 41, 100, 600),
        ],
    )
    sci_payload.update(mtf._OPTION_IS_INPUTS)
    bundles["sci"] = list(
        css.generate_dossier(_normalise_civil(sci_payload), base / "sci").docx_paths
    )

    # SCI IRIS : 1 PM + 1 PP (statuts SCI IRIS, même chemin signature « A <lieu>, le <date> »).
    bundles["sci_iris"] = list(
        css.generate_dossier(
            _normalise_civil(
                mtf._civil_base(
                    "SCI IRIS",
                    "sci_iris",
                    [mtf._pm(40, 1, 40, 400), mtf._pp("Alice", "Martin", 60, 41, 100, 600)],
                )
            ),
            base / "sci_iris",
        ).docx_paths
    )

    # MICRO HOLDING : statuts micro-holding (SPFPL morale + praticien physique, signataires côte à
    # côte) — branche de rendu distincte de SCI/SCM/SCS, jamais exercée avant la 4e passe Akainu.
    micro_payload = micro_tests._payload()
    micro_payload.update(
        {
            "signataire_nom_pere": "Pierre Durand",
            "signataire_nom_mere": "Anne Durand",
            "signataire_adresse_num": "1",
            "signataire_adresse_voie": "rue Exemple",
            "signataire_adresse_cp": "75000",
            "signataire_adresse_ville": "Paris",
            "signataire_fonction": "gerant",
        }
    )
    bundles["micro_holding"] = list(
        css.generate_dossier(micro_payload, base / "micro_holding").docx_paths
    )

    # SELARL cession médicale : acte + compromis cession cabinet + appel de fonds + avenant bail
    # (table) + courrier SDE cession SCM (lettre signée) + demande inscription + statuts SELARL.
    scenario = dataclasses.replace(
        build_selarl_scenario("selarl_medecin_cession_cabinet_medical"),
        siege_voie="avenue de Breteuil",
    )
    if scenario.scm_cession_context is not None:
        scenario.scm_cession_context.cedant.nationalite = "française"
        scenario.scm_cession_context.cedant.situation_maritale = "marié"
    bundles["selarl_cession"] = list(
        generate_selarl_dossier(scenario, base / "selarl_cession").docx_paths
    )

    # SELARL cession SCM (dentiste) : porte le courrier SDE cession SCM (lettre signée sans ancre).
    scenario_scm = dataclasses.replace(
        build_selarl_scenario("selarl_dentiste_cession_scm"), siege_voie="avenue de Breteuil"
    )
    if scenario_scm.scm_cession_context is not None:
        scenario_scm.scm_cession_context.cedant.nationalite = "française"
        scenario_scm.scm_cession_context.cedant.situation_maritale = "marié"
    bundles["selarl_cession_scm"] = list(
        generate_selarl_dossier(scenario_scm, base / "selarl_cession_scm").docx_paths
    )

    # SELARL MULTI-associés (≥3 membres) : statuts SELARL en mode multi (« Fait à » + liste des
    # signataires + mention « Lu et approuvé ») — c'est le cas où les NOMS des signataires
    # doivent porter keepNext (même bug que le PV B1).
    bundles["selarl_multi"] = list(
        generate_selarl_dossier(_selarl_multi_input(), base / "selarl_multi").docx_paths
    )

    return bundles


# Structures NEUVES apportées par la convergence @All (doivent exercer ≥1 doc signé chacune).
# 4e passe Akainu (2026-07-16) : SCI, SCI IRIS et MICRO_HOLDING (statuts civils à branche de rendu
# signature distincte) ajoutés — sans eux le garde n'était pas exhaustif.
_REQUIRED_STRUCTURES: tuple[str, ...] = (
    "sas",
    "spfpl_apport",
    "selas_uni_medecin",
    "selas_uni_dentiste",
    "selarl_multi",
    "sci",
    "sci_iris",
    "micro_holding",
)


@pytest.fixture(scope="module")
def sample_bundles(tmp_path_factory: pytest.TempPathFactory) -> dict[str, list[Path]]:
    return _generate_bundles(tmp_path_factory.mktemp("kan36_signature_cohesion"))


def test_signed_documents_apply_kan36_protection(
    sample_bundles: dict[str, list[Path]],
) -> None:
    """Tout document signé de TOUTES les structures porte la protection KAN-36 réelle."""
    all_paths = [path for paths in sample_bundles.values() for path in paths]
    checked = 0
    unprotected: list[str] = []
    for path in all_paths:
        if path.name in KNOWN_SIGNATURE_GAPS:
            continue
        if not _is_signed(path):
            continue
        checked += 1
        if not _has_kan36_protection(path):
            unprotected.append(path.name)

    # Garde-fou anti-garde-vide : l'échantillon exerce un GRAND volume de docs signés (≈80).
    assert checked >= 60, (
        f"échantillon de documents signés trop maigre ({checked}) — la garde n'exerce plus rien."
    )
    assert not unprotected, (
        "Documents SIGNÉS sans protection KAN-36 (bloc signature scindable sur deux pages) : "
        + ", ".join(sorted(set(unprotected)))
        + ". Câbler keep_final_signature_block_together (bloc paragraphe) "
        "ou keep_signature_block_together (grille en table)."
    )


def test_guard_exercises_every_structure(sample_bundles: dict[str, list[Path]]) -> None:
    """Chaque structure recensée exerce ≥1 document signé ; les structures NEUVES sont présentes."""
    for structure, paths in sample_bundles.items():
        signed = [p for p in paths if _is_signed(p)]
        assert signed, f"structure {structure!r} n'exerce aucun document signé (recensement cassé)."
    missing = [s for s in _REQUIRED_STRUCTURES if s not in sample_bundles]
    assert not missing, f"structures neuves absentes du recensement : {missing}"


def test_pv_agrement_plusieurs_associes_b1_closing_block_kept_together(
    tmp_path: Path,
) -> None:
    """Cas B1 reproduit : PV agrément SPFPL avec 9 associés présents (noms longs).

    De la clôture (« … dressé le présent procès-verbal … signé après lecture … »)
    jusqu'à la fin, TOUS les paragraphes portent keepNext (le tout dernier exclu par
    construction du helper) — le bloc clôture + signataires reste sur une seule page.
    """
    import test_lot_05_spfpl_agrement_info as t

    from sydel_doc_engine.domain.models import AssocieCible
    from sydel_doc_engine.generators.lot_05.pv_agrement_cession_spfpl_plusieurs_associes import (
        PvAgrementCessionSpfplPlusieursAssociesGenerator,
    )

    ctx = t._plural_context()
    ctx.associes_cible = [
        AssocieCible(
            civilite_affichage="Docteur",
            prenom=f"Jean-Baptiste{i}",
            nom=f"de la Rochefoucauld-Montmorency{i}",
            nb_parts_avant=10,
            nb_parts_apres=10,
            plage_parts=f"{i * 10 + 1} a {i * 10 + 10}",
            est_present_ou_represente=True,
        )
        for i in range(9)
    ]

    output = PvAgrementCessionSpfplPlusieursAssociesGenerator().generate(ctx, tmp_path)
    paras = Document(str(output)).paragraphs

    closing = next(
        (
            i
            for i, p in enumerate(paras)
            if "signé après lecture" in p.text or "dressé le présent procès-verbal" in p.text
        ),
        None,
    )
    assert closing is not None, "clôture du PV (signé après lecture) introuvable"

    tail = paras[closing:-1]  # le helper pose keepNext sur [start:-1] (dernier exclu)
    assert tail, "aucun paragraphe entre la clôture et la fin du document"
    not_kept = [
        closing + offset
        for offset, p in enumerate(tail)
        if not p.paragraph_format.keep_with_next
    ]
    assert not not_kept, (
        f"paragraphes {not_kept} (clôture -> signataires) sans keepNext : le bloc signature "
        "du PV à 9 associés est scindable sur deux pages (régression cas B1)."
    )

    # 9 associés présents => au moins 9 paragraphes solidarisés (clôture + lignes de noms).
    kept = sum(1 for p in paras if p.paragraph_format.keep_with_next)
    assert kept >= 9, f"attendu >= 9 paragraphes keepNext (9 associés présents), obtenu {kept}"
