"""Tests micro holding (DOC-047) — VRAI modele Albane 2026-06-29.

Micro holding = « societe civile de portefeuille A CAPITAL VARIABLE » (26 articles, modele
distinct du SCI fourni par Albane). L'objet social (art. 2) est VERBATIM dans le modele
(plus de variante A/B). Ces tests verifient la FIDELITE au modele Albane (corps legal
byte-exact via les donnees de reference Berte/Gosset), le capital variable (separateur de
milliers a point, max = 10x le minimum, accord des lettres), l'absence de token residuel et
le bundle de creation.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from docx import Document

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    CentreImpots,
    Company,
    DocumentGenerationContext,
    DossierOptions,
    Person,
    Signature,
    StatutsCivilsApport,
    StatutsCivilsAssocie,
    StatutsCivilsCapitalDepot,
    StatutsCivilsContext,
    StatutsCivilsParts,
    StatutsCivilsRepresentant,
)
from sydel_doc_engine.front_app import civil_statuts_slice as cs
from sydel_doc_engine.front_app.type_registry import registered_type_by_key
from sydel_doc_engine.generators.lot_01.autorisation_domiciliation import (
    AutorisationDomiciliationGenerator,
)
from sydel_doc_engine.generators.lot_04.statuts_micro_holding import (
    StatutsMicroHoldingGenerator,
)
from sydel_doc_engine.generators.lot_05.lettre_option_is import LettreOptionIsGenerator
from sydel_doc_engine.registry.catalog import ALL_STRUCTURES


def _docx_text(path: Path) -> str:
    document = Document(path)
    texts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs if p.text)
    return "\n".join(texts)


def _ctx_berte() -> DocumentGenerationContext:
    """Contexte de reference = l'exemple Berte/Gosset du modele Albane (SPFPL + praticienne)."""
    spfpl = StatutsCivilsAssocie(
        type_personne="personne_morale",
        denomination="SPFPL DU DR JESSICA GOSSET",
        forme_juridique=(
            "Société de Participations Financières de Profession Libérale par actions "
            "simplifiée de Médecins"
        ),
        capital_social="1.020",
        siege=Address(adresse_affichee="20 rue Isabey, 54000 NANCY"),
        ville_rcs="NANCY",
        numero_rcs="981 102 072",
        representant=StatutsCivilsRepresentant(
            civilite_affichage="Madame", prenom="Jessica", nom="GOSSET", fonction="Président"
        ),
        apport=StatutsCivilsApport(montant="1010", montant_lettres="mille dix"),
        parts=StatutsCivilsParts(nb=1010, nb_lettres="mille dix"),
        est_signataire=True,
    )
    gosset = StatutsCivilsAssocie(
        type_personne="personne_physique",
        genre=Gender.FEMININ,
        civilite_affichage="Madame",
        prenoms="Jessica",
        nom="GOSSET",
        date_naissance="1er décembre 1978",
        ville_naissance="CHAMBRAY LES TOURS",
        departement_naissance="37",
        nationalite="française",
        situation_maritale=(
            "Mariée sous le régime de la séparation de biens avec Monsieur Nicolas BERTE"
        ),
        adresse_personnelle_affichee="19 rue Sainte Catherine, 54000 NANCY",
        apport=StatutsCivilsApport(montant="10", montant_lettres="dix"),
        parts=StatutsCivilsParts(nb=10, nb_lettres="dix"),
        est_signataire=True,
    )
    return DocumentGenerationContext(
        structure="MICRO_HOLDING",
        personne_signataire=Person(
            genre=Gender.FEMININ, civilite="Madame", prenom="Jessica", nom="GOSSET"
        ),
        signature=Signature(lieu="Nancy", date=date(2026, 5, 22)),
        societe=Company(
            denomination="Micro holding famille Berte",
            forme_sociale="Société Civile",
            siege=Address(
                num_voie="19",
                voie="rue Sainte Catherine",
                cp="54000",
                ville="NANCY",
                adresse_affichee="19 rue Sainte Catherine, 54000 NANCY",
            ),
            ville_rcs="NANCY",
        ),
        statuts_civils=StatutsCivilsContext(
            type="micro_holding",
            forme_sociale="Société Civile",
            mention_capital_variable="à capital variable",
            capital_social="1.020",
            capital_social_lettres="mille vingt",
            capital_maximal="10.200",
            capital_maximal_lettres="dix mille deux cents",
            capital_autorise="10.200",
            capital_autorise_lettres="dix mille deux cents",
            nb_parts_total=1020,
            nb_parts_total_lettres="mille vingt",
            valeur_nominale_part="1",
            valeur_nominale_part_lettres="un",
            date_cloture_premier_exercice="31 décembre 2026",
            capital_depot=StatutsCivilsCapitalDepot(
                banque_nom="HSBC", banque_adresse="48 Rue Saint-Jean, 54000 NANCY"
            ),
            associes=[spfpl, gosset],
        ),
    )


def _payload(capital_min: int = 1020) -> dict:
    def _phys(
        prenom: str, nom: str, nb: int, debut: int, fin: int, montant: int
    ) -> StatutsCivilsAssocie:
        return StatutsCivilsAssocie(
            type_personne="personne_physique",
            genre=Gender.MASCULIN,
            civilite_affichage="Monsieur",
            prenom=prenom,
            prenoms=prenom,
            nom=nom,
            date_naissance="1er janvier 1980",
            ville_naissance="Paris",
            departement_naissance="75",
            nationalite="française",
            situation_maritale="célibataire",
            adresse_personnelle_affichee="1 rue Exemple, 75000 Paris",
            # DNC par associe (Rafael 2026-07-09) : filiation requise pour CHAQUE associe.
            nom_pere=f"Pierre {nom}",
            nom_mere=f"Anne {nom}",
            apport=StatutsCivilsApport(montant=str(montant), montant_lettres=str(montant)),
            parts=StatutsCivilsParts(nb=nb, nb_lettres=str(nb), debut=debut, fin=fin),
        )

    return {
        "structure": "MICRO_HOLDING",
        "statuts_type": "micro_holding",
        "denomination": "MA MICRO HOLDING",
        "capital_social": str(capital_min),
        "nb_parts_total": 100,
        "valeur_nominale_part": "",
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "ville_rcs": "Paris",
        "banque_nom": "BANQUE",
        "banque_adresse": "1 rue Banque, 75009 Paris",
        "date_cloture_premier_exercice": "31 decembre 2026",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 15),
        "decision_date": date(2026, 5, 10),
        "associes": [
            _phys("Jean", "Durand", 40, 1, 40, capital_min * 40 // 100),
            _phys("Alice", "Martin", 60, 41, 100, capital_min * 60 // 100),
        ],
        "gerant_index": 0,
    }


# --- enregistrement du type --------------------------------------------------


def test_micro_holding_structure_registered() -> None:
    assert "MICRO_HOLDING" in ALL_STRUCTURES
    assert "MICRO_HOLDING" in cs.CIVIL_TYPE_BY_STRUCTURE
    statuts_type, doc_code = cs.CIVIL_TYPE_BY_STRUCTURE["MICRO_HOLDING"]
    assert statuts_type == "micro_holding"
    assert doc_code == "DOC-047"


def test_micro_holding_registered_in_type_registry() -> None:
    rt = registered_type_by_key("micro_holding_v1")
    assert rt.structure == "MICRO_HOLDING"
    assert rt.slice_module == "sydel_doc_engine.front_app.micro_holding_slice"
    assert rt.generation_enabled is True


# --- fidelite au modele Albane (corps legal byte-exact) ----------------------

# Lignes DISTINCTIVES du modele Albane « societe civile de portefeuille » (verbatim attendu
# dans la sortie generee avec les donnees de reference Berte/Gosset). Couvrent l'objet social
# (art. 2), la forme civile (art. 1), la gerance avec gerant suppleant (art. 17), la
# variabilite du capital (art. 8) et la signature electronique YouSign (art. 26) — autant de
# marqueurs ABSENTS de l'ancien clone SCI.
_VERBATIM_ALBANE = [
    "Ont établi ainsi qu’il suit les statuts d’une Société civile de portefeuille",
    "D’acquérir et de gérer notamment des biens immobiliers, mobiliers, créances et placements",
    "La prise de participation par achat, souscription, apport, fusion de tous biens mobiliers",
    "ARTICLE 8 – VARIABILITE DU CAPITAL",
    "Sans aucun vote, le gérant suppléant sera par défaut, le conjoint ou partenaire de PACS",
    "ARTICLE 26 - CONVENTION SUR LA PREUVE – SIGNATURE ELECTRONIQUE",
    "un procédé de signature électronique avancée mis en œuvre par un prestataire de services",
]


def test_micro_holding_fidelity_legal_body(tmp_path: Path) -> None:
    out = StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path)
    text = _docx_text(out)
    for ligne in _VERBATIM_ALBANE:
        assert ligne in text, f"ligne Albane absente du rendu : {ligne!r}"
    # L'ancien objet « clone SCI » NE DOIT PLUS apparaitre (preuve du recablage).
    assert "biens et droits immobiliers" not in text
    # Aucun placeholder source residuel.
    assert "[" not in text and "]" not in text


def test_micro_holding_capital_variable_lines(tmp_path: Path) -> None:
    out = StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path)
    text = _docx_text(out)
    # Capital variable : minimum / maximum / effectif (lettres en MAJUSCULES, montant a point).
    assert "Le capital social est variable." in text
    assert "Le capital social minimal est fixé à MILLE VINGT (1.020€)." in text
    assert "Le capital social maximal est fixé à DIX MILLE DEUX CENTS EUROS (10.200€)." in text
    assert "Le capital social est divisé en 1020 parts de 1€ (UN EURO) chacune." in text
    # Comparution morale (SPFPL) + repartition par associe.
    assert "- SPFPL DU DR JESSICA GOSSET" in text
    assert "La SPFPL DU DR JESSICA GOSSET apporte la somme de mille dix euros" in text


def test_micro_holding_comparution_nom_usage_epouse(tmp_path: Path) -> None:
    # Q4/MH-épouse (Albane 2026-07-01, « comme les modeles ») : quand un NOM DE NAISSANCE distinct
    # est saisi, la comparution physique porte le nom d'usage « <maiden> épouse <nom> » (modele
    # micro holding « Madame Jessica GOSSET épouse BERTE, » — SANS virgule avant « épouse »).
    # Convention : nom = nom d'usage (apport/signature) ; nom_naissance = maiden (comparution).
    ctx = _ctx_berte()
    physique = next(
        a for a in ctx.statuts_civils.associes if a.type_personne == "personne_physique"
    )
    physique.nom = "BERTE"  # nom d'usage marital
    physique.nom_naissance = "GOSSET"  # nom de naissance (maiden)
    out = StatutsMicroHoldingGenerator().generate(ctx, tmp_path)
    text = _docx_text(out)
    assert "Madame Jessica GOSSET épouse BERTE," in text


def test_micro_holding_comparution_sans_nom_naissance_inchangee(tmp_path: Path) -> None:
    # Non-regression Q4 : SANS nom_naissance distinct, la comparution reste « <civ> <prenoms>
    # <nom>, » (byte-identique, aucune mention « épouse »).
    ctx = _ctx_berte()  # gosset : nom="GOSSET", nom_naissance=None
    out = StatutsMicroHoldingGenerator().generate(ctx, tmp_path)
    text = _docx_text(out)
    assert "Madame Jessica GOSSET," in text
    assert "épouse" not in text


def test_micro_holding_signature_date_longue_et_civilite_abregee(tmp_path: Path) -> None:
    # MH signature (Albane 2026-07-01, « comme les modeles ») : zone signature du modele
    # Statuts_Micro_holding.docx = « Fait à Nancy » / « Le 22 mai 2026 » (date LONGUE) puis
    # « Mme Jessica GOSSET \t...\t SPFPL DU DR JESSICA GOSSET » (civilite ABREGEE « Mme »).
    out = StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path)
    text = _docx_text(out)
    assert "Fait à Nancy" in text
    assert "Le 22 mai 2026" in text  # forme longue (date signature = 22/05/2026)
    assert "Le 22/05/2026" not in text  # plus la date courte dans la zone signature
    assert "Mme Jessica GOSSET" in text  # civilite abregee
    # La zone signature ne porte plus la civilite pleine (« Madame Jessica GOSSET » suivi d'une
    # tabulation = ligne signataire) ; la comparution garde « Madame » ailleurs, non testee ici.
    assert "Madame Jessica GOSSET\t" not in text


# --- capital variable, max = 10x min (separateur a point), accords ------------


def test_capital_dot_format_and_max_ten_times() -> None:
    # Le slice formate le capital avec un separateur de milliers A POINT et calcule le
    # maximum = 10x le minimum saisi (« 1.020 » -> « 10.200 »).
    sc = cs.build_generation_context(_payload(capital_min=1020)).statuts_civils
    assert sc.capital_social == "1.020"
    assert sc.capital_maximal == "10.200"
    assert sc.capital_autorise == "10.200"
    assert sc.capital_social_lettres == "mille vingt"
    assert sc.capital_maximal_lettres == "dix mille deux cents"
    assert sc.mention_capital_variable == "a capital variable"


@pytest.mark.parametrize(
    ("capital_min", "max_digits", "max_lettres"),
    [
        (1500, "15.000", "quinze mille"),
        (8000, "80.000", "quatre-vingt mille"),  # « vingt » invariable devant « mille »
        (30000, "300.000", "trois cent mille"),  # « cent » invariable devant « mille »
        (28000, "280.000", "deux cent quatre-vingt mille"),
        (100_000_000, "1.000.000.000", "un milliard"),  # franchit le milliard
    ],
)
def test_capital_autorise_lettres_accords(
    capital_min: int, max_digits: str, max_lettres: str
) -> None:
    sc = cs.build_generation_context(_payload(capital_min=capital_min)).statuts_civils
    assert sc.capital_autorise == max_digits
    assert sc.capital_autorise_lettres == max_lettres
    assert "mille millions" not in sc.capital_autorise_lettres


def test_capital_max_ligne_dans_docx_reel(tmp_path: Path) -> None:
    out = StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path)
    text = _docx_text(out)
    assert "Le capital social est variable" in text
    assert "DIX MILLE DEUX CENTS EUROS (10.200€)" in text


def test_micro_holding_inherits_model_form_and_overwrites_client_footer(tmp_path: Path) -> None:
    # R1 (Albane 2026-06-30) : la micro holding herite la GEOMETRIE / mise en page du modele
    # Albane (A4, taille 12 pt) au lieu du profil SYDEL Letter US.
    # MAJ police (Albane 2026-07-01 : « police Times au lieu de Roboto, notre charte ») : on garde
    # la geometrie ET la taille du modele, mais on IMPOSE la police de la CHARTE (Roboto). Le
    # modele est en Times New Roman ; on ne le reproduit plus (c'etait l'ecart signale). Le modele
    # porte aussi un footer CLIENT (« ...Berte ») qui DOIT etre ecrase par la denomination.
    from docx import Document
    from docx.shared import Cm

    out = StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path)
    document = Document(out)
    section = document.sections[0]
    normal = document.styles["Normal"]

    assert normal.font.name == "Roboto"  # charte SYDEL (imposee), plus le Times du modele
    # Albane 2026-07-02 (« tout est en police 12 au lieu de 10 ») : la TAILLE de la charte (10 pt)
    # est desormais imposee sur les defaults (le « Normal » du modele etait a 12 pt -> les runs
    # re-emis heritaient 12 pt). Plus la taille du modele : la taille de la charte SYDEL.
    assert normal.font.size is not None and normal.font.size.pt == 10.0
    assert abs(section.page_height - Cm(29.7)) < Cm(0.1)  # A4 du modele conserve
    assert abs(section.page_width - Cm(21.59)) > Cm(0.1)  # pas Letter US SYDEL
    footer_text = " | ".join(p.text for p in section.footer.paragraphs if p.text)
    assert "Micro holding famille Berte - Statuts constitutifs" in footer_text
    # M2 (Akainu 2026-06-30) : pas de paragraphe vide d'amorce en tete -> body[0] = le titre.
    assert document.paragraphs[0].text.strip() != ""


def test_micro_holding_police_10pt_rendue_et_espacements_albane(tmp_path: Path) -> None:
    # Retour Albane 2026-07-02 (verrous Akainu m2 + M1) :
    #  (m2) TOUT le corps RENDU est a 10 pt — pas seulement le style « Normal ». Un run a taille
    #       explicite != 10 (regression future) doit casser (verbatim « TOUT est en 12 »).
    #  (M1) Les 3 espacements demandes sont presents (paragraphe vide, space_after = 6 pt) —
    #       sinon un add_spacer supprime/deplace regresse SILENCIEUSEMENT sur un retour explicite.
    from docx.shared import Pt

    document = Document(StatutsMicroHoldingGenerator().generate(_ctx_berte(), tmp_path))

    def _all_runs(doc):
        for paragraph in doc.paragraphs:
            yield from paragraph.runs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield from paragraph.runs

    # (m2) aucun run (corps + tables) ne rend a une taille explicite != 10 pt (herite => Normal=10).
    for run in _all_runs(document):
        size = run.font.size
        assert size is None or size.pt == 10.0, f"run a {size.pt} pt (attendu 10) : {run.text!r}"

    # (M1) les 3 spacers Albane (paragraphe vide, 6 pt) sont presents ET l'un precede « ARTICLE 1 ».
    paras = document.paragraphs
    spacers = [
        i
        for i, p in enumerate(paras)
        if not p.text.strip() and p.paragraph_format.space_after == Pt(6)
    ]
    assert len(spacers) >= 3, f"attendu >= 3 spacers Albane (vide, 6 pt), trouve {len(spacers)}"
    art1 = next(i for i, p in enumerate(paras) if p.text.strip().upper().startswith("ARTICLE 1 "))
    assert art1 - 1 in spacers, "pas de spacer 6 pt juste avant « ARTICLE 1 »"
    # espace apres la comparution : la derniere ligne « Demeurant … » est suivie d'un spacer 6 pt.
    demeurant = [i for i, p in enumerate(paras) if p.text.strip().startswith("Demeurant")]
    assert demeurant, "ligne « Demeurant » (fin de comparution associe) introuvable"
    assert demeurant[-1] + 1 in spacers, "pas de spacer 6 pt apres l'adresse du dernier associe"


def test_micro_holding_footer_no_client_residue_for_other_dossier(tmp_path: Path) -> None:
    # M1 (Akainu 2026-06-30) : le footer du modele micro a 3 paragraphes dont [1] = « Statuts
    # Societe Micro holding famille Berte » (NOM CLIENT). Ecraser uniquement [0] laissait
    # « Berte » FUITER dans tout dossier. Avec une denomination NON-Berte, le nom client du
    # modele ne doit apparaitre NULLE PART, et le footer doit porter la SEULE denomination.
    ctx = _ctx_berte()
    ctx.societe.denomination = "HOLDING TEST DUPONT ZZZ"
    document = Document(StatutsMicroHoldingGenerator().generate(ctx, tmp_path))

    fragments = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                fragments.extend(p.text for p in cell.paragraphs)
    for sect in document.sections:
        fragments.extend(p.text for p in sect.footer.paragraphs)
        fragments.extend(p.text for p in sect.header.paragraphs)
    full_doc = "\n".join(fragments)

    assert "Berte" not in full_doc, "residu du nom client modele a fuite dans le dossier"
    footer_text = " | ".join(p.text for p in document.sections[0].footer.paragraphs if p.text)
    assert footer_text == "HOLDING TEST DUPONT ZZZ - Statuts constitutifs"


# --- bundle de creation = statuts + tronc commun civil + PV + option IS -------


def test_bundle_is_statuts_plus_civil_tronc_commun() -> None:
    plan = cs.build_civil_plan(_payload())
    codes = set(plan.document_codes)
    # Statuts micro holding + tronc commun civil (DNC, domiciliation, procuration) + PV gerant.
    assert "DOC-047" in codes  # statuts micro holding
    assert "DOC-001" in codes  # declaration non-condamnation
    assert "DOC-002" in codes  # autorisation de domiciliation
    assert "DOC-003" in codes  # procuration
    assert "DOC-004" in codes  # PV nomination gerant
    # Pas de satellites SCM sur la micro holding.
    assert "DOC-025" not in codes
    assert "DOC-026" not in codes
    # Sans option IS cochee, la lettre d'option IS (DOC-022) n'est PAS dans le bundle.
    assert "DOC-022" not in codes


def _payload_option_is() -> dict:
    payload = _payload()
    payload.update(
        {
            "option_is": True,
            "siren": "En cours d’immatriculation",
            "impots_service": "SIE MEURTHE-ET-MOSELLE",
            "impots_adresse_ligne_1": "CITE ADMINISTRATIVE",
            "impots_adresse_ligne_2": "45 RUE SAINTE CATHERINE",
            "impots_cp": "54000",
            "impots_ville": "NANCY",
        }
    )
    return payload


def test_bundle_includes_lettre_option_is_when_option_is_on() -> None:
    # Mail Albane 2026-06-29 : la lettre d'option IS (DOC-022) entre dans le bundle micro
    # holding quand l'option IS est cochee.
    codes = set(cs.build_civil_plan(_payload_option_is()).document_codes)
    assert "DOC-022" in codes
    assert "DOC-047" in codes


# --- lettre option IS micro holding : fidelite au modele Albane ---------------


def _ctx_lettre_is() -> DocumentGenerationContext:
    spfpl = StatutsCivilsAssocie(
        type_personne="personne_morale",
        denomination="SPFPL DU DR JESSICA GOSSET",
        siege=Address(adresse_affichee="20 rue Isabey, 54000 NANCY"),
        apport=StatutsCivilsApport(montant="1010", montant_lettres="mille dix"),
        parts=StatutsCivilsParts(nb=1010, nb_lettres="mille dix"),
    )
    gosset = StatutsCivilsAssocie(
        type_personne="personne_physique",
        genre=Gender.FEMININ,
        civilite_affichage="Madame",
        prenom="Jessica",
        prenoms="Jessica",
        nom="GOSSET",
        adresse_personnelle_affichee="19 rue Sainte Catherine, 54000 NANCY",
        apport=StatutsCivilsApport(montant="10", montant_lettres="dix"),
        parts=StatutsCivilsParts(nb=10, nb_lettres="dix", qualite_associe="gérante"),
    )
    return DocumentGenerationContext(
        structure="MICRO_HOLDING",
        dossier_options=DossierOptions(option_is=True),
        personne_signataire=Person(
            genre=Gender.FEMININ, civilite="Madame", prenom="Jessica", nom="GOSSET"
        ),
        signature=Signature(lieu="Nancy", date=date(2023, 12, 6)),
        societe=Company(
            denomination="Micro holding famille Berte",
            siege=Address(adresse_affichee="19 rue Sainte Catherine, 54000 NANCY"),
            siren="En cours d’immatriculation",
            capital_social="1020",
        ),
        impots=CentreImpots(
            service="SIE MEURTHE-ET-MOSELLE",
            adresse_ligne_1="CITE ADMINISTRATIVE",
            adresse_ligne_2="45 RUE SAINTE CATHERINE",
            cp="54000",
            ville="NANCY",
        ),
        statuts_civils=StatutsCivilsContext(
            type="micro_holding",
            capital_social="1020",
            nb_parts_total=1020,
            associes=[gosset, spfpl],
        ),
    )


def test_lettre_option_is_micro_holding_fidelity(tmp_path: Path) -> None:
    out = LettreOptionIsGenerator().generate(_ctx_lettre_is(), tmp_path)
    text = _docx_text(out)
    assert "Centre des Finances Publiques" in text
    assert "Objet : Demande d'option pour le régime de l'impôt sur les sociétés" in text
    assert "la société civile dont vous trouverez la description ci-après opte" in text
    # Tableau d'identification : associe morale + physique au format Albane.
    assert "Micro holding famille Berte" in text
    assert "En cours d’immatriculation" in text
    assert (
        "Madame Jessica GOSSET, demeurant 19 rue Sainte Catherine, 54000 NANCY, gérante, "
        "détenant 10 parts." in text
    )
    assert (
        "La société SPFPL DU DR JESSICA GOSSET, ayant son siège social au 20 rue Isabey, "
        "54000 NANCY, détenant 1010 parts." in text
    )
    assert "Le gérant" in text


def test_option_is_gerant_recoit_qualite_gerant() -> None:
    # Akainu m1 (2026-06-29) : dans le bundle REEL (slice), l'associe physique GERANT de la
    # lettre option IS doit recevoir « gerant/gerante » (modele Albane « ..., gerante, ... »),
    # pas « associe(e) ». Verifie le comportement du slice de production (pas un fixture force).
    ctx = cs.build_generation_context(_payload_option_is())
    associes = ctx.statuts_civils.associes
    # gerant_index = 0 (1er associe physique) -> « gerant » (masculin) ; l'autre -> « associee/e ».
    assert associes[0].parts.qualite_associe == "gérant"
    assert associes[1].parts.qualite_associe in {"associé", "associée"}


def test_domiciliation_micro_holding_mention_capital_variable(tmp_path: Path) -> None:
    # Akainu M1 (2026-06-29) : l'autorisation de domiciliation du bundle micro holding doit
    # porter la mention capital variable du modele Albane (« a capital variable au capital
    # minimum de X € et au capital effectif de X € »), pas le generique « au capital de X euros ».
    ctx = DocumentGenerationContext(
        structure="MICRO_HOLDING",
        personne_signataire=Person(
            genre=Gender.MASCULIN, civilite="Monsieur", prenom="Jérémie", nom="BERDAH"
        ),
        signature=Signature(lieu="Paris", date=date(2025, 10, 29)),
        societe=Company(
            denomination="Micro holding famille Berte",
            capital="1.020",
            siege=Address(num_voie="54", voie="rue Ordener", cp="75018", ville="PARIS"),
        ),
    )
    out = AutorisationDomiciliationGenerator().generate(ctx, tmp_path)
    text = _docx_text(out)
    assert (
        "à capital variable au capital minimum de 1.020 € et au capital effectif de 1.020 €"
        in text
    )
    assert "au capital de 1.020 euros" not in text
