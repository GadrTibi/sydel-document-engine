from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from streamlit.testing.v1 import AppTest

from sydel_doc_engine.app.front_shell import PROTOTYPE_TOOL_LABELS
from sydel_doc_engine.app.single_document_mode import (
    UNIT_STATUS_MANUAL_ONLY,
    UNIT_STATUS_NOT_SUPPORTED,
    UNIT_STATUS_SUPPORTED,
    SingleDocumentAssociateInput,
    SingleDocumentInput,
    build_single_document_context,
    field_specs_for_document,
    sample_single_document_input,
    single_document_choices,
    validate_single_document_input,
)
from sydel_doc_engine.app.ui_runtime import (
    generate_docx_files_for_document_codes,
    generate_zip_file,
)


def test_single_document_mode_is_visible_next_to_existing_streamlit_modes() -> None:
    app_source = Path("src/sydel_doc_engine/app/streamlit_app.py").read_text(encoding="utf-8")

    assert "PROTOTYPE_TOOL_LABELS" in app_source
    assert "Assistant metier prototype" in PROTOTYPE_TOOL_LABELS
    assert "Document unitaire" in PROTOTYPE_TOOL_LABELS
    assert "Technique / diagnostic" in PROTOTYPE_TOOL_LABELS


def test_streamlit_single_document_mode_renders_document_selector() -> None:
    app = AppTest.from_file("src/sydel_doc_engine/app/streamlit_app.py").run(timeout=120)

    app.radio[0].set_value("Prototype / outils de test")
    app.run(timeout=120)
    app.radio[1].set_value("Document unitaire")
    app.run(timeout=120)

    assert app.selectbox(key="single_document_choice").label == "Document a tester"


def test_single_document_choices_mark_supported_and_unsupported_selarl_docs() -> None:
    choices = single_document_choices(
        "SELARL",
        {
            "profession": "medecin",
            "site_distinct": False,
            "scm_cession": False,
            "regime_communautaire": False,
            "derogation": False,
            "cession": False,
        },
    )

    by_code = {choice.document_code: choice for choice in choices}

    assert by_code["DOC-001"].status == UNIT_STATUS_SUPPORTED
    assert by_code["DOC-002"].status == UNIT_STATUS_SUPPORTED
    assert by_code["DOC-003"].status == UNIT_STATUS_SUPPORTED
    assert by_code["DOC-004"].status == UNIT_STATUS_SUPPORTED
    assert by_code["DOC-017"].status == UNIT_STATUS_NOT_SUPPORTED


def test_single_document_choices_keep_manual_documents_out_of_generation() -> None:
    choices = single_document_choices(
        "SELARL",
        {
            "profession": "chirurgien_dentiste",
            "site_distinct": True,
            "scm_cession": False,
            "regime_communautaire": False,
            "derogation": False,
            "cession": False,
        },
    )

    manual = next(choice for choice in choices if choice.document_code is None)

    assert manual.document_key == "site_distinct_cd94_sel"
    assert manual.status == UNIT_STATUS_MANUAL_ONLY


def test_single_document_field_specs_are_scoped_to_selected_document() -> None:
    doc_001_fields = {field.key for field in field_specs_for_document("DOC-001")}
    doc_002_fields = {field.key for field in field_specs_for_document("DOC-002")}

    assert "personne_date_naissance" in doc_001_fields
    assert "societe_denomination" not in doc_001_fields
    assert "domiciliation_adresse_affichee" in doc_002_fields
    assert "personne_adresse_num_voie" not in doc_002_fields


def test_single_document_validation_detects_missing_doc_004_associate_parts() -> None:
    data = SingleDocumentInput(
        **{
            **sample_single_document_input("DOC-004").__dict__,
            "associes": (
                SingleDocumentAssociateInput(
                    civilite_affichage="Monsieur",
                    prenom="Jean",
                    nom="Durand",
                    nb_parts=None,
                ),
            ),
        }
    )

    missing = validate_single_document_input(data)

    assert "associes[1].nb_parts" in missing


def test_single_document_generates_docx_and_zip_for_doc_001(tmp_path: Path) -> None:
    data = sample_single_document_input("DOC-001")
    ctx = build_single_document_context(data)

    docx_paths = generate_docx_files_for_document_codes(ctx, tmp_path, ("DOC-001",))
    zip_path = generate_zip_file(tmp_path, docx_paths)

    assert [path.name for path in docx_paths] == ["declaration_non_condamnation.docx"]
    assert "[" not in _docx_text(docx_paths[0])
    assert "]" not in _docx_text(docx_paths[0])
    with ZipFile(zip_path) as archive:
        assert {"declaration_non_condamnation.docx", "manifest.json"}.issubset(
            set(archive.namelist())
        )


def test_single_document_generates_docx_for_doc_004(tmp_path: Path) -> None:
    data = sample_single_document_input("DOC-004")
    ctx = build_single_document_context(data)

    docx_paths = generate_docx_files_for_document_codes(ctx, tmp_path, ("DOC-004",))

    assert [path.name for path in docx_paths] == ["pv_nomination_gerant.docx"]


def _docx_text(path: Path) -> str:
    document = Document(path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    table_text = [
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    ]
    return "\n".join([*paragraph_text, *table_text])
