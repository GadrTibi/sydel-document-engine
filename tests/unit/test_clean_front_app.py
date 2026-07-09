from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from unicodedata import normalize

from docx import Document
from streamlit.testing.v1 import AppTest

from sydel_doc_engine.domain.models import (
    StatutsCivilsApport,
    StatutsCivilsAssocie,
    StatutsCivilsParts,
)
from sydel_doc_engine.front_app.data_entry import build_clean_data_entry
from sydel_doc_engine.front_app.dossier_selection import dossier_type_by_label
from sydel_doc_engine.front_app.generation import build_clean_generation_plan
from sydel_doc_engine.front_app.legacy_boundary import legacy_boundary_items
from sydel_doc_engine.front_app.routing import clean_front_routes
from sydel_doc_engine.front_app.selarl_slice import (
    PROFESSION_DENTISTE,
    PROFESSION_MEDECIN,
    build_generation_context,
    generate_selarl_dossier,
    selected_selarl_document_codes,
)
from sydel_doc_engine.orchestrator.service import DocumentOrchestrator
from sydel_doc_engine.registry.catalog import build_seed_catalog
from sydel_doc_engine.scenarios.selarl import (
    build_selarl_scenario,
    cession_fixture_for_profession,
    scm_cession_fixture,
)


def test_clean_front_routes_are_minimal() -> None:
    assert [route.label for route in clean_front_routes()] == [
        "Type de dossier",
        "Donnees a saisir",
        "Generation",
    ]


def test_clean_front_selarl_slice_is_generable_for_medecin() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = _valid_selarl_input(PROFESSION_MEDECIN)

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is True
    assert plan.status == "ready"
    assert plan.target_engine_adapter == "front_app.selarl_slice"
    assert plan.document_codes == (
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-017",
    )
    assert "DOC-006" not in plan.document_codes


def test_clean_front_selarl_slice_switches_statuts_for_dentiste() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = _valid_selarl_input(PROFESSION_DENTISTE)

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is True
    assert "DOC-016" in plan.document_codes
    assert "DOC-017" not in plan.document_codes


def test_clean_front_selarl_slice_adds_regime_batch_only_for_regime() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = _valid_selarl_input(PROFESSION_MEDECIN, regime_communautaire=True)

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is True
    assert "DOC-005" in plan.document_codes
    assert "DOC-006" in plan.document_codes
    assert any(
        row.doc_code == "DOC-006" and row.status == "generable"
        for row in plan.document_rows
    )


def test_clean_front_selarl_medecin_regime_derives_conjoint_only_when_active() -> None:
    standard_ctx = build_generation_context(_valid_selarl_input(PROFESSION_MEDECIN))
    regime_ctx = build_generation_context(
        _valid_selarl_input(PROFESSION_MEDECIN, regime_communautaire=True)
    )

    assert standard_ctx.conjoint is None
    assert standard_ctx.regime_communautaire is None
    assert regime_ctx.conjoint is not None
    assert regime_ctx.conjoint.prenom == "Claire"
    assert regime_ctx.conjoint.nom == "Martin"
    assert regime_ctx.conjoint.adresse_perso is not None
    assert regime_ctx.conjoint.adresse_perso.adresse_affichee == "10 rue Test, 75001 Paris"
    assert regime_ctx.statuts_sel is not None
    assert regime_ctx.statuts_sel.overlay == "selarl_medecin"
    assert regime_ctx.regime_communautaire is not None
    assert regime_ctx.regime_communautaire.date_courrier_avertissement == date.today()
    assert regime_ctx.regime_communautaire.renonciation is not None
    assert (
        regime_ctx.regime_communautaire.renonciation.nombre_exemplaires_lettres
        == "quatre"
    )


def test_clean_front_selarl_regime_does_not_require_conjoint_address() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    kwargs = _valid_selarl_kwargs(PROFESSION_MEDECIN, regime_communautaire=True)
    data_entry = build_clean_data_entry(dossier_type, **kwargs)

    plan = build_clean_generation_plan(dossier_type, data_entry)
    ctx = build_generation_context(data_entry)

    assert plan.can_generate is True
    assert ctx.conjoint is not None
    assert ctx.conjoint.adresse_perso is not None
    assert ctx.conjoint.adresse_perso.adresse_affichee == "10 rue Test, 75001 Paris"


def test_clean_front_selarl_medecin_separation_de_biens_generates_statuts(
    tmp_path: Path,
) -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = _valid_selarl_input(
        PROFESSION_MEDECIN,
        married_separation=True,
    )

    plan = build_clean_generation_plan(dossier_type, data_entry)
    ctx = build_generation_context(data_entry)
    generated = generate_selarl_dossier(data_entry, tmp_path / "selarl-medecin-separation")
    statuts_path = next(
        path for path in generated.docx_paths if path.name.lower().startswith("statuts")
    )
    statuts_text = _docx_text(statuts_path)

    assert plan.can_generate is True
    assert "DOC-005" not in plan.document_codes
    assert "DOC-006" not in plan.document_codes
    assert ctx.conjoint is not None
    assert ctx.conjoint.adresse_perso is None
    assert ctx.associes[0].conjoint is not None
    assert (
        "marié sous le régime de la séparation de biens avec Madame Claire Martin"
        in statuts_text
    )


def test_clean_front_selarl_medecin_separation_de_biens_blocks_without_conjoint() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    kwargs = _valid_selarl_kwargs(PROFESSION_MEDECIN, married_separation=True)
    kwargs.update(
        {
            "conjoint_civilite": "",
            "conjoint_prenom": "",
            "conjoint_nom": "",
        }
    )
    data_entry = build_clean_data_entry(dossier_type, **kwargs)

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is False
    assert any("conjoint" in blocker.casefold() for blocker in plan.blockers)


def test_clean_front_selarl_regime_ui_never_exposes_conjoint_address_fields() -> None:
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    # Retours client 2026-06-11 : le regime de la communaute est derive de la
    # situation matrimoniale (plus de case a cocher dediee).
    app.selectbox(key="selarl_situation_maritale").set_value(
        "Marié(e) sous le régime légal / communauté"
    )
    app.run(timeout=120)

    conjoint_address_labels = [
        widget.label
        for widget in app.text_input
        if "conjoint" in widget.label.casefold()
        and "adresse" in widget.label.casefold()
    ]
    conjoint_address_keys = [
        str(widget.key)
        for widget in app.text_input
        if "conjoint_adresse" in str(widget.key)
    ]

    assert conjoint_address_labels == []
    assert conjoint_address_keys == []


def test_clean_front_selarl_ui_exposes_ordre_connecteur_selector() -> None:
    # M2 (Akainu, 2026-06-30, parite SELAS) : le formulaire SELARL expose le selecteur
    # de connecteur grammatical (de / du / des) avant le departement, sinon « du Calvados »
    # est inatteignable (le SELARL restait bloque sur « de »).
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    connecteur_box = next(
        (w for w in app.selectbox if str(w.key) == "selarl_ordre_connecteur"),
        None,
    )
    assert connecteur_box is not None
    assert [str(o) for o in connecteur_box.options] == ["de", "du", "des"]


def test_clean_front_selarl_slice_blocks_out_of_scope_cases() -> None:
    # La cession est desormais SUPPORTEE quand les donnees cession sont fournies
    # (cession_context). Demander la cession (flag) sans donnees reste bloque.
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = build_clean_data_entry(
        dossier_type,
        **{
            **_valid_selarl_kwargs(PROFESSION_MEDECIN),
            "cession": True,
        },
    )

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is False
    assert "Cession demandee mais donnees cession manquantes." in plan.blockers


def test_clean_front_selarl_cession_cabinet_medical_generates_acte(tmp_path: Path) -> None:
    # Cession avec donnees (scenario fige) -> l'acte de cession cabinet medical est generable.
    data = build_selarl_scenario("selarl_medecin_cession_cabinet_medical")
    dossier_type = dossier_type_by_label("SELARL creation V1")

    plan = build_clean_generation_plan(dossier_type, data)

    assert plan.can_generate is True
    assert "DOC-009" in plan.document_codes
    assert "DOC-007" in plan.document_codes
    # Appel de fonds (DOC-008) = document commun « Si cession », present aussi en medical.
    assert "DOC-008" in plan.document_codes

    result = generate_selarl_dossier(data, tmp_path)
    names = {path.name for path in result.docx_paths}
    assert "acte_cession_cabinet_medical.docx" in names
    assert "avenant_contrat_bail.docx" in names
    assert "appel_fond_sel.docx" in names


def test_clean_front_selarl_cession_cabinet_dentaire_generates_full_pack(tmp_path: Path) -> None:
    # Cession dentaire : acte (DOC-011) + avenant bail (DOC-007) + appel de fonds (DOC-008).
    data = build_selarl_scenario("selarl_dentiste_cession_cabinet_dentaire")
    dossier_type = dossier_type_by_label("SELARL creation V1")

    plan = build_clean_generation_plan(dossier_type, data)

    assert plan.can_generate is True
    assert {"DOC-011", "DOC-008", "DOC-007"}.issubset(set(plan.document_codes))

    result = generate_selarl_dossier(data, tmp_path)
    names = {path.name for path in result.docx_paths}
    assert "acte_cession_cabinet_dentaire.docx" in names
    assert "appel_fond_sel.docx" in names
    assert "avenant_contrat_bail.docx" in names


def test_clean_front_selarl_cession_scm_generates_scm_docs(tmp_path: Path) -> None:
    # Cession de parts de SCM : PV AGE (DOC-031) + courrier SDE (DOC-032) + acte (DOC-033).
    data = build_selarl_scenario("selarl_dentiste_cession_scm")
    dossier_type = dossier_type_by_label("SELARL creation V1")

    plan = build_clean_generation_plan(dossier_type, data)

    assert plan.can_generate is True
    assert {"DOC-031", "DOC-032", "DOC-033"}.issubset(set(plan.document_codes))

    result = generate_selarl_dossier(data, tmp_path)
    names = {path.name for path in result.docx_paths}
    assert "pv_age_cession_parts_scm.docx" in names
    assert "courrier_sde_cession_scm.docx" in names
    assert "acte_cession_parts_scm.docx" in names


def test_cession_context_medical_selects_acte_bail_appel_fonds() -> None:
    # (a) cession_context medical + bail -> DOC-009 (acte medical) + DOC-007 (bail)
    # + DOC-008 (appel de fonds).
    cession, bail = cession_fixture_for_profession(PROFESSION_MEDECIN)
    data_entry = _valid_selarl_input(
        PROFESSION_MEDECIN,
        cession=True,
        cession_context=cession,
        bail_context=bail,
    )

    codes = selected_selarl_document_codes(data_entry)

    assert "DOC-009" in codes
    assert "DOC-007" in codes
    assert "DOC-008" in codes
    assert "DOC-011" not in codes


def test_cession_context_dentaire_selects_acte_and_appel_fonds() -> None:
    # (a) cession_context dentaire -> DOC-011 (acte dentaire) + DOC-008 (appel de fonds).
    cession, bail = cession_fixture_for_profession(PROFESSION_DENTISTE)
    data_entry = _valid_selarl_input(
        PROFESSION_DENTISTE,
        cession=True,
        cession_context=cession,
        bail_context=bail,
    )

    codes = selected_selarl_document_codes(data_entry)

    assert "DOC-011" in codes
    assert "DOC-008" in codes
    assert "DOC-009" not in codes


def test_scm_cession_context_selects_scm_docs() -> None:
    # (c) scm_cession_context -> DOC-031 / DOC-032 / DOC-033.
    data_entry = _valid_selarl_input(
        PROFESSION_DENTISTE,
        scm=True,
        scm_cession_context=scm_cession_fixture(),
    )

    codes = selected_selarl_document_codes(data_entry)

    assert {"DOC-031", "DOC-032", "DOC-033"}.issubset(set(codes))


def test_scm_flag_without_data_is_blocked() -> None:
    # Le garde-fou reste : SCM coche sans donnees -> bloque (pas de generation muette).
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = build_clean_data_entry(
        dossier_type,
        **{**_valid_selarl_kwargs(PROFESSION_MEDECIN), "scm": True},
    )

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is False
    assert any("SCM" in blocker for blocker in plan.blockers)


def test_clean_front_ui_prefill_generates_cession_medical_without_residual_tokens(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # (b) Chemin UI complet : profession medecin, bouton de test (active la cession +
    # prereremplit selarl_cession_*), generation -> acte cession medical + appel de fonds
    # + bail, sans aucun token [xxx] residuel.
    _assert_ui_prefill_cession_generates(
        tmp_path,
        monkeypatch,
        profession_label="Medecin",
        expected_doc="acte_cession_cabinet_medical.docx",
    )


def test_clean_front_ui_prefill_generates_cession_dentaire_without_residual_tokens(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # (b) Chemin UI complet pour le cabinet dentaire (acte dentaire + salaries).
    _assert_ui_prefill_cession_generates(
        tmp_path,
        monkeypatch,
        profession_label="Chirurgien-dentiste",
        expected_doc="acte_cession_cabinet_dentaire.docx",
    )


def _assert_ui_prefill_cession_generates(
    tmp_path: Path,
    monkeypatch,
    *,
    profession_label: str,
    expected_doc: str,
) -> None:
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-cession")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="selarl_profession").set_value(profession_label)
    app = app.run(timeout=180)

    app.button(key="clean_generate_test_data").click()
    app = app.run(timeout=180)

    assert app.checkbox(key="selarl_cession").value is True
    assert app.checkbox(key="selarl_scm").value is True
    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_generate_dossier").disabled is False

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=180)

    download_labels = [item.label for item in app.get("download_button")]
    assert f"Telecharger {expected_doc}" in download_labels
    assert "Telecharger appel_fond_sel.docx" in download_labels
    assert "Telecharger avenant_contrat_bail.docx" in download_labels
    assert "Telecharger pv_age_cession_parts_scm.docx" in download_labels

    generated = app.session_state["clean_generated_dossier"]
    combined_text = "\n".join(
        _docx_text(Path(path)) for path in generated["docx_paths"]
    )
    assert "[" not in combined_text
    assert "]" not in combined_text


def test_clean_front_ui_prefill_selas_uni_medecin_generates(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Retour Rafael 2026-06-25 (#4) : le bouton « donnees de test » de la SELAS unipersonnelle
    # medecin pre-remplit un dossier COHERENT et GENERABLE en un clic (avant : aucun bouton).
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-uni")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(
        "SELAS unipersonnelle medecin creation V1"
    )
    app = app.run(timeout=180)

    app.button(key="clean_test_data_SELAS_uni_medecin").click()
    app = app.run(timeout=180)

    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_typed_generate_dossier").disabled is False

    app.button(key="clean_typed_generate_dossier").click()
    app = app.run(timeout=180)

    generated = app.session_state[shell.TYPED_GENERATED_STATE_KEY]
    combined_text = "\n".join(
        _docx_text(Path(path)) for path in generated["docx_paths"]
    )
    assert "SELAS EXEMPLE" in combined_text
    # Akainu M1/M2 (retour #1) : duree forcee « 99 ans » -> le template porte deja « années »,
    # le rendu doit etre « fixée à 99 années » (JAMAIS le doublon « 99 ans années » ni
    # « 99 ans ans », ni le placeholder brut).
    assert "fixée à 99 années" in combined_text
    assert "99 ans années" not in combined_text
    assert "99 ans ans" not in combined_text
    assert "[duree_societe]" not in combined_text
    # Marie sous communaute -> clause conjoint rendue (avec Madame Alice Durand).
    assert "Alice Durand" in combined_text
    # R5 (Albane) : destinataire FORME LONGUE « Conseil départemental de l’Ordre des
    # <profession_pluriel> <connecteur> <departement> » (profession PUIS departement, forme
    # du modele d'Albane, confirmee Rafael/Albane 2026-07-01 « comme les modeles »).
    assert "Conseil départemental de l’Ordre des médecins du Rhone" in combined_text
    # Dossier propre : aucun token/placeholder residuel.
    assert "[" not in combined_text
    assert "]" not in combined_text


def test_clean_front_ui_prefill_selas_uni_dentiste_generates(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Retour Rafael #5 : « la SELAS unipers dentiste = la pluripersonnelle dentiste mais
    # avec un seul associe ». Le bouton « donnees de test » pre-remplit un dossier
    # COHERENT et GENERABLE en un clic, profession chirurgien-dentiste (statuts DOC-046).
    from _accents import assert_no_unaccented_french

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-uni-dent")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(
        "SELAS unipersonnelle dentiste creation V1"
    )
    app = app.run(timeout=180)

    app.button(key="clean_test_data_SELAS_uni_dentiste").click()
    app = app.run(timeout=180)

    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_typed_generate_dossier").disabled is False

    app.button(key="clean_typed_generate_dossier").click()
    app = app.run(timeout=180)

    generated = app.session_state[shell.TYPED_GENERATED_STATE_KEY]
    # R10 (Rafael 2026-07-07) : le fichier statuts porte la denomination (prefill
    # « SELAS EXEMPLE »), plus le nom fixe « statuts_selas_dentiste.docx ».
    statuts_text = next(
        _docx_text(Path(path))
        for path in generated["docx_paths"]
        if Path(path).name == "Statuts SELAS EXEMPLE.docx"
    )
    # Corpus dentiste (verbatim du modele source dentiste pluri, uni-fie).
    assert "chirurgien-dentiste" in statuts_text
    assert "société d'exercice libéral par actions simplifiée" in statuts_text
    # Uni-fication : un seul soussigne, un seul apporteur, « A établi » (singulier).
    assert "LE SOUSSIGNE" in statuts_text
    assert "A établi ainsi qu’il suit" in statuts_text
    # Rafael 2026-07-09 « supprimer partout » : civilite CIVILE, plus « Docteur ».
    assert "- Monsieur Jean Durand, apporte" in statuts_text
    assert "Docteur" not in statuts_text
    # Article President dentiste GENERIQUE (jamais nomme) preserve.
    assert "Article 19 - Président de La société" in statuts_text
    # Marie sous communaute -> clause conjoint rendue (avec Madame Alice Durand).
    assert "Alice Durand" in statuts_text
    # Texte from-scratch dentiste : entierement accentue (garde-fou centralise).
    assert_no_unaccented_french(statuts_text)
    # Dossier propre : aucun token/placeholder residuel.
    assert "[" not in statuts_text
    assert "]" not in statuts_text


def test_clean_front_ui_prefill_micro_holding_generates(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Levier anti-Rafael : le bouton « donnees de test » de la Micro holding (societe civile
    # a capital variable) pre-remplit un dossier COHERENT et GENERABLE en un clic (avant :
    # aucun bouton pour ce type). Bundle = statuts (DOC-047) + tronc commun + PV gerant +
    # lettre d'option IS (DOC-022, activee par le prefill pour la demontrer).
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-micro-holding")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("Micro holding creation V1")
    app = app.run(timeout=180)

    # Le bouton n'apparait QUE si un prefill est enregistre pour la structure.
    app.button(key="clean_test_data_MICRO_HOLDING").click()
    app = app.run(timeout=180)

    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_typed_generate_dossier").disabled is False

    app.button(key="clean_typed_generate_dossier").click()
    app = app.run(timeout=180)

    generated = app.session_state[shell.TYPED_GENERATED_STATE_KEY]
    names = {Path(path).name for path in generated["docx_paths"]}
    # Statuts micro holding + tronc commun + PV gerant + lettre d'option IS (prefill IS actif).
    assert "Statuts MICRO HOLDING EXEMPLE.docx" in names  # R10 : denomination au nom
    assert "pv_nomination_gerant.docx" in names
    assert "lettre_option_is.docx" in names
    combined_text = "\n".join(_docx_text(Path(path)) for path in generated["docx_paths"])
    assert "MICRO HOLDING EXEMPLE" in combined_text
    # Societe civile A CAPITAL VARIABLE : la mention doit etre rendue.
    assert "capital variable" in combined_text
    # Dossier propre : aucun token/placeholder residuel.
    assert "[" not in combined_text
    assert "]" not in combined_text


def test_clean_front_ui_prefill_sasu_holding_generates(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Levier anti-Rafael : le bouton « donnees de test » de la SASU Holding (SAS
    # unipersonnelle, holding patrimoniale generaliste) pre-remplit un dossier COHERENT
    # et GENERABLE en un clic (avant : aucun bouton). Bundle 6 pieces : statuts (DOC-048)
    # + tronc commun + PV remuneration president (DOC-049) + liste souscripteurs (DOC-050).
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-sasu-holding")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(
        "SASU Holding (holding patrimoniale) creation V1"
    )
    app = app.run(timeout=180)

    app.button(key="clean_test_data_SASU_HOLDING").click()
    app = app.run(timeout=180)

    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_typed_generate_dossier").disabled is False

    app.button(key="clean_typed_generate_dossier").click()
    app = app.run(timeout=180)

    generated = app.session_state[shell.TYPED_GENERATED_STATE_KEY]
    names = {Path(path).name for path in generated["docx_paths"]}
    # 6 pieces : statuts + tronc commun (DNC/domic/procuration) + PV remu president + souscripteurs.
    assert "Statuts SASU HOLDING EXEMPLE.docx" in names  # R10 : denomination au nom
    assert "pv_remuneration_president_sasu_holding.docx" in names
    assert "liste_souscripteurs_sasu_holding.docx" in names
    assert len(generated["docx_paths"]) == 6
    combined_text = "\n".join(_docx_text(Path(path)) for path in generated["docx_paths"])
    assert "SASU HOLDING EXEMPLE" in combined_text
    assert "Jean Durand" in combined_text
    # Dossier propre : aucun token/placeholder residuel.
    assert "[" not in combined_text
    assert "]" not in combined_text


def test_selas_uni_dentiste_generator_matches_dentiste_model_wording(
    tmp_path: Path,
) -> None:
    # Fidelite : le statuts SELAS uni dentiste reprend le wording VERBATIM du modele
    # source dentiste pluri (`Statuts_SELAS_dentiste_pluri_modele.docx`), uni-fie. On
    # verifie quelques ancres litterales propres au corpus dentiste (et absentes du
    # corpus medecin), plus l'absence de placeholder et l'accentuation complete.
    from _accents import assert_no_unaccented_french

    from sydel_doc_engine.front_app import selas_uni_dentiste_slice as sd

    payload: dict[str, object] = {
        "denomination": "SELAS DENTAIRE TEST",
        "capital_social": "1000",
        "nb_actions_total": 100,
        "valeur_nominale_action": "10",
        "duree": "99 ans",
        "ville_rcs": "Rennes",
        "lieu_exercice_adresse": "9 rue du Centre, 35000 Rennes",
        "siege_num": "9",
        "siege_voie": "rue du Centre",
        "siege_cp": "35000",
        "siege_ville": "Rennes",
        "banque_nom": "BANQUE TEST",
        "banque_adresse": "8 place Test, 35000 Rennes",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 décembre",
        "exercice_cloture": "31 décembre 2026",
        "civilite": "Monsieur",
        "prenom": "Jean",
        "nom": "Durand",
        "date_naissance": date(1980, 1, 1),
        "ville_naissance": "Rennes",
        "departement_naissance": "35",
        "nationalite": "française",
        "titre_affichage": "Docteur",
        "adresse_num": "31",
        "adresse_voie": "boulevard Test",
        "adresse_cp": "35700",
        "adresse_ville": "Rennes",
        "nom_pere": "Pierre Durand",
        "nom_mere": "Anne Durand",
        "situation_maritale": "celibataire",
        # Pour un celibataire, le formulaire derive regime_matrimonial = « celibataire »
        # (regime_matrimonial_from_status) — non vide, donc generable sans conjoint.
        "regime_matrimonial": "celibataire",
        "regime_communautaire": False,
        "conjoint_civilite": "",
        "conjoint_prenom": "",
        "conjoint_nom": "",
        "departement_ordre": "Ille et Vilaine",
        "connecteur_departement": "de",
        "ordre_president_feminin": False,
        "mandataire_prenom": "Marc",
        "mandataire_nom": "Sydel",
        "numero_ordre": "79630",
        "numero_rpps": "10101676194",
        "ordre_ville": "Rennes",
        "ordre_cp": "35000",
        "ordre_adresse_ligne_1": "1 rue Ordre",
        "signature_lieu": "Rennes",
        "signature_date": date(2026, 9, 22),
    }
    plan = sd.build_selas_uni_dentiste_plan(payload)
    assert plan.can_generate, plan.blockers
    assert plan.document_codes[0] == "DOC-046"

    ctx = sd.build_generation_context(payload)
    from sydel_doc_engine.generators.lot_04.statuts_selas_dentiste import (
        StatutsSelasDentisteGenerator,
    )

    out_path = StatutsSelasDentisteGenerator().generate(ctx, tmp_path / "dent")
    text = _docx_text(out_path)

    # Ancres litterales DENTISTE (corpus dentiste pluri, absentes du corpus medecin).
    assert "La société est constituée en Société d'Exercice Libéral par Actions Simplifiée." in text
    assert "par les articles R. 4113-1 et suivants du Code de la santé publique" in text
    assert "La société a pour objet l’exercice en commun de la profession de chirurgien-dentiste." in text  # noqa: E501
    assert "l’agrément de celle-ci par le Conseil Départemental de l’Ordre des Chirurgiens-dentistes" in text  # noqa: E501
    assert "Article 19 - Président de La société" in text
    assert "Article 20 - Directeur Général" in text
    # Uni-fication : associe unique, un seul apporteur / attributaire.
    assert "A la constitution de la Société, l’associé unique a fait les apports suivants" in text
    # Rafael 2026-07-09 « supprimer partout » : apport ET repartition en civilite CIVILE, plus
    # « Docteur »/« Le Docteur » (le titre professionnel n'apparait plus dans la sortie).
    assert "- Monsieur Jean Durand, apporte mille euros" in text
    assert "- Monsieur Jean Durand, cent actions" in text
    assert "Docteur" not in text
    # Retour Albane « mise en forme » 2.3 : le mot « euro(s) » (accorde sur la FIGURE)
    # figure apres la valeur nominale en lettres. VN=10 (front-app reel -> lettres « dix »,
    # SANS euro) -> le token « [euro_nominal_word] » du modele dentiste ajoute « euros ».
    assert "100 actions de 10 € (dix euros) chacune" in text
    assert "(dix) chacune" not in text  # garde-fou : le mot euro NE doit PAS manquer
    assert "euros euros" not in text and "euro euros" not in text  # pas de double euro
    # Celibataire : la clause matrimoniale rend juste « celibataire » (pas de conjoint).
    assert "marié" not in text.replace("non marié", "")
    # Propre + accentue.
    assert "[" not in text and "]" not in text
    assert_no_unaccented_french(text)


def test_clean_front_cession_libre_dates_have_accented_months(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # LIVE-03 : les dates a SAISIE LIBRE du sous-formulaire cession/bail
    # (date du bail, date d'origine de propriete, date de naissance vendeur)
    # doivent ressortir avec les mois ACCENTUES dans le DOCX, meme si l'operateur
    # saisit « aout » sans accent. Le generateur reste un echo fidele : la
    # re-accentuation se fait EN AMONT (front, _render_cession_form).
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-cession-accents")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="selarl_profession").set_value("Medecin")
    app = app.run(timeout=180)

    app.button(key="clean_generate_test_data").click()
    app = app.run(timeout=180)
    assert app.checkbox(key="selarl_cession").value is True

    # Saisies LIBRES non accentuees (« 1er aout 2021 ») sur les dates du bail et de
    # l'origine de propriete (toutes deux passees par _cession_date, comme la date de
    # naissance vendeur a saisie libre — meme helper).
    app.text_input(key="selarl_cession_bail_date_bail").set_value("1er aout 2021")
    app.text_input(key="selarl_cession_cabinet_origine_date").set_value("3 aout 2019")
    app = app.run(timeout=180)

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=180)
    assert [e.value for e in app.error] == []

    generated = app.session_state["clean_generated_dossier"]
    combined = "\n".join(_docx_text(Path(path)) for path in generated["docx_paths"])

    # Mois ACCENTUE present (date du bail + date d'origine de propriete).
    assert "août" in combined
    # Plus aucune occurrence du mois « aout » NON accentue (toutes re-accentuees).
    assert re.search(r"\baout\b", combined) is None


def test_clean_front_ui_creation_only_unchanged(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # (d) Sans cession ni SCM : la creation seule genere toujours les 6 documents,
    # inchangee par le cablage du sous-formulaire.
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = _valid_selarl_input(PROFESSION_MEDECIN)

    plan = build_clean_generation_plan(dossier_type, data_entry)
    generated = generate_selarl_dossier(data_entry, tmp_path / "creation-only")

    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-017",
    )
    assert data_entry.cession_context is None
    assert data_entry.scm_cession_context is None
    assert len(generated.docx_paths) == 6


def test_clean_front_cession_form_returns_none_without_flag() -> None:
    # Le sous-formulaire ne rend rien et renvoie (None, None) quand la cession
    # n'est pas demandee -> aucun expander parasite dans le wizard de base.
    from sydel_doc_engine.front_app import shell

    cession_context, bail_context = shell._render_cession_form(
        False,
        PROFESSION_MEDECIN,
        praticien={},
        societe={},
        ordre={},
        generation={},
    )
    scm_context = shell._render_scm_cession_form(
        False, praticien={}, societe={}, profession_label="", ordre={}
    )

    assert cession_context is None
    assert bail_context is None
    assert scm_context is None


def test_scm_cession_derive_apres_cession_deux_associes() -> None:
    # §4.1 : derivation deterministe de l'apres-cession a partir des presents.
    # 2 presents (Paul 40 / Jean 60 = 100). Jean (cedant) cede 20 parts (plage
    # 81 à 100) a la SEL acquereur. Attendu : Paul inchange (40), Jean reduit
    # (40, plage residuelle 41 à 80), SEL entrante (20, plage 81 à 100).
    # N4 (2026-07-01) : plages avec « à » accentue (« comme les modeles »).
    from sydel_doc_engine.domain.models import (
        ScmCessionAssocie,
        ScmCessionPartsAttribution,
    )
    from sydel_doc_engine.front_app import shell

    presents = [
        ScmCessionAssocie(
            type_personne="personne_physique", civilite_affichage="Monsieur",
            prenom="Paul", nom="Bernard",
            parts=ScmCessionPartsAttribution(nb=40, plage="1 à 40"),
        ),
        ScmCessionAssocie(
            type_personne="personne_physique", civilite_affichage="Monsieur",
            prenom="Jean", nom="Dupont",
            parts=ScmCessionPartsAttribution(nb=60, plage="41 à 100"),
        ),
    ]
    cedant = {"prenom": "Jean", "nom": "Dupont"}
    cessionnaire = {"denomination": "SELARL CABINET DUPONT", "forme_juridique": "SELARL"}
    parts_cedees = {"nb": 20, "plage": "81 à 100"}

    apres = shell._derive_scm_apres_cession(presents, cedant, cessionnaire, parts_cedees)

    assert len(apres) == 3
    paul, jean, sel = apres
    assert (paul.prenom, paul.nom, paul.parts.nb, paul.parts.plage) == (
        "Paul", "Bernard", 40, "1 à 40",
    )
    # Cedant reduit : 60 - 20 = 40, plage residuelle = complement (41 à 80).
    assert (jean.prenom, jean.nom, jean.parts.nb, jean.parts.plage) == (
        "Jean", "Dupont", 40, "41 à 80",
    )
    # SEL acquereur entrante (personne morale) avec les parts cedees.
    assert sel.type_personne == "personne_morale"
    assert sel.denomination == "SELARL CABINET DUPONT"
    assert (sel.parts.nb, sel.parts.plage) == (20, "81 à 100")
    # Sommes coherentes : 40 + 40 + 20 = 100 = total des presents.
    assert sum(a.parts.nb for a in apres) == sum(p.parts.nb for p in presents)

    # signataires_pv derive des presents (libelle court).
    signataires = shell._derive_scm_signataires_pv(presents)
    assert signataires == ["M. Paul Bernard", "M. Jean Dupont"]


def test_scm_cession_cedant_cede_toutes_ses_parts_sort_de_lapres() -> None:
    # §4.1 : si le cedant cede TOUTES ses parts -> il disparait de l'apres-cession,
    # remplace par la SEL acquereur (pas de cedant a 0 part).
    from sydel_doc_engine.domain.models import (
        ScmCessionAssocie,
        ScmCessionPartsAttribution,
    )
    from sydel_doc_engine.front_app import shell

    presents = [
        ScmCessionAssocie(
            type_personne="personne_physique", civilite_affichage="Madame",
            prenom="Anne", nom="Martin",
            parts=ScmCessionPartsAttribution(nb=50, plage="1 a 50"),
        ),
        ScmCessionAssocie(
            type_personne="personne_physique", civilite_affichage="Monsieur",
            prenom="Jean", nom="Dupont",
            parts=ScmCessionPartsAttribution(nb=50, plage="51 a 100"),
        ),
    ]
    cedant = {"prenom": "Jean", "nom": "Dupont"}
    cessionnaire = {"denomination": "SELAS CABINET", "forme_juridique": "SELAS"}
    parts_cedees = {"nb": 50, "plage": "51 a 100"}

    apres = shell._derive_scm_apres_cession(presents, cedant, cessionnaire, parts_cedees)

    noms = [(a.type_personne, a.nom or a.denomination) for a in apres]
    assert ("personne_physique", "Dupont") not in noms  # cedant sorti
    assert ("personne_physique", "Martin") in noms  # autre associe garde
    assert ("personne_morale", "SELAS CABINET") in noms  # SEL entrante
    assert sum(a.parts.nb for a in apres) == 100


def test_clean_front_selarl_cession_compromis_generates(tmp_path: Path) -> None:
    # Compromis de cession : médical (DOC-010) et dentaire (DOC-012), même moteur que l'acte.
    dossier_type = dossier_type_by_label("SELARL creation V1")
    for scenario, expected_doc, filename in (
        (
            "selarl_medecin_cession_compromis_medical",
            "DOC-010",
            "compromis_cession_cabinet_medical.docx",
        ),
        (
            "selarl_dentiste_cession_compromis_dentaire",
            "DOC-012",
            "compromis_cession_cabinet_dentaire.docx",
        ),
    ):
        data = build_selarl_scenario(scenario)
        plan = build_clean_generation_plan(dossier_type, data)
        assert plan.can_generate is True
        assert expected_doc in plan.document_codes
        result = generate_selarl_dossier(data, tmp_path / scenario)
        names = {path.name for path in result.docx_paths}
        assert filename in names


def test_clean_front_selarl_context_selects_only_expected_engine_docs() -> None:
    ctx = build_generation_context(_valid_selarl_input(PROFESSION_DENTISTE))
    selected = DocumentOrchestrator(build_seed_catalog()).select_documents_for_context(ctx)
    selected_codes = {document.doc_id for document in selected}

    assert {"DOC-001", "DOC-002", "DOC-003", "DOC-004", "DOC-034", "DOC-016"}.issubset(
        selected_codes
    )
    assert "DOC-017" not in selected_codes
    assert "DOC-005" not in selected_codes
    assert "DOC-006" not in selected_codes


def test_clean_front_selarl_context_derives_hidden_ux_values() -> None:
    ctx = build_generation_context(_valid_selarl_input(PROFESSION_MEDECIN))

    assert ctx.personne_signataire.genre.value == "masculin"
    assert ctx.personne_signataire.date_naissance == date(1984, 4, 12)
    assert ctx.personne_signataire.ville_naissance == "Paris"
    assert ctx.personne_signataire.ville_naissance_article_au is False
    assert ctx.societe.capital_social_lettres == "mille"
    assert ctx.capital is not None
    assert ctx.capital.valeur_nominale_titre == "10"
    assert ctx.capital.nombre_titres_total_lettres == "cent"
    assert ctx.capital.valeur_nominale_titre_lettres == "dix"
    assert ctx.reunion is not None
    assert ctx.reunion.date_lettres == "vingt-six mai deux mille vingt-six"
    assert ctx.reunion.heure is None
    assert ctx.reunion.president is not None
    assert ctx.reunion.president.civilite_president_seance == "Monsieur"
    assert ctx.reunion.president.prenom_president_seance == "Jean"
    assert ctx.reunion.president.nom_personne_seance == "Martin"
    assert ctx.signature.prestataire_signature_electronique == "Yousign"
    assert ctx.gerance is not None
    # R5 (Albane 2026-07-07) : seuils groupes par 3 des la construction du contexte.
    assert ctx.gerance.seuil_achat_materiel == "5 000"
    assert ctx.gerance.seuil_emprunt == "10 000"
    assert ctx.mandataire is not None
    assert ctx.mandataire.cabinet == "SYDEL"
    assert ctx.exercice_social is not None
    assert ctx.exercice_social.lieux[0].adresse_affichee == "20 avenue du Siege, 75002 Paris"
    assert ctx.associes[0].nb_parts == 100
    # Ticket 2.2 : sans 2e lieu saisi, exactement UN lieu (le siege) — inchange.
    assert len(ctx.exercice_social.lieux) == 1


def test_clean_front_selarl_second_lieu_appends_lieux_1() -> None:
    # Ticket 2.2 ADDITIF : un 2e lieu (nom + adresse) ajoute lieux[1] sans toucher
    # lieux[0] (toujours le siege). Le contexte porte alors DEUX lieux.
    ctx = build_generation_context(
        _valid_selarl_input(
            PROFESSION_MEDECIN,
            second_lieu_exercice_nom="Cabinet secondaire",
            second_lieu_exercice_adresse="20 rue Bleue, 75009 Paris",
        )
    )
    assert ctx.exercice_social is not None
    assert len(ctx.exercice_social.lieux) == 2
    assert ctx.exercice_social.lieux[0].adresse_affichee == "20 avenue du Siege, 75002 Paris"
    assert ctx.exercice_social.lieux[1].nom == "Cabinet secondaire"
    assert ctx.exercice_social.lieux[1].adresse_affichee == "20 rue Bleue, 75009 Paris"


def test_clean_front_selarl_exercice_debut_accentuates_month() -> None:
    # LIVE-03 / MINEUR 5 (re-Akainu T4) : le debut d'exercice est un text_input LIBRE ;
    # un mois saisi sans accent (« 1er aout ») doit ressortir accentue (« 1er août »),
    # comme fin/cloture. On verrouille la re-accentuation EN AMONT cote selarl_slice.
    ctx = build_generation_context(
        _valid_selarl_input(PROFESSION_MEDECIN, exercice_debut="1er aout")
    )
    assert ctx.exercice_social is not None
    assert ctx.exercice_social.debut == "1er août"


def test_clean_front_selarl_partial_second_lieu_ignored() -> None:
    # Ticket 2.2 : nom OU adresse seul -> pas de 2e lieu cote front (un seul lieu).
    ctx = build_generation_context(
        _valid_selarl_input(
            PROFESSION_MEDECIN,
            second_lieu_exercice_nom="Cabinet secondaire",
        )
    )
    assert ctx.exercice_social is not None
    assert len(ctx.exercice_social.lieux) == 1


def test_clean_front_selarl_accepts_french_date_strings_outside_streamlit_range() -> None:
    dossier_type = dossier_type_by_label("SELARL creation V1")
    data_entry = build_clean_data_entry(
        dossier_type,
        **{
            **_valid_selarl_kwargs(PROFESSION_MEDECIN),
            "date_naissance": "31/12/1974",
            "signature_date": "27/05/2026",
            "decision_date": "27/05/2026",
        },
    )

    plan = build_clean_generation_plan(dossier_type, data_entry)

    assert plan.can_generate is True
    assert data_entry.date_naissance == date(1974, 12, 31)
    assert data_entry.signature_date == date(2026, 5, 27)
    assert data_entry.decision_date == date(2026, 5, 27)


def test_clean_front_selarl_maps_birth_city_article_au() -> None:
    ctx = build_generation_context(
        _valid_selarl_input(
            PROFESSION_MEDECIN,
            ville_naissance="Bourget",
            ville_naissance_article_au=True,
        )
    )

    assert ctx.personne_signataire.ville_naissance == "Bourget"
    assert ctx.personne_signataire.ville_naissance_article_au is True


def test_clean_front_selarl_generation_smoke(tmp_path: Path) -> None:
    generated = generate_selarl_dossier(
        _valid_selarl_input(PROFESSION_MEDECIN),
        tmp_path / "selarl-medecin",
    )

    assert len(generated.docx_paths) == 6
    assert generated.zip_path.exists()
    assert {path.name for path in generated.docx_paths} >= {
        # O24-02 : la DNC porte le nom du gerant (Martin) dans son nom de fichier.
        "declaration_non_condamnation_Martin.docx",
        "autorisation_domiciliation.docx",
        "procuration.docx",
        "pv_nomination_gerant.docx",
        "demande_inscription_ordre.docx",
        # Retour Albane 2026-06-10 : intitulé du doc = « Statuts {dénomination} ».
        "Statuts SELARL MARTIN.docx",
    }
    combined_text = "\n".join(_docx_text(path) for path in generated.docx_paths)
    assert "SELARL SELARL" not in combined_text
    assert "Société d’exercice libéral à responsabilité limitée de médecin" in combined_text
    # R5 (Albane) : destinataire FORME LONGUE « Conseil départemental de l’Ordre des
    # <profession_pluriel> <connecteur> <departement> » (profession PUIS departement).
    # Retour fonctionnel 9.2 (Albane 2026-07-06) : le departement est affiche par son NOM
    # (« Paris »), pas par son numero (« 75 ») — propage a tous les types via le generateur partage.
    assert "Conseil départemental de l’Ordre des médecins de Paris" in combined_text
    assert "Au capital de 1 000 euros" in combined_text
    assert "Au capital de 1000" not in combined_text
    assert " medecin" not in combined_text


def test_clean_front_selarl_medecin_regime_communautaire_generation_smoke(
    tmp_path: Path,
) -> None:
    generated = generate_selarl_dossier(
        _valid_selarl_input(PROFESSION_MEDECIN, regime_communautaire=True),
        tmp_path / "selarl-medecin-regime-communautaire",
    )

    names = {path.name for path in generated.docx_paths}
    assert len(generated.docx_paths) == 8
    assert generated.zip_path.exists()
    assert names == {
        "declaration_non_condamnation_Martin.docx",  # O24-02 : DNC nommee par le gerant
        "autorisation_domiciliation.docx",
        "procuration.docx",
        "pv_nomination_gerant.docx",
        "demande_inscription_ordre.docx",
        "Statuts SELARL MARTIN.docx",
        "lettre_renonciation_associe.docx",
        "lettre_avertissement_conjoint.docx",
    }
    assert "Statuts SEL CHIRURGIEN.docx" not in names

    combined_text = "\n".join(_docx_text(path) for path in generated.docx_paths)
    ascii_text = _ascii_text(combined_text)
    assert "[" not in combined_text
    assert "]" not in combined_text
    assert "SELARL SELARL" not in combined_text
    # Retour Albane 2026-06-10 : la procuration porte desormais le RCS/SIREN + tel SYDEL.
    assert "RCS PARIS 788 531 432" in combined_text
    assert "0153814303" in combined_text
    assert "Société d’exercice libéral à responsabilité limitée de médecin" in combined_text
    assert "Au capital de 1 000 €" in combined_text
    assert "Au capital de 1000" not in combined_text
    assert " medecin" not in combined_text
    assert f"Par courrier en date du {date.today():%d/%m/%Y}" in combined_text
    assert "euros dependant de notre communaute." in ascii_text
    assert "regime de communaute" not in ascii_text
    # A3 (Albane 2026-06-26) : le destinataire conjoint porte desormais son PRENOM.
    assert "Madame Claire Martin" in combined_text
    assert "Madame Martin\n" not in combined_text
    assert "10 rue Test" in combined_text
    assert "30 rue Conjoint" not in combined_text


def test_clean_front_legacy_boundary_is_explicit() -> None:
    items = legacy_boundary_items()
    decisions = {item.decision for item in items}
    components = {item.component for item in items}

    assert {"reused", "legacy_reference", "ignored_by_clean_front", "remove_later"}.issubset(
        decisions
    )
    assert "src/sydel_doc_engine/front_data/" in components
    assert "src/sydel_doc_engine/app/streamlit_app.py" in components


def test_clean_front_entrypoint_does_not_import_legacy_screens() -> None:
    source = "\n".join(
        Path(path).read_text(encoding="utf-8")
        for path in (
            "src/sydel_doc_engine/front_app/app.py",
            "src/sydel_doc_engine/front_app/shell.py",
            "src/sydel_doc_engine/front_app/selarl_slice.py",
        )
    )

    assert "business_wizard" not in source
    assert "single_document_mode" not in source
    assert "streamlit_app" not in source
    assert "front_internal_tool" not in source
    assert "Technique / diagnostic" not in source
    assert "Document unitaire" not in source
    assert "Debug interne" not in source


def test_clean_front_streamlit_surface_is_not_legacy() -> None:
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    assert [item.value for item in app.subheader] == [
        "Type de dossier",
        "Donnees a saisir",
        "Generation",
    ]
    assert app.selectbox(key="clean_dossier_type").label == "Type de dossier"
    assert app.selectbox(key="clean_dossier_type").value == "SELARL creation V1"
    assert app.selectbox(key="selarl_profession").label == "Profession"
    assert not any(str(widget.key) == "selarl_case_mode" for widget in app.selectbox)
    assert app.button(key="clean_generate_test_data").label == "Generer des donnees de test"
    assert app.button(key="clean_generate_dossier").disabled is True
    assert app.button(key="selarl_signature_date_today").label == "Aujourd'hui"
    assert len(app.radio) == 0
    assert len(app.table) == 0
    assert len(app.expander) == 0
    assert not any(item.label == "Outils internes" for item in app.checkbox)

    visible_labels = {
        *[item.label for item in app.text_input],
        *[item.label for item in app.number_input],
        *[item.label for item in app.selectbox],
        *[item.label for item in app.date_input],
    }
    # Rafael 2026-07-02 : le calendrier `st.date_input` (propose en option depuis R3) a ete
    # RETIRE — il s'affichait comme un champ date VIDE en double (« tout seul ») au-dessus de
    # chaque champ date, sur tous les formulaires. La surface n'expose donc PLUS aucun
    # `st.date_input` ; chaque date = un champ texte JJ/MM/AAAA + un bouton « Aujourd'hui ».
    assert len(app.date_input) == 0
    assert "Genre" not in visible_labels
    assert "Titre affichage" not in visible_labels
    assert "Capital social en lettres" not in visible_labels
    assert "Nombre de parts en lettres" not in visible_labels
    assert "Valeur nominale d'une part (€)" not in visible_labels
    assert "Valeur nominale en lettres" not in visible_labels
    assert "Date reunion en lettres" not in visible_labels
    assert "Heure de decision" not in visible_labels
    assert "Regime matrimonial" not in visible_labels
    assert "Prestataire signature electronique" not in visible_labels
    assert "Seuil achat materiel" not in visible_labels
    assert "Seuil emprunt" not in visible_labels
    assert "Lieu exercice" not in visible_labels
    assert "Civilite mandataire" not in visible_labels
    assert "Civilite conjoint" not in visible_labels
    assert app.selectbox(key="selarl_nationalite_choice").label == "Nationalite"
    assert "Portugaise" in app.selectbox(key="selarl_nationalite_choice").options
    assert app.selectbox(key="selarl_situation_maritale").label == "Situation matrimoniale"
    # Retours client 2026-06-11 (ticket 1.2) : les regimes matrimoniaux sont des
    # options de la situation matrimoniale ; plus de case a cocher dediee.
    assert "Marié(e) sous le régime légal / communauté" in app.selectbox(
        key="selarl_situation_maritale"
    ).options
    assert "Marié(e) sous le régime de la séparation de biens" in app.selectbox(
        key="selarl_situation_maritale"
    ).options
    assert not any(
        str(widget.key) == "selarl_regime_communautaire" for widget in app.checkbox
    )
    # Ticket 1.1 : libelle « Cession de fonds liberal ».
    assert any(
        widget.label == "Cession de fonds liberal"
        for widget in app.checkbox
        if str(widget.key) == "selarl_cession"
    )
    # O24-03 : adresse personnelle sur UNE ligne (plus de champ voie/CP/Ville separes).
    assert (
        app.text_input(key="selarl_adresse_ligne").label
        == "Adresse personnelle (N° et voie, CP Ville)"
    )
    for stray in ("selarl_adresse_voie", "selarl_adresse_cp", "selarl_adresse_ville"):
        assert not any(str(widget.key) == stray for widget in app.text_input), stray
    # Ticket 1.7 : dates d'exercice preremplies dynamiquement (cloture N+1).
    assert app.text_input(key="selarl_exercice_debut").value == "1er janvier"
    assert app.text_input(key="selarl_exercice_fin").value == "31 décembre"
    assert (
        app.text_input(key="selarl_exercice_cloture_premier").value
        == f"31 décembre {date.today().year + 1}"
    )
    assert (
        app.text_input(key="selarl_departement_ordre").label
        == "Departement d'inscription a l'ordre"
    )


def test_clean_front_streamlit_no_longer_exposes_multi_associes_case() -> None:
    # SELARL = unipersonnelle (decision Gad 2026-06-04) : aucun selecteur de cas multi,
    # aucun champ de sous-formulaire multi-associes ne doit subsister dans le wizard.
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    assert not any(str(widget.key) == "selarl_case_mode" for widget in app.selectbox)
    multi_keys = [
        str(widget.key)
        for group in (app.number_input, app.selectbox, app.text_input)
        for widget in group
        if "doc004" in str(widget.key)
    ]
    assert multi_keys == []


def test_clean_front_streamlit_random_data_button_prefills_generable_case(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "random-data")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    app.button(key="clean_generate_test_data").click()
    app = app.run(timeout=120)

    assert app.text_input(key="selarl_dossier_reference").value.startswith("TEST-SELARL-")
    assert app.number_input(key="selarl_capital_social").value > 0
    assert app.number_input(key="selarl_nb_parts_total").value > 0
    assert app.button(key="clean_generate_dossier").disabled is False
    assert not any("Blocage" in item.value for item in app.caption)

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=120)

    assert any(
        item.label == "Telecharger le dossier ZIP"
        for item in app.get("download_button")
    )


def test_clean_front_streamlit_generation_exposes_download_buttons(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "streamlit-downloads")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    _fill_valid_streamlit_selarl_form(app)
    app = app.run(timeout=120)

    assert app.button(key="clean_generate_dossier").disabled is False

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=120)

    download_labels = [item.label for item in app.get("download_button")]
    assert download_labels == [
        "Telecharger le dossier ZIP",
        "Telecharger declaration_non_condamnation_Martin.docx",
        "Telecharger autorisation_domiciliation.docx",
        "Telecharger procuration.docx",
        "Telecharger pv_nomination_gerant.docx",
        "Telecharger demande_inscription_ordre.docx",
        "Telecharger Statuts SELARL MARTIN.docx",
    ]


def _membre_bernard() -> StatutsCivilsAssocie:
    """Membre additionnel SELARL multi VALIDE (donnees DOC-001 completes — DNC par
    associe, Rafael 2026-07-09 : filiation requise pour chaque membre physique)."""
    return StatutsCivilsAssocie(
        type_personne="personne_physique",
        civilite_affichage="Madame",
        prenom="Lea",
        nom="Bernard",
        profession="medecin",
        date_naissance="03/04/1985",
        ville_naissance="Lyon",
        departement_naissance="69",
        nationalite="française",
        situation_maritale="celibataire",
        adresse_personnelle_affichee="8 rue Centrale, 69001 Lyon",
        nom_pere="Laurent Bernard",
        nom_mere="Julie Bernard",
        ordre_departemental="69",
        numero_ordre="ORD-999",
        numero_rpps="20000000002",
        apport=StatutsCivilsApport(montant="400", montant_lettres="quatre cents"),
        parts=StatutsCivilsParts(nb=40, nb_lettres="quarante"),
    )


def test_clean_front_selarl_multi_associes_generates_statuts(tmp_path: Path) -> None:
    # Retours V3 2026-06-17 (SELARL multi-associes) — ADDITIF. Praticien (60 parts)
    # + 1 membre additionnel personne physique (40 parts) = 100 parts (= capital).
    dossier_type = dossier_type_by_label("SELARL creation V1")
    membre = _membre_bernard()
    data = _valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        praticien_apport="600",
        membres_additionnels=(membre,),
    )

    plan = build_clean_generation_plan(dossier_type, data)
    assert plan.can_generate is True, plan.blockers

    result = generate_selarl_dossier(data, tmp_path)
    statuts = next(p for p in result.docx_paths if "Statuts" in p.name)
    text = _docx_text(statuts)

    assert "LES SOUSSIGNÉS" in text
    assert "Madame Lea Bernard" in text
    assert "1° Jean Martin, détenant 60 parts ;" in text
    assert "2° Lea Bernard, détenant 40 parts." in text


def test_clean_front_selarl_multi_dnc_une_par_associe(tmp_path: Path) -> None:
    # DNC par associe (Rafael 2026-07-09) : en multi-associes, UNE declaration de
    # non-condamnation PAR associe personne physique — praticien (Martin, DNC de
    # l'orchestrateur renommee O24-02) + membre additionnel (Bernard, generee et
    # nommee par son nom). 2 associes -> 2 documents.
    data = _valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        praticien_apport="600",
        membres_additionnels=(_membre_bernard(),),
    )
    result = generate_selarl_dossier(data, tmp_path / "selarl-multi-dnc")
    names = {p.name for p in result.docx_paths}
    assert "declaration_non_condamnation_Martin.docx" in names  # praticien
    assert "declaration_non_condamnation_Bernard.docx" in names  # membre
    assert "declaration_non_condamnation.docx" not in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 2
    # La DNC du membre porte SON identite (signataire = Lea Bernard) et SA filiation.
    dnc_bernard = _docx_text(
        next(
            p
            for p in result.docx_paths
            if p.name == "declaration_non_condamnation_Bernard.docx"
        )
    )
    assert "Madame Lea Bernard" in dnc_bernard
    assert "Laurent Bernard" in dnc_bernard
    assert "Julie Bernard" in dnc_bernard
    assert "Jean Martin" not in dnc_bernard


def test_clean_front_selarl_multi_membre_sans_filiation_blocks() -> None:
    # DNC par associe (Rafael 2026-07-09) : membre sans filiation -> la generation
    # de SA declaration de non-condamnation crasherait ; on bloque en amont.
    membre = _membre_bernard().model_copy(update={"nom_pere": None, "nom_mere": None})
    data = _valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        praticien_apport="600",
        membres_additionnels=(membre,),
    )
    dossier_type = dossier_type_by_label("SELARL creation V1")
    plan = build_clean_generation_plan(dossier_type, data)
    assert plan.can_generate is False
    assert any("nom du pere du membre 2" in b for b in plan.blockers), plan.blockers
    assert any("nom de la mere du membre 2" in b for b in plan.blockers), plan.blockers


def test_clean_front_selarl_multi_membre_sans_ordre_blocks() -> None:
    # Dogfood 2026-06-22 : membre additionnel sans inscription a l'ordre (departement /
    # numero / RPPS) passait la validation puis crashait a la generation. Doit bloquer.
    membre = StatutsCivilsAssocie(
        type_personne="personne_physique",
        civilite_affichage="Madame",
        prenom="Lea",
        nom="Bernard",
        profession="medecin",
        date_naissance="03/04/1985",
        ville_naissance="Lyon",
        departement_naissance="69",
        nationalite="française",
        situation_maritale="celibataire",
        adresse_personnelle_affichee="8 rue Centrale, 69001 Lyon",
        apport=StatutsCivilsApport(montant="400", montant_lettres="quatre cents"),
        parts=StatutsCivilsParts(nb=40, nb_lettres="quarante"),
    )  # SANS ordre_departemental / numero_ordre / numero_rpps
    data = _valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        praticien_apport="600",
        membres_additionnels=(membre,),
    )
    dossier_type = dossier_type_by_label("SELARL creation V1")
    plan = build_clean_generation_plan(dossier_type, data)
    assert plan.can_generate is False
    assert any(("ordre" in b.lower() or "rpps" in b.lower()) for b in plan.blockers)


def test_clean_front_selarl_multi_associes_blocks_incoherent_total(tmp_path: Path) -> None:
    # Somme des parts (praticien 60 + membre 30 = 90) != capital (100 parts) -> bloque.
    dossier_type = dossier_type_by_label("SELARL creation V1")
    membre = StatutsCivilsAssocie(
        type_personne="personne_physique",
        civilite_affichage="Madame",
        prenom="Lea",
        nom="Bernard",
        parts=StatutsCivilsParts(nb=30, nb_lettres="trente"),
    )
    data = _valid_selarl_input(
        PROFESSION_MEDECIN,
        dossier_unipersonnel=False,
        praticien_nb_parts=60,
        membres_additionnels=(membre,),
    )

    plan = build_clean_generation_plan(dossier_type, data)

    assert plan.can_generate is False
    assert any("somme des parts" in blocker for blocker in plan.blockers)


def test_clean_front_selarl_unipersonnel_unchanged_without_membres(tmp_path: Path) -> None:
    # Garde-fou retro-compat : sans membre additionnel, le dossier reste unipersonnel
    # et byte-equivalent au parcours historique (entete singulier).
    data = _valid_selarl_input(PROFESSION_MEDECIN)
    ctx = build_generation_context(data)
    assert ctx.dossier_options.associe_unique is True
    assert ctx.statuts_sel.membres == []

    result = generate_selarl_dossier(data, tmp_path)
    statuts = next(p for p in result.docx_paths if "Statuts" in p.name)
    text = _docx_text(statuts)
    assert "LE SOUSSIGNE" in text
    assert "LES SOUSSIGNÉS" not in text


def _valid_selarl_input(
    profession: str,
    *,
    regime_communautaire: bool = False,
    married_separation: bool = False,
    **overrides: object,
):
    dossier_type = dossier_type_by_label("SELARL creation V1")
    values = _valid_selarl_kwargs(
        profession,
        regime_communautaire=regime_communautaire,
        married_separation=married_separation,
    )
    values.update(overrides)
    return build_clean_data_entry(
        dossier_type,
        **values,
    )


def _valid_selarl_kwargs(
    profession: str,
    *,
    regime_communautaire: bool = False,
    married_separation: bool = False,
) -> dict[str, object]:
    is_married = regime_communautaire or married_separation
    return {
        "dossier_reference": "B-SELARL-001",
        "profession": profession,
        "dossier_unipersonnel": True,
        "regime_communautaire": regime_communautaire,
        "civilite": "Monsieur",
        "prenom": "Jean",
        "nom": "Martin",
        "date_naissance": date(1984, 4, 12),
        "ville_naissance": "Paris",
        "departement_naissance": "75",
        "nationalite": "française",
        "nom_pere": "Pierre Martin",
        "nom_mere": "Anne Martin",
        "adresse_num_voie": "10",
        "adresse_voie": "rue Test",
        "adresse_cp": "75001",
        "adresse_ville": "Paris",
        "situation_maritale": "marie" if is_married else "celibataire",
        "regime_matrimonial": (
            "regime de communaute"
            if regime_communautaire
            else "separation de biens"
            if married_separation
            else ""
        ),
        "numero_ordre": "ORD-123",
        "numero_rpps": "10000000001",
        "departement_ordre": "75",
        "denomination": "SELARL MARTIN",
        "capital_social": "1000",
        "nb_parts_total": 100,
        "valeur_nominale_part": "10",
        "siege_num_voie": "20",
        "siege_voie": "avenue du Siege",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "ville_rcs": "Paris",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "ordre_cp": "75008",
        "ordre_ville": "Paris",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 26),
        "decision_date": date(2026, 5, 26),
        "depot_banque_nom": "Banque Test",
        "depot_banque_adresse": "30 boulevard Banque, 75009 Paris",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 decembre",
        "exercice_cloture_premier": "31 decembre 2026",
        **(
            {
                "conjoint_civilite": "Madame",
                "conjoint_prenom": "Claire",
                "conjoint_nom": "Martin",
            }
            if profession == PROFESSION_DENTISTE or regime_communautaire or married_separation
            else {}
        ),
    }


def _docx_text(path: Path) -> str:
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(text for text in texts if text)


def _ascii_text(value: str) -> str:
    return normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")


def _fill_valid_streamlit_selarl_form(app: AppTest) -> None:
    values = {
        "selarl_dossier_reference": "B-SELARL-DOWNLOAD-TEST",
        "selarl_prenom": "Jean",
        "selarl_nom": "Martin",
        "selarl_date_naissance": "31/12/1974",
        "selarl_ville_naissance": "Paris",
        "selarl_departement_naissance": "75",
        "selarl_nom_pere": "Pierre Martin",
        "selarl_nom_mere": "Anne Martin",
        "selarl_numero_ordre": "ORD-123",
        "selarl_numero_rpps": "10000000001",
        # O24-03 : adresse perso + siege sur UNE ligne (le slice reparse num/voie/cp/ville).
        "selarl_adresse_ligne": "10 rue Test, 75001 Paris",
        "selarl_denomination": "SELARL MARTIN",
        "selarl_ville_rcs": "Paris",
        "selarl_siege_ligne": "20 avenue du Siege, 75002 Paris",
        "selarl_departement_ordre": "75",
        # O24-03 : adresse de l'ordre sur UNE ligne (le slice reparse ligne_1/cp/ville).
        "selarl_ordre_adresse_ligne": "1 rue de l'Ordre, 75008 Paris",
        "selarl_signature_lieu": "Paris",
        "selarl_signature_date": "27/05/2026",
        "selarl_depot_banque_nom": "Banque Test",
        "selarl_depot_banque_adresse": "30 boulevard Banque, 75009 Paris",
        "selarl_exercice_debut": "1er janvier",
        "selarl_exercice_fin": "31 decembre",
        "selarl_exercice_cloture_premier": "31 decembre 2026",
    }
    for key, value in values.items():
        app.text_input(key=key).set_value(value)
    app.checkbox(key="selarl_ville_naissance_article_au").set_value(False)
    app.number_input(key="selarl_capital_social").set_value(1000)
    app.number_input(key="selarl_nb_parts_total").set_value(100)


# ---------------------------------------------------------------------------
# Retours client 2026-06-11 — ticket SELARL dentiste unipersonnelle + cession
# ---------------------------------------------------------------------------


def test_regime_matrimonial_derivations_cover_new_options() -> None:
    from sydel_doc_engine.front_app.field_derivations import (
        MATRIMONIAL_STATUS_PRESETS,
        regime_communautaire_from_status,
        regime_matrimonial_from_status,
    )

    # Ticket 1.2 : les quatre regimes maries sont proposes.
    married = [item for item in MATRIMONIAL_STATUS_PRESETS if item.startswith("Marié")]
    assert len(married) == 4

    # Seul le regime legal / communaute declenche DOC-005/DOC-006.
    assert regime_communautaire_from_status("Marié(e) sous le régime légal / communauté")
    assert not regime_communautaire_from_status(
        "Marié(e) sous le régime de la séparation de biens"
    )
    assert not regime_communautaire_from_status(
        "Marié(e) sous le régime de la communauté universelle"
    )
    assert not regime_communautaire_from_status(
        "Marié(e) sous le régime de la participation aux acquêts"
    )
    assert not regime_communautaire_from_status("Célibataire")

    # Le regime injecte dans les statuts suit l'option choisie.
    assert (
        regime_matrimonial_from_status(
            "Marié(e) sous le régime de la communauté universelle", False
        )
        == "communaute universelle"
    )
    assert (
        regime_matrimonial_from_status(
            "Marié(e) sous le régime de la participation aux acquêts", False
        )
        == "participation aux acquets"
    )
    assert (
        regime_matrimonial_from_status(
            "Marié(e) sous le régime de la séparation de biens", False
        )
        == "separation de biens"
    )


def test_split_numero_voie_extracts_leading_number() -> None:
    from sydel_doc_engine.front_app.field_derivations import split_numero_voie

    assert split_numero_voie("10 rue Test") == ("10", "rue Test")
    assert split_numero_voie("10 bis avenue Foch") == ("10 bis", "avenue Foch")
    assert split_numero_voie("Lieu-dit Les Pins") == ("", "Lieu-dit Les Pins")
    assert split_numero_voie("") == ("", "")


def test_clean_front_dentiste_sans_salarie_facultatifs_vides_genere(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Cas EXACT du ticket client : dentiste unipersonnel + cession de fonds
    # liberal, AUCUN salarie, champs facultatifs vides (adresse banque, loyer,
    # CA/resultat, infos bancaires cession) -> la generation N'EST PAS bloquee,
    # la phrase salaries est supprimee, aucun token residuel.
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "dentiste-sans-salarie")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="selarl_profession").set_value("Chirurgien-dentiste")
    app = app.run(timeout=180)

    _fill_valid_streamlit_selarl_form(app)
    # Marie sous le regime de la separation de biens : aucune logique
    # documentaire supplementaire (pas de DOC-005/006), conjoint requis.
    app.selectbox(key="selarl_situation_maritale").set_value(
        "Marié(e) sous le régime de la séparation de biens"
    )
    app = app.run(timeout=180)
    app.text_input(key="selarl_conjoint_prenom").set_value("Claire")
    app.text_input(key="selarl_conjoint_nom").set_value("Martin")
    # Adresse banque NON renseignee (ticket 3.2).
    app.text_input(key="selarl_depot_banque_adresse").set_value("")
    app.checkbox(key="selarl_cession").set_value(True)
    app = app.run(timeout=180)

    # Cession : prix total uniquement ; tout le reste reste vide (3.1).
    app.text_input(key="selarl_cession_prix_total").set_value("250 000")
    app = app.run(timeout=180)

    assert app.checkbox(key="selarl_cession_salaries_aucun").value is True
    assert not any("Blocage" in item.value for item in app.caption)
    assert app.button(key="clean_generate_dossier").disabled is False

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=180)

    assert [e.value for e in app.error] == []
    download_labels = [item.label for item in app.get("download_button")]
    assert "Telecharger acte_cession_cabinet_dentaire.docx" in download_labels
    assert "Telecharger appel_fond_sel.docx" in download_labels
    # Pas de documents du regime de la communaute (ticket 1.2).
    assert "Telecharger lettre_renonciation_associe.docx" not in download_labels
    assert "Telecharger lettre_avertissement_conjoint.docx" not in download_labels

    generated = app.session_state["clean_generated_dossier"]
    combined = "\n".join(_docx_text(Path(path)) for path in generated["docx_paths"])
    assert "[" not in combined
    assert "]" not in combined
    acte = next(
        Path(path)
        for path in generated["docx_paths"]
        if "cession_cabinet_dentaire" in Path(path).name
    )
    acte_text = _docx_text(acte)
    # Aucun salarie -> phrase supprimee, pas de « Neant » (tickets 2.12 / 3.3).
    assert "contrats de travail de" not in acte_text
    assert "Néant" not in acte_text
    # Vendeur = associe unique repris automatiquement (ticket 2.1).
    assert "Jean" in acte_text
    assert "Martin" in acte_text


def _scm_acte_text_for_situation(
    tmp_path: Path,
    monkeypatch,
    *,
    situation_preset: str,
    subdir: str,
) -> str:
    """Bout-en-bout O24-11 : SELARL dentiste + SCM via le SLICE reel (AppTest).

    Clique le bouton de donnees de test (prefill cession + SCM + `selarl_scm`),
    surcharge la situation matrimoniale + le conjoint de facon deterministe, genere,
    et renvoie le texte de l'acte de cession de parts SCM.
    """
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / subdir)
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="selarl_profession").set_value("Chirurgien-dentiste")
    app = app.run(timeout=180)

    app.button(key="clean_generate_test_data").click()
    app = app.run(timeout=180)

    # Surcharge deterministe : civilite (genre fixe), situation matrimoniale ciblee
    # + conjoint nomme. La personne de test est aleatoire (genre M/F), on la fige.
    app.selectbox(key="selarl_civilite").set_value("Monsieur")
    app.selectbox(key="selarl_situation_maritale").set_value(situation_preset)
    app = app.run(timeout=180)
    app.text_input(key="selarl_conjoint_prenom").set_value("Claire")
    app.text_input(key="selarl_conjoint_nom").set_value("Dupont")
    app = app.run(timeout=180)

    assert app.checkbox(key="selarl_scm").value is True
    assert app.button(key="clean_generate_dossier").disabled is False

    app.button(key="clean_generate_dossier").click()
    app = app.run(timeout=180)
    assert [e.value for e in app.error] == []

    generated = app.session_state["clean_generated_dossier"]
    acte = next(
        Path(path)
        for path in generated["docx_paths"]
        if "acte_cession_parts_scm" in Path(path).name
    )
    return _docx_text(acte)


def test_scm_cession_acte_keeps_regime_matrimonial_separation(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # O24-11 (BLOQUANT re-Akainu T4) : l'acte de cession de parts SCM doit conserver
    # le REGIME matrimonial accentue (pas le « marie » collapse, non accentue). Le
    # modele n'a qu'un placeholder [situation_maritale_cedant] : le slice y injecte le
    # libelle complet « marié sous le régime de séparation de biens avec <conjoint> ».
    acte_text = _scm_acte_text_for_situation(
        tmp_path,
        monkeypatch,
        situation_preset="Marié(e) sous le régime de la séparation de biens",
        subdir="scm-o2411-separation",
    )
    # Statut accentue/genre (« marié » ou « mariée » selon la personne de test aleatoire)
    # + regime accentue + conjoint, dans une seule clause.
    assert (
        "marié sous le régime de séparation de biens avec Madame Claire Dupont" in acte_text
        or "mariée sous le régime de séparation de biens avec Madame Claire Dupont"
        in acte_text
    )
    # Pas de fuite de la valeur collapsee brute / non accentuee.
    assert "marie sous le" not in acte_text


def test_scm_cession_acte_keeps_regime_matrimonial_communaute_universelle(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # O24-11 : second regime (communaute universelle) pour prouver que ce n'est pas
    # un cas unique cable en dur.
    acte_text = _scm_acte_text_for_situation(
        tmp_path,
        monkeypatch,
        situation_preset="Marié(e) sous le régime de la communauté universelle",
        subdir="scm-o2411-universelle",
    )
    assert (
        "marié sous le régime de communauté universelle avec Madame Claire Dupont"
        in acte_text
        or "mariée sous le régime de communauté universelle avec Madame Claire Dupont"
        in acte_text
    )
    assert "marie sous le" not in acte_text


def test_scm_cedant_situation_display_covers_both_paths(monkeypatch) -> None:
    # O24-11 : le helper qui reconstruit le libelle complet est partage par les DEUX
    # chemins (SELARL prefix=selarl ; SELAS multi prefix=selas). On verrouille qu'il lit
    # bien la cle de session `{prefix}_situation_maritale` (libelle BRUT) + le genre +
    # la valeur collapsee, pour les 3 regimes + un non-marie.
    from sydel_doc_engine.domain.enums import Gender
    from sydel_doc_engine.front_app import shell

    class _FakeSt:
        def __init__(self) -> None:
            self.session_state: dict[str, object] = {}

    fake = _FakeSt()
    monkeypatch.setattr(shell, "st", fake)

    cases = [
        ("selarl", "Marié(e) sous le régime de la séparation de biens", Gender.MASCULIN,
         "marié sous le régime de séparation de biens"),
        ("selas", "Marié(e) sous le régime de la communauté universelle", Gender.FEMININ,
         "mariée sous le régime de communauté universelle"),
        ("selarl", "Marié(e) sous le régime légal / communauté", Gender.MASCULIN,
         "marié sous le régime de communauté réduite aux acquêts"),
    ]
    for prefix, brut, genre, expected in cases:
        fake.session_state = {f"{prefix}_situation_maritale": brut}
        praticien = {"situation_maritale": "marie", "genre": genre}
        assert (
            shell._scm_cedant_situation_maritale_display(praticien, prefix=prefix)
            == expected
        )

    # Non marie : juste le statut accentue, jamais « sous le régime de » (et jamais vide).
    fake.session_state = {"selas_situation_maritale": "Divorcé(e)"}
    praticien = {"situation_maritale": "divorce", "genre": Gender.FEMININ}
    assert (
        shell._scm_cedant_situation_maritale_display(praticien, prefix="selas")
        == "divorcée"
    )


def test_scm_cession_acte_keeps_regime_matrimonial_communaute_legale(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # O24-11 : 3e regime (communaute legale -> « communauté réduite aux acquêts »).
    acte_text = _scm_acte_text_for_situation(
        tmp_path,
        monkeypatch,
        situation_preset="Marié(e) sous le régime légal / communauté",
        subdir="scm-o2411-legale",
    )
    assert (
        "marié sous le régime de communauté réduite aux acquêts avec Madame Claire Dupont"
        in acte_text
        or "mariée sous le régime de communauté réduite aux acquêts avec Madame Claire Dupont"
        in acte_text
    )
    assert "marie sous le" not in acte_text


def test_clean_front_signature_lieu_seeded_from_siege_ville(
    tmp_path: Path,
    monkeypatch,
) -> None:
    # Ticket 1.6 : le lieu de signature est preremplie avec la ville du siege
    # social et reste modifiable.
    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "lieu-signature")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    # O24-03 : siege sur UNE ligne ; la ville (« Lyon ») est parsee de la ligne unique.
    app.text_input(key="selarl_siege_ligne").set_value("20 avenue du Siege, 69002 Lyon")
    app = app.run(timeout=120)

    assert app.text_input(key="selarl_signature_lieu").value == "Lyon"

    app.text_input(key="selarl_signature_lieu").set_value("Paris")
    app = app.run(timeout=120)
    assert app.text_input(key="selarl_signature_lieu").value == "Paris"


def test_clean_front_statuts_render_new_marriage_regimes(tmp_path: Path) -> None:
    # Ticket 1.2 : communaute universelle et participation aux acquets sont
    # proposes, rendus correctement dans les statuts, SANS DOC-005/DOC-006.
    for slug, regime, expected in (
        ("universelle", "communaute universelle", "la communauté universelle"),
        ("acquets", "participation aux acquets", "la participation aux acquêts"),
    ):
        data_entry = _valid_selarl_input(
            PROFESSION_MEDECIN,
            married_separation=True,
            regime_matrimonial=regime,
        )
        dossier_type = dossier_type_by_label("SELARL creation V1")
        plan = build_clean_generation_plan(dossier_type, data_entry)
        assert plan.can_generate is True
        assert "DOC-005" not in plan.document_codes
        assert "DOC-006" not in plan.document_codes

        generated = generate_selarl_dossier(data_entry, tmp_path / slug)
        statuts_path = next(
            path for path in generated.docx_paths if path.name.lower().startswith("statuts")
        )
        statuts_text = _docx_text(statuts_path)
        assert (
            f"marié sous le régime de {expected} avec Madame Claire Martin"
            in statuts_text
        )


def test_clean_front_banque_adresse_vide_ne_bloque_pas(tmp_path: Path) -> None:
    # Ticket 3.2 : adresse de banque non renseignee -> generation NON bloquee,
    # zone vide a completer dans les statuts medecin.
    data_entry = _valid_selarl_input(PROFESSION_MEDECIN, depot_banque_adresse="")
    dossier_type = dossier_type_by_label("SELARL creation V1")

    plan = build_clean_generation_plan(dossier_type, data_entry)
    assert plan.can_generate is True
    assert not any("banque" in blocker.casefold() for blocker in plan.blockers)

    generated = generate_selarl_dossier(data_entry, tmp_path / "sans-adresse-banque")
    statuts_path = next(
        path for path in generated.docx_paths if path.name.lower().startswith("statuts")
    )
    statuts_text = _docx_text(statuts_path)
    assert "[" not in statuts_text
    assert "]" not in statuts_text


def test_clean_front_demande_ordre_presidente_and_conseiller(tmp_path: Path) -> None:
    # Retour Albane 2026-06-10 (demande d'inscription a l'ordre) :
    # - « Madame la Présidente » si la présidente de l'ordre est une femme ;
    # - nom du conseiller (mandataire) adaptable au lieu de Jordan ELBAZ en dur ;
    # - CP + ville de l'adresse perso sur une ligne séparée.
    from dataclasses import replace

    base = build_selarl_scenario("selarl_medecin_simple")

    default_gen = generate_selarl_dossier(base, tmp_path / "defaut")
    default_doc = next(
        p for p in default_gen.docx_paths if "demande_inscription_ordre" in p.name
    )
    default_text = _docx_text(default_doc)
    assert "Monsieur le Président" in default_text
    assert "Jordan ELBAZ" in default_text

    custom = replace(
        base,
        ordre_president_feminin=True,
        mandataire_prenom="Sophie",
        mandataire_nom="MARTIN",
    )
    custom_gen = generate_selarl_dossier(custom, tmp_path / "custom")
    custom_doc = next(
        p for p in custom_gen.docx_paths if "demande_inscription_ordre" in p.name
    )
    paragraphs = [p.text for p in Document(custom_doc).paragraphs]
    custom_text = "\n".join(paragraphs)
    assert "Madame la Présidente" in custom_text
    assert "Monsieur le Président" not in custom_text
    assert "Sophie MARTIN" in custom_text
    assert "Jordan ELBAZ" not in custom_text
    # Adresse perso sur deux lignes (rue / CP ville).
    assert "10 rue Test" in paragraphs
    assert "75001 Paris" in paragraphs


def test_clean_front_selarl_pv_decision_reunion_signature_divergent_b1(tmp_path: Path) -> None:
    # B1/SCS2/SU3 (Albane 2026-06-25, lock Akainu) : SELARL via le CHEMIN REEL
    # (build_clean_data_entry -> generate_selarl_dossier). Ville signature != siege + date decision
    # d'une ANNEE differente : le PV DOIT forcer signature=siege et decision=reunion=signature (pas
    # de contradiction interne). Sans le fix data_entry.py:115, reunion_date_lettres tirait de
    # decision_date -> regression non protegee par le lock multi_type (qui ne couvre pas SELARL).
    from docx import Document

    dossier_type = dossier_type_by_label("SELARL creation V1")
    kwargs = dict(_valid_selarl_kwargs(PROFESSION_MEDECIN))
    kwargs.update(
        {
            "siege_ville": "Lyon",
            "signature_lieu": "VilleSignatureDivergente",
            "signature_date": date(2026, 5, 26),
            "decision_date": date(2024, 3, 7),
        }
    )
    data = build_clean_data_entry(dossier_type, **kwargs)
    result = generate_selarl_dossier(data, tmp_path)
    for path in result.docx_paths:
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts += [p.text for p in cell.paragraphs]
        text = "\n".join(parts)
        assert "VilleSignatureDivergente" not in text, f"{path.name} ville signature (SU3)"
        assert "sept mars" not in text, f"{path.name} date decision (B1)"
        assert "07/03/2024" not in text, f"{path.name} date decision chiffres"
        assert "deux mille vingt-quatre" not in text, f"{path.name} annee (B1)"
