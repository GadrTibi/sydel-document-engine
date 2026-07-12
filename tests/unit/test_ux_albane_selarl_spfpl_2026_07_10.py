"""Retours Albane 2026-07-10 — volet FRONT + orchestration/bundle (SELARL + SPFPL cession).

Couvre :
  O1  — choix « Cabinet cree / achete » ouvert au medical ET au dentaire ; le mode pilote
        le libelle de la date et la visibilite du precedent proprietaire + prix.
  O2  — elements corporels derives par defaut = prix total - elements incorporels.
  F2  — la reference dossier ne bloque plus l'edition.
  F3  — la banque du depot des fonds ne bloque plus l'edition.
  F1  — les champs SCM ne sont plus pre-remplis ; vides -> bloques avec un message clair.
  MD1 — la cession SELARL produit l'acte ET le compromis ensemble.
  Procuration SEL — la cession SPFPL genere une 2e procuration au nom de la societe cible.
"""

from __future__ import annotations

from test_multi_type_front import _spfpl_payload

from sydel_doc_engine.domain.models import ScmCessionContext
from sydel_doc_engine.front_app.selarl_slice import (
    generate_selarl_dossier,
    validate_selarl_input,
)
from sydel_doc_engine.front_app.spfpl_slice import generate_dossier as spfpl_generate_dossier
from sydel_doc_engine.scenarios.selarl import build_selarl_scenario


# --------------------------------------------------------------------------- O1
def test_o1_origine_choice_present_and_mode_drives_labels(tmp_path, monkeypatch) -> None:
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-o1")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    app.checkbox(key="selarl_cession").set_value(True)
    app = app.run(timeout=180)

    origine = app.selectbox(key="selarl_cession_cabinet_origine_mode")
    assert set(origine.options) == {
        "Cabinet cree par le vendeur",
        "Cabinet achete par le vendeur",
    }

    # Le champ date est un text_input labellise (date_input_with_today).
    def text_labels() -> list[str]:
        return [str(w.label) for w in app.text_input]

    # Mode « cree » : libelle « creation », pas de precedent proprietaire ni de prix d'origine.
    origine.set_value("Cabinet cree par le vendeur")
    app = app.run(timeout=180)
    assert any("Date de creation du cabinet" in lbl for lbl in text_labels())
    assert not any("precedent proprietaire" in lbl.lower() for lbl in text_labels())

    # Mode « achete » : libelle « acquisition » + precedent proprietaire + prix d'origine.
    app.selectbox(key="selarl_cession_cabinet_origine_mode").set_value(
        "Cabinet achete par le vendeur"
    )
    app = app.run(timeout=180)
    assert any("Date d'acquisition du cabinet" in lbl for lbl in text_labels())
    assert any("precedent proprietaire" in lbl.lower() for lbl in text_labels())
    assert any("Prix d'origine" in lbl for lbl in text_labels())


def test_o1_dentaire_acte_supports_created_cabinet(tmp_path) -> None:
    """La clause « cree » est bien rendue par le modele DENTAIRE (raison d'ouvrir le choix
    aux deux professions)."""
    from docx import Document

    data = build_selarl_scenario("selarl_dentiste_cession_cabinet_dentaire")
    cabinet = data.cession_context.cabinet.model_copy(
        update={
            "origine_propriete_mode": "cree",
            "date_origine_propriete": "3 janvier 2015",
            "prix_origine_propriete": None,
            "precedent_proprietaire": None,
        }
    )
    cession = data.cession_context.model_copy(update={"cabinet": cabinet})
    import dataclasses

    data = dataclasses.replace(data, cession_context=cession)
    result = generate_selarl_dossier(data, tmp_path / "dentaire-cree")
    acte = next(p for p in result.docx_paths if "acte_cession_cabinet_dentaire" in p.name)
    text = "\n".join(par.text for par in Document(str(acte)).paragraphs)
    assert "régulièrement créés le 3 janvier 2015" in text
    assert "acquis auprès de" not in text


# --------------------------------------------------------------------------- O2
def test_o2_corporel_derived_from_total_minus_incorporel(tmp_path, monkeypatch) -> None:
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-o2")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    app.checkbox(key="selarl_cession").set_value(True)
    app = app.run(timeout=180)

    def set_text(key: str, value: str) -> None:
        for widget in app.text_input:
            if str(widget.key) == key:
                widget.set_value(value)
                return
        raise KeyError(key)

    set_text("selarl_cession_prix_total", "300 000")
    set_text("selarl_cession_prix_incorporels", "250 000")
    app = app.run(timeout=180)

    corporels = next(
        w for w in app.text_input
        if str(w.key) == "selarl_cession_prix_corporels"
    )
    assert corporels.value.replace(" ", " ").replace("\xa0", " ") == "50 000"


# ----------------------------------------------------------------------- F2 / F3
def test_f2_f3_reference_and_banque_not_blocking() -> None:
    data = build_selarl_scenario("selarl_medecin_simple")
    import dataclasses

    data = dataclasses.replace(data, dossier_reference="", depot_banque_nom="")
    blockers = validate_selarl_input(data)
    assert not any("Reference dossier" in b for b in blockers)
    assert not any("Banque du depot" in b for b in blockers)


# --------------------------------------------------------------------------- F1
def test_f1_blank_scm_fields_block_with_clear_messages() -> None:
    data = build_selarl_scenario("selarl_medecin_simple")
    import dataclasses

    blank_scm = ScmCessionContext.model_validate(
        {"scm_cedee": {}, "parts_cedees": {}, "prix": {}}
    )
    data = dataclasses.replace(data, scm=True, scm_cession_context=blank_scm)
    blockers = validate_selarl_input(data)
    for needle in (
        "denomination de la SCM cedee",
        "siege de la SCM cedee",
        "capital social de la SCM cedee",
        "nombre de parts cedees",
        "prix global",
    ):
        assert any(needle in b for b in blockers), needle


def test_f1_fixture_scm_scenario_still_generable() -> None:
    """Le jeu FIXTURE (donnees completes) passe les nouveaux blocages."""
    data = build_selarl_scenario("selarl_dentiste_cession_scm")
    assert not any("Cession de parts SCM" in b for b in validate_selarl_input(data))


# -------------------------------------------------------------------------- MD1
def test_md1_selarl_cession_bundle_has_acte_and_compromis(tmp_path) -> None:
    for scenario, acte_name, compromis_name in (
        (
            "selarl_medecin_cession_cabinet_medical",
            "acte_cession_cabinet_medical.docx",
            "compromis_cession_cabinet_medical.docx",
        ),
        (
            "selarl_dentiste_cession_cabinet_dentaire",
            "acte_cession_cabinet_dentaire.docx",
            "compromis_cession_cabinet_dentaire.docx",
        ),
    ):
        data = build_selarl_scenario(scenario)
        result = generate_selarl_dossier(data, tmp_path / scenario)
        names = {p.name for p in result.docx_paths}
        assert acte_name in names
        assert compromis_name in names


# ---------------------------------------------------------------- Procuration SEL
def test_spfpl_cession_generates_procuration_for_sel_target(tmp_path) -> None:
    from docx import Document

    payload = _spfpl_payload("SPFPL cession")
    result = spfpl_generate_dossier(payload, tmp_path / "spfpl-cession")
    names = {p.name for p in result.docx_paths}
    assert "procuration.docx" in names
    assert "procuration_SEL.docx" in names

    def company_line(name: str) -> str:
        doc = Document(str(tmp_path / "spfpl-cession" / name))
        return " ".join(par.text for par in doc.paragraphs)

    # La 2e procuration nomme la SOCIETE CIBLE (SEL), distincte de la SPFPL.
    assert "SPFPL MARTIN" in company_line("procuration.docx")
    sel_text = company_line("procuration_SEL.docx")
    assert "SPFPL MARTIN" not in sel_text
    assert "MARTIN" in sel_text  # denomination de la SEL cible (SELARL/SELAS CABINET MARTIN)


def test_spfpl_apport_does_not_generate_sel_procuration(tmp_path) -> None:
    """La 2e procuration SEL est propre a la CESSION ; l'apport ne la produit pas."""
    payload = _spfpl_payload("SPFPL apport")
    result = spfpl_generate_dossier(payload, tmp_path / "spfpl-apport")
    names = {p.name for p in result.docx_paths}
    assert "procuration.docx" in names
    assert "procuration_SEL.docx" not in names
