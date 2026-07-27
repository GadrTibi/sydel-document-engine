from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from _accents import assert_no_unaccented_french
from docx import Document

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    CapitalSouscripteur,
    CapitalSouscription,
    CessionBanque,
    DepotFonds,
    DocumentGenerationContext,
    DossierOptions,
    Person,
    Signature,
    SocieteSpfpl,
)
from sydel_doc_engine.generators.lot_05.attestation_capital_souscripteurs_selas import (
    AttestationCapitalSouscripteursSelasGenerator,
)


def _base_context() -> DocumentGenerationContext:
    return DocumentGenerationContext(
        structure="SELAS",
        dossier_options=DossierOptions(),
        personne_signataire=Person(
            genre=Gender.MASCULIN,
            civilite="Monsieur",
            prenom="Alain",
            nom="Fedorowsky",
        ),
        signature=Signature(lieu="Rennes", date=date(2026, 6, 15)),
        societe_spfpl=SocieteSpfpl(
            denomination="SELAS DES BOIS",
            forme_sociale="Société d'exercice libérale par actions simplifiée",
            capital_social="1.000",
            nb_actions_total=1000,
            valeur_nominale_action="1",
            profession="Médecins",
            siege=Address(
                num_voie="31B",
                voie="Boulevard de Sévigné",
                cp="35700",
                ville="Rennes",
                adresse_affichee="31B Boulevard de Sévigné, 35700 Rennes",
            ),
        ),
        depot_fonds=DepotFonds(
            banque=CessionBanque(
                nom="Crédit Agricole",
                adresse_affichee="2 rue de la Monnaie, 35000 Rennes",
            )
        ),
        capital_souscription=CapitalSouscription(
            nb_actions_total=1000,
            valeur_nominale_action="1",
            apports_numeraire_montant="1.000",
            president=CapitalSouscripteur(
                civilite_affichage="Docteur",
                genre=Gender.MASCULIN,
                prenom="Alain",
                nom="Fedorowsky",
            ),
            # R3 durci (Rafael 2026-07-07) : le titre « Docteur » pose en civilite est
            # rendu CIVIL et accorde au genre du souscripteur (Claire = FEMININ pour
            # verrouiller « à Madame »).
            souscripteurs=[
                CapitalSouscripteur(
                    civilite_affichage="Docteur",
                    genre=Gender.MASCULIN,
                    prenom="Alain",
                    nom="Fedorowsky",
                    nb_actions=510,
                ),
                CapitalSouscripteur(
                    civilite_affichage="Docteur",
                    genre=Gender.FEMININ,
                    prenom="Claire",
                    nom="Martin",
                    nb_actions=490,
                ),
            ],
        ),
    )


def _docx_text(path: Path) -> str:
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(text for text in texts if text)


def _assert_clean(text: str) -> None:
    assert "[" not in text
    assert "]" not in text


def test_attestation_selas_generates_multi_subscriber_wording(tmp_path: Path) -> None:
    output_path = AttestationCapitalSouscripteursSelasGenerator().generate(
        _base_context(),
        tmp_path,
    )

    text = _docx_text(output_path)

    assert output_path.name == "attestation_capital_souscripteurs_selas.docx"
    assert "SELAS DES BOIS" in text
    assert "Société d'exercice libérale par Actions simplifiées de Médecins" in text
    assert "Au capital de 1.000 euros" in text
    assert "En cours d'immatriculation" in text
    assert "Liste des souscripteurs" in text
    assert (
        "Je soussigné Monsieur Alain Fedorowsky, Président de la Société SELAS DES BOIS, "
        "atteste que le capital de ladite société est réparti de la manière suivante :"
    ) in text
    assert "Capital social : 1.000 € en numéraire" in text
    assert "Nombre d'actions: 1000 actions d'un montant de 1 euro chacune" in text
    # R3 durci (Rafael 2026-07-07, siloing pluripersonnel) : repartition + apport
    # rendent la civilite CIVILE accordee au genre du souscripteur — supersede les
    # anciens verrous « au Dr X » / « Le Docteur X a fait un apport » (A26-45/49).
    assert "510 actions attribuées à Monsieur Alain Fedorowsky," in text
    assert "490 actions attribuées à Madame Claire Martin," in text
    # KAN-46 (Rafael) : l'ADRESSE de la banque est reportee, pas juste le nom.
    assert (
        "Capital social de 1.000 € entièrement libéré et déposé dans les livres de la "
        "banque Crédit Agricole, 2 rue de la Monnaie, 35000 Rennes"
    ) in text
    assert "Monsieur Alain Fedorowsky a fait un apport de 510 euros en numéraire." in text
    assert "Madame Claire Martin a fait un apport de 490 euros en numéraire." in text
    # R3 (Albane 2026-07-07) : « Docteur » n'est pas une civilité — le slot
    # « par le Président, __ » rend la civilité CIVILE (accord au genre du
    # president, repli signataire).
    # KAN-46 (Rafael) : point final apres le nom dans la derniere phrase.
    assert (
        "certifié exact, sincère et véritable par le Président, Monsieur Alain Fedorowsky."
    ) in text
    assert "Docteur" not in text
    assert "au Dr " not in text
    assert "Fait à Rennes" in text
    assert "Le 15/06/2026" in text
    _assert_clean(text)
    assert_no_unaccented_french(text)


def test_attestation_selas_three_subscribers(tmp_path: Path) -> None:
    # R3 durci (Rafael 2026-07-07) : civilite CIVILE partout ; Hugo SANS genre
    # (appelant legacy) -> masculin par defaut, comme derive_gender_from_civilite.
    ctx = _base_context()
    ctx.capital_souscription.souscripteurs = [
        CapitalSouscripteur(civilite_affichage="Docteur", genre=Gender.MASCULIN,
                            prenom="Alain", nom="Fedorowsky", nb_actions=400),
        CapitalSouscripteur(civilite_affichage="Docteur", genre=Gender.FEMININ,
                            prenom="Claire", nom="Martin", nb_actions=400),
        CapitalSouscripteur(civilite_affichage="Docteur", prenom="Hugo", nom="Bernard",
                            nb_actions=200),
    ]

    output_path = AttestationCapitalSouscripteursSelasGenerator().generate(ctx, tmp_path)
    text = _docx_text(output_path)

    assert "400 actions attribuées à Monsieur Alain Fedorowsky," in text
    assert "400 actions attribuées à Madame Claire Martin," in text
    assert "200 actions attribuées à Monsieur Hugo Bernard," in text
    assert "Monsieur Hugo Bernard a fait un apport de 200 euros en numéraire." in text
    assert "Docteur" not in text
    _assert_clean(text)


def test_attestation_selas_rejects_action_total_mismatch(tmp_path: Path) -> None:
    ctx = _base_context()
    ctx.capital_souscription.souscripteurs[1].nb_actions = 100  # 510 + 100 != 1000

    with pytest.raises(ValueError, match="répartition des souscripteurs"):
        AttestationCapitalSouscripteursSelasGenerator().generate(ctx, tmp_path)


def test_attestation_selas_rejects_missing_bank(tmp_path: Path) -> None:
    ctx = _base_context()
    ctx.depot_fonds = None

    with pytest.raises(ValueError, match="depot_fonds.banque"):
        AttestationCapitalSouscripteursSelasGenerator().generate(ctx, tmp_path)


def test_attestation_selas_rejects_wrong_structure(tmp_path: Path) -> None:
    ctx = _base_context()
    ctx.structure = "SAS"

    with pytest.raises(ValueError, match="SELAS"):
        AttestationCapitalSouscripteursSelasGenerator().generate(ctx, tmp_path)
