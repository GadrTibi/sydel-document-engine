"""Tests statuts SASU Holding (nouveau type, modele officiel Albane 2026-06-29).

SASU Holding = SAS unipersonnelle, holding patrimoniale GENERALISTE (objet participations,
pas de profession reglementee, 20 articles, president + vice-presidents). DISTINCT de la
« SAS / SPFPL medecins » du moteur (conservee). Generateur = token-replacement sur le modele
officiel. Ces tests verifient : tous les tokens remplis (zero residuel), le corps legal
present verbatim, la comparution + l'accord de genre de l'associe unique.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from docx import Document

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    Company,
    DocumentGenerationContext,
    ExerciceSocial,
    Person,
    Signature,
    StatutsSasuHoldingContext,
)
from sydel_doc_engine.generators.lot_04.statuts_sasu_holding import (
    StatutsSasuHoldingGenerator,
)


def _docx_text(path: Path) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


def _ctx(*, genre: Gender = Gender.MASCULIN) -> DocumentGenerationContext:
    civilite = "Madame" if genre == Gender.FEMININ else "Monsieur"
    return DocumentGenerationContext(
        structure="SASU_HOLDING",
        personne_signataire=Person(
            genre=genre,
            civilite=civilite,
            prenom="Malo",
            nom="LE GUEN",
            date_naissance=date(1990, 1, 1),
            ville_naissance="Lorient",
            nationalite="française",
            adresse_perso=Address(
                num_voie="5", voie="Allée de la Clarté", cp="56700", ville="KERVIGNAC"
            ),
        ),
        signature=Signature(lieu="KERVIGNAC", date=date(2025, 11, 13)),
        societe=Company(
            denomination="MLG",
            forme_sociale="Société par actions simplifiée unipersonnelle",
            capital="1000",
            siege=Address(num_voie="5", voie="Allée de la Clarté", cp="56700", ville="KERVIGNAC"),
        ),
        exercice_social=ExerciceSocial(
            debut="1er janvier",
            fin="31 décembre",
            date_cloture_premier_exercice="le 31 décembre 2026",
        ),
        statuts_sasu_holding=StatutsSasuHoldingContext(
            forme_sociale="Société par actions simplifiée unipersonnelle",
            capital_social="1000",
            capital_social_lettres="mille",
            nb_actions=10000,
            nom_banque="HSBC",
            qualite_associe="Associé unique et Président",
            fonction_dirigeant="Président",
        ),
    )


# Lignes DISTINCTIVES du modele SASU Holding generaliste (verbatim attendu) — ABSENTES de la
# SPFPL medecins (preuve que c'est bien le bon acte, pas le clone medecin).
_VERBATIM_ALBANE = [
    "La société est une société par Actions Simplifiée régie par les lois et règlements",
    "La participation de la société, par tous moyens, à toutes entreprises ou sociétés",
    "Article 11 : Sortie des actionnaires",
    "Article 13 : Direction de la société",
    "Sans aucun vote, le Président suppléant sera par défaut, le conjoint séparé de corps",
    "Article 20 : Contestations",
]


def test_sasu_holding_fidelity_legal_body(tmp_path: Path) -> None:
    text = _docx_text(StatutsSasuHoldingGenerator().generate(_ctx(), tmp_path))
    for ligne in _VERBATIM_ALBANE:
        assert ligne in text, f"ligne Albane absente : {ligne!r}"
    # Aucune mention SPFPL / medecin / Ordre (preuve : holding generaliste, pas SPFPL medecins).
    assert "Participations Financières" not in text
    assert "Ordre des Médecins" not in text
    # Aucun token source residuel.
    assert "[" not in text and "]" not in text


def test_sasu_holding_tokens_filled(tmp_path: Path) -> None:
    text = _docx_text(StatutsSasuHoldingGenerator().generate(_ctx(), tmp_path))
    assert "MLG" in text  # denomination
    assert "Société par actions simplifiée unipersonnelle" in text  # forme
    assert "Monsieur Malo LE GUEN" in text  # associe unique
    assert "10000 actions" in text  # nb actions
    assert "mille (1000)" in text  # capital lettres + chiffre
    assert "Fait à KERVIGNAC, le 13 novembre 2025" in text  # signature
    assert "Associé unique et Président" in text  # qualite
    # Comparution reconstruite (proxy P022 casse) : identite complete propre.
    assert (
        # Rafael/Albane 2026-07-09 : « demeurant [adresse] » -> « demeurant au [adresse] ».
        "né le 1er janvier 1990 à Lorient, de nationalité française, demeurant au "
        "5 Allée de la Clarté" in text
    )


def test_sasu_holding_feminin_accord(tmp_path: Path) -> None:
    # Associee unique feminin -> « La soussignée », « née le » (le SASU Holding n'impose PAS
    # le masculin, contrairement a la SPFPL medecins du moteur).
    text = _docx_text(StatutsSasuHoldingGenerator().generate(_ctx(genre=Gender.FEMININ), tmp_path))
    assert "La soussignée" in text
    assert "née le 1er janvier 1990" in text
    assert "Le soussigné :" not in text


def test_sasu_holding_requires_structure(tmp_path: Path) -> None:
    ctx = _ctx()
    ctx.structure = "SAS"
    with pytest.raises(ValueError, match="SASU_HOLDING"):
        StatutsSasuHoldingGenerator().generate(ctx, tmp_path)


def test_sasu_holding_sans_surlignage_et_annexe_nouvelle_page(tmp_path: Path) -> None:
    # Rafael 2026-07-09 : (1) AUCUN surlignage dans le document genere (le modele
    # source Albane porte un run surligne d'edition « MLG » que le token-replacement
    # preservait) ; (2) l'annexe demarre TOUJOURS en debut de nouvelle page
    # (page_break_before sur le titre « ANNEXE »).
    from docx.oxml.ns import qn

    document = Document(StatutsSasuHoldingGenerator().generate(_ctx(), tmp_path))

    surlignes = [
        run.text
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run._element.find(qn("w:rPr")) is not None
        and run._element.find(qn("w:rPr")).find(qn("w:highlight")) is not None
    ]
    assert not surlignes, f"runs surlignes residuels : {surlignes!r}"

    annexe = next(
        p for p in document.paragraphs if p.text.strip().upper().startswith("ANNEXE")
    )
    assert annexe.paragraph_format.page_break_before is True
