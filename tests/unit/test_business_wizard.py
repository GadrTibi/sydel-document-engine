from __future__ import annotations

import unicodedata
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from sydel_doc_engine.app import business_wizard
from sydel_doc_engine.app.business_wizard import (
    STATUS_CONTEXT_INCOMPLETE,
    STATUS_GENERABLE,
    STATUS_MANUAL_ONLY,
    STATUS_NOT_IMPLEMENTED,
    BusinessWizardInput,
    business_document_table_rows,
    evaluate_business_wizard,
    sample_business_wizard_input,
    selarl_ui_address_labels,
    selarl_ui_block_visibility,
    selarl_ui_condition_specs,
    selarl_ui_document_specs,
    selarl_ui_reuse_rules,
)
from sydel_doc_engine.app.ui_runtime import (
    generate_docx_files_for_document_codes,
    generate_zip_file,
)
from sydel_doc_engine.orchestrator.service import DocumentOrchestrator
from sydel_doc_engine.registry.catalog import build_seed_catalog


def test_business_wizard_uses_case_catalog(monkeypatch) -> None:
    calls: list[object] = []

    def fake_get_expected_documents(case_input: object) -> list[object]:
        calls.append(case_input)
        return []

    monkeypatch.setattr(
        business_wizard,
        "get_expected_documents",
        fake_get_expected_documents,
    )

    evaluate_business_wizard(
        BusinessWizardInput(structure="SCI", sci_iris=False, option_is=False)
    )

    assert calls


def test_business_wizard_builds_minimal_valid_context() -> None:
    validation = evaluate_business_wizard(sample_business_wizard_input())

    assert validation.context is not None
    assert validation.can_generate_docx is True
    assert validation.missing_fields == ()
    assert validation.inconsistencies == ()
    assert validation.context.structure == "SCI"
    assert validation.generatable_document_codes == (
        "DOC-001",
        "DOC-003",
        "DOC-002",
        "DOC-004",
    )


def test_business_wizard_validation_detects_missing_fields() -> None:
    data = BusinessWizardInput(structure="SCI")

    validation = evaluate_business_wizard(data)

    assert validation.can_generate_docx is False
    assert "conditions.sci_iris" in validation.missing_fields
    assert "conditions.option_is" in validation.missing_fields
    assert "personne_signataire.civilite" in validation.missing_fields
    assert "societe.denomination" in validation.missing_fields
    assert "signature.date" in validation.missing_fields
    assert any(row.status == "blocked_missing_fields" for row in validation.document_rows)


def test_business_wizard_document_rows_show_expected_documents() -> None:
    validation = evaluate_business_wizard(sample_business_wizard_input())

    rows = business_document_table_rows(validation)

    assert [row["code document"] for row in rows] == [
        "DOC-020",
        "DOC-001",
        "DOC-003",
        "DOC-002",
        "DOC-004",
    ]
    assert rows[0]["statut"] == "Contexte incomplet pour génération V2"
    assert {row["statut"] for row in rows[1:]} == {"Générable"}


def test_business_wizard_context_uses_orchestrator_selection() -> None:
    validation = evaluate_business_wizard(sample_business_wizard_input())
    assert validation.context is not None
    orchestrator = DocumentOrchestrator(build_seed_catalog())

    selected = orchestrator.select_documents_for_context(validation.context)

    assert [document.doc_id for document in selected] == [
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
    ]


def test_business_wizard_sci_option_is_displays_doc_022() -> None:
    validation = evaluate_business_wizard(
        replace(sample_business_wizard_input(), option_is=True)
    )

    assert "DOC-022" in _row_codes(validation)
    option_is_row = _row_by_code(validation, "DOC-022")
    assert option_is_row.status == STATUS_CONTEXT_INCOMPLETE


def test_business_wizard_sci_iris_displays_doc_021_and_not_doc_020() -> None:
    validation = evaluate_business_wizard(
        replace(sample_business_wizard_input(), sci_iris=True)
    )

    assert "DOC-021" in _row_codes(validation)
    assert "DOC-020" not in _row_codes(validation)


def test_business_wizard_selarl_regime_communautaire_displays_doc_005_and_doc_006() -> None:
    validation = evaluate_business_wizard(
        _case_data(
            "SELARL",
            profession="medecin",
            site_distinct=False,
            scm_cession=False,
            regime_communautaire=True,
            derogation=False,
            cession=False,
        )
    )

    assert "DOC-005" in _row_codes(validation)
    assert "DOC-006" in _row_codes(validation)


def test_business_wizard_selarl_site_distinct_displays_manual_document() -> None:
    validation = evaluate_business_wizard(
        _case_data(
            "SELARL",
            profession="chirurgien_dentiste",
            site_distinct=True,
            scm_cession=False,
            regime_communautaire=False,
            derogation=False,
            cession=False,
        )
    )

    manual_row = next(row for row in validation.document_rows if row.document_code is None)

    assert manual_row.document_key == "site_distinct_cd94_sel"
    assert manual_row.status == STATUS_MANUAL_ONLY


def test_business_wizard_selas_derogation_displays_not_implemented_document() -> None:
    validation = evaluate_business_wizard(
        _case_data(
            "SELAS",
            profession="medecin",
            scm=False,
            regime_communautaire=False,
            derogation=True,
            cession=False,
        )
    )

    row = next(
        row
        for row in validation.document_rows
        if row.document_key == "derogation_cumul_selarl_salariee"
    )

    assert row.status == STATUS_NOT_IMPLEMENTED


def test_business_wizard_spfpl_cession_displays_agrement_by_associate_count() -> None:
    unique = evaluate_business_wizard(
        _case_data(
            "SPFPL cession",
            regime_communautaire=False,
            associe_unique=True,
            cession_actions=False,
        )
    )
    several = evaluate_business_wizard(
        _case_data(
            "SPFPL cession",
            regime_communautaire=False,
            associe_unique=False,
            cession_actions=False,
        )
    )

    assert "DOC-038" in _row_codes(unique)
    assert "DOC-039" not in _row_codes(unique)
    assert "DOC-039" in _row_codes(several)
    assert "DOC-038" not in _row_codes(several)


def test_business_wizard_manual_and_not_implemented_documents_are_not_sent_to_generation() -> None:
    validation = evaluate_business_wizard(
        _case_data(
            "SELARL",
            profession="medecin",
            site_distinct=True,
            scm_cession=False,
            regime_communautaire=False,
            derogation=True,
            cession=False,
        )
    )

    excluded_statuses = {STATUS_MANUAL_ONLY, STATUS_NOT_IMPLEMENTED}
    excluded_rows = [row for row in validation.document_rows if row.status in excluded_statuses]

    assert excluded_rows
    assert all(
        row.document_code not in validation.generatable_document_codes
        for row in excluded_rows
    )
    assert all(
        _row_by_code(validation, code).status == STATUS_GENERABLE
        for code in validation.generatable_document_codes
    )


def test_streamlit_technical_mode_remains_accessible() -> None:
    app_source = Path("src/sydel_doc_engine/app/streamlit_app.py").read_text(encoding="utf-8")

    assert "Technique / diagnostic" in app_source
    assert "Assistant metier" in app_source


def test_streamlit_business_mode_exposes_sci_and_selarl() -> None:
    dossier_types = {item.structure for item in business_wizard.business_dossier_types()}

    assert {"SCI", "SELARL"}.issubset(dossier_types)


def test_selarl_ui_conditions_are_available() -> None:
    assert {spec.key for spec in selarl_ui_condition_specs()} == {
        "profession",
        "site_distinct",
        "scm_cession",
        "regime_communautaire",
        "derogation",
        "cession",
        "cabinet_type",
    }


def test_streamlit_selarl_path_uses_business_wording() -> None:
    app_source = Path("src/sydel_doc_engine/app/streamlit_app.py").read_text(encoding="utf-8")
    banned = "professionnel " + "principal"

    assert "Dirigeant / pharmacien" not in app_source
    assert banned not in app_source.casefold()
    assert "Ecran 3 - Fiche Client" in app_source
    assert "Praticien" in app_source


def test_selarl_ui_address_labels_are_qualified() -> None:
    ambiguous_labels = {"adresse", "numero", "voie", "ville", "code postal"}
    labels = selarl_ui_address_labels()

    assert labels
    assert all(label.casefold().strip() not in ambiguous_labels for label in labels)
    assert any("adresse personnelle du praticien" in label.casefold() for label in labels)
    assert any("adresse du siege social" in _plain(label) for label in labels)
    assert any("adresse de la banque" in label.casefold() for label in labels)


def test_selarl_docs_006_013_014_have_required_ui_statuses() -> None:
    validation = evaluate_business_wizard(
        _case_data(
            "SELARL",
            profession="medecin",
            site_distinct=False,
            scm_cession=False,
            regime_communautaire=True,
            derogation=True,
            cession=False,
        )
    )

    assert any("vraie V2" in note for note in _row_by_code(validation, "DOC-006").notes)
    assert _row_by_code(validation, "DOC-013").status == STATUS_MANUAL_ONLY
    assert _row_by_code(validation, "DOC-014").status == STATUS_MANUAL_ONLY
    assert "DOC-013" not in validation.generatable_document_codes
    assert "DOC-014" not in validation.generatable_document_codes


def test_selarl_ui_reuse_rules_cover_main_deduplications() -> None:
    rule_keys = {rule.key for rule in selarl_ui_reuse_rules()}

    assert {
        "signataire_is_associe_1",
        "gerant_is_professional",
        "signataire_is_professional",
        "mandataire_is_signataire",
        "selarl_is_acquirer",
        "selarl_is_scm_transferee",
        "domiciliation_is_registered_office",
    }.issubset(rule_keys)


def test_selarl_ui_document_specs_stay_aligned_with_catalog_statuses() -> None:
    spec_codes = {document.document_code for document in selarl_ui_document_specs()}
    validation = evaluate_business_wizard(
        _case_data(
            "SELARL",
            profession="medecin",
            site_distinct=True,
            scm_cession=True,
            regime_communautaire=True,
            derogation=True,
            cession=True,
            cabinet_type="medical",
        )
    )
    row_codes = {row.document_code for row in validation.document_rows}

    assert {"DOC-006", "DOC-013", "DOC-014"}.issubset(spec_codes)
    assert {"DOC-006", "DOC-013", "DOC-014"}.issubset(row_codes)


def test_selarl_ui_block_visibility_masks_inactive_specific_blocks() -> None:
    data = _case_data(
        "SELARL",
        profession="medecin",
        site_distinct=False,
        scm_cession=False,
        regime_communautaire=False,
        derogation=False,
        cession=False,
    )

    visibility = selarl_ui_block_visibility(data)

    assert visibility["societe"] is True
    assert visibility["regime_conjoint"] is False
    assert visibility["scm"] is False
    assert visibility["cession_cabinet"] is False
    assert visibility["bail"] is False
    assert visibility["banque_financement"] is False


def test_business_wizard_generates_docx_and_zip_without_residual_placeholders(
    tmp_path: Path,
) -> None:
    validation = evaluate_business_wizard(sample_business_wizard_input())
    assert validation.context is not None

    docx_paths = generate_docx_files_for_document_codes(
        validation.context,
        tmp_path,
        validation.generatable_document_codes,
    )
    zip_path = generate_zip_file(tmp_path, docx_paths)

    assert len(docx_paths) == 4
    for docx_path in docx_paths:
        text = _docx_text(docx_path)
        assert "[" not in text
        assert "]" not in text
    with ZipFile(zip_path) as archive:
        names = set(archive.namelist())
    assert "manifest.json" in names
    assert "declaration_non_condamnation.docx" in names


def _case_data(structure: str, **conditions: object) -> BusinessWizardInput:
    return replace(sample_business_wizard_input(), structure=structure, **conditions)


def _row_codes(validation) -> list[str]:
    return [row.document_code for row in validation.document_rows if row.document_code]


def _row_by_code(validation, code: str):
    return next(row for row in validation.document_rows if row.document_code == code)


def _plain(value: str) -> str:
    return (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
        .casefold()
    )


def _docx_text(path: Path) -> str:
    document = Document(path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    table_text = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraph_text, *table_text])
