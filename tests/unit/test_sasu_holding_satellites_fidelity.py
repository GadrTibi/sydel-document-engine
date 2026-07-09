"""Tests de FIDELITE des satellites SASU Holding generalistes (modeles Albane 2026-06-29).

PV remuneration president (DOC-049) + liste des souscripteurs (DOC-050) : generateurs
GENERALISTES dedies, byte-fideles aux modeles officiels Albane
(`docs/review/albane_sas_2026-06-29/`), DISTINCTS des satellites SPFPL medecins DOC-023 /
DOC-024 (conserves, verrouilles SPFPL medecins). Dossier de reference : SAS MLG / Malo LE GUEN,
president, 5 Allee de la Clarte 56700 KERVIGNAC, 10000 actions, 1000 euros.

On assert les lignes VERBATIM des modeles Albane (corps + tableau + signature) et l'ABSENCE de
tout wording SPFPL / medecin / profession reglementee (preuve : holding generaliste, genre libre).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    Company,
    DocumentGenerationContext,
    DossierOptions,
    ExerciceSocial,
    Person,
    Signature,
    StatutsSasuHoldingContext,
)
from sydel_doc_engine.generators.lot_01.autorisation_domiciliation import (
    AutorisationDomiciliationGenerator,
)
from sydel_doc_engine.generators.lot_01.procuration import ProcurationGenerator
from sydel_doc_engine.generators.lot_05.liste_souscripteurs_sasu_holding import (
    ListeSouscripteursSasuHoldingGenerator,
)
from sydel_doc_engine.generators.lot_05.pv_remuneration_president_sasu_holding import (
    PvRemunerationPresidentSasuHoldingGenerator,
)

# Ponctuation fine francaise des modeles Albane (apostrophe courbe U+2019, espace insecable
# U+00A0) — la sortie doit etre byte-fidele a ces caracteres.
_APOS = chr(0x2019)
_NBSP = chr(0x00A0)


def _ctx(*, genre: Gender = Gender.MASCULIN) -> DocumentGenerationContext:
    civilite = "Madame" if genre == Gender.FEMININ else "Monsieur"
    adresse = Address(
        num_voie="5", voie="Allée de la Clarté", cp="56700", ville="KERVIGNAC",
        adresse_affichee="5 Allée de la Clarté, 56700 KERVIGNAC",
    )
    return DocumentGenerationContext(
        structure="SASU_HOLDING",
        dossier_options=DossierOptions(associe_unique=True),
        personne_signataire=Person(
            genre=genre,
            civilite=civilite,
            prenom="Malo",
            nom="LE GUEN",
            adresse_perso=adresse,
            adresse_personnelle_affichee="5 Allée de la Clarté, 56700 KERVIGNAC",
            fonction_dirigeant="Président",
        ),
        signature=Signature(lieu="KERVIGNAC", date=date(2025, 11, 13)),
        societe=Company(
            denomination="MLG",
            forme_sociale="Société par actions simplifiée unipersonnelle",
            capital="1000",
            capital_social="1000",
            siege=adresse,
            ville_rcs="KERVIGNAC",
        ),
        exercice_social=ExerciceSocial(
            debut="1er janvier",
            fin="31 décembre",
            date_cloture_premier_exercice="le 31 décembre 2026",
        ),
        statuts_sasu_holding=StatutsSasuHoldingContext(
            forme_sociale="Société par actions simplifiée unipersonnelle",
            capital_social="1000",
            capital_social_lettres="mille",
            nb_actions=10000,
            nom_banque="HSBC",
            qualite_associe="Associé unique et Président",
            fonction_dirigeant="Président",
        ),
    )


def _paragraphs(path: Path) -> list[str]:
    return [p.text for p in Document(path).paragraphs]


def _all_text(path: Path) -> str:
    document = Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _table_rows(path: Path) -> list[list[str]]:
    document = Document(path)
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}tbl"):
            table = Table(child, document)
            return [[cell.text for cell in row.cells] for row in table.rows]
    raise AssertionError("aucun tableau trouve dans le document")


def _body_in_order(path: Path) -> list[str]:
    """Texte du corps dans l'ordre (paragraphes + cellules de tableau aplaties)."""
    document = Document(path)
    out: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            out.append(Paragraph(child, document).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                out.extend(cell.text for cell in row.cells)
    return out


# --- Autorisation de domiciliation (tronc commun, chemin SASU_HOLDING) -------------------


def test_domiciliation_sasu_holding_fidelity(tmp_path: Path) -> None:
    # Gate Akainu : la domiciliation SASU porte la forme « de la SAS <denom> » (coherent avec
    # la procuration) + « dans les locaux situes a <adresse>, » (une holding n'est pas un
    # cabinet). Structure-aware SASU_HOLDING (autres types inchanges).
    path = AutorisationDomiciliationGenerator().generate(_ctx(), tmp_path)
    text = _all_text(path)
    assert "de la SAS MLG" in text
    assert "dans les locaux situés à" in text
    assert ", pour une durée indéterminée" in text
    assert "du cabinet" not in text  # plus de wording cabinet


def test_procuration_sasu_holding_no_legal_effect_line(tmp_path: Path) -> None:
    # Gate Akainu / source : le modele Procuration_SAS.docx d'Albane n'a PAS la ligne « Fait
    # pour servir et valoir ce que de droit. » (present seulement dans le modele micro holding)
    # -> omise pour la SASU Holding. Forme abregee « SAS » portee par le slice (forme_affichage).
    ctx = _ctx()
    # Le slice SASU pose la forme abregee « SAS » sur la Company (forme_sociale + affichage).
    ctx.societe.forme_sociale = "SAS"
    ctx.societe.forme_sociale_affichage = "SAS"
    text = _all_text(ProcurationGenerator().generate(ctx, tmp_path))
    assert "de la SAS MLG" in text
    assert "Fait pour servir et valoir ce que de droit" not in text


# --- PV remuneration president (DOC-049) -------------------------------------------------


def test_pv_remuneration_president_fidelity(tmp_path: Path) -> None:
    path = PvRemunerationPresidentSasuHoldingGenerator().generate(_ctx(), tmp_path)
    paragraphs = _paragraphs(path)
    all_text = _all_text(path)

    # Titre encadre (verbatim modele Albane), date longue francaise.
    assert "PROCES-VERBAL DES DECISIONS" in all_text
    assert f"DE L{_APOS}ASSOCIE UNIQUE" in all_text
    assert "DU 13 novembre 2025" in all_text

    # Corps verbatim (lignes du modele Albane) : apostrophe courbe + NBSP + « jusqu'au 31 ... »
    # (sans « le » : pas de double article — gate Akainu B1).
    expected_lines = [
        "Monsieur Malo LE GUEN",
        "Demeurant 5 Allée de la Clarté, 56700 KERVIGNAC",
        "Associé unique et Président de la SASU en cours de formation.",
        # Rafael 2026-07-09 (R12) : majuscule en debut de phrase (la ligne suit un
        # point) — supersede la minuscule du modele (precedent « tout le texte doit
        # etre correct », Rafael 2026-07-07).
        f"A pris la décision suivante{_NBSP}: ",
        "Fixation de la rémunération du Président",
        "DECISION UNIQUE",
        (
            f"Monsieur LE GUEN, associé unique, décide qu{_APOS}il ne percevra aucune "
            "rémunération au titre de son mandat de Président, à compter de son immatriculation, "
            f"et ce, jusqu{_APOS}au 31 décembre 2026 inclus, date de la clôture du premier "
            "exercice social."
        ),
        (
            "Il pourra donc prétendre au remboursement sur justification de ses frais de "
            "représentation et de déplacement."
        ),
        (
            f"De tout ce que dessus, l{_APOS}associé unique a dressé et signé le présent "
            "procès-verbal."
        ),
        "Fait à KERVIGNAC en trois exemplaires ",
        "________________",
        "Malo LE GUEN",
    ]
    for ligne in expected_lines:
        assert ligne in paragraphs, f"ligne PV absente : {ligne!r}"

    # Aucun wording SPFPL medecins (preuve : satellite generaliste, pas le DOC-023 verrouille).
    for interdit in (
        "Participations Financières",
        "Profession Libérale",
        "médecin",
        "Dr ",
    ):
        assert interdit not in all_text, f"wording SPFPL/medecin inattendu : {interdit!r}"


def test_pv_remuneration_president_requires_structure(tmp_path: Path) -> None:
    ctx = _ctx()
    ctx.structure = "SAS"
    with pytest.raises(ValueError, match="SASU_HOLDING"):
        PvRemunerationPresidentSasuHoldingGenerator().generate(ctx, tmp_path)


def test_pv_remuneration_president_feminin(tmp_path: Path) -> None:
    # Genre libre : une associee unique femme passe sans la garde « masculin » du DOC-023.
    path = PvRemunerationPresidentSasuHoldingGenerator().generate(
        _ctx(genre=Gender.FEMININ), tmp_path
    )
    paragraphs = _paragraphs(path)
    assert "Madame Malo LE GUEN" in paragraphs
    assert (
        f"Madame LE GUEN, associé unique, décide qu{_APOS}il ne percevra aucune rémunération"
        in "\n".join(paragraphs)
    )


# --- Liste des souscripteurs (DOC-050) ---------------------------------------------------


def test_liste_souscripteurs_fidelity(tmp_path: Path) -> None:
    path = ListeSouscripteursSasuHoldingGenerator().generate(_ctx(), tmp_path)
    body = _body_in_order(path)
    rows = _table_rows(path)

    # En-tete « Etat des souscriptions / De la SAS <denom> » (forme courte verbatim modele).
    assert "Etat des souscriptions " in body
    assert "De la SAS MLG" in body

    # Tableau : en-tetes verbatim.
    assert rows[0] == [
        "Noms, prénoms et adresse des souscripteurs",
        f"Nombre d{_APOS}actions souscrites",
        "Montant des souscriptions",
    ]
    # Ligne souscripteur unique (nb actions formate a POINT comme le modele : « 10.000 »).
    assert rows[1] == ["Monsieur Malo LE GUEN", "10.000 ", "1000 "]
    # Ligne TOTAL verbatim.
    assert rows[2] == ["TOTAL", "10.000 actions", "1000 euros"]

    # Signature : « Fait à <lieu> » / « Le <jj/mm/aaaa> » / « <prenom> <nom> » (espace de tete).
    assert "Fait à KERVIGNAC" in body
    assert "Le 13/11/2025" in body
    assert " Malo LE GUEN" in body

    # Aucun wording SPFPL / apport en nature (satellite generaliste).
    all_text = _all_text(path)
    for interdit in ("Participations Financières", "Apports en nature", "médecin"):
        assert interdit not in all_text, f"wording inattendu : {interdit!r}"


def test_liste_souscripteurs_requires_structure(tmp_path: Path) -> None:
    ctx = _ctx()
    ctx.structure = "SAS"
    with pytest.raises(ValueError, match="SASU_HOLDING"):
        ListeSouscripteursSasuHoldingGenerator().generate(ctx, tmp_path)
