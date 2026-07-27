from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from sydel_doc_engine.domain.enums import Gender
from sydel_doc_engine.domain.models import (
    Address,
    RegimeCommunautaireAssocie,
    StatutsCivilsApport,
    StatutsCivilsAssocie,
    StatutsCivilsParts,
    StatutsCivilsRepresentant,
)
from sydel_doc_engine.front_app import (
    civil_statuts_slice as css,
)
from sydel_doc_engine.front_app import (
    sas_slice,
    selas_multi_slice,
    spfpl_slice,
)
from sydel_doc_engine.front_app.dossier_selection import (
    dossier_type_by_label,
    dossier_type_labels,
)
from sydel_doc_engine.front_app.type_registry import registered_types


def _docx_text(path: Path) -> str:
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs if p.text)
    for section in document.sections:
        texts.extend(p.text for p in section.footer.paragraphs if p.text)
    return "\n".join(texts)


def _assert_clean(text: str) -> None:
    assert "[" not in text
    assert "]" not in text
    # SCS4 (Akainu m1 2026-06-25) : le marqueur « (À COMPLÉTER : …) » est volontairement
    # SANS crochets (pour échapper au garde-fou source) -> un champ ctx oublié le shippe
    # silencieusement (vécu : DOC-006 SCS « (À COMPLÉTER : forme_sociale_abregee) »). On le
    # scanne donc explicitement, pour TOUS les documents.
    assert "À COMPLÉTER" not in text, "marqueur (À COMPLÉTER : …) résiduel"
    assert "A COMPLETER" not in text.upper().replace("À", "A").replace("É", "E"), (
        "marqueur (A COMPLETER : …) résiduel"
    )
    # O24-01 (onglet 24) : les 2 items frais de cabinet de création (« lettre de mission » /
    # « acompte des honoraires ») sont retirés de l'annexe de TOUS les statuts, partout.
    low = text.lower()
    assert "lettre de mission" not in low, "O24-01 lettre de mission à retirer"
    assert "acompte des honoraires" not in low, "O24-01 acompte honoraires à retirer"


# --- Registre + deroulante ----------------------------------------------------


def test_registry_exposes_all_ready_types() -> None:
    labels = dossier_type_labels()
    assert labels[0] == "SELARL"  # SELARL TOUJOURS en premier (defaut)
    structures = {item.structure for item in registered_types()}
    assert {
        "SELARL",
        "SCM",
        "SCI",
        "SCI IRIS",
        "SCS",
        # Micro holding (Albane 2026-06-26) : societe civile a capital variable.
        "MICRO_HOLDING",
        "SAS",
        # SASU Holding (Albane 2026-06-29) : SAS unipersonnelle, holding patrimoniale
        # generaliste, DISTINCTE de la SAS / SPFPL medecins.
        "SASU_HOLDING",
        "SPFPL cession",
        "SPFPL apport",
        "SELAS",
        "SELAS uni medecin",
        "SELAS uni dentiste",
    } == structures


def test_selarl_option_unchanged() -> None:
    option = dossier_type_by_label("SELARL")
    assert option.key == "selarl_v1"
    assert option.structure == "SELARL"
    assert option.generation_enabled is True
    assert option.slice_module is None  # chemin historique dedie


# --- Helpers de construction d'associes civils --------------------------------


def _pp(prenom, nom, nb, debut, fin, apport, role=None, prof="chirurgien-dentiste"):
    return StatutsCivilsAssocie(
        type_personne="personne_physique",
        role_statutaire=role,
        genre=Gender.MASCULIN,
        civilite_affichage="Monsieur",
        prenom=prenom,
        prenoms=prenom,
        nom=nom,
        profession=prof,
        date_naissance="1 janvier 1980",
        ville_naissance="Paris",
        departement_naissance="75",
        nationalite="francaise",
        situation_maritale="celibataire",
        adresse_personnelle_affichee="1 rue Exemple, 75000 Paris",
        # DNC par associe (Rafael 2026-07-09) : filiation requise pour CHAQUE associe.
        nom_pere=f"Pierre {nom}",
        nom_mere=f"Anne {nom}",
        apport=StatutsCivilsApport(montant=str(apport), montant_lettres=str(apport)),
        parts=StatutsCivilsParts(
            nb=nb,
            nb_lettres=str(nb),
            plage_affichee=f"{debut} a {fin}",
            debut=debut,
            fin=fin,
            qualite_associe=role,
        ),
    )


def _pm(nb, debut, fin, apport, prof="chirurgien-dentiste"):
    return StatutsCivilsAssocie(
        type_personne="personne_morale",
        denomination="SEL IRIS",
        forme_juridique="SELARL",
        profession=prof,
        capital_social="1 000 euros",
        siege=Address(adresse_affichee="2 rue Pro, 75000 Paris"),
        numero_rcs="900 000 001",
        ville_rcs="Paris",
        representant=StatutsCivilsRepresentant(
            civilite_affichage="Monsieur",
            prenom="Jean",
            nom="Durand",
            fonction="gerant",
        ),
        apport=StatutsCivilsApport(montant=str(apport), montant_lettres=str(apport)),
        parts=StatutsCivilsParts(nb=nb, nb_lettres=str(nb), debut=debut, fin=fin),
    )


def _civil_base(structure, statuts_type, associes):
    payload = {
        "structure": structure,
        "statuts_type": statuts_type,
        "denomination": f"{structure} EXEMPLE",
        "forme_sociale": "societe civile",
        "capital_social": "1000",
        "nb_parts_total": 100,
        "valeur_nominale_part": "10",
        "duree_societe": "99",
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "ville_rcs": "Paris",
        "banque_nom": "BANQUE",
        "banque_adresse": "1 rue Banque, 75009 Paris",
        "date_cloture_premier_exercice": "31 decembre 2026",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 15),
        "associes": associes,
        # Documents communs (DNC / procuration / PV gerant) : signataire = 1er physique.
        "signataire_nom_pere": "Pierre Durand",
        "signataire_nom_mere": "Anne Durand",
        "signataire_adresse_num": "1",
        "signataire_adresse_voie": "rue Exemple",
        "signataire_adresse_cp": "75000",
        "signataire_adresse_ville": "Paris",
        "signataire_fonction": "gerant",
        "signataire_titre": "Docteur",
        "decision_date": date(2026, 5, 15),
    }
    if structure == "SCM":
        payload.update(
            {
                "ordre_conseil": "Conseil departemental",
                "ordre_departement": "75",
                "ordre_adresse_ligne_1": "1 rue de l'Ordre",
                "ordre_cp": "75008",
                "ordre_ville": "Paris",
                "ordre_numero": "ORD-1",
                # Satellites SCM (pacte + liste depenses, generes a 2 associes).
                "pacte_ville_tribunal": "Paris",
                "societe_numero_rcs": "en cours de constitution",
            }
        )
    return payload


# --- Generation par type (smoke, 0 placeholder residuel) ----------------------


def _assert_bundle_clean(generated, expected_names) -> None:
    # O24-02 (onglet 24) : la DNC est renommee avec le nom du dirigeant. On normalise
    # vers le nom generique pour le controle de COMPLETUDE du bundle ; le nommage par
    # dirigeant est verifie par les tests DNC dedies (SELARL clean / SELAS dnc multi).
    names = {
        "declaration_non_condamnation.docx"
        if p.name.startswith("declaration_non_condamnation")
        else p.name
        for p in generated.docx_paths
    }
    assert expected_names <= names
    for path in generated.docx_paths:
        _assert_clean(_docx_text(path))


# Tronc commun present dans CHAQUE bundle de creation.
_TRONC_DOCS = {
    "declaration_non_condamnation.docx",
    "autorisation_domiciliation.docx",
    "procuration.docx",
    "pv_nomination_dirigeant.docx",
}


def test_sci_slice_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-020", "DOC-001", "DOC-002", "DOC-003", "DOC-004")
    generated = css.generate_dossier(payload, tmp_path / "sci")
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI EXEMPLE.docx"})


# Centre des impots requis par la lettre d'option IS (DOC-022) ; saisies utilisateur.
_OPTION_IS_INPUTS = {
    "option_is": True,
    "siren": "900 000 001",
    "impots_service": "SIE",
    "impots_centre": "Centre des Finances Publiques",
    "impots_adresse_ligne_1": "1 rue des Impots",
    "impots_adresse_ligne_2": "BP 100",
    "impots_cp": "75002",
    "impots_ville": "Paris",
}


def test_sci_option_is_off_keeps_base_bundle(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload["option_is"] = False
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert "DOC-022" not in plan.document_codes
    generated = css.generate_dossier(payload, tmp_path / "sci-is-off")
    names = {p.name for p in generated.docx_paths}
    assert "lettre_option_is.docx" not in names


def test_sci_option_is_on_adds_lettre_option_is(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload.update(_OPTION_IS_INPUTS)
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-020",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-022",
    )
    generated = css.generate_dossier(payload, tmp_path / "sci-is-on")
    _assert_bundle_clean(
        generated,
        _TRONC_DOCS | {"Statuts SCI EXEMPLE.docx", "lettre_option_is.docx"},
    )


def test_sci_iris_option_is_on_adds_lettre_option_is(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload.update(_OPTION_IS_INPUTS)
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert "DOC-022" in plan.document_codes
    generated = css.generate_dossier(payload, tmp_path / "iris-is-on")
    _assert_bundle_clean(
        generated,
        _TRONC_DOCS | {"Statuts SCI IRIS EXEMPLE.docx", "lettre_option_is.docx"},
    )


def test_sci_option_is_on_requires_centre_des_impots() -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload["option_is"] = True  # toggle ON mais aucune saisie centre des impots
    plan = css.build_civil_plan(payload)
    # KAN-2 (Rafael 2026-07-14, rejeté 2×) : un centre des impôts non renseigné ne BLOQUE PLUS ->
    # génération possible, le manque devient un AVERTISSEMENT (zone « (À COMPLÉTER : …) »).
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("option IS" in w for w in plan.warnings)


def test_scm_slice_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCM",
        "scm",
        [_pm(70, 1, 70, 700), _pp("Alice", "Martin", 30, 71, 100, 300)],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    # SCM a 2 associes : le bundle inclut les satellites pacte (DOC-026) + liste
    # des depenses communes (DOC-030), decision Rafael 2026-06-08.
    assert plan.document_codes == (
        "DOC-025",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-030",
        "DOC-026",
    )
    generated = css.generate_dossier(payload, tmp_path / "scm")
    _assert_bundle_clean(
        generated,
        _TRONC_DOCS
        | {
            "Statuts SCM EXEMPLE.docx",
            "demande_inscription_ordre.docx",
            "pacte_associes_scm.docx",
            "liste_depenses_communes_scm.docx",
        },
    )


# Documents inter-SEL (opt-in) : identite de la SEL de chaque associe + telephone +
# parametres du contrat de frais communs / reglement interieur. Personnes derivees des
# associes (pas de ressaisie). Forme sociale IDENTIQUE pour les 2 SEL (contrainte source).
_SCM_INTER_SEL_INPUTS = {
    "inter_sel_active": True,
    "inter_sel_forme": "SELARL",
    "inter_sel_titre": "Docteur",
    "inter_sel_parties": [
        {
            "denomination": "SEL DOCTEUR DURAND",
            "capital": "1 000 euros",
            "siege": "1 rue des Soins, 75001 Paris",
            "ville_rcs": "Paris",
            "numero_rcs": "900 000 001",
            "telephone": "01 00 00 00 01",
        },
        {
            "denomination": "SEL DOCTEUR MARTIN",
            "capital": "1 000 euros",
            "siege": "2 rue des Soins, 75002 Paris",
            "ville_rcs": "Paris",
            "numero_rcs": "900 000 002",
            "telephone": "01 00 00 00 02",
        },
    ],
    "inter_sel_locaux": "10 rue de la Paix, 75002 Paris",
    "inter_sel_date_effet": "1er janvier 2027",
    "inter_sel_seuil": "1 500 euros",
    "inter_sel_annee_ref": "2027",
    "inter_sel_date_fin_gestion": "31 decembre 2027",
    "inter_sel_date_attribution": "1er janvier",
}

_SCM_INTER_SEL_DOCS = {"contrat_frais_communs.docx", "reglement_interieur_scm.docx"}


def test_scm_inter_sel_adds_frais_communs_reglement(tmp_path: Path) -> None:
    # Opt-in actif + 2 associes physiques -> le bundle ajoute le contrat de frais
    # communs (DOC-027) et le reglement interieur (DOC-028).
    payload = _civil_base(
        "SCM",
        "scm",
        [_pp("Jean", "Durand", 50, 1, 50, 500), _pp("Alice", "Martin", 50, 51, 100, 500)],
    )
    payload.update(_SCM_INTER_SEL_INPUTS)
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-025",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-030",
        "DOC-026",
        "DOC-027",
        "DOC-028",
    )
    generated = css.generate_dossier(payload, tmp_path / "scm-inter-sel")
    _assert_bundle_clean(
        generated,
        _TRONC_DOCS
        | {
            "Statuts SCM EXEMPLE.docx",
            "pacte_associes_scm.docx",
            "liste_depenses_communes_scm.docx",
        }
        | _SCM_INTER_SEL_DOCS,
    )


def test_scm_inter_sel_off_keeps_base_satellites(tmp_path: Path) -> None:
    # Inter-SEL explicitement inactif (bascule decochee) -> pas de docs inter-SEL, bundle
    # satellites de base inchange. NB : depuis R22-03/04 le defaut UI est ACTIF ; ce test
    # couvre le cas ou l'utilisateur decoche (SCM sans SEL distincte).
    payload = _civil_base(
        "SCM",
        "scm",
        [_pp("Jean", "Durand", 50, 1, 50, 500), _pp("Alice", "Martin", 50, 51, 100, 500)],
    )
    payload["inter_sel_active"] = False
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert "DOC-027" not in plan.document_codes
    assert "DOC-028" not in plan.document_codes
    generated = css.generate_dossier(payload, tmp_path / "scm-no-inter-sel")
    names = {p.name for p in generated.docx_paths}
    assert "contrat_frais_communs.docx" not in names
    assert "reglement_interieur_scm.docx" not in names


def test_scm_inter_sel_blocks_with_personne_morale(tmp_path: Path) -> None:
    # Opt-in actif mais un associe est une personne morale -> blocage (la SEL de chaque
    # associe physique est la partie ; une PM ne fournit pas de praticien-representant).
    payload = _civil_base(
        "SCM",
        "scm",
        [_pm(50, 1, 50, 500), _pp("Alice", "Martin", 50, 51, 100, 500)],
    )
    payload.update(_SCM_INTER_SEL_INPUTS)
    plan = css.build_civil_plan(payload)
    # KAN-2 : plus de blocage — l'avertissement signale que les 2 associés doivent être des
    # personnes physiques pour les documents inter-SEL ; la génération reste possible.
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("personnes physiques" in w for w in plan.warnings)


def test_civil_capital_not_divisible_by_parts_now_generates(tmp_path: Path) -> None:
    # N1 (Rafael/Vincent 2026-06-24) : la valeur nominale PEUT etre decimale (regle ratifiee).
    # Un capital non divisible par le nb de parts (1000/3) ne bloque PLUS ; l'arrondi au centime
    # (calculate_nominal_value) evite la decimale infinie. Ancienne garde de divisibilite retiree.
    payload = _civil_base("SCI", "sci", [_pp("Jean", "Durand", 3, 1, 3, 1000)])
    payload["capital_social"] = "1000"
    payload["nb_parts_total"] = 3
    plan = css.build_civil_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_civil_capital_divisible_by_parts_ok(tmp_path: Path) -> None:
    # Contre-epreuve : 999 / 3 = 333 (entier) -> pas de blocage de divisibilite.
    payload = _civil_base("SCI", "sci", [_pp("Jean", "Durand", 3, 1, 3, 999)])
    payload["capital_social"] = "999"
    payload["nb_parts_total"] = 3
    plan = css.build_civil_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_selas_capital_not_divisible_by_actions_now_generates() -> None:
    # N1 (Rafael/Vincent 2026-06-24) : valeur nominale decimale autorisee -> SELAS multi avec
    # capital non divisible par le nb d'actions ne bloque PLUS (ancienne garde O24-05 retiree).
    payload = _selas_payload()
    payload["capital_social"] = "1000"
    payload["nb_actions_total"] = 3
    # Repartition coherente (somme = 3) pour isoler l'absence de blocker de divisibilite.
    payload["associes"][0].nb_actions = 2
    payload["associes"][1].nb_actions = 1
    plan = selas_multi_slice.build_selas_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_selas_capital_divisible_by_actions_ok() -> None:
    # Contre-epreuve : 999 / 3 = 333 (entier) -> pas de blocage de divisibilite.
    payload = _selas_payload()
    payload["capital_social"] = "999"
    payload["nb_actions_total"] = 3
    payload["associes"][0].nb_actions = 2
    payload["associes"][1].nb_actions = 1
    plan = selas_multi_slice.build_selas_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_spfpl_capital_not_divisible_by_actions_now_generates() -> None:
    # N1 (Rafael/Vincent 2026-06-24) : valeur nominale decimale autorisee -> SPFPL avec capital
    # non divisible par le nb d'actions ne bloque PLUS (ancienne garde O24-05 retiree).
    payload = _spfpl_payload("SPFPL cession")
    payload["capital_social"] = "1000"
    payload["nb_actions_total"] = 7
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_spfpl_capital_divisible_by_actions_ok() -> None:
    # Contre-epreuve : 60000 / 600 = 100 (entier) -> pas de blocage de divisibilite.
    payload = _spfpl_payload("SPFPL cession")
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert not any("divisible" in b for b in plan.blockers)


def test_valeur_nominale_decimale_rendue_n1() -> None:
    # N1 (Rafael/Vincent 2026-06-24) : la valeur nominale PEUT etre decimale, PARTOUT. La FIGURE
    # est arrondie au centime + format FR (« 1,25 »).
    # Containment 2026-07-06 (re-architecture) : `number_words_from_value` N'EST PLUS
    # decimal-aware globalement (cause de doubles-euros en cascade). Un decimal -> la FIGURE en
    # format FR (« 0,01 »), sure pour TOUTE recomposition « lettres + euro » separee ; l'entier
    # reste des mots NUS (« dix »). La phrase monetaire « un centime d'euro » (Albane 7.5) est
    # LOCALISEE au SEUL point de composition `grammar.montant_lettres_avec_unite` (teste dans
    # test_grammar.py + les tests SEL exercice ci-dessous).
    from decimal import Decimal

    from sydel_doc_engine.front_app.field_derivations import (
        calculate_nominal_value,
        number_words_from_value,
    )

    assert calculate_nominal_value("1000", 800) == "1,25"
    assert calculate_nominal_value("1000", 3) == "333,33"  # arrondi : plus de decimale infinie
    assert calculate_nominal_value("999", 3) == "333"  # entier reste propre
    # Containment : decimal -> FIGURE (jamais la phrase monetaire, qui doublerait « euro »).
    assert number_words_from_value(Decimal("1.25")) == "1,25"
    assert number_words_from_value("0,01") == "0,01"
    assert number_words_from_value("0,50") == "0,5"
    assert number_words_from_value("2,50") == "2,5"
    assert number_words_from_value(10) == "dix"  # entier inchange (mots nus, sans unite)

    # monetary_words_from_value : le batisseur monetaire dedie (formes standard FR).
    from sydel_doc_engine.front_app.field_derivations import monetary_words_from_value

    assert monetary_words_from_value("1") == "un euro"
    assert monetary_words_from_value("100") == "cent euros"
    assert monetary_words_from_value("0,01") == "un centime d’euro"
    assert monetary_words_from_value("0,50") == "cinquante centimes d’euro"
    assert monetary_words_from_value("2,50") == "deux euros et cinquante centimes"
    assert monetary_words_from_value("1,01") == "un euro et un centime"

    # prix_lettres_from_value : le PRIX n'est PAS ratifie en lettres (contrairement a la valeur
    # nominale 7.5) -> entier = mots nus, DECIMAL = FIGURE (jamais la phrase monetaire, qui
    # doublerait « euro » a la recomposition « <lettres> + euro_word »). Akainu M1/M2 2026-07-06.
    from sydel_doc_engine.front_app.field_derivations import prix_lettres_from_value

    assert prix_lettres_from_value("1") == "un"
    assert prix_lettres_from_value("1000") == "mille"
    assert prix_lettres_from_value("0,01") == "0,01"  # DECIMAL -> figure, PAS « un centime d'euro »
    assert prix_lettres_from_value("2,5") == "2,5"
    assert prix_lettres_from_value("") == ""

    # _euro_amount_letters (prix cabinet SELARL, 3e surface) : recompose « <lettres> <euro> »
    # -> sur un DECIMAL, jamais de double « euro » ni de phrase monetaire (regression Akainu M1).
    from sydel_doc_engine.front_app.shell import _euro_amount_letters

    assert _euro_amount_letters("100") == "cent euros"
    assert _euro_amount_letters("1") == "un euro"
    assert "euro euro" not in _euro_amount_letters("0,01")
    assert "centime d’euro" not in _euro_amount_letters("2,50")


def test_selas_ajout_associe_preserve_les_precedents_n5() -> None:
    # N5 (Albane 2026-06-24) : ajouter un associe ne doit PAS effacer les champs des precedents.
    # Cause : st.rerun() du bouton "Ajouter" se declenchait AVANT le rendu de la boucle associes
    # -> Streamlit garbage-collectait l'etat des widgets non instancies -> champs effaces. Fix :
    # retrait du st.rerun() + re-lecture du count (la boucle rend directement le nouveau nombre).
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(w for w in app.number_input if str(w.key) == "selas_associe_0_nb_actions").set_value(60)
    next(w for w in app.number_input if str(w.key) == "selas_associe_1_nb_actions").set_value(40)
    app = app.run(timeout=180)
    next(b for b in app.button if str(b.key) == "selas_add").click()
    app = app.run(timeout=180)
    a0 = next(w for w in app.number_input if str(w.key) == "selas_associe_0_nb_actions").value
    a1 = next(w for w in app.number_input if str(w.key) == "selas_associe_1_nb_actions").value
    assert a0 == 60 and a1 == 40  # champs des associes precedents preserves apres ajout


def test_civil_repeater_ajout_associe_preserve_les_precedents_n5() -> None:
    # N5 propagation (Albane 2026-06-24) : le repeater PARTAGE (civil/SCI/SCM via
    # render_associe_repeater) ne doit pas non plus effacer les champs des associes precedents
    # a l'ajout. Meme cause/fix que le SELAS (retrait st.rerun). Verrouille le 2e des 3 repeaters.
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SCM")
    app = app.run(timeout=180)
    ti = {str(w.key): w for w in app.text_input}
    ti["scm_associe_0_profession"].set_value("Medecin")
    ti["scm_associe_1_profession"].set_value("Dentiste")
    app = app.run(timeout=180)
    next(b for b in app.button if str(b.key) == "scm_add").click()
    app = app.run(timeout=180)
    p0 = next(w for w in app.text_input if str(w.key) == "scm_associe_0_profession").value
    p1 = next(w for w in app.text_input if str(w.key) == "scm_associe_1_profession").value
    assert p0 == "Medecin" and p1 == "Dentiste"  # repeater partage : precedents preserves


def test_selarl_multi_ajout_associe_preserve_les_precedents_n5() -> None:
    # N5 propagation (Albane 2026-06-24) : SELARL multi-associes (_render_selarl_membres) ne doit
    # pas effacer les champs des membres precedents a l'ajout. Bug identique (st.rerun premature)
    # reste INTACT sur ce chemin (Akainu M1). Verrouille le 3e des 3 repeaters.
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    # Passer en mode multi-associes : decocher « Dossier unipersonnel » (defaut coche).
    next(w for w in app.checkbox if str(w.key) == "selarl_dossier_unipersonnel").set_value(False)
    app = app.run(timeout=180)
    next(w for w in app.number_input if str(w.key) == "selarl_membre_0_nb_parts").set_value(30)
    app = app.run(timeout=180)
    next(b for b in app.button if str(b.key) == "selarl_membres_add").click()
    app = app.run(timeout=180)
    v0 = next(w for w in app.number_input if str(w.key) == "selarl_membre_0_nb_parts").value
    assert v0 == 30  # le membre precedent garde sa valeur apres ajout


def test_civil_capital_zero_blocks() -> None:
    # Dogfood 2026-06-22 : « 0 » passait la garde de presence -> capital 0. Doit bloquer.
    payload = _civil_base(
        "SCM", "scm", [_pp("Jean", "Durand", 50, 1, 50, 0), _pp("Alice", "Martin", 50, 51, 100, 0)]
    )
    payload["capital_social"] = "0"
    plan = css.build_civil_plan(payload)
    # KAN-2 : un capital à 0 ne BLOQUE PLUS -> avertissement (zone à compléter à la main).
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("apital" in w and "superieur a zero" in w for w in plan.warnings)


def test_civil_missing_banque_adresse_blocks() -> None:
    # Dogfood 2026-06-22 : le generateur exige l'adresse de la banque -> bloquer avant crash.
    payload = _civil_base("SCI", "sci", [_pp("Jean", "Durand", 100, 1, 100, 1000)])
    payload["banque_adresse"] = ""
    plan = css.build_civil_plan(payload)
    # KAN-2 : adresse de banque manquante -> avertissement (marqueur), plus de blocage.
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("anque" in w and "dresse" in w for w in plan.warnings)


def test_civil_morale_associe_without_representant_blocks() -> None:
    # Dogfood 2026-06-22 : associe personne morale sans representant -> can_generate=True
    # puis ValueError au generateur. Doit bloquer proprement.
    pm = _pm(50, 1, 50, 500)
    pm.representant = None
    payload = _civil_base("SCM", "scm", [pm, _pp("Alice", "Martin", 50, 51, 100, 500)])
    plan = css.build_civil_plan(payload)
    # KAN-2 : un représentant manquant ne BLOQUE PLUS -> avertissement (marqueur au générateur).
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("representant" in w.lower() for w in plan.warnings)


def test_scs_commanditaire_gerant_blocks() -> None:
    # Source NotebookLM : seul le commandite gere ; un commanditaire designe gerant -> bloque.
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
        ],
    )
    payload["gerant_index"] = 1  # le commanditaire designe gerant
    plan = css.build_civil_plan(payload)
    # KAN-2 : le choix « commanditaire gérant » ne BLOQUE PLUS la génération -> avertissement métier
    # (la règle « seul le commandité gère » est rappelée, mais le document reste générable).
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("commandite" in w.lower() and "gerant" in w.lower() for w in plan.warnings)


def test_scs_slice_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-019", "DOC-001", "DOC-002", "DOC-003", "DOC-004", "DOC-LSS-SCS"
    )
    generated = css.generate_dossier(payload, tmp_path / "scs")
    _assert_bundle_clean(
        generated, _TRONC_DOCS | {"Statuts SCS EXEMPLE.docx", "liste_souscripteurs_scs.docx"}
    )


def test_scs4_associe_marie_communaute_genere_doc005_006(tmp_path: Path) -> None:
    # SCS4 (Albane 2026-06-25) : la SCS reprend le bloc regime matrimonial de la SELAS
    # pluri -> un associe SCS marie sous communaute legale genere DOC-005 (renonciation)
    # + DOC-006 (avertissement), per-associe comme la SELAS (avant : capture morte, le
    # bloc UI captait le conjoint/regime mais AUCUN document n'etait emis — BLOQUANT Akainu).
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
        ],
    )
    payload["associes"][0].regime_communautaire_associe = _regime_associe("Paule", "Durand")
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" in plan.document_codes
    assert "DOC-006" in plan.document_codes
    generated = css.generate_dossier(payload, tmp_path / "scs-regime")
    names = {p.name for p in generated.docx_paths}
    assert _REGIME_DOCS <= names
    # SCS4 (Akainu m1) : le CONTENU des DOC-005/006 doit etre propre (le test ne se
    # contente plus de la presence des fichiers — un placeholder « (À COMPLÉTER : …) »
    # y avait echappe). On scanne chaque document du couple regime.
    for path in generated.docx_paths:
        if path.name in _REGIME_DOCS:
            _assert_clean(_docx_text(path))


def test_scs5_souscripteurs_ordre_et_personne_morale(tmp_path: Path) -> None:
    # SCS5 (Akainu M1+M2 2026-06-25) : (M1) l'ordre des souscripteurs suit l'ordre des associes
    # (pas d'inversion LIFO des 3 associes) ; (M2) un associe personne morale est LISTE et compte
    # dans le TOTAL (sinon souscripteur omis + total != capital).
    # --- M1 : 3 associes a valeurs distinctes, ordre preserve ---
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Un", "Aaa", 50, 1, 50, 500, role="commandite"),
            _pp("Deux", "Bbb", 30, 51, 80, 300, role="commanditaire"),
            _pp("Trois", "Ccc", 20, 81, 100, 200, role="commanditaire"),
        ],
    )
    generated = css.generate_dossier(payload, tmp_path / "scs-ordre")
    text = _names_in(generated, "liste_souscripteurs_scs.docx")
    pos = [text.index(n) for n in ("Un Aaa", "Deux Bbb", "Trois Ccc")]
    assert pos == sorted(pos), f"ordre des souscripteurs inverse : {pos}"
    # --- M2 : associe personne morale liste + TOTAL juste ---
    pm = _pm(40, 61, 100, 400)
    pm.role_statutaire = "commanditaire"
    payload_pm = _civil_base(
        "SCS",
        "scs",
        [_pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"), pm],
    )
    generated_pm = css.generate_dossier(payload_pm, tmp_path / "scs-pm")
    text_pm = _names_in(generated_pm, "liste_souscripteurs_scs.docx")
    assert "SEL IRIS" in text_pm  # la PM est listee
    assert "100 parts" in text_pm  # TOTAL = somme PP (60) + PM (40), pas seulement les PP
    _assert_clean(text_pm)


def test_scs5_liste_souscripteurs_parts_president_civilite(tmp_path: Path) -> None:
    # SCS5 (Albane 2026-06-25 ; arbitrage Rafael « adapte les termes a la SCS (parts/president) ») :
    # la SCS genere une LISTE DES SOUSCRIPTEURS adaptee — « parts » (jamais « actions »),
    # « President » conserve, civilite civile M./Mme (SP2), une ligne de table par associe + TOTAL.
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
        ],
    )
    generated = css.generate_dossier(payload, tmp_path / "scs-souscr")
    names = {p.name for p in generated.docx_paths}
    assert "liste_souscripteurs_scs.docx" in names
    text = _names_in(generated, "liste_souscripteurs_scs.docx")
    # Termes SCS : « parts », jamais « actions ».
    assert "souscription de parts" in text
    assert "parts souscrites" in text
    assert "actions" not in text
    # President conserve (Rafael), civilite civile M./Mme (SP2), pas « Docteur ».
    assert "Président" in text
    assert "Monsieur Jean Durand" in text
    assert "Monsieur Alice Martin" in text
    assert "Docteur" not in text
    # Table dynamique : 1 ligne par associe + TOTAL des parts.
    assert "100 parts" in text
    _assert_clean(text)


def test_scs_sans_associe_marie_ne_genere_pas_doc005_006(tmp_path: Path) -> None:
    # SCS4 non-regression : une SCS sans associe marie sous communaute NE genere PAS
    # le couple regime (bundle de base inchange).
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire"),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert "DOC-005" not in plan.document_codes
    assert "DOC-006" not in plan.document_codes
    generated = css.generate_dossier(payload, tmp_path / "scs-no-regime")
    names = {p.name for p in generated.docx_paths}
    assert not (_REGIME_DOCS & names)


def test_sci_iris_slice_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-021", "DOC-001", "DOC-002", "DOC-003", "DOC-004")
    generated = css.generate_dossier(payload, tmp_path / "iris")
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI IRIS EXEMPLE.docx"})


def test_sci_iris_pv_forme_is_sci_not_internal_key(tmp_path: Path) -> None:
    # Coherence 2026-06-22 : l'en-tete du PV affichait « SCI IRIS » (cle interne) comme
    # forme au lieu de la forme reelle. Il dit desormais « SCI » ; la denomination reste
    # « SCI IRIS EXEMPLE ».
    payload = _civil_base(
        "SCI IRIS", "sci_iris", [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)]
    )
    generated = css.generate_dossier(payload, tmp_path / "iris-pv")
    pv = next(p for p in generated.docx_paths if p.name == "pv_nomination_dirigeant.docx")
    paras = [par.text.strip() for par in Document(pv).paragraphs if par.text.strip()]
    assert "SCI" in paras
    assert "SCI IRIS" not in paras  # plus la cle interne comme ligne de forme


def test_sci_standard_allows_personne_morale(tmp_path: Path) -> None:
    # Ratifie Rafael 2026-06-08 : une SCI classique peut avoir une societe comme
    # associee (SCI -> micro-holding -> SPFPL). Le moteur ne bloque plus ; l'identite
    # morale est rendue proprement (meme machinerie que SCM / SCI IRIS).
    payload = _civil_base(
        "SCI",
        "sci",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-020", "DOC-001", "DOC-002", "DOC-003", "DOC-004")
    generated = css.generate_dossier(payload, tmp_path / "sci_pm")
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI EXEMPLE.docx"})
    statuts_text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "Statuts SCI EXEMPLE.docx")
    )
    # L'identite morale de l'associe societe apparait dans les statuts.
    assert "SEL IRIS" in statuts_text


def test_civil_blocks_incoherent_parts_sum() -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 50, 41, 90, 500)],
    )
    plan = css.build_civil_plan(payload)
    # KAN-2 : une somme de parts incohérente ne BLOQUE PLUS -> avertissement (le document reflète
    # la saisie ; la cohérence est signalée mais n'empêche pas la génération).
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("Somme des parts" in w for w in plan.warnings)


# --- A2 : preuve multi-associes N (3 et 5) ------------------------------------
# Le socle multi-N est DEJA cable (repeater 1..6, generateurs iterent sur la
# liste d'associes). Ces tests PROUVENT qu'il tient a 3 et 5 ; ils ne corrigent
# rien. Si l'un casse -> bug reel a remonter, pas a patcher a l'aveugle.


def _names_in(generated, docx_name: str) -> str:
    return _docx_text(next(p for p in generated.docx_paths if p.name == docx_name))


def test_sci_three_associes_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [
            _pp("Jean", "Durand", 40, 1, 40, 400),
            _pp("Alice", "Martin", 30, 41, 70, 300),
            _pp("Paul", "Petit", 30, 71, 100, 300),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-020", "DOC-001", "DOC-002", "DOC-003", "DOC-004")
    generated = css.generate_dossier(payload, tmp_path / "sci3")
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI EXEMPLE.docx"})
    # Les 3 associes apparaissent dans les statuts ET sont listes au PV nomination.
    statuts_text = _names_in(generated, "Statuts SCI EXEMPLE.docx")
    pv_text = _names_in(generated, "pv_nomination_dirigeant.docx")
    for nom in ("Durand", "Martin", "Petit"):
        assert nom in statuts_text
        assert nom in pv_text


def test_sci_five_associes_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCI",
        "sci",
        [
            _pp("Jean", "Durand", 20, 1, 20, 200),
            _pp("Alice", "Martin", 20, 21, 40, 200),
            _pp("Paul", "Petit", 20, 41, 60, 200),
            _pp("Marie", "Robert", 20, 61, 80, 200),
            _pp("Luc", "Bernard", 20, 81, 100, 200),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    generated = css.generate_dossier(payload, tmp_path / "sci5")
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI EXEMPLE.docx"})
    statuts_text = _names_in(generated, "Statuts SCI EXEMPLE.docx")
    for nom in ("Durand", "Martin", "Petit", "Robert", "Bernard"):
        assert nom in statuts_text


def test_scs_three_associes_generates_clean(tmp_path: Path) -> None:
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 40, 1, 40, 400, role="commandite"),
            _pp("Alice", "Martin", 30, 41, 70, 300, role="commanditaire"),
            _pp("Paul", "Petit", 30, 71, 100, 300, role="commanditaire"),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-019", "DOC-001", "DOC-002", "DOC-003", "DOC-004", "DOC-LSS-SCS"
    )
    generated = css.generate_dossier(payload, tmp_path / "scs3")
    _assert_bundle_clean(
        generated, _TRONC_DOCS | {"Statuts SCS EXEMPLE.docx", "liste_souscripteurs_scs.docx"}
    )
    statuts_text = _names_in(generated, "Statuts SCS EXEMPLE.docx")
    for nom in ("Durand", "Martin", "Petit"):
        assert nom in statuts_text


def test_sci_iris_three_associes_generates_clean(tmp_path: Path) -> None:
    # 1 personne morale (contrainte IRIS) + 2 personnes physiques.
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [
            _pm(40, 1, 40, 400),
            _pp("Alice", "Martin", 30, 41, 70, 300),
            _pp("Paul", "Petit", 30, 71, 100, 300),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-021", "DOC-001", "DOC-002", "DOC-003", "DOC-004")
    generated = css.generate_dossier(payload, tmp_path / "iris3")
    # Les groupes de resultat IRIS sont produits sans token residuel pour 3 plages.
    _assert_bundle_clean(generated, _TRONC_DOCS | {"Statuts SCI IRIS EXEMPLE.docx"})


def test_scm_three_associes_drops_satellites(tmp_path: Path) -> None:
    # Les satellites SCM (pacte DOC-026 + liste depenses DOC-030) sont verrouilles
    # a EXACTEMENT 2 associes (modeles batis pour 2). A 3, ils disparaissent ; le
    # bundle de base reste genere proprement. Generaliser >2 = FLAG Rafael.
    payload = _civil_base(
        "SCM",
        "scm",
        [
            _pp("Jean", "Durand", 40, 1, 40, 400),
            _pp("Alice", "Martin", 30, 41, 70, 300),
            _pp("Paul", "Petit", 30, 71, 100, 300),
        ],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True
    assert "DOC-026" not in plan.document_codes
    assert "DOC-030" not in plan.document_codes
    assert plan.document_codes == (
        "DOC-025",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
    )
    generated = css.generate_dossier(payload, tmp_path / "scm3")
    names = {p.name for p in generated.docx_paths}
    assert "pacte_associes_scm.docx" not in names
    assert "liste_depenses_communes_scm.docx" not in names
    _assert_bundle_clean(
        generated,
        _TRONC_DOCS | {"Statuts SCM EXEMPLE.docx", "demande_inscription_ordre.docx"},
    )


# --- A3/A4 civils : selection du gerant + DNC sous le gerant -------------------


def test_civil_gerant_selectable_via_index() -> None:
    # Le gerant civil peut etre un autre associe physique que le 1er.
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload["gerant_index"] = 1  # Alice Martin gerante
    ctx = css.build_generation_context(payload)
    assert ctx.personne_signataire.nom == "Martin"


def test_civil_gerant_defaults_to_first_physique_skipping_morale() -> None:
    # SCI IRIS : associe 0 = personne morale -> le gerant par defaut = 1er associe
    # PHYSIQUE (index 1), jamais la personne morale.
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    ctx = css.build_generation_context(payload)  # pas de gerant_index
    assert ctx.personne_signataire.nom == "Martin"


def test_civil_gerant_index_morale_falls_back() -> None:
    # Un gerant_index pointant la personne morale retombe sur le 1er physique.
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    payload["gerant_index"] = 0  # pointe la personne morale -> fallback index 1
    ctx = css.build_generation_context(payload)
    assert ctx.personne_signataire.nom == "Martin"


def _sas_payload():
    return {
        "denomination": "SPFPL MARTIN",
        "siege": "10 rue de la Paix, 75002 Paris",
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "capital_social": "12000",
        "nb_actions_total": 120,
        "valeur_nominale_action": "100",
        "apports_nature_montant": "10000",
        "apports_numeraire_montant": "2000",
        "civilite": "Docteur",
        "prenom": "Camille",
        "nom": "Martin",
        "genre": Gender.MASCULIN,
        "qualification_principale": "Medecin cardiologue",
        "date_naissance": "2 janvier 1980",
        "date_naissance_iso": date(1980, 1, 2),
        "ville_naissance": "Paris",
        "departement_naissance": "75",
        "nationalite": "francaise",
        "regime_matrimonial": "la communaute legale",
        "adresse": "5 rue Royale, 75008 Paris",
        "adresse_num": "5",
        "adresse_voie": "rue Royale",
        "adresse_cp": "75008",
        "adresse_ville": "Paris",
        "nom_pere": "Pierre Martin",
        "nom_mere": "Anne Martin",
        "conjoint_civilite": "Madame",
        "conjoint_prenom": "Alice",
        "conjoint_nom": "Martin",
        "ordre_departement": "Paris",
        "numero_ordre": "12345",
        "numero_rpps": "10000000001",
        "cible_denomination": "SELARL CABINET MARTIN",
        "cible_forme": "SELARL",
        "cible_siege": "12 avenue des Ternes, 75017 Paris",
        "cible_ville_rcs": "Paris",
        "cible_numero_rcs": "900 000 001",
        "apport_nb_parts": 50,
        "banque_nom": "BANQUE EXEMPLE",
        "signature_lieu": "Paris",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 decembre",
        "date_cloture": "31 decembre 2026",
        "signature_date": date(2026, 5, 14),
    }


def test_sas_slice_generates_clean(tmp_path: Path) -> None:
    payload = _sas_payload()
    plan = sas_slice.build_sas_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-015", "DOC-001", "DOC-002", "DOC-003", "DOC-024", "DOC-023")
    generated = sas_slice.generate_dossier(payload, tmp_path / "sas")
    _assert_bundle_clean(
        generated,
        {
            "Statuts SPFPL MARTIN.docx",
            "declaration_non_condamnation.docx",
            "autorisation_domiciliation.docx",
            "procuration.docx",
            "attestation_capital_liste_souscripteurs_sas.docx",
            "pv_remuneration_president.docx",
        },
    )


def test_sas_bundle_celibataire_sans_fuite(tmp_path: Path) -> None:
    # R0702-02 / Akainu n2 : bundle SAS COMPLET pour un actionnaire CELIBATAIRE (menu complet). La
    # comparification repliee -> AUCUN document du bundle ne fuit « <statut> avec … », marqueur
    # « À COMPLÉTER » ni token residuel ; validation non bloquante (regime/conjoint non exiges).
    payload = _sas_payload()
    payload["situation_maritale"] = "Célibataire"
    payload["regime_matrimonial"] = ""
    payload["conjoint_civilite"] = ""
    payload["conjoint_prenom"] = ""
    payload["conjoint_nom"] = ""
    plan = sas_slice.build_sas_plan(payload)
    assert plan.can_generate is True
    assert not any("conjoint" in b.lower() or "Regime matrimonial" in b for b in plan.blockers)
    generated = sas_slice.generate_dossier(payload, tmp_path / "sas-celib")
    seen = False
    for path in generated.docx_paths:
        text = _docx_text(path)
        assert "À COMPLÉTER" not in text, f"marqueur À COMPLÉTER dans {path.name}"
        assert (
            "Célibataire avec" not in text
        ), f"« Célibataire avec » (conjoint fantome) dans {path.name}"
        assert "[" not in text and "]" not in text, f"token residuel dans {path.name}"
        if "Célibataire" in text:
            seen = True
    assert seen, "« Célibataire » absent de tout le bundle SAS"


def test_sas_live03_accentuates_exercice_months(tmp_path: Path) -> None:
    """LIVE-03 : un mois saisi sans accent (« 31 decembre ») ressort accentue.

    Les champs de date d'exercice sont des text_input LIBRES : une saisie sans accent
    partirait verbatim dans le DOCX. La re-accentuation se fait EN AMONT (front), pas dans
    le generateur (echo fidele du modele). Le payload SAS feed « 31 decembre » / « 31
    decembre 2026 » -> la sortie DOCX doit contenir « décembre » accentue, jamais « decembre ».
    """
    payload = dict(_sas_payload())
    # debut accentue comme fin/cloture (oubli releve par re-Akainu T4) : « 1er aout » -> « août ».
    payload["exercice_debut"] = "1er aout"
    payload["exercice_fin"] = "31 decembre"
    payload["date_cloture"] = "31 decembre 2026"
    generated = sas_slice.generate_dossier(payload, tmp_path / "sas_live03")
    full_text = "\n".join(_docx_text(p) for p in generated.docx_paths)
    assert "décembre" in full_text, "LIVE-03 : le mois doit ressortir accentue"
    assert "août" in full_text, "LIVE-03 : le debut d'exercice doit ressortir accentue"
    assert (
        "31 decembre" not in full_text
    ), "LIVE-03 : aucune occurrence sans accent ne doit subsister"
    assert (
        "1er aout" not in full_text
    ), "LIVE-03 : le debut d'exercice sans accent ne doit pas subsister"


def test_sas_live03_accentuates_date_naissance(tmp_path: Path) -> None:
    """LIVE-03 (date_naissance, SAS) : la date de naissance a saisie LIBRE
    (« 1er aout 1980 ») ressort accentuee dans les statuts SAS. Verrouille la
    re-accentuation EN AMONT (sas_slice.build_generation_context), le generateur
    restant un echo fidele."""
    payload = dict(_sas_payload())
    payload["date_naissance"] = "1er aout 1980"
    generated = sas_slice.generate_dossier(payload, tmp_path / "sas_naissance_live03")
    full_text = "\n".join(_docx_text(p) for p in generated.docx_paths)
    assert "août 1980" in full_text, "LIVE-03 : le mois de naissance doit ressortir accentue"
    assert "1er aout 1980" not in full_text, "LIVE-03 : la saisie sans accent ne doit pas subsister"


def test_spfpl_live03_accentuates_date_naissance(tmp_path: Path) -> None:
    """LIVE-03 (date_naissance, SPFPL) : meme garde que SAS pour les statuts SPFPL.
    Le champ accepte un libelle textuel (« 2 decembre 1979 ») -> mois accentue en
    sortie ; une saisie ISO resterait intacte (aucun nom de mois)."""
    payload = dict(_spfpl_payload("SPFPL cession"))
    payload["date_naissance"] = "2 decembre 1979"
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl_naissance_live03")
    full_text = "\n".join(_docx_text(p) for p in generated.docx_paths)
    assert "décembre 1979" in full_text, "LIVE-03 : le mois de naissance SPFPL doit etre accentue"
    assert "2 decembre 1979" not in full_text, "LIVE-03 : la saisie sans accent doit disparaitre"


def test_scm_live03_accentuates_cloture_month(tmp_path: Path) -> None:
    """LIVE-03 (type civil) : la cloture « 31 decembre 2026 » saisie sans accent ressort
    accentuee dans le DOCX civil. Verrouille la re-accentuation EN AMONT cote
    civil_statuts_slice (date_cloture_premier_exercice), independamment de SAS."""
    payload = _civil_base(
        "SCM",
        "scm",
        [_pp("Jean", "Durand", 60, 1, 60, 600), _pp("Alice", "Martin", 40, 61, 100, 400)],
    )
    payload["date_cloture_premier_exercice"] = "31 decembre 2026"
    generated = css.generate_dossier(payload, tmp_path / "scm_live03")
    full_text = "\n".join(_docx_text(p) for p in generated.docx_paths)
    assert "décembre" in full_text, "LIVE-03 : le mois doit ressortir accentue (civil)"
    assert (
        "31 decembre" not in full_text
    ), "LIVE-03 : aucune occurrence sans accent ne doit subsister (civil)"


def test_scm_inter_sel_live03_accentuates_three_free_dates(tmp_path: Path) -> None:
    """LIVE-03 (BLOQUANT re-Akainu T4) : les 3 dates inter-SEL a saisie libre
    (date_effet / date_fin_gestion / date_attribution) ressortent accentuees dans
    contrat_frais_communs.docx et reglement_interieur_scm.docx. On force les 3 mois
    accentuables (aout / fevrier / decembre) saisis SANS accent."""
    payload = _civil_base(
        "SCM",
        "scm",
        [_pp("Jean", "Durand", 50, 1, 50, 500), _pp("Alice", "Martin", 50, 51, 100, 500)],
    )
    payload.update(_SCM_INTER_SEL_INPUTS)
    payload["inter_sel_date_effet"] = "1er aout 2027"
    payload["inter_sel_date_fin_gestion"] = "28 fevrier 2028"
    payload["inter_sel_date_attribution"] = "31 decembre 2027"

    generated = css.generate_dossier(payload, tmp_path / "scm-inter-sel-live03")
    targets = {
        p.name: _docx_text(p)
        for p in generated.docx_paths
        if p.name in {"contrat_frais_communs.docx", "reglement_interieur_scm.docx"}
    }
    assert set(targets) == {"contrat_frais_communs.docx", "reglement_interieur_scm.docx"}
    contrat = targets["contrat_frais_communs.docx"]
    reglement = targets["reglement_interieur_scm.docx"]
    # date_effet -> contrat_frais_communs
    assert "août" in contrat
    assert "aout" not in contrat
    # date_fin_gestion + date_attribution -> reglement_interieur_scm
    assert "février" in reglement
    assert "décembre" in reglement
    assert "fevrier" not in reglement
    assert "decembre" not in reglement


def test_sas_apports_sum_mismatch_warns_not_blocks() -> None:
    # KAN-2 (Rafael, rejeté 2×) : « tous les documents doivent pouvoir être générés, même sans
    # aucun champ ». La génération n'est PLUS JAMAIS bloquée -> l'incohérence somme apports !=
    # capital devient un AVERTISSEMENT (à corriger à la main), plus un blocage. Miroir du contrat
    # spfpl (build_spfpl_plan : can_generate=True, blockers=(), tout écart -> warnings).
    payload = dict(_sas_payload())
    payload["apports_numeraire_montant"] = "999999"  # 10000 + 999999 != 12000
    plan = sas_slice.build_sas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("apports" in w.lower() and "capital" in w.lower() for w in plan.warnings)


def test_sas_civilite_genre_incoherent_warns_not_blocks() -> None:
    # KAN-2 (Rafael, rejeté 2×) : plus AUCUN blocage. Une civilité genrée incohérente avec le genre
    # (« Madame » + masculin) reste SIGNALÉE en avertissement (à corriger à la main), mais ne bloque
    # plus la génération. « Docteur » reste neutre.
    payload = dict(_sas_payload())
    payload["civilite"] = "Madame"
    payload["genre"] = Gender.MASCULIN
    plan = sas_slice.build_sas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("incoherent" in w.lower() for w in plan.warnings)


def _spfpl_payload(structure):
    operation = "apport" if "apport" in structure else "cession"
    return {
        "structure": structure,
        "operation": operation,
        "is_apport": operation == "apport",
        "cession_data": {
            "nb_cedees": 60,
            "prix_unitaire": "1000",
            "plage_cedee": "41 a 100",
            "cible_forme_complete": "societe d'exercice liberal a responsabilite limitee",
            "cible_siege_num": "12",
            "cible_siege_voie": "avenue des Ternes",
            "cible_siege_cp": "75017",
            "cible_siege_ville": "Paris",
            "associes": [
                {
                    "civilite": "Docteur", "prenom": "Camille", "nom": "Martin",
                    "avant": 70, "apres": 10, "plage": "1 a 10",
                },
                {
                    "civilite": "Docteur", "prenom": "Louise", "nom": "Bernard",
                    "avant": 30, "apres": 30, "plage": "11 a 40",
                },
            ],
        },
        "denomination": "SPFPL MARTIN",
        "siege": "10 rue de la Paix, 75002 Paris",
        "capital_social": "60000",
        "nb_actions_total": 600,
        "valeur_nominale_action": "100",
        # §14.2 : civilite = civilite CIVILE (M./Mme) ; le titre pro « Docteur »
        # est porte par titre_affichage (automatique pour une SPFPL dentiste).
        "civilite": "Monsieur",
        "titre_affichage": "Docteur",
        "prenom": "Camille",
        "prenoms": "Camille Andre",
        "nom": "Martin",
        "genre": Gender.MASCULIN,
        # §6.6 : profession SAISIE de l'associe unique (defaut UI « chirurgien-dentiste »). KAN-2 /
        # M3 : plus de repli hardcode cote moteur -> un dossier COMPLET la fournit explicitement
        # (sinon marqueur « (À COMPLÉTER : profession de l'associé unique) », comportement teste
        # a part). Requis pour un bundle nominal propre (associe unique = cedant/president).
        "profession_associe_unique": "chirurgien-dentiste",
        "date_naissance": "02/01/1980",
        "ville_naissance": "Paris",
        "departement_naissance": "75",
        "nationalite": "francaise",
        "regime_matrimonial": "la communaute legale",
        "adresse": "5 rue Royale, 75008 Paris",
        "adresse_num": "5",
        "adresse_voie": "rue Royale",
        "adresse_cp": "75008",
        "adresse_ville": "Paris",
        "nom_pere": "Pierre Martin",
        "nom_mere": "Anne Martin",
        "decision_date": date(2026, 5, 14),
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "ville_rcs": "Paris",
        "conjoint_civilite": "Madame",
        "conjoint_prenom": "Alice",
        "conjoint_nom": "Martin",
        "ordre_departement": "Paris",
        "numero_ordre": "12345",
        "numero_rpps": "10000000001",
        "ordre_conseil": "Conseil departemental",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "ordre_cp": "75008",
        "ordre_ville": "Paris",
        "banque_nom": "BANQUE EXEMPLE",
        "banque_adresse": "1 boulevard Haussmann, 75009 Paris",
        "apport_montant": "60000",
        "apport_nb_parts": 60,
        "apport_plage": "41 a 100",
        "apport_valeur_globale": "60000",
        "cible_denomination": "SELARL CABINET MARTIN",
        "cible_siege": "12 avenue des Ternes, 75017 Paris",
        "cible_ville_rcs": "Paris",
        "cible_numero_rcs": "900 000 001",
        "cible_forme": "SELARL",
        "cible_profession": "chirurgien-dentiste",
        "cible_capital": "10000",
        "cible_nb_parts": 100,
        "cible_valeur_part": "100",
        # Operation apport (DOC-041/042/043) : detail des titres + organes de controle.
        "apport_nature_titres": "parts sociales",
        "apport_valeur_par_titre": "1000",
        "commissaire_denomination": "CAA EXPERTISE",
        "commissaire_forme": "SAS",
        "commissaire_capital": "1 000 euros",
        "commissaire_siege": "1 rue Scheffer, 75016 Paris",
        "commissaire_ville_rcs": "Paris",
        "commissaire_numero_rcs": "948 483 730",
        "commissaire_rep_civilite": "Monsieur",
        "commissaire_rep_prenom": "Nabil",
        "commissaire_rep_nom": "Saidi",
        "evaluateur_denomination": "EVAL CONSEIL",
        "evaluateur_forme": "SAS",
        "evaluateur_capital": "1 000 euros",
        "evaluateur_siege": "1 rue Scheffer, 75016 Paris",
        "evaluateur_ville_rcs": "Paris",
        "evaluateur_numero_rcs": "948 483 730",
        "evaluateur_rep_civilite": "Madame",
        "evaluateur_rep_prenom": "Eva",
        "evaluateur_rep_nom": "Lemoine",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 decembre",
        "date_cloture": "31 decembre 2026",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 14),
    }


_SPFPL_BUNDLE_TRONC = {
    "declaration_non_condamnation.docx",
    "autorisation_domiciliation.docx",
    "procuration.docx",
    "pv_nomination_dirigeant.docx",
    "demande_inscription_ordre.docx",
}

# Documents d'operation apport (DOC-037/041/042/043) ajoutes au bundle de creation.
# Rafael 2026-07-09 : la note d'information (DOC-037) porte cession ET apport — elle
# manquait au bundle apport.
_SPFPL_APPORT_DOCS = {
    "note_information.docx",
    "contrat_apport_spfpl.docx",
    "attestation_capital_liste_souscripteurs.docx",
    "attestation_commissaire_apports.docx",
}

# Documents d'operation cession : note d'info (DOC-037) + PV d'agrement plusieurs
# associes (DOC-039, car le payload de test a 2 associes cible) + acte (DOC-040).
_SPFPL_CESSION_DOCS = {
    "note_information.docx",
    "pv_agrement_cession_spfpl_plusieurs_associes.docx",
    "acte_cession_parts_spfpl.docx",
    # Retour Albane 11 : attestation capital / liste des souscripteurs, variante cession
    # (capital en numeraire). Le bundle cession n'en produisait aucune (DOC-051).
    "attestation_capital_liste_souscripteurs_cession.docx",
}


def test_spfpl_cession_slice_generates_clean(tmp_path: Path) -> None:
    payload = _spfpl_payload("SPFPL cession")
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    # Creation + documents d'operation cession (note + PV agrement plusieurs + acte +
    # attestation capital cession DOC-051, retour Albane 11).
    assert plan.document_codes == (
        "DOC-035",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-037",
        "DOC-039",
        "DOC-040",
        "DOC-051",
    )
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-cession")
    _assert_bundle_clean(
        generated,
        _SPFPL_BUNDLE_TRONC | {"Statuts SPFPL MARTIN.docx"} | _SPFPL_CESSION_DOCS,
    )


@pytest.mark.parametrize(
    ("genre", "civilite", "attendu"),
    [
        (Gender.MASCULIN, "Monsieur", "marié"),
        (Gender.FEMININ, "Madame", "mariée"),
    ],
)
def test_spfpl_cession_acte_situation_cedant_accentuee(
    tmp_path: Path, genre, civilite, attendu
) -> None:
    # O24-11 (MINEUR 2b) : l'acte de cession SPFPL rend cedant.situation_maritale
    # verbatim. Le slice posait « marie » nu (non accentue, non accorde) -> on prouve
    # bout-en-bout que la situation ressort ACCENTUEE et accordee au genre, sur les
    # deux genres, et qu'aucun « marie » nu ne fuit.
    payload = _spfpl_payload("SPFPL cession")
    payload["genre"] = genre
    payload["civilite"] = civilite
    generated = spfpl_slice.generate_dossier(payload, tmp_path / f"spfpl-cession-{attendu}")
    acte = next(
        path for path in generated.docx_paths
        if "acte_cession_parts_spfpl" in path.name
    )
    acte_text = _docx_text(acte)
    assert f"{attendu} avec Madame Alice Martin" in acte_text
    # Le « marie » NON accentue (statut nu, suivi d'une virgule ou « avec ») ne doit
    # plus apparaitre. NB : le reste de la phrase reste un echo fidele non accentue
    # (« ne le », « a Paris ») — seul le statut matrimonial est accorde/accentue.
    assert "marie avec" not in acte_text
    assert "marie sous" not in acte_text


@pytest.mark.parametrize(
    ("genre", "civilite", "statut"),
    [(Gender.MASCULIN, "Monsieur", "marié"), (Gender.FEMININ, "Madame", "mariée")],
)
@pytest.mark.parametrize(
    ("regime", "regime_affiche"),
    [
        ("la communaute legale", "la communauté légale"),
        ("la separation de biens", "la séparation de biens"),
        ("la communaute universelle", "la communauté universelle"),
        ("la participation aux acquets", "la participation aux acquêts"),
    ],
)
def test_spfpl_cession_comparution_marie_ligne_complete(
    tmp_path: Path, genre, civilite, statut, regime, regime_affiche
) -> None:
    # Akainu m1 (2026-07-02) : byte-fidelite du cas MARIE sur les 4 REGIMES x 2 GENRES.
    # La comparution des statuts de cession rend EXACTEMENT « <statut accorde> sous le régime
    # de <regime> avec <conjoint> » (wording propre a la SPFPL : « avec », pas « époux/épouse
    # de »). Assertion de ligne COMPLETE (pas une sous-chaine d'un seul regime/genre).
    # M1 (Akainu 2026-07-06) : 7.2 exige une MAJUSCULE en tete de chaque element du soussigne
    # -> la ligne matrimoniale commence desormais par « Marié »/« Mariée » (capitalisee).
    # B1 (Akainu doc-entier 2026-07-09) : le regime sort ACCENTUE (« la communauté légale »…),
    # avec une SEULE preposition « de » (plus l'echo brut non accentue du preset).
    payload = _spfpl_payload("SPFPL cession")
    payload["genre"] = genre
    payload["civilite"] = civilite
    payload["regime_matrimonial"] = regime
    slug = f"{statut}-{regime.split()[-1]}"
    generated = spfpl_slice.generate_dossier(payload, tmp_path / f"spfpl-cession-marie-{slug}")
    statuts = next(
        p for p in generated.docx_paths if p.name == "Statuts SPFPL MARTIN.docx"
    )
    text = _docx_text(statuts)
    statut_capitalise = statut[0].upper() + statut[1:]  # « marié » -> « Marié »
    expected = f"{statut_capitalise} sous le régime de {regime_affiche} avec Madame Alice Martin"
    assert expected in text, f"comparution mariee attendue absente ({slug}) : {expected!r}"
    # M1 : jamais la version minuscule en tete de la ligne matrimoniale.
    assert f"\n{statut} sous le régime" not in text
    # B1 : jamais la preposition doublee ni le regime non accentue.
    assert "sous le régime de la communaute" not in text
    assert "de regime de" not in text


@pytest.mark.parametrize("structure", ["SPFPL cession", "SPFPL apport"])
def test_spfpl_comparution_celibataire_sans_conjoint(tmp_path: Path, structure) -> None:
    # Akainu m2 (round 1) + M1/m1 (round 2, 2026-07-02) : la SPFPL utilise le menu complet
    # « Situation matrimoniale » — sur les DEUX operations (cession ET apport), servies par le
    # meme formulaire. Un CELIBATAIRE rend juste « célibataire » dans la comparution, SANS
    # « avec <conjoint> » ni marqueur « À COMPLÉTER » residuel, dans AUCUN document du bundle
    # (scan exhaustif ET sur les 2 operations — sinon siloing : l'apport fuyait alors que la
    # cession etait propre, d'ou 840 verts trompeurs au round 1).
    payload = _spfpl_payload(structure)
    payload["situation_maritale"] = "célibataire"
    payload["regime_matrimonial"] = ""
    payload["conjoint_civilite"] = ""
    payload["conjoint_prenom"] = ""
    payload["conjoint_nom"] = ""
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert not any(
        "conjoint" in b.lower() or "Regime matrimonial" in b for b in plan.blockers
    )
    slug = structure.split()[-1]
    generated = spfpl_slice.generate_dossier(payload, tmp_path / f"spfpl-{slug}-celib")
    seen_celibataire = False
    for path in generated.docx_paths:
        text = _docx_text(path)
        assert "À COMPLÉTER" not in text, f"marqueur À COMPLÉTER dans {path.name} ({structure})"
        assert (
            "célibataire avec" not in text
        ), f"« célibataire avec » (conjoint fantome) dans {path.name} ({structure})"
        assert "[" not in text and "]" not in text, f"token residuel dans {path.name} ({structure})"
        if "célibataire" in text:
            seen_celibataire = True
    assert seen_celibataire, f"« célibataire » absent de tout le bundle {structure}"


def test_spfpl_form_situation_menu_complet_et_conjoint_conditionnel() -> None:
    # Retour Rafael 2026-07-02 : le formulaire SPFPL expose DESORMAIS le menu complet
    # « Situation matrimoniale » (comme les autres types), plus l'ancien « Regime
    # matrimonial » restreint aux regimes maries ; les champs conjoint sont
    # conditionnels (affiches uniquement pour un marie).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app.field_derivations import MATRIMONIAL_STATUS_PRESETS

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SPFPL dentistes — cession")
    app = app.run(timeout=120)

    situation_box = next(
        s for s in app.selectbox if str(s.key) == "spfpl_cession_situation"
    )
    assert situation_box.label == "Situation matrimoniale"
    assert list(situation_box.options) == list(MATRIMONIAL_STATUS_PRESETS)
    # Plus d'ancien selecteur « Regime matrimonial » (menu married-only).
    assert not any(str(s.key) == "spfpl_cession_regime_label" for s in app.selectbox)
    assert "Regime matrimonial" not in [s.label for s in app.selectbox]

    # Defaut = Celibataire (1er preset) -> champs conjoint MASQUES.
    conjoint_keys = {str(w.key) for w in app.text_input if "conjoint" in str(w.key)}
    assert "spfpl_cession_conjoint_prenom" not in conjoint_keys
    assert "spfpl_cession_conjoint_nom" not in conjoint_keys

    # Choisir un statut MARIE -> les champs conjoint APPARAISSENT.
    situation_box.set_value("Marié(e) sous le régime légal / communauté")
    app = app.run(timeout=120)
    conjoint_keys = {str(w.key) for w in app.text_input if "conjoint" in str(w.key)}
    assert "spfpl_cession_conjoint_prenom" in conjoint_keys
    assert "spfpl_cession_conjoint_nom" in conjoint_keys


def test_sas_form_situation_menu_complet_et_conjoint_conditionnel() -> None:
    # R0702-02 (Gad 2026-07-02) : le formulaire SAS (SPFPL medecins forme SAS) expose DESORMAIS le
    # MEME menu complet « Situation matrimoniale » que les autres types, plus l'ancien champ regime
    # en TEXTE LIBRE ni le « marie » verrouille. Champs conjoint conditionnels (marie uniquement).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app.field_derivations import MATRIMONIAL_STATUS_PRESETS

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SPFPL médecins (forme SAS)")
    app = app.run(timeout=120)

    situation_box = next(s for s in app.selectbox if str(s.key) == "sas_situation")
    assert situation_box.label == "Situation matrimoniale"
    assert list(situation_box.options) == list(MATRIMONIAL_STATUS_PRESETS)
    # Plus d'ancien champ « Regime matrimonial » en texte libre.
    assert not any(str(w.key) == "sas_regime_matrimonial" for w in app.text_input)
    assert "Regime matrimonial (ex: la communaute legale)" not in [w.label for w in app.text_input]

    # Defaut = Celibataire (1er preset) -> champs conjoint MASQUES.
    conjoint_keys = {str(w.key) for w in app.text_input if "conjoint" in str(w.key)}
    assert "sas_conjoint_prenom" not in conjoint_keys
    assert "sas_conjoint_nom" not in conjoint_keys

    # Choisir un statut MARIE -> les champs conjoint APPARAISSENT.
    situation_box.set_value("Marié(e) sous le régime légal / communauté")
    app = app.run(timeout=120)
    conjoint_keys = {str(w.key) for w in app.text_input if "conjoint" in str(w.key)}
    assert "sas_conjoint_prenom" in conjoint_keys
    assert "sas_conjoint_nom" in conjoint_keys


def test_spfpl_cession_missing_cible_forme_generates_with_gap() -> None:
    # KAN-2 (Rafael 2026-07-13) : la forme complete de la cible est un champ TEXTE -> son
    # absence ne BLOQUE plus la generation. Elle sort en « (À COMPLÉTER : …) » (required_text,
    # a completer a la main) et est surfacee en AVERTISSEMENT, jamais en blocage. Seuls les
    # manques STRUCTURELS / NUMERIQUES (nb actions/parts, capital, prix, dates) bloquent encore.
    payload = _spfpl_payload("SPFPL cession")
    payload["cession_data"] = {**payload["cession_data"], "cible_forme_complete": ""}
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("forme" in w.lower() and "cible" in w.lower() for w in plan.warnings)


@pytest.mark.parametrize(
    "flow, field",
    [
        # Champs TEXTE.
        ("SPFPL cession", "nom"),
        ("SPFPL cession", "banque_nom"),
        ("SPFPL cession", "cible_denomination"),
        ("SPFPL cession", "ordre_ville"),
        ("SPFPL apport", "nom"),
        ("SPFPL apport", "nationalite"),
        ("SPFPL apport", "commissaire_denomination"),
        ("SPFPL apport", "evaluateur_denomination"),
        # Champs DATE et NUMÉRIQUES — ceux dont ce lot RETIRE la garde, donc les seuls qui
        # prouvent quelque chose (gate Akainu M2 : les 13 cas d'origine n'étaient que du texte,
        # ils passaient au vert en évitant exactement le périmètre du ticket).
        ("SPFPL cession", "date_naissance"),
        ("SPFPL cession", "signature_date"),
        ("SPFPL cession", "nb_actions_total"),
        ("SPFPL cession", "cible_nb_parts"),
        ("SPFPL cession", "capital_social"),
        ("SPFPL cession", "cession_data.nb_cedees"),
        ("SPFPL cession", "cession_data.prix_unitaire"),
        ("SPFPL cession", "cession_data.associes"),
        ("SPFPL apport", "signature_date"),
        ("SPFPL apport", "nb_actions_total"),
        ("SPFPL apport", "cible_nb_parts"),
        ("SPFPL apport", "apport_nb_parts"),
    ],
)
def test_spfpl_generates_with_single_missing_field(
    flow: str, field: str, tmp_path: Path
) -> None:
    # KAN-2 (Rafael 2026-07-14) : « tous les documents doivent pouvoir être générés, même si je ne
    # remplis aucun champ ». Vider UN SEUL champ — texte, date OU numérique — ne doit JAMAIS faire
    # échouer la génération : la zone sort « (À COMPLÉTER : …) » à compléter à la main, et le
    # livrable ne contient jamais de crochets « [ ] » (règle R1).
    #
    # Gate Akainu : le plan ne doit pas MENTIR. Un `can_generate=True` suivi d'un crash est le
    # défaut exact qui a fait rejeter ce ticket -> chaque cas GÉNÈRE réellement (M3), et les
    # champs date/numériques sont couverts (M2), pas seulement le texte.
    payload = _spfpl_payload(flow)
    if "." in field:  # champ imbriqué du sous-formulaire cession
        parent, enfant = field.split(".", 1)
        bloc = dict(payload[parent])  # type: ignore[arg-type]
        bloc[enfant] = [] if enfant == "associes" else ("" if enfant == "prix_unitaire" else 0)
        payload[parent] = bloc
    else:
        assert _vider_champ(payload, field), f"{field} n'est pas un champ videable"

    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True, f"{flow}/{field} : le plan ne doit jamais bloquer"
    assert plan.blockers == (), f"{flow}/{field} : aucun blocage attendu"
    spfpl_slice.generate_dossier(payload, tmp_path)
    docs = list(tmp_path.rglob("*.docx"))
    assert docs, f"{flow}/{field} : le bundle doit se générer malgré le champ manquant"
    texte = " ".join(p.text for d in docs for p in Document(d).paragraphs)
    assert "[" not in texte and "]" not in texte, f"{flow}/{field} : marqueur à crochets interdit"


# Clés de ROUTAGE : le type de dossier / l'opération, choisis dans un menu — jamais « non
# remplis ». Tout le reste est un champ que l'utilisateur peut laisser vide.
_ROUTAGE_KEYS = {"structure", "operation", "is_apport"}


def _docx_texte_complet(path: Path) -> str:
    """Texte INTÉGRAL d'un .docx : corps + TABLEAUX + en-têtes/pieds.

    Akainu M4 (2026-07-15) : `Document(d).paragraphs` ne descend PAS dans les cellules ni dans
    les sections -> le filet R1 (« aucun crochet livré ») était aveugle aux tableaux de
    répartition, qui sont précisément là où vivent les quantités. Un filet troué ne prouve rien.
    """
    document = Document(path)
    morceaux: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                morceaux.extend(p.text for p in cell.paragraphs)
    for section in document.sections:
        for part in (section.header, section.footer):
            morceaux.extend(p.text for p in part.paragraphs)
    return " ".join(t for t in morceaux if t)


# Akainu B1 : une QUANTITÉ de titres affirmée à zéro dans une PHRASE. On code l'INTENTION
# (« un compte de titres non renseigné ne s'affirme jamais »), pas les 3 tournures qu'Akainu a
# citées — sinon la 4e formulation repasserait au vert (leçon R3 « Docteur », payée 3 fois).
# Le lookbehind écarte la QUEUE d'un nombre plus grand (« 10 parts », « 100 actions »).
# NB — les TABLEAUX de répartition avant/après sont hors périmètre : « 0 » y est une valeur
# RÉELLE (la holding détient 0 part AVANT la cession). D'où le ciblage « 0 <mot> » en toutes
# lettres dans le fil du texte, et non toute cellule valant « 0 ».
_QUANTITE_NULLE = re.compile(
    r"(?<![\d,.])0\s+(?:parts?|actions?|titres?)\b", re.IGNORECASE
)


def _quantites_nulles_affirmees(texte: str) -> list[str]:
    return [m.group(0) for m in _QUANTITE_NULLE.finditer(texte)]


def _vider_champ(payload: dict[str, object], key: str) -> bool:
    """Vide UN champ du payload comme le ferait un utilisateur qui ne le remplit pas.

    KAN-2 / gate Akainu (M1) : les objets `date` doivent être vidés EUX AUSSI. Un premier jet ne
    vidait que str/int/dict/list -> `signature_date` restait valorisée et le test « formulaire
    entièrement vide » esquivait précisément le champ qui faisait crasher. `0` est la valeur d'un
    `number_input(min_value=0)` non rempli ; `None` celle d'une date effacée ou invalide.

    Renvoie False pour les clés de routage et les booléens (une case cochée est toujours dans un
    état) — rien à vider.
    """
    if key in _ROUTAGE_KEYS:
        return False
    value = payload[key]
    if isinstance(value, bool):
        return False
    if isinstance(value, str):
        payload[key] = ""
    elif isinstance(value, date):
        payload[key] = None
    elif isinstance(value, int):
        payload[key] = 0
    elif isinstance(value, dict):
        payload[key] = {"associes": []} if key == "cession_data" else {}
    elif isinstance(value, list):
        payload[key] = []
    else:
        payload[key] = None
    return True


@pytest.mark.parametrize("flow", ["SPFPL cession", "SPFPL apport"])
def test_spfpl_formulaire_entierement_vide_genere(flow: str, tmp_path: Path) -> None:
    # KAN-2 (Rafael 2026-07-14), verbatim : « Tous les documents doivent pouvoir être générés,
    # MÊME SI JE NE REMPLIS AUCUN CHAMP dans le générateur. » C'est LE test du ticket : on ne
    # garde que les clés de ROUTAGE (le type de dossier, choisi dans un menu — jamais vide) et
    # on vide tout le reste. Le bundle complet doit sortir, zones en « (À COMPLÉTER : …) »,
    # sans aucun blocage ni crochet « [ ] ».
    payload = _spfpl_payload(flow)
    for key in list(payload):
        _vider_champ(payload, key)

    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True, f"{flow} : un formulaire vide doit rester générable"
    assert plan.blockers == (), f"{flow} : aucun blocage ne doit subsister"

    spfpl_slice.generate_dossier(payload, tmp_path)
    docs = list(tmp_path.rglob("*.docx"))
    assert len(docs) >= 10, f"{flow} : le bundle complet doit sortir ({len(docs)} docs)"
    texte = " ".join(_docx_texte_complet(d) for d in docs)
    assert "À COMPLÉTER" in texte
    assert "[" not in texte and "]" not in texte
    # Akainu B1 (2026-07-15) : generer sans blocage ne suffit pas — un acte SIGNABLE qui AFFIRME
    # « 0 parts sociales » est FAUX, et pire que le blocage que le client a fait retirer. Une
    # quantite de titres non renseignee doit sortir en marqueur, jamais en zero affirme.
    fautes = _quantites_nulles_affirmees(texte)
    assert not fautes, (
        f"{flow} : quantite(s) de titres affirmee(s) a 0 sur un formulaire VIDE "
        f"(le champ n'est pas 'zero', il est 'non rempli') -> {sorted(set(fautes))[:6]}"
    )


def test_spfpl_missing_field_marks_zone_a_completer(tmp_path: Path) -> None:
    # KAN-2 : le mécanisme de marqueur « (À COMPLÉTER : …) » fonctionne (date de naissance vide ->
    # zone visible dans le PV de nomination), sans jamais de crochets.
    payload = _spfpl_payload("SPFPL cession")
    payload["date_naissance"] = ""
    spfpl_slice.generate_dossier(payload, tmp_path)
    texte = " ".join(
        p.text for d in tmp_path.rglob("*.docx") for p in Document(d).paragraphs
    )
    assert "À COMPLÉTER" in texte
    assert "[" not in texte and "]" not in texte


def test_spfpl_nb_actions_zero_ne_bloque_plus(tmp_path: Path) -> None:
    # KAN-2 (Rafael 2026-07-14) : « Tous les documents doivent pouvoir être générés, MÊME SI je
    # ne remplis AUCUN champ. » -> nb_actions = 0 (la valeur d'un `number_input` non rempli) ne
    # BLOQUE plus ; le manque est surfacé en avertissement.
    #
    # Gate Akainu (M3) : ce test n'assertait QUE le plan — or c'est précisément ce que ce ticket a
    # cassé (le plan annonce « générable » puis la génération lève). Un test de non-blocage qui ne
    # génère pas ne prouve rien -> on génère réellement.
    payload = _spfpl_payload("SPFPL cession")
    payload["nb_actions_total"] = 0
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("actions" in w.lower() and "zero" in w.lower() for w in plan.warnings)
    spfpl_slice.generate_dossier(payload, tmp_path)
    assert list(tmp_path.rglob("*.docx"))


def test_spfpl_apport_slice_generates_clean(tmp_path: Path) -> None:
    payload = _spfpl_payload("SPFPL apport")
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    # L'apport complete le bundle de creation par ses 4 documents d'operation :
    # note d'information DOC-037 (Rafael 2026-07-09 : cession ET apport) + contrat
    # d'apport DOC-041 + attestations capital DOC-042 / commissaire DOC-043.
    assert plan.document_codes == (
        "DOC-036",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-037",
        "DOC-041",
        "DOC-042",
        "DOC-043",
    )
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-apport")
    _assert_bundle_clean(
        generated,
        _SPFPL_BUNDLE_TRONC | {"Statuts SPFPL MARTIN.docx"} | _SPFPL_APPORT_DOCS,
    )
    # Rafael 2026-07-09 : la note d'information (DOC-037) sort en APPORT avec la
    # repartition APRES operation DERIVEE — le fondateur garde (100 - 60) parts sans
    # numerotation inventee, le holding recoit les 60 parts apportees (plage saisie).
    note = _docx_text(
        next(p for p in generated.docx_paths if p.name == "note_information.docx")
    )
    assert "prévoit de recevoir en apport en nature" in note
    assert "Après ledit apport" in note
    assert "Monsieur Camille Martin, titulaire de 40 parts sociales" in note
    assert "SPFPL MARTIN, titulaire de 60 parts sociales, numérotées de 41 à 100" in note


_REGIME_DOCS = {
    "lettre_renonciation_associe.docx",
    "lettre_avertissement_conjoint.docx",
}


def test_spfpl_regime_off_keeps_base_bundle(tmp_path: Path) -> None:
    payload = _spfpl_payload("SPFPL cession")
    payload["regime_communautaire"] = False
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" not in plan.document_codes
    assert "DOC-006" not in plan.document_codes
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-regime-off")
    names = {p.name for p in generated.docx_paths}
    assert not (_REGIME_DOCS & names)


def test_spfpl_cession_regime_on_adds_regime_docs(tmp_path: Path) -> None:
    payload = _spfpl_payload("SPFPL cession")
    payload["regime_communautaire"] = True
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-035",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-005",
        "DOC-006",
        "DOC-037",
        "DOC-039",
        "DOC-040",
        "DOC-051",
    )
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-cession-regime")
    _assert_bundle_clean(
        generated,
        _SPFPL_BUNDLE_TRONC
        | {"Statuts SPFPL MARTIN.docx"}
        | _SPFPL_CESSION_DOCS
        | _REGIME_DOCS,
    )


def test_spfpl_apport_regime_on_adds_regime_docs(tmp_path: Path) -> None:
    payload = _spfpl_payload("SPFPL apport")
    payload["regime_communautaire"] = True
    plan = spfpl_slice.build_spfpl_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" in plan.document_codes
    assert "DOC-006" in plan.document_codes
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-apport-regime")
    _assert_bundle_clean(
        generated,
        _SPFPL_BUNDLE_TRONC | {"Statuts SPFPL MARTIN.docx"} | _SPFPL_APPORT_DOCS | _REGIME_DOCS,
    )
    # Rafael 2026-07-09 : plus JAMAIS de double unite dans les lettres de regime —
    # avertissement conjoint « soixante mille (60 000) euros » (et non « soixante
    # mille euros (60 000) euros ») ; renonciation « 60 000 (soixante mille) euros ».
    avertissement = _docx_text(
        next(p for p in generated.docx_paths if p.name == "lettre_avertissement_conjoint.docx")
    )
    assert "de soixante mille (60 000) euros" in avertissement
    assert "euros (60 000) euros" not in avertissement
    renonciation = _docx_text(
        next(p for p in generated.docx_paths if p.name == "lettre_renonciation_associe.docx")
    )
    assert "60 000 (soixante mille) euros" in renonciation
    assert "(soixante mille euros) euros" not in renonciation


# --- Regressions « variables mal injectees » (audit 2026-06-09) ---------------


def test_spfpl_apport_capital_not_duplicated(tmp_path: Path) -> None:
    # Le capital (art. 8) ne doit plus etre injecte en double (« 60000 € 60000euros »).
    payload = _spfpl_payload("SPFPL apport")
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-apport-cap")
    text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "Statuts SPFPL MARTIN.docx")
    )
    assert "60000euros" not in text  # plus de « euros » colle
    assert "€ 60000" not in text  # plus de montant duplique apres le €
    # Albane 2026-07-07 (statuts SPFPL) : montants GROUPES -> la forme propre est desormais
    # « 60 000 euros » (en-tete), « 60000 » nu n'apparait plus dans les statuts.
    assert "60 000 euros" in text  # forme propre groupee
    assert "60000" not in text


# --- §14.2 : SPFPL dentiste, « Docteur » automatique (hors deroulante civilite) -


def test_spfpl_titre_docteur_automatic_civilite_is_civil() -> None:
    # §14.2 : la civilite portee au contexte est la civilite CIVILE (M./Mme), et le
    # titre pro « Docteur » est applique automatiquement (titre_affichage), sans
    # dependre de la valeur choisie dans la deroulante civilite.
    payload = _spfpl_payload("SPFPL cession")
    assert payload["civilite"] == "Monsieur"  # civilite civile, plus « Docteur »
    ctx = spfpl_slice.build_generation_context(payload)
    assert ctx.personne_signataire.civilite == "Monsieur"
    assert ctx.personne_signataire.titre_affichage == "Docteur"
    # Le dirigeant / president gardent la civilite CIVILE (mutualise SELARL/SELAS).
    assert ctx.dirigeant_nomine.civilite_affichage == "Monsieur"


def test_spfpl_demande_ordre_uses_civil_civilite(tmp_path: Path) -> None:
    # R3 durci (Rafael 2026-07-07, « partout ») — SUPERSEDE §14.2 : la demande
    # d'inscription a l'ordre rend la civilite CIVILE (« Monsieur Camille Martin »),
    # plus le titre « Docteur » (titre_affichage reste porte au contexte mais ne
    # sort plus sur ce document).
    payload = _spfpl_payload("SPFPL cession")
    generated = spfpl_slice.generate_dossier(payload, tmp_path / "spfpl-titre-ordre")
    text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "demande_inscription_ordre.docx")
    )
    assert "Monsieur Camille Martin" in text
    assert "Docteur" not in text
    assert "Dr " not in text


def test_spfpl_form_civilite_drops_docteur_and_genre_selector() -> None:
    # §14.2 : la deroulante civilite ne propose plus « Docteur » (M./Mme seulement)
    # et le selecteur « Genre civil » redondant a disparu.
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SPFPL dentistes — cession")
    app = app.run(timeout=120)

    civilite_box = next(s for s in app.selectbox if str(s.key) == "spfpl_cession_civilite")
    assert "Docteur" not in list(civilite_box.options)
    assert set(civilite_box.options) == {"Monsieur", "Madame"}
    assert not any(str(s.key) == "spfpl_cession_genre_label" for s in app.selectbox)


def test_sas_attestation_no_double_docteur(tmp_path: Path) -> None:
    # Le titre « Docteur » ne doit plus etre injecte en double
    # (« Le Docteur Docteur Camille Martin »).
    payload = _sas_payload()
    generated = sas_slice.generate_dossier(payload, tmp_path / "sas-titre")
    text = _docx_text(
        next(
            p
            for p in generated.docx_paths
            if p.name == "attestation_capital_liste_souscripteurs_sas.docx"
        )
    )
    assert "Docteur Docteur" not in text
    # R3 Rafael 2026-07-07 : Docteur retiré partout (supersede A26-45/49) — la phrase
    # d'apport rend la civilité CIVILE, plus de titre « Le Docteur X a fait ».
    assert "Monsieur Camille Martin a fait la totalité des apports en nature." in text
    assert "Docteur" not in text


def _selas_payload():
    phys = StatutsCivilsAssocie(
        type_personne="personne_physique",
        genre=Gender.FEMININ,
        civilite_affichage="Madame",
        prenom="Claire",
        prenoms="Claire",
        nom="Durand",
        date_naissance="1 janvier 1980",
        ville_naissance="Lyon",
        departement_naissance="69",
        nationalite="française",
        profession="Docteur",
        situation_maritale="célibataire",
        adresse_personnelle_affichee="10 rue Exemple, 69000 Lyon",
        qualification_principale="qualifiée en médecine générale",
        ordre_departemental="Rhône",
        numero_ordre="69-12345",
        numero_rpps="10100000001",
        qualite_capital="associée exerçante",
        nb_actions=75,
        nb_actions_lettres="soixante-quinze",
        # DNC par associe (Rafael 2026-07-09) : filiation requise pour CHAQUE associe.
        nom_pere="Pierre Durand",
        nom_mere="Anne Durand",
        apport=StatutsCivilsApport(montant="750", montant_lettres="sept cent cinquante"),
    )
    mor = StatutsCivilsAssocie(
        type_personne="personne_morale",
        denomination="SOCIETE CIVILE EXEMPLE",
        forme_juridique="Société Civile, à capital variable",
        capital_social="1 020",
        siege=Address(adresse_affichee="2 avenue Pro, 69000 Lyon"),
        ville_rcs="Lyon",
        numero_rcs="900 000 001",
        representant=StatutsCivilsRepresentant(
            civilite_affichage="Monsieur",
            prenom="Paul",
            nom="Martin",
            fonction="gérant",
        ),
        qualite_capital="associée non exerçante",
        nb_actions=25,
        nb_actions_lettres="vingt-cinq",
        apport=StatutsCivilsApport(montant="250", montant_lettres="deux cent cinquante"),
    )
    return {
        "denomination": "SELAS EXEMPLE",
        "siege": "5 place du Centre, 69000 Lyon",
        "siege_num": "5",
        "siege_voie": "place du Centre",
        "siege_cp": "69000",
        "siege_ville": "Lyon",
        "profession_reglementee": "médecin",
        "profession_reglementee_pluriel": "médecins",
        "capital_social": "1000",
        "nb_actions_total": 100,
        "valeur_nominale_action": "10",
        "ville_rcs": "Lyon",
        "adresse_lieu_exercice": "5 place du Centre, 69000 Lyon",
        "banque_nom": "BANQUE EXEMPLE",
        "banque_adresse": "1 rue Banque, 69009 Lyon",
        "date_cloture": "31 décembre 2026",
        "signature_lieu": "Lyon",
        "signature_date": date(2026, 5, 15),
        "associes": [phys, mor],
        # Documents communs : signataire = president (1er physique).
        "signataire_nom_pere": "Pierre Durand",
        "signataire_nom_mere": "Anne Durand",
        "signataire_adresse_num": "10",
        "signataire_adresse_voie": "rue Exemple",
        "signataire_adresse_cp": "69000",
        "signataire_adresse_ville": "Lyon",
        "signataire_nationalite": "française",
        "signataire_titre": "Docteur",
        "signataire_date_naissance": date(1980, 1, 1),
        "ordre_conseil": "Conseil departemental",
        "ordre_departement": "Rhône",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "ordre_cp": "69002",
        "ordre_ville": "Lyon",
        "ordre_numero": "69-12345",
    }


def test_selas_multi_slice_generates_clean(tmp_path: Path) -> None:
    payload = _selas_payload()
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == ("DOC-044", "DOC-001", "DOC-002", "DOC-003", "DOC-004", "DOC-034")
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas")
    _assert_bundle_clean(generated, _SELAS_BUNDLE_NAMES)


def test_selas_dnc_filename_carries_dirigeant_name(tmp_path: Path) -> None:
    # #2 (onglet 24) : la declaration de non-condamnation porte le NOM DU DIRIGEANT
    # (president, ici « Durand ») dans son nom de fichier. L'associe personne MORALE
    # n'a pas de DNC (le document declare une personne physique) -> 1 seule DNC ici.
    payload = _selas_payload()
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-dnc")
    names = {p.name for p in generated.docx_paths}
    assert "declaration_non_condamnation_Durand.docx" in names
    assert "declaration_non_condamnation.docx" not in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


# --- DNC = GERANT UNIQUEMENT (Albane, Direction Juridique, 2026-07-09) ---------
# La declaration de non-condamnation ne se genere QUE pour les GERANTS (en SELAS =
# les dirigeants : President + DG / DG Associe). Un associe NON dirigeant n'a PAS de
# DNC. SUPERSEDE le retour Rafael du matin « 1 DNC par associe » (+ verrou R11) :
# Rafael a confirme « Albane a raison ».


def test_selas_dnc_gerants_uniquement_pas_les_associes(tmp_path: Path) -> None:
    # 3 associes physiques, SEUL le president (Durand, defaut) est dirigeant : les
    # 2 autres associes (Martin, Petit) NON dirigeants ne recoivent PAS de DNC.
    # Payload nu (sans dirigeants_nomines) -> seul le president est gerant.
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 40),
            _selas_phys("Paul", "Martin", 30),
            _selas_phys("Marie", "Petit", 30),
        ]
    )
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-dnc-n")
    names = {p.name for p in generated.docx_paths}
    assert "declaration_non_condamnation_Durand.docx" in names  # president (gerant)
    assert "declaration_non_condamnation_Martin.docx" not in names  # non dirigeant
    assert "declaration_non_condamnation_Petit.docx" not in names  # non dirigeant
    assert "declaration_non_condamnation.docx" not in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


def test_selas_associe_non_dirigeant_sans_filiation_ne_bloque_pas(tmp_path: Path) -> None:
    # Albane 2026-07-09 : un associe physique NON dirigeant sans filiation ne bloque
    # PLUS (il n'a plus de DNC). « Je n'ai volontairement pas complete les parents de
    # mon 2e associe car il n'est pas gerant, mais ca bloque » -> corrige.
    associe_sans_filiation = _selas_phys("Paul", "Martin", 40).model_copy(
        update={"nom_pere": None, "nom_mere": None}
    )
    payload = _selas_payload_n(
        [_selas_phys("Claire", "Durand", 60), associe_sans_filiation]
    )
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True, plan.blockers
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-nofil")
    names = {p.name for p in generated.docx_paths}
    # Seule la DNC du president (gerant) est generee ; Martin non dirigeant -> aucune.
    assert "declaration_non_condamnation_Durand.docx" in names
    assert "declaration_non_condamnation_Martin.docx" not in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


def test_civil_dnc_gerant_uniquement(tmp_path: Path) -> None:
    # SCI a 2 associes physiques -> UNE seule DNC, celle du GERANT (Durand, 1er
    # physique, via l'orchestrateur renommee O24-02). Albane 2026-07-09 : l'associe
    # NON gerant (Martin) ne recoit PAS de DNC (supersede « 1 DNC par associe »).
    payload = _civil_base(
        "SCI",
        "sci",
        [
            _pp("Jean", "Durand", 40, 1, 40, 400),
            _pp("Alice", "Martin", 60, 41, 100, 600),
        ],
    )
    generated = css.generate_dossier(payload, tmp_path / "sci-dnc")
    names = {p.name for p in generated.docx_paths}
    assert "declaration_non_condamnation_Durand.docx" in names  # gerant
    assert "declaration_non_condamnation_Martin.docx" not in names  # non gerant
    assert "declaration_non_condamnation.docx" not in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


def test_civil_dnc_trois_associes_scs_gerant_uniquement(tmp_path: Path) -> None:
    # SCS (commandite gerant + 2 commanditaires) -> UNE seule DNC, celle du gerant
    # (Durand, commandite, 1er physique). Les commanditaires non gerants -> pas de DNC.
    payload = _civil_base(
        "SCS",
        "scs",
        [
            _pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
            _pp("Alice", "Martin", 20, 61, 80, 200, role="commanditaire"),
            _pp("Marc", "Petit", 20, 81, 100, 200, role="commanditaire"),
        ],
    )
    generated = css.generate_dossier(payload, tmp_path / "scs-dnc")
    names = {p.name for p in generated.docx_paths}
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1
    assert "declaration_non_condamnation_Durand.docx" in names  # gerant (commandite)
    assert "declaration_non_condamnation_Martin.docx" not in names
    assert "declaration_non_condamnation_Petit.docx" not in names


def test_civil_dnc_personne_morale_sans_dnc(tmp_path: Path) -> None:
    # SCI IRIS (1 personne morale + 1 physique) : la personne morale n'a PAS de
    # DNC -> exactement 1 declaration, celle de l'associe physique (gerant).
    payload = _civil_base(
        "SCI IRIS",
        "sci_iris",
        [_pm(40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    generated = css.generate_dossier(payload, tmp_path / "sci-iris-dnc")
    names = {p.name for p in generated.docx_paths}
    assert "declaration_non_condamnation_Martin.docx" in names
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


def test_civil_associe_non_gerant_sans_filiation_ne_bloque_pas(tmp_path: Path) -> None:
    # Albane 2026-07-09 : un associe physique NON gerant sans filiation ne bloque PLUS
    # (il n'a plus de DNC). « Je n'ai volontairement pas complete les parents de mon
    # 2e associe car il n'est pas gerant, mais ca bloque » -> corrige.
    associe_sans_filiation = _pp("Alice", "Martin", 60, 41, 100, 600).model_copy(
        update={"nom_pere": None, "nom_mere": None}
    )
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), associe_sans_filiation],
    )
    plan = css.build_civil_plan(payload)
    assert plan.can_generate is True, plan.blockers
    generated = css.generate_dossier(payload, tmp_path / "sci-nofil")
    names = {p.name for p in generated.docx_paths}
    assert "declaration_non_condamnation_Durand.docx" in names  # gerant
    assert sum(1 for n in names if n.startswith("declaration_non_condamnation")) == 1


def test_selas_ordre_conseil_derive_sans_champ_libelle(tmp_path: Path) -> None:
    # R5 (2026-06-18) : le « Conseil departemental » n'est plus saisi ; le libelle
    # de la demande d'inscription est derive du departement + connecteur. Sans
    # ordre_conseil dans le payload (champ supprime), la generation reste propre
    # et le libelle est reconstruit cote moteur.
    payload = _selas_payload()
    payload.pop("ordre_conseil", None)  # champ supprime du formulaire
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True  # plus de blocage sur ordre_conseil
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-ordre-derive")
    text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "demande_inscription_ordre.docx")
    )
    # Connecteur par defaut « de » (ordre_connecteur absent).
    # R5 (Albane) : FORME LONGUE « Conseil départemental de l’Ordre des <profession_pluriel>
    # <connecteur> <departement> » (profession PUIS departement, forme du modele d'Albane).
    assert "Conseil départemental de l’Ordre des médecins de Rhône" in text


def test_selas_ordre_connecteur_du(tmp_path: Path) -> None:
    # R6 (2026-06-18) : le connecteur grammatical « du » remplace « de » avant le
    # departement de l'ordre dans le libelle derive (gestion de l'accord).
    payload = _selas_payload()
    payload.pop("ordre_conseil", None)
    payload["ordre_connecteur"] = "du"
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-ordre-du")
    text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "demande_inscription_ordre.docx")
    )
    # R5 (Albane) : FORME LONGUE profession (pluriel) PUIS connecteur « du » + departement,
    # comme le modele d'Albane (« des médecins du Rhône »). Confirme 2026-07-01.
    assert "Conseil départemental de l’Ordre des médecins du Rhône" in text
    assert "Conseil départemental de Rhône" not in text


def test_selas_incoherent_actions_sum_warns_but_generates() -> None:
    # KAN-2 (Rafael, rejeté 2×) : « Tous les documents doivent pouvoir être générés. » Une somme
    # d'actions incohérente NE BLOQUE PLUS la génération ; elle est signalée en AVERTISSEMENT.
    payload = _selas_payload()
    payload["associes"][1].nb_actions = 10  # 75 + 10 != 100
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    # R3 (2026-06-18) : message reformule façon Rafael (explicite, sans jargon) -> avertissement.
    assert any("Problème de calcul" in w for w in plan.warnings)
    assert any("ne correspond pas au nombre total d'actions" in w for w in plan.warnings)


# Saisies conjoint requises par DOC-005 / DOC-006 (regime communautaire).
_SELAS_REGIME_INPUTS = {
    "regime_communautaire": True,
    "conjoint_civilite": "Monsieur",
    "conjoint_genre": Gender.MASCULIN,
    "conjoint_prenom": "Paul",
    "conjoint_nom": "Durand",
    "regime_matrimonial": "la communaute legale",
}


def test_selas_regime_off_keeps_base_bundle(tmp_path: Path) -> None:
    payload = _selas_payload()
    payload["regime_communautaire"] = False
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" not in plan.document_codes
    assert "DOC-006" not in plan.document_codes
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-regime-off")
    names = {p.name for p in generated.docx_paths}
    assert not (_REGIME_DOCS & names)


def test_selas_regime_on_adds_regime_docs(tmp_path: Path) -> None:
    payload = _selas_payload()
    payload.update(_SELAS_REGIME_INPUTS)
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.document_codes == (
        "DOC-044",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-005",
        "DOC-006",
    )
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-regime-on")
    _assert_bundle_clean(generated, _SELAS_BUNDLE_NAMES | _REGIME_DOCS)


def _regime_associe(
    prenom: str,
    nom: str,
    conjoint_civilite: str = "Monsieur",
    conjoint_genre: Gender = Gender.MASCULIN,
) -> RegimeCommunautaireAssocie:
    return RegimeCommunautaireAssocie(
        actif=True,
        regime_matrimonial="la communaute legale",
        conjoint_civilite=conjoint_civilite,
        conjoint_genre=conjoint_genre,
        conjoint_prenom=prenom,
        conjoint_nom=nom,
    )


def test_selas_regime_par_associe_active_les_docs(tmp_path: Path) -> None:
    # R7 (2026-06-18) : le regime communautaire est porte PAR associe physique.
    # Toggle global INACTIF mais le 1er associe (president) marie sous communaute
    # -> DOC-005/006 generes (chemin per-associe), sans toggle global.
    payload = _selas_payload()
    payload["regime_communautaire"] = False  # pas de toggle global
    payload["associes"][0].regime_communautaire_associe = _regime_associe("Paul", "Durand")
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" in plan.document_codes
    assert "DOC-006" in plan.document_codes
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-regime-assoc")
    names = {p.name for p in generated.docx_paths}
    assert _REGIME_DOCS <= names


def test_selas_regime_par_associe_valide_le_conjoint(tmp_path: Path) -> None:
    # R7 : un associe marie sous communaute sans conjoint complet.
    # KAN-2 (Rafael, rejeté 2×) : ce manque NE BLOQUE PLUS -> avertissement, génération possible.
    payload = _selas_payload()
    payload["regime_communautaire"] = False
    payload["associes"][0].regime_communautaire_associe = RegimeCommunautaireAssocie(
        actif=True
    )  # conjoint + regime manquants
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("conjoint" in w.lower() for w in plan.warnings)


def _selas_payload_deux_maries():
    """Payload SELAS avec 2 associes physiques maries sous communaute (R7).

    Associe 0 = Claire Durand (conjoint Paul Durand) ; associe 1 = Marc Petit
    (conjoint Sophie Petit). On remplace la personne morale du payload de base
    par un 2e physique valide (etat civil complet) pour avoir deux maries."""
    payload = _selas_payload()
    payload["regime_communautaire"] = False
    phys2 = payload["associes"][0].model_copy(
        update={
            "genre": Gender.MASCULIN,
            "civilite_affichage": "Monsieur",
            "prenom": "Marc",
            "prenoms": "Marc",
            "nom": "Petit",
            "nb_actions": 25,
            "nb_actions_lettres": "vingt-cinq",
            "qualite_capital": "associé exerçant",
            "apport": StatutsCivilsApport(montant="250", montant_lettres="deux cent cinquante"),
            # Conjointe de Marc : Sophie Petit (civilite feminine coherente).
            "regime_communautaire_associe": _regime_associe(
                "Sophie", "Petit", conjoint_civilite="Madame", conjoint_genre=Gender.FEMININ
            ),
        }
    )
    payload["associes"][0].regime_communautaire_associe = _regime_associe("Paul", "Durand")
    payload["associes"][0].nb_actions = 75
    payload["associes"][1] = phys2
    return payload


def test_selas_deux_associes_maries_flag_multiplication(tmp_path: Path) -> None:
    # R7 : 2 associes physiques maries sous communaute -> warning explicite qu'un
    # couple DOC-005/006 sera genere POUR CHACUN (generation par-personne).
    payload = _selas_payload_deux_maries()
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert any(
        "2 associés mariés" in w and "pour chacun" in w for w in plan.warnings
    )


def test_selas_deux_maries_generent_deux_couples_distincts(tmp_path: Path) -> None:
    # R7 : 2 associes maries -> 2 renonciations + 2 avertissements, fichiers de
    # noms distincts (suffixes par nom d'associe), chacun avec le BON conjoint.
    payload = _selas_payload_deux_maries()
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-2maries")
    names = {p.name for p in generated.docx_paths}
    # Deux couples nommes par associe (le couple historique de nom fixe NE doit
    # PAS apparaitre quand il y a plusieurs maries).
    assert "lettre_renonciation_associe_Durand.docx" in names
    assert "lettre_avertissement_conjoint_Durand.docx" in names
    assert "lettre_renonciation_associe_Petit.docx" in names
    assert "lettre_avertissement_conjoint_Petit.docx" in names
    assert "lettre_renonciation_associe.docx" not in names
    assert "lettre_avertissement_conjoint.docx" not in names
    renonciation_count = sum(1 for n in names if n.startswith("lettre_renonciation_associe"))
    avertissement_count = sum(1 for n in names if n.startswith("lettre_avertissement_conjoint"))
    assert renonciation_count == 2
    assert avertissement_count == 2
    # Le BON conjoint dans le BON fichier (renonciation = signature du conjoint).
    durand_renonciation = _docx_text(
        next(p for p in generated.docx_paths if p.name == "lettre_renonciation_associe_Durand.docx")
    )
    petit_renonciation = _docx_text(
        next(p for p in generated.docx_paths if p.name == "lettre_renonciation_associe_Petit.docx")
    )
    assert "Paul Durand" in durand_renonciation
    assert "Sophie Petit" not in durand_renonciation
    assert "Sophie Petit" in petit_renonciation
    assert "Paul Durand" not in petit_renonciation
    # L'apporteur (renoncant) est bien CHAQUE associe dans son propre fichier.
    assert "Claire Durand" in durand_renonciation
    assert "Marc Petit" in petit_renonciation
    # Avertissement : le conjoint destinataire (civilite + nom) et l'apporteur
    # (civilite prenom nom dans la mention) correspondent a CHAQUE associe.
    durand_avert = _docx_text(
        next(
            p
            for p in generated.docx_paths
            if p.name == "lettre_avertissement_conjoint_Durand.docx"
        )
    )
    petit_avert = _docx_text(
        next(
            p
            for p in generated.docx_paths
            if p.name == "lettre_avertissement_conjoint_Petit.docx"
        )
    )
    # Apporteur correct par fichier (mention manuscrite « ... par {apporteur} »).
    assert "Madame Claire Durand" in durand_avert
    assert "Monsieur Marc Petit" in petit_avert
    # A3 (Albane 2026-06-26) : le conjoint destinataire porte son PRENOM
    # (« Monsieur Paul Durand » / « Madame Sophie Petit »), plus jamais sans prenom.
    assert "Monsieur Paul Durand" in durand_avert
    assert "Madame Sophie Petit" in petit_avert
    assert "Monsieur Durand\n" not in durand_avert
    assert "Madame Petit\n" not in petit_avert
    # A1 (Albane 2026-06-26) : l'entete societe affiche la PROFESSION REGLEMENTEE
    # (« médecin »), jamais le titre « Docteur ».
    assert "de médecin" in durand_avert
    assert "de Docteur" not in durand_avert
    assert "de Docteur" not in petit_avert
    # A5 (Albane 2026-06-26) : montant = apport INDIVIDUEL du renoncant (750 pour
    # Durand, 250 pour Petit), pas le capital total (1 000 / 1 020). A4 : la mention
    # s'accorde au CONJOINT signataire (Durand -> conjoint Paul MASCULIN « informé » ;
    # Petit -> conjointe Sophie FEMININ « informée »).
    assert "été informé de l’apport de 750 euros" in durand_avert
    assert "informée de l’apport de 750" not in durand_avert
    assert "de l’apport de 1 000 euros" not in durand_avert
    assert "de l’apport de 1 020 euros" not in durand_avert
    assert "été informée de l’apport de 250 euros" in petit_avert
    assert "de l’apport de 1 000 euros" not in petit_avert
    # R3 (Albane 2026-06-26) : renonciation utilise aussi l'apport individuel.
    assert "en apportant 750 (sept cent cinquante) euros" in durand_renonciation
    assert "en apportant 250 (deux cent cinquante) euros" in petit_renonciation
    assert "1 000" not in durand_renonciation
    # R2 (Albane 2026-06-26) : plus de mention « exemplaires » dans la renonciation.
    assert "exemplaires" not in durand_renonciation
    assert "exemplaires" not in petit_renonciation
    _assert_bundle_clean(generated, {_SELAS_STATUTS_NAME})


def test_selas_deux_maries_adresses_foyer_distinctes(tmp_path: Path) -> None:
    # R7 (2026-06-18) : chaque associe marie porte SON adresse de foyer structuree
    # dans son propre avertissement (DOC-006), plus de repli sur l'adresse du
    # president (residu d63393c). Le president habite « 10 rue Exemple, 69000 Lyon ».
    payload = _selas_payload_deux_maries()
    payload["associes"][0].adresse_personnelle = Address(
        num_voie="3",
        voie="rue des Lilas",
        cp="33000",
        ville="Bordeaux",
        adresse_affichee="3 rue des Lilas, 33000 Bordeaux",
    )
    payload["associes"][1].adresse_personnelle = Address(
        num_voie="7",
        voie="avenue du Parc",
        cp="44000",
        ville="Nantes",
        adresse_affichee="7 avenue du Parc, 44000 Nantes",
    )
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-2foyers")
    durand_avert = _docx_text(
        next(
            p
            for p in generated.docx_paths
            if p.name == "lettre_avertissement_conjoint_Durand.docx"
        )
    )
    petit_avert = _docx_text(
        next(
            p
            for p in generated.docx_paths
            if p.name == "lettre_avertissement_conjoint_Petit.docx"
        )
    )
    # Chaque avertissement porte SON foyer (bloc destinataire conjoint).
    assert "rue des Lilas" in durand_avert
    assert "33000 Bordeaux" in durand_avert
    assert "avenue du Parc" in petit_avert
    assert "44000 Nantes" in petit_avert
    # Pas de fuite croisee entre les deux foyers.
    assert "avenue du Parc" not in durand_avert
    assert "44000 Nantes" not in durand_avert
    assert "rue des Lilas" not in petit_avert
    assert "33000 Bordeaux" not in petit_avert
    # Plus de repli sur l'adresse PERSONNELLE du president (residu d63393c).
    # Le president habite « 10 rue Exemple, 69000 Lyon » ; « rue Exemple » ne doit
    # plus apparaitre comme adresse du foyer dans les avertissements per-associe.
    # (Le siege social de la societe contient « 69000 Lyon » et reste legitime en
    # en-tete expediteur, donc on ne l'asserte pas.)
    assert "rue Exemple" not in durand_avert
    assert "rue Exemple" not in petit_avert


def test_selas_un_seul_marie_garde_nom_fixe(tmp_path: Path) -> None:
    # R7 : un SEUL associe marie (chemin per-associe) -> couple unique au nom
    # FIXE historique (byte-identique avec le toggle global), pas de suffixe.
    payload = _selas_payload()
    payload["regime_communautaire"] = False
    payload["associes"][0].regime_communautaire_associe = _regime_associe("Paul", "Durand")
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-1marie")
    names = {p.name for p in generated.docx_paths}
    assert "lettre_renonciation_associe.docx" in names
    assert "lettre_avertissement_conjoint.docx" in names
    # Pas de fichier suffixe pour un seul marie.
    assert not any(n.startswith("lettre_renonciation_associe_") for n in names)
    renonciation = _docx_text(
        next(p for p in generated.docx_paths if p.name == "lettre_renonciation_associe.docx")
    )
    assert "Paul Durand" in renonciation
    assert "Claire Durand" in renonciation


def test_selas_zero_marie_aucun_doc_regime(tmp_path: Path) -> None:
    # R7 : aucun associe marie + toggle global inactif -> aucun document regime.
    payload = _selas_payload()
    payload["regime_communautaire"] = False
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-0marie")
    names = {p.name for p in generated.docx_paths}
    assert not (_REGIME_DOCS & names)
    assert not any(n.startswith("lettre_renonciation_associe") for n in names)
    assert not any(n.startswith("lettre_avertissement_conjoint") for n in names)


def test_selas_regime_on_warns_for_conjoint() -> None:
    # KAN-2 (Rafael, rejeté 2×) : régime ON sans saisies conjoint NE BLOQUE PLUS -> avertissement.
    payload = _selas_payload()
    payload["regime_communautaire"] = True  # toggle ON sans saisies conjoint
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("regime communautaire" in w.casefold() for w in plan.warnings)


# --- A2 SELAS : preuve multi-associes N (3 et 5, borne moteur 5) --------------


def _selas_phys(prenom, nom, nb_actions, qualite="associé exerçant"):
    """Associe SELAS personne physique exercante, tous champs requis remplis."""
    return StatutsCivilsAssocie(
        type_personne="personne_physique",
        genre=Gender.MASCULIN,
        civilite_affichage="Monsieur",
        prenom=prenom,
        prenoms=prenom,
        nom=nom,
        date_naissance="1 janvier 1980",
        ville_naissance="Lyon",
        departement_naissance="69",
        nationalite="française",
        profession="Docteur",
        situation_maritale="célibataire",
        adresse_personnelle_affichee="10 rue Exemple, 69000 Lyon",
        qualification_principale="qualifié en médecine générale",
        ordre_departemental="Rhône",
        numero_ordre="69-12345",
        numero_rpps="10100000001",
        qualite_capital=qualite,
        nb_actions=nb_actions,
        nb_actions_lettres=str(nb_actions),
        # DNC par associe (Rafael 2026-07-09) : filiation requise pour CHAQUE associe.
        nom_pere=f"Pierre {nom}",
        nom_mere=f"Anne {nom}",
        apport=StatutsCivilsApport(
            montant=str(nb_actions * 10), montant_lettres=str(nb_actions * 10)
        ),
    )


def _selas_payload_n(associes):
    """Payload SELAS de base, avec une liste d'associes custom (somme actions auto)."""
    payload = _selas_payload()
    payload["associes"] = associes
    payload["nb_actions_total"] = sum(int(a.nb_actions or 0) for a in associes)
    return payload


# O24-02 (onglet 24) : la DNC porte le nom du dirigeant. La COMPLETUDE du bundle se
# verifie au nom generique (cf. _assert_bundle_clean qui normalise) ; le nommage par
# dirigeant (Durand president + Martin DG) est verifie par le test DNC multi dedie.
# ST1 (Albane 2026-06-26) : le statuts porte la denomination dans son nom
# (« Statuts SELAS EXEMPLE.docx » pour les payloads de denomination « SELAS EXEMPLE »).
_SELAS_STATUTS_NAME = "Statuts SELAS EXEMPLE.docx"
_SELAS_BUNDLE_NAMES = {
    _SELAS_STATUTS_NAME,
    "declaration_non_condamnation.docx",
    "autorisation_domiciliation.docx",
    "procuration.docx",
    "pv_nomination_dirigeant.docx",
    "demande_inscription_ordre.docx",
}


def test_selas_three_associes_generates_clean(tmp_path: Path) -> None:
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 40),
            _selas_phys("Paul", "Martin", 30),
            _selas_phys("Marie", "Petit", 30),
        ]
    )
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    # ANO-045 : dossier ENTIEREMENT physique (somme actions == total) -> attestable ->
    # l'attestation souscripteurs (DOC-045) est incluse au bundle.
    assert plan.document_codes == (
        "DOC-044",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-045",
    )
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas3")
    _assert_bundle_clean(
        generated,
        _SELAS_BUNDLE_NAMES | {"attestation_capital_souscripteurs_selas.docx"},
    )
    statuts_text = _docx_text(
        next(p for p in generated.docx_paths if p.name == _SELAS_STATUTS_NAME)
    )
    for nom in ("Durand", "Martin", "Petit"):
        assert nom in statuts_text


def test_selas_five_associes_generates_clean(tmp_path: Path) -> None:
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 20),
            _selas_phys("Paul", "Martin", 20),
            _selas_phys("Marie", "Petit", 20),
            _selas_phys("Luc", "Robert", 20),
            _selas_phys("Anne", "Bernard", 20),
        ]
    )
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    # ANO-045 : dossier entierement physique -> attestable -> DOC-045 au bundle.
    assert "DOC-045" in plan.document_codes
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas5")
    _assert_bundle_clean(
        generated,
        _SELAS_BUNDLE_NAMES | {"attestation_capital_souscripteurs_selas.docx"},
    )
    statuts_text = _docx_text(
        next(p for p in generated.docx_paths if p.name == _SELAS_STATUTS_NAME)
    )
    for nom in ("Durand", "Martin", "Petit", "Robert", "Bernard"):
        assert nom in statuts_text


# --- ANO-045 : attestation souscripteurs SELAS (DOC-045) ----------------------


_ATTESTATION_SELAS_NAME = "attestation_capital_souscripteurs_selas.docx"


def test_selas_multi_all_physical_generates_attestation_souscripteurs(
    tmp_path: Path,
) -> None:
    """ANO-045 : un dossier SELAS multi entierement PHYSIQUE (somme actions == total)
    produit l'attestation souscripteurs, bien remplie (une ligne de repartition + une
    d'apport par souscripteur), header profession au pluriel capitalise."""
    # R3 durci (Rafael 2026-07-07) : associee FEMININE pour verrouiller l'accord de
    # la civilite civile (« à Madame Claire Durand »), Paul reste masculin.
    claire = _selas_phys("Claire", "Durand", 60)
    claire.civilite_affichage = "Madame"
    claire.genre = Gender.FEMININ
    payload = _selas_payload_n([claire, _selas_phys("Paul", "Martin", 40)])
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert "DOC-045" in plan.document_codes
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-att")
    attestation = next(
        p for p in generated.docx_paths if p.name == _ATTESTATION_SELAS_NAME
    )
    text = _docx_text(attestation)
    assert "Liste des souscripteurs" in text
    # Header : profession reglementee au pluriel, capitalisee (parite modele Albane).
    assert "par Actions simplifiées de Médecins" in text
    # R3 durci (Rafael 2026-07-07, siloing) : repartition + apport rendent la
    # civilite CIVILE accordee au genre du souscripteur — supersede les anciens
    # verrous « au Dr X » / « Le Docteur X a fait un apport » de cette variante
    # pluripersonnelle (l'unipersonnelle etait deja civile).
    assert "60 actions attribuées à Madame Claire Durand," in text
    assert "40 actions attribuées à Monsieur Paul Martin," in text
    assert "Madame Claire Durand a fait un apport de 600 euros en numéraire." in text
    assert "Monsieur Paul Martin a fait un apport de 400 euros en numéraire." in text
    # R3 (Albane 2026-07-07) + R3 durci : le slot « par le Président, __ » rend la
    # civilite CIVILE accordee au genre du PRESIDENT (ici Claire, feminine).
    assert (
        "certifié exact, sincère et véritable par le Président, Madame Claire Durand"
        in text
    )
    assert "Docteur" not in text
    assert "au Dr " not in text
    assert "[" not in text and "]" not in text


def test_selas_multi_with_morale_omits_attestation_souscripteurs(
    tmp_path: Path,
) -> None:
    """ANO-045 : un dossier SELAS multi comportant un associe PERSONNE MORALE n'est PAS
    attestable (le modele ne represente que des souscripteurs physiques) -> DOC-045
    absent du plan ET du bundle, sans crash du generateur."""
    payload = _selas_payload()  # 1 physique + 1 personne morale
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True
    assert "DOC-045" not in plan.document_codes
    generated = selas_multi_slice.generate_dossier(payload, tmp_path / "selas-att-morale")
    names = {p.name for p in generated.docx_paths}
    assert _ATTESTATION_SELAS_NAME not in names


def test_selas_multi_incoherent_actions_omits_attestation() -> None:
    """ANO-045 : une somme d'actions != total (dossier bloque par ailleurs) n'est pas
    attestable -> DOC-045 jamais annonce au plan."""
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 60),
            _selas_phys("Paul", "Martin", 40),
        ]
    )
    payload["nb_actions_total"] = 120  # 60 + 40 != 120
    plan = selas_multi_slice.build_selas_plan(payload)
    assert "DOC-045" not in plan.document_codes


def _selas_uni_payload() -> dict[str, object]:
    """Payload minimal generable d'une SELAS unipersonnelle (medecin / dentiste)."""
    return {
        "denomination": "SELAS EXEMPLE",
        "capital_social": "1000",
        "nb_actions_total": 100,
        "valeur_nominale_action": "10",
        "duree": "99 ans",
        "ville_rcs": "Rennes",
        "lieu_exercice_adresse": "5 place du Centre, 35000 Rennes",
        "siege_num": "5",
        "siege_voie": "place du Centre",
        "siege_cp": "35000",
        "siege_ville": "Rennes",
        "banque_nom": "BANQUE EXEMPLE",
        "banque_adresse": "1 rue Banque, 35000 Rennes",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 décembre",
        "exercice_cloture": "31 décembre 2027",
        "civilite": "Monsieur",
        "prenom": "Alain",
        "nom": "Fedorowsky",
        "date_naissance": date(1980, 1, 1),
        "ville_naissance": "Rennes",
        "departement_naissance": "35",
        "nationalite": "française",
        "titre_affichage": "Docteur",
        "adresse_num": "10",
        "adresse_voie": "rue Exemple",
        "adresse_cp": "35000",
        "adresse_ville": "Rennes",
        "nom_pere": "Pierre",
        "nom_mere": "Anne",
        "situation_maritale": "celibataire",
        "regime_matrimonial": "celibataire",
        "regime_communautaire": False,
        "conjoint_civilite": "",
        "conjoint_prenom": "",
        "conjoint_nom": "",
        "departement_ordre": "Ille-et-Vilaine",
        "connecteur_departement": "de",
        "ordre_president_feminin": False,
        "mandataire_prenom": "Jordan",
        "mandataire_nom": "ELBAZ",
        "numero_ordre": "35-1",
        "numero_rpps": "10100000001",
        "ordre_ville": "Rennes",
        "ordre_cp": "35000",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "signature_lieu": "Rennes",
        "signature_date": date(2026, 5, 15),
    }


def test_selas_uni_medecin_generates_attestation_souscripteurs(tmp_path: Path) -> None:
    """ANO-045 : la SELAS unipersonnelle medecin est toujours attestable (associe
    physique unique) -> DOC-045 au bundle, un unique souscripteur = le president."""
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as sm

    payload = _selas_uni_payload()
    plan = sm.build_selas_uni_medecin_plan(payload)
    assert plan.can_generate is True
    assert "DOC-045" in plan.document_codes
    generated = sm.generate_dossier(payload, tmp_path / "selas-uni-med")
    attestation = next(
        p for p in generated.docx_paths if p.name == _ATTESTATION_SELAS_NAME
    )
    text = _docx_text(attestation)
    assert "par Actions simplifiées de Médecins" in text
    # R3 durci (Rafael 2026-07-07) : civilite CIVILE (titre « Docteur » du payload
    # converti, accord au genre du signataire) — plus de « au Dr » / « Le Docteur ».
    assert "100 actions attribuées à Monsieur Alain Fedorowsky," in text
    # R5 (Rafael 2026-07-09) : montant groupe des 4 chiffres (« 1 000 euros », plus « 1000 »).
    assert (
        "Monsieur Alain Fedorowsky a fait un apport de 1 000 euros en numéraire."
        in text
    )
    assert "Docteur" not in text
    assert "[" not in text and "]" not in text


def test_selas_uni_dentiste_generates_attestation_souscripteurs(tmp_path: Path) -> None:
    """ANO-045 : la SELAS unipersonnelle dentiste est toujours attestable -> DOC-045
    au bundle, header profession dentiste au pluriel capitalise."""
    from sydel_doc_engine.front_app import selas_uni_dentiste_slice as sd

    payload = _selas_uni_payload()
    plan = sd.build_selas_uni_dentiste_plan(payload)
    assert plan.can_generate is True
    assert "DOC-045" in plan.document_codes
    generated = sd.generate_dossier(payload, tmp_path / "selas-uni-dent")
    attestation = next(
        p for p in generated.docx_paths if p.name == _ATTESTATION_SELAS_NAME
    )
    text = _docx_text(attestation)
    assert "par Actions simplifiées de Chirurgiens-dentistes" in text
    # R3 durci (Rafael 2026-07-07) : civilite CIVILE, plus de « au Dr ».
    assert "100 actions attribuées à Monsieur Alain Fedorowsky," in text
    assert "Docteur" not in text
    assert "[" not in text and "]" not in text


# --- A1 : bornes du repeater nommees + alignees sur le moteur -----------------


def test_repeater_bounds_are_named_and_aligned() -> None:
    # A1 : les bornes sont des constantes nommees (pas de litteraux disperses) et
    # la borne SELAS front est ALIGNEE sur le generateur (un ecart silencieux
    # casserait la generation). Ce test verrouille les valeurs.
    from sydel_doc_engine.generators.lot_04.statuts_selas_multi import (
        MAX_ASSOCIES,
        MIN_ASSOCIES,
    )

    assert css.CIVIL_NB_MIN_BY_STRUCTURE == {
        "SCI": 1,
        "SCM": 1,
        # Micro holding : 1 associe minimum (societe civile a capital variable).
        "MICRO_HOLDING": 1,
        "SCI IRIS": 2,
        "SCS": 2,
    }
    assert css.CIVIL_NB_MAX_ASSOCIES == 6
    assert (selas_multi_slice.SELAS_NB_MIN, selas_multi_slice.SELAS_NB_MAX) == (
        MIN_ASSOCIES,
        MAX_ASSOCIES,
    )


# --- A3 : selection du dirigeant (president) SELAS ----------------------------


def test_selas_president_defaults_to_first_physique() -> None:
    # Sans choix UI (payload sans president_index) : le 1er associe physique est
    # president — comportement historique, non-regression.
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 40),
            _selas_phys("Paul", "Martin", 30),
            _selas_phys("Marie", "Petit", 30),
        ]
    )
    ctx = selas_multi_slice.build_generation_context(payload)
    assert ctx.dirigeant_nomine.ref_associe_index == 0
    assert ctx.statuts_selas_multi.president.ref_associe_index == 0
    assert ctx.dirigeant_nomine.nom == "Durand"


def test_selas_president_selectable() -> None:
    # Le dirigeant peut etre un autre associe physique (le 3e ici).
    payload = _selas_payload_n(
        [
            _selas_phys("Claire", "Durand", 40),
            _selas_phys("Paul", "Martin", 30),
            _selas_phys("Marie", "Petit", 30),
        ]
    )
    payload["president_index"] = 2
    ctx = selas_multi_slice.build_generation_context(payload)
    assert ctx.dirigeant_nomine.ref_associe_index == 2
    assert ctx.statuts_selas_multi.president.ref_associe_index == 2
    assert ctx.dirigeant_nomine.nom == "Petit"


def test_selas_president_index_morale_falls_back_to_physique() -> None:
    # Un index pointant une personne morale ne doit JAMAIS atteindre le moteur
    # (President = personne physique) : on retombe sur le 1er physique.
    morale = _selas_payload()["associes"][1]
    morale.nb_actions = 40
    payload = _selas_payload_n([_selas_phys("Claire", "Durand", 60), morale])
    payload["president_index"] = 1  # pointe la personne morale
    ctx = selas_multi_slice.build_generation_context(payload)
    assert ctx.dirigeant_nomine.ref_associe_index == 0
    plan = selas_multi_slice.build_selas_plan(payload)
    assert plan.can_generate is True


# --- Surface UI Streamlit -----------------------------------------------------


def test_front_dropdown_lists_all_types_with_selarl_default() -> None:
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)

    selector = app.selectbox(key="clean_dossier_type")
    assert selector.value == "SELARL"
    assert list(selector.options) == [
        "SELARL",
        "SCM",
        "SCI",
        "SCI IRIS",
        "SCS",
        "Micro holding",
        "SPFPL médecins (forme SAS)",
        "SASU Holding (holding patrimoniale)",
        "SPFPL dentistes — cession",
        "SPFPL dentistes — apport",
        "SELAS pluripersonnelle",
        "SELAS unipersonnelle médecin",
        "SELAS unipersonnelle dentiste",
    ]
    # Surface SELARL inchangee : aucun expander sur le defaut.
    assert len(app.expander) == 0


def test_front_routes_to_sci_slice_and_generates(tmp_path: Path, monkeypatch) -> None:
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-sci")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=180)

    def set_text(key: str, value: str) -> None:
        for widget in app.text_input:
            if str(widget.key) == key:
                widget.set_value(value)
                return
        raise KeyError(key)

    def set_number(key: str, value: int) -> None:
        for widget in app.number_input:
            if str(widget.key) == key:
                widget.set_value(value)
                return
        raise KeyError(key)

    # forme_sociale derivee (§18.1), valeur nominale calculee (§18.2), duree figee
    # (§18.3), lieu signature = ville siege (§18.4) : ces champs ne sont plus saisis.
    society = {
        "sci_denomination": "SCI EXEMPLE",
        # O24-03 : siege sur UNE ligne (plus de champs No/Voie/CP/Ville separes).
        "sci_siege_adresse": "10 rue de la Paix, 75002 Paris",
        "sci_ville_rcs": "Paris",
        "sci_banque_nom": "BANQUE",
        "sci_banque_adresse": "1 rue Banque, 75009 Paris",
        "sci_date_cloture_premier_exercice": "31 decembre 2026",
        # Documents communs (hors identite du gerant : fonction/titre).
        "sci_signataire_fonction": "gerant",
        "sci_signataire_titre": "Docteur",
    }
    for key, value in society.items():
        set_text(key, value)
    # Capital social = number_input (parite gold) -> set via set_number.
    set_number("sci_capital_social", 1000)
    set_text("sci_signature_date", "15/05/2026")
    app = app.run(timeout=180)

    # Nationalite = deroulant (§SCREEN-1, cle _choice) ; adresse personnelle
    # structuree (§18.5) ; profession non demandee hors SCM (§18.6) ; parts debut/fin
    # derivees (§18.5) ; DNC = noms des parents uniquement.
    associes_text = {
        "sci_associe_0_prenom": "Jean",
        "sci_associe_0_nom": "Durand",
        "sci_associe_0_date_naissance": "1 janvier 1980",
        "sci_associe_0_ville_naissance": "Paris",
        "sci_associe_0_departement_naissance": "75",
        # O24-03 : adresse personnelle sur UNE ligne.
        "sci_associe_0_adresse": "1 rue Exemple, 75000 Paris",
        "sci_associe_0_apport_montant": "400",
        # DNC du gerant (associe 0, coche « Dirigeant » par defaut) saisie sous lui.
        "sci_associe_0_sig_nom_pere": "Pierre Durand",
        "sci_associe_0_sig_nom_mere": "Anne Durand",
        "sci_associe_1_prenom": "Alice",
        "sci_associe_1_nom": "Martin",
        "sci_associe_1_date_naissance": "2 fevrier 1982",
        "sci_associe_1_ville_naissance": "Lyon",
        "sci_associe_1_departement_naissance": "69",
        # O24-03 : adresse personnelle sur UNE ligne.
        "sci_associe_1_adresse": "2 rue Exemple, 69000 Lyon",
        # DNC par associe (Rafael 2026-07-09) : filiation de CHAQUE associe.
        "sci_associe_1_sig_nom_pere": "Pierre Martin",
        "sci_associe_1_sig_nom_mere": "Anne Martin",
        "sci_associe_1_apport_montant": "600",
    }
    for key, value in associes_text.items():
        set_text(key, value)
    # KAN-37 (Rafael 2026-07-27) : la situation matrimoniale est un MENU DÉROULANT (plus un champ
    # texte) — on selectionne le preset « Célibataire » sur chaque associe.
    _set_selectbox(app, "sci_associe_0_situation", "Célibataire")
    _set_selectbox(app, "sci_associe_1_situation", "Célibataire")
    associes_numbers = {
        "sci_associe_0_nb_titres": 40,
        "sci_associe_1_nb_titres": 60,
        "sci_nb_parts_total": 100,
    }
    for key, value in associes_numbers.items():
        set_number(key, value)
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)

    generate_button.click()
    app = app.run(timeout=180)

    download_labels = [item.label for item in app.get("download_button")]
    assert "Telecharger Statuts SCI EXEMPLE.docx" in download_labels
    assert "Telecharger le dossier ZIP" in download_labels


def test_sci_unipersonnel_apport_derived_from_capital(
    tmp_path: Path, monkeypatch
) -> None:
    """§1 (Albane) — VALEUR / DERIVATION : une societe UNIPERSONNELLE (un seul associe)
    ne SAISIT PLUS le montant d'apport de l'associe unique — il vaut FORCEMENT le capital
    social (pas de double saisie). Verrou de bout en bout sur le FLUX REEL (repeater) :
    - le champ « Apport (montant) » de l'associe unique n'existe PAS (widget absent) ;
    - les statuts SCI rendent le capital comme apport (« La somme de mille euros, ») ;
    - une seule ligne « La somme de » (pas d'apport separe qui doublerait le montant)."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-sci-uni")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    # Passer en UNIPERSONNEL : retirer le 2e associe seede par les donnees de test.
    # Retour Albane 2026-07-07 : le bouton global « Retirer un associe » (dernier
    # seulement) est remplace par un bouton PAR LIGNE -> on retire la ligne index 1.
    next(b for b in app.button if str(b.key) == "sci_remove_associe_1").click()
    app = app.run(timeout=180)

    # §1 : le champ d'apport de l'associe unique a DISPARU (repris auto du capital).
    assert "sci_associe_0_apport_montant" not in [str(w.key) for w in app.text_input]

    def set_number(key: str, value: int) -> None:
        next(w for w in app.number_input if str(w.key) == key).set_value(value)

    # Aligner le total de parts sur l'associe unique restant + capital = 1000 €.
    parts_unique = next(
        w for w in app.number_input if str(w.key) == "sci_associe_0_nb_titres"
    ).value
    set_number("sci_nb_parts_total", int(parts_unique))
    set_number("sci_capital_social", 1000)
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)
    generate_button.click()
    app = app.run(timeout=180)

    text = _dir_docx_text(tmp_path / "ui-sci-uni")
    # L'apport de l'associe unique = capital (1000 €) rendu en lettres.
    assert "La somme de mille euros," in text
    # Pas de double saisie : une seule ligne d'apport (le capital), jamais deux.
    assert text.count("La somme de mille euros,") == 1


@pytest.mark.parametrize(
    "label, statuts_name",
    [
        # R10 (Rafael 2026-07-07) : TOUS les statuts portent la denomination du prefill
        # dans le nom de fichier (« Statuts <denomination>.docx »).
        ("SCM", "Statuts SCM DES DOCTEURS EXEMPLE.docx"),
        ("SCI", "Statuts SCI EXEMPLE.docx"),
        ("SCI IRIS", "Statuts SCI IRIS EXEMPLE.docx"),
        ("SCS", "Statuts SCS EXEMPLE.docx"),
        # Micro holding (Albane 2026-06-26) : societe civile a capital variable, socle civil.
        ("Micro holding", "Statuts MICRO HOLDING EXEMPLE.docx"),
        ("SPFPL médecins (forme SAS)", "Statuts SPFPL MARTIN.docx"),
        # SASU Holding (Albane 2026-06-29) : SAS unipersonnelle generaliste, slice dedie.
        ("SASU Holding (holding patrimoniale)", "Statuts SASU HOLDING EXEMPLE.docx"),
        ("SPFPL dentistes — cession", "Statuts SPFPL MARTIN.docx"),
        ("SPFPL dentistes — apport", "Statuts SPFPL MARTIN.docx"),
        ("SELAS pluripersonnelle", "Statuts SELAS EXEMPLE.docx"),
    ],
)
def test_typed_test_data_button_generates(
    tmp_path: Path, monkeypatch, label: str, statuts_name: str
) -> None:
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-civil")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(label)
    app = app.run(timeout=180)

    # Le bouton "donnees de test" remplit le dossier sans saisie manuelle.
    test_data_button = next(b for b in app.button if "test_data" in str(b.key))
    test_data_button.click()
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)

    generate_button.click()
    app = app.run(timeout=180)

    download_labels = [item.label for item in app.get("download_button")]
    assert f"Telecharger {statuts_name}" in download_labels
    assert "Telecharger le dossier ZIP" in download_labels


def _dir_docx_text(directory: Path) -> str:
    """Concatene le texte de tous les .docx generes sous un repertoire (artifacts UI)."""
    return "\n".join(_docx_text(p) for p in directory.rglob("*.docx"))


def _set_text_widget(app, key: str, value: str) -> None:
    for widget in app.text_input:
        if str(widget.key) == key:
            widget.set_value(value)
            return
    raise KeyError(key)


def _set_selectbox(app, key: str, value: str) -> None:
    # KAN-37 : la situation matrimoniale (entre autres) est un menu déroulant, plus un champ texte.
    for widget in app.selectbox:
        if str(widget.key) == key:
            widget.set_value(value)
            return
    raise KeyError(key)


def test_sci_repeater_live03_accentuates_date_naissance(tmp_path: Path, monkeypatch) -> None:
    """LIVE-03 (date_naissance via le REPEATER, SCI) : une date de naissance saisie sans
    accent (« 1er aout 1980 ») dans le repeater d'associes ressort accentuee dans les
    statuts SCI (« Statuts SCI EXEMPLE.docx »). Exerce le code REEL du formulaire
    (associe_repeater._render_personne_physique -> accentuate_french_months), pas un
    associe pre-construit."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-sci-live03")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Force une date de naissance LIBRE sans accent sur le 1er associe (repeater).
    _set_text_widget(app, "sci_associe_0_date_naissance", "1er aout 1980")
    app = app.run(timeout=180)

    generate_button = next(b for b in app.button if str(b.key) == "clean_typed_generate_dossier")
    assert generate_button.disabled is False
    generate_button.click()
    app = app.run(timeout=180)

    text = _dir_docx_text(tmp_path / "ui-sci-live03")
    assert "août 1980" in text, "LIVE-03 : le mois de naissance (repeater SCI) doit etre accentue"
    assert "1er aout 1980" not in text, "LIVE-03 : la saisie sans accent ne doit pas subsister"


def test_selas_multi_live03_accentuates_date_naissance(tmp_path: Path, monkeypatch) -> None:
    """LIVE-03 (date_naissance, SELAS multi) : la date de naissance LIBRE du 1er associe
    SELAS ressort accentuee dans statuts_selas_multi.docx. Exerce
    selas_multi_slice._render_personne_physique (accentuate_french_months)."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-live03")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    _set_text_widget(app, "selas_associe_0_date_naissance", "1er aout 1980")
    app = app.run(timeout=180)

    generate_button = next(b for b in app.button if str(b.key) == "clean_typed_generate_dossier")
    assert generate_button.disabled is False
    generate_button.click()
    app = app.run(timeout=180)

    text = _dir_docx_text(tmp_path / "ui-selas-live03")
    assert "août 1980" in text, "LIVE-03 : le mois de naissance (SELAS multi) doit etre accentue"
    assert "1er aout 1980" not in text, "LIVE-03 : la saisie sans accent ne doit pas subsister"


def test_selarl_membre_additionnel_live03_accentuates_date_naissance(
    tmp_path: Path, monkeypatch
) -> None:
    """LIVE-03 (date_naissance, MEMBRE SELARL additionnel) : la date de naissance LIBRE
    d'un associe additionnel (forme multi-associes) ressort accentuee dans les statuts.
    Exerce shell._render_one_selarl_membre (text_input + accentuate via _accentuate_date_value).
    Le 1er associe (praticien) utilise un date_input -> objet date, deja accentue par le moteur ;
    seuls les membres additionnels saisissent la date en TEXTE libre."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selarl-live03")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Passer en multi-associes : decocher « Dossier unipersonnel » fait apparaitre le
    # bloc des associes additionnels.
    app.checkbox(key="selarl_dossier_unipersonnel").set_value(False)
    app = app.run(timeout=180)

    # Parts du praticien (60) ; le membre additionnel prend les 40 restantes (capital
    # test = 100 parts). Un seul membre est rendu par defaut (selarl_membres_count=1).
    def set_number(key: str, value: int) -> None:
        for widget in app.number_input:
            if str(widget.key) == key:
                widget.set_value(value)
                return
        raise KeyError(key)

    # Le total de parts du jeu de test est ALEATOIRE (100/200/500/1000) : on le fige a
    # 100 pour un partage deterministe (praticien 60 + membre 40 = total).
    set_number("selarl_nb_parts_total", 100)
    app = app.run(timeout=180)
    set_number("selarl_praticien_nb_parts", 60)
    _set_text_widget(app, "selarl_praticien_apport", "600")
    set_number("selarl_membre_0_nb_parts", 40)
    _set_text_widget(app, "selarl_membre_0_apport", "400")
    _set_text_widget(app, "selarl_membre_0_prenom", "Lea")
    _set_text_widget(app, "selarl_membre_0_nom", "Bernard")
    # Date de naissance LIBRE sans accent -> doit ressortir accentuee.
    _set_text_widget(app, "selarl_membre_0_date_naissance", "1er aout 1985")
    _set_text_widget(app, "selarl_membre_0_ville_naissance", "Lyon")
    _set_text_widget(app, "selarl_membre_0_dep_naissance", "69")
    _set_text_widget(app, "selarl_membre_0_nationalite", "française")
    _set_selectbox(app, "selarl_membre_0_situation", "Célibataire")
    _set_text_widget(app, "selarl_membre_0_profession", "medecin")
    _set_text_widget(app, "selarl_membre_0_adresse", "8 rue Centrale, 69001 Lyon")
    # DNC par associe (Rafael 2026-07-09) : filiation du membre requise (sa DNC).
    _set_text_widget(app, "selarl_membre_0_sig_nom_pere", "Laurent Bernard")
    _set_text_widget(app, "selarl_membre_0_sig_nom_mere", "Julie Bernard")
    _set_text_widget(app, "selarl_membre_0_ordre_dep", "69")
    _set_text_widget(app, "selarl_membre_0_numero_ordre", "ORD-999")
    _set_text_widget(app, "selarl_membre_0_numero_rpps", "20000000002")
    app = app.run(timeout=180)

    # La SELARL passe par le parcours legacy (cle de bouton « clean_generate_dossier »,
    # distincte du parcours typé « clean_typed_generate_dossier »).
    generate_button = next(b for b in app.button if str(b.key) == "clean_generate_dossier")
    assert generate_button.disabled is False, [c.value for c in app.caption if "Blocage" in c.value]
    generate_button.click()
    app = app.run(timeout=180)

    text = _dir_docx_text(tmp_path / "ui-selarl-live03")
    assert "août 1985" in text, "LIVE-03 : la date du membre additionnel SELARL doit etre accentuee"
    assert "1er aout 1985" not in text, "LIVE-03 : la saisie sans accent ne doit pas subsister"


def test_selas_madame_la_presidente_toggle_drives_destinataire() -> None:
    # Parite gold (RAF-003 / Albane 2026-06-10) : la SELAS doit, comme la SELARL,
    # adresser « Madame la Presidente » quand la presidente de l'ordre est une femme
    # (demande d'inscription DOC-034), au lieu de « Monsieur le President » en dur.
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    base = _selas_payload()
    assert sms.build_generation_context(base).ordre.destinataire_appel == (
        "Monsieur le Président"
    )
    feminin = {**_selas_payload(), "ordre_president_feminin": True}
    assert sms.build_generation_context(feminin).ordre.destinataire_appel == (
        "Madame la Présidente"
    )


def test_selas_cession_codes_flow_into_plan_and_orchestrator() -> None:
    # Completude V2 : la SELAS reutilise les sous-formulaires cession SELARL
    # valides (prefix='selas'). Des qu'un contexte cession/bail/scm est saisi,
    # les codes des documents de cession doivent entrer dans le bundle (plan ET
    # codes confies a l'orchestrateur), comme la SELARL.
    from sydel_doc_engine.domain.models import (
        BailContext,
        CessionContext,
        ScmCessionContext,
    )
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    payload = {
        "cession_context": CessionContext.model_validate(
            {"etape": "compromis", "type_cabinet": "dentaire"}
        ),
        "bail_context": BailContext(),
        "scm_cession_context": ScmCessionContext(),
    }
    plan = sms._selas_document_codes(payload)
    orch = sms._orchestrator_codes(payload)
    for code in ("DOC-012", "DOC-007", "DOC-031", "DOC-032", "DOC-033"):
        assert code in plan, f"{code} absent du plan SELAS"
        assert code in orch, f"{code} absent des codes orchestrateur SELAS"
    # ANO-008 : l'appel de fonds (DOC-008) est restreint a la SELARL cote moteur et
    # non clairement requis pour la SELAS au canon -> il ne doit PAS etre inscrit au
    # plan SELAS (sinon annonce mais jamais genere). Question metier flaggee (Albane).
    assert "DOC-008" not in plan
    # Sans cession : aucun code cession ajoute (bundle de creation inchange).
    assert sms._cession_codes({}) == ()


def test_selas_cession_acquereur_forme_sociale_corrigee_en_selas() -> None:
    # B1 (fidelite gold, machine 2026-06-22) : le sous-formulaire de cession PARTAGE
    # code l'acquereur en « SELARL » (hardcode unipersonnel shell.py:2125). En SELAS,
    # l'acquereur EST la SELAS creee -> les actes/compromis de cession doivent afficher
    # « SELAS », pas « SELARL ». Post-correction SELAS-only (gold SELARL intact).
    from sydel_doc_engine.domain.models import CessionAcquereur, CessionContext
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    payload = {
        **_selas_payload(),
        "cession_context": CessionContext(
            etape="compromis",
            type_cabinet="dentaire",
            acquereur=CessionAcquereur(forme_sociale="SELARL"),
        ),
    }
    ctx = sms.build_generation_context(payload)
    assert ctx.cession is not None and ctx.cession.acquereur is not None
    assert ctx.cession.acquereur.forme_sociale == "SELAS"


def test_selas_cession_masque_le_bloc_acquereur(tmp_path: Path, monkeypatch) -> None:
    # #15 (onglet 24) : en SELAS, l'acquereur EST la societe en cours de creation ->
    # aucun champ de saisie acquereur (RCS / SIRET / dates) ne doit apparaitre dans le
    # formulaire de cession (entierement derive de la fiche societe).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-cess")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert not any("cession_acquereur" in k for k in keys), (
        "Le bloc acquereur ne doit pas etre saisi en SELAS (#15)."
    )


def test_selas_cession_warns_ca_et_resultat_des_exercices() -> None:
    # #13 (onglet 24) : en cession SELAS, le CA et le resultat des exercices sont attendus.
    # KAN-2 (Rafael, rejeté 2×) : un exercice au CA/resultat vide NE BLOQUE PLUS la génération ->
    # il est signalé en AVERTISSEMENT (zone à compléter à la main).
    from sydel_doc_engine.domain.models import CessionContext, CessionExercice
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    payload = {
        **_selas_payload(),
        "cession_context": CessionContext(
            exercices=[CessionExercice(periode="2024", chiffre_affaires="", resultat="")]
        ),
    }
    plan = sms.build_selas_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("chiffre d'affaires de l'exercice" in w for w in plan.warnings)
    assert any("resultat de l'exercice" in w for w in plan.warnings)

    payload_ok = {
        **_selas_payload(),
        "cession_context": CessionContext(
            exercices=[
                CessionExercice(periode="2024", chiffre_affaires="200 000", resultat="50 000")
            ]
        ),
    }
    assert not any("exercice" in w for w in sms.build_selas_plan(payload_ok).warnings)


def test_selas_cession_genere_acte_et_compromis_ensemble() -> None:
    # #14 (onglet 24) : en SELAS, l'acte ET le compromis du cabinet sont produits
    # ENSEMBLE pour le type de cabinet, quelle que soit l'etape saisie.
    from sydel_doc_engine.domain.models import CessionContext
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    codes_dent = sms._cession_codes(
        {"cession_context": CessionContext(etape="compromis", type_cabinet="dentaire")}
    )
    assert "DOC-011" in codes_dent and "DOC-012" in codes_dent  # acte + compromis dentaire
    codes_med = sms._cession_codes(
        {"cession_context": CessionContext(etape="acte", type_cabinet="medical")}
    )
    assert "DOC-009" in codes_med and "DOC-010" in codes_med  # acte + compromis medical


def test_selas_cession_generate_dossier_emits_all_cession_files(
    tmp_path: Path, monkeypatch
) -> None:
    # TACHE A (cablage cession SELAS) : PREUVE de bout en bout que generate_dossier
    # SELAS EMET reellement les documents de cession quand les contextes sont saisis.
    # Avant le cablage, le dossier levait « document est obligatoire pour
    # CODE-CESSION-CAB-001 » (acte/compromis/bail sans `document=`) et n'emettait
    # JAMAIS DOC-031/032/033 (scm_cession.variante_structure='selarl' != 'SELAS' ->
    # plan menteur). On reutilise les GENERATEURS de cession SELARL (memes fichiers)
    # comme le ferait le sous-formulaire partage, MAIS avec un acquereur
    # REPRESENTATIF de la SELAS : en SELAS reel l'acquereur EST la societe creee,
    # dont la denomination est derivee de la fiche societe (« SELAS EXEMPLE »), PAS
    # « SELARL CABINET DURAND ». Un acquereur prefixe 'SELARL' validait un scenario
    # impossible et produisait « Pour la SELAS SELARL CABINET DURAND » (Akainu
    # MAJEUR 1, 2026-06-24).
    from sydel_doc_engine.front_app import selas_multi_slice as sms
    from sydel_doc_engine.front_app import shell
    from sydel_doc_engine.scenarios.selarl import (
        PROFESSION_MEDECIN,
        cession_fixture_for_profession,
        scm_cession_fixture,
    )

    cession, bail = cession_fixture_for_profession(PROFESSION_MEDECIN)
    # Acquereur REPRESENTATIF SELAS : denomination derivee de la fiche societe (sans
    # prefixe 'SELARL') ; la forme reelle ('SELAS') sera de toute facon imposee par
    # la post-correction SELAS-only de build_generation_context. La fixture SELARL
    # sert seulement de squelette d'etat civil/financement valide.
    acquereur_selas = cession.acquereur.model_copy(
        update={"denomination_societe": "SELAS EXEMPLE", "forme_sociale": "SELAS"}
    )
    cession = cession.model_copy(update={"acquereur": acquereur_selas})
    scm = scm_cession_fixture()
    # La fixture SCM SELARL porte variante_structure='selarl' ; on PROUVE que le
    # cablage SELAS la realigne (build_generation_context normalise -> 'selas').
    assert scm.variante_structure == "selarl"
    # O24-11 (MINEUR 2) : le modele SCM n'a qu'UN placeholder situation_maritale ; le
    # front y injecte le libelle COMPLET ACCENTUE via le VRAI helper du shell. On
    # DERIVE donc la valeur de _scm_cedant_situation_maritale_display(prefix='selas')
    # — exactement le code que le sous-formulaire SELAS appelle — au lieu de coder en
    # dur une valeur que le helper ne produit JAMAIS (« communauté légale »). Ainsi
    # ce test e2e exerce REELLEMENT le chemin O24-11 et regresse si le helper casse.
    class _FakeSt:
        def __init__(self) -> None:
            self.session_state = {
                "selas_situation_maritale": (
                    "Marie(e) sous le regime de la communaute reduite aux acquets"
                )
            }

    monkeypatch.setattr(shell, "st", _FakeSt())
    situation_complete = shell._scm_cedant_situation_maritale_display(
        {"situation_maritale": "marie", "genre": Gender.MASCULIN}, prefix="selas"
    )
    # Le helper produit le libelle accentue+accorde reel (jamais « communauté légale »).
    assert situation_complete == "marié sous le régime de communauté réduite aux acquêts"
    scm.cedant.situation_maritale = situation_complete

    payload = {
        **_selas_payload(),
        "cession_context": cession,
        "bail_context": bail,
        "scm_cession_context": scm,
    }

    plan = sms.build_selas_plan(payload)
    assert plan.can_generate is True, plan.reason
    # Le plan ANNONCE bien les codes de cession (cabinet medical acte+compromis,
    # bail, SCM DOC-031/032/033).
    for code in ("DOC-009", "DOC-010", "DOC-007", "DOC-031", "DOC-032", "DOC-033"):
        assert code in plan.document_codes, f"{code} absent du plan SELAS"

    generated = sms.generate_dossier(payload, tmp_path / "selas-cession")
    names = {path.name for path in generated.docx_paths}
    # PREUVE : tous les fichiers de cession sont EMIS (plus de plan menteur, plus de
    # crash « document est obligatoire »).
    expected_cession_files = {
        "acte_cession_cabinet_medical.docx",
        "compromis_cession_cabinet_medical.docx",
        "avenant_contrat_bail.docx",
        "pv_age_cession_parts_scm.docx",
        "courrier_sde_cession_scm.docx",
        "acte_cession_parts_scm.docx",
    }
    missing = expected_cession_files - names
    assert not missing, f"Fichiers de cession SELAS manquants : {sorted(missing)}"

    # MAJEUR 1 : l'acte ET le compromis du cabinet affichent la forme de l'acquereur
    # en « SELAS » (denomination + ligne « <forme> au capital de » + signature), et
    # ne contiennent JAMAIS le doublon auto-contradictoire « SELAS SELARL ».
    for fichier in (
        "acte_cession_cabinet_medical.docx",
        "compromis_cession_cabinet_medical.docx",
    ):
        path = next(p for p in generated.docx_paths if p.name == fichier)
        texte = _docx_text(path)
        assert "SELAS EXEMPLE" in texte, f"{fichier} : denomination acquereur SELAS attendue"
        assert "SELAS au capital de" in texte, (
            f"{fichier} : la forme de l'acquereur doit ressortir « SELAS au capital de »"
        )
        assert "SELAS SELARL" not in texte, (
            f"{fichier} : doublon de forme « SELAS SELARL » (acquereur auto-contradictoire)"
        )
        # Aucune fuite residuelle de « SELARL » pour l'acquereur (denomination ou forme).
        assert "SELARL au capital de" not in texte, (
            f"{fichier} : forme acquereur figee « SELARL au capital de » (fidelite SELAS)"
        )

    # O24-11 : l'acte SCM emis affiche « marié sous le régime de <regime> avec
    # <conjoint> » ACCENTUE (le generateur ajoute « avec <conjoint> » via mentions_conjoint).
    acte_scm = next(p for p in generated.docx_paths if p.name == "acte_cession_parts_scm.docx")
    texte = _docx_text(acte_scm)
    assert (
        "marié sous le régime de communauté réduite aux acquêts avec Madame Claire Dupont"
        in texte
    ), "L'acte SCM doit afficher la situation matrimoniale accentuee + conjoint (O24-11)."


class _StScmStub:
    """Mock Streamlit minimal pour piloter `_render_scm_cession_form` hors AppTest.

    Chaque widget lit/ecrit dans un `session_state` partage : `text_input` /
    `number_input` / `selectbox` / `checkbox` renvoient la valeur de session (ou un
    defaut). Les conteneurs (`columns`) renvoient le stub lui-meme : `_cession_text`
    appelle `container.text_input(label, key=...)`, donc le meme objet suffit. Les
    expanders sont des context managers no-op. Couvre le maillon UI du sous-formulaire
    SCM, aujourd'hui non exerce (les tests existants n'appellent la fonction qu'avec
    scm=False)."""

    def __init__(self, session_state: dict[str, object]) -> None:
        self.session_state = session_state
        self.warnings: list[str] = []

    # --- conteneurs / mise en page ---
    def columns(self, spec):
        n = spec if isinstance(spec, int) else len(spec)
        return tuple(self for _ in range(n))

    class _Expander:
        def __enter__(self):
            return None

        def __exit__(self, *exc):
            return False

    def expander(self, *_a, **_k):
        return self._Expander()

    # --- widgets ---
    def text_input(self, _label, *, key=None, disabled=False, value=None, **_k):
        if value is not None:
            return value
        return self.session_state.get(key, "")

    def number_input(self, _label, *, key=None, min_value=None, step=None, **_k):
        if key in self.session_state:
            return self.session_state[key]
        return min_value if min_value is not None else 0

    def selectbox(self, _label, options, *, key=None, **_k):
        if key in self.session_state:
            return self.session_state[key]
        return options[0] if options else None

    def checkbox(self, _label, *, key=None, **_k):
        return bool(self.session_state.get(key, False))

    def text_area(self, _label, *, key=None, **_k):
        return self.session_state.get(key, "")

    def date_input(self, _label, *, key=None, value=None, **_k):
        if key in self.session_state:
            return self.session_state[key]
        return value

    def button(self, *_a, **_k):
        return False

    # --- divers no-op ---
    def markdown(self, *_a, **_k):
        return None

    def caption(self, *_a, **_k):
        return None

    def info(self, *_a, **_k):
        return None

    def warning(self, message, *_a, **_k):
        self.warnings.append(str(message))


def test_render_scm_cession_form_selas_injecte_situation_accentuee(monkeypatch) -> None:
    # MINEUR 2 (ideal) : couvre le MAILLON UI `_render_scm_cession_form(prefix='selas')`
    # — non exerce par les tests existants. On pilote le sous-formulaire avec un stub
    # Streamlit minimal et on PROUVE que le ScmCessionContext RETOURNE par le formulaire
    # porte deja le libelle de situation maritale du cedant COMPLET et ACCENTUE
    # (statut/genre + « sous le régime de » + regime), via le vrai helper du shell. Si
    # le maillon UI cessait d'injecter ce libelle, ce test regresserait (le e2e seul ne
    # couvrait pas l'appel direct du formulaire).
    from sydel_doc_engine.front_app import shell

    session_state = {
        # Total des presents = nb_parts_total de la fixture (300) -> aucun warning de
        # coherence et un signataires_pv derive proprement.
        "selas_cession_scm_presents_count": 1,
        "selas_cession_scm_present_0_civilite": "Monsieur",
        "selas_cession_scm_present_0_prenom": "Jean",
        "selas_cession_scm_present_0_nom": "Dupont",
        "selas_cession_scm_present_0_nb_parts": "300",
        "selas_cession_scm_present_0_plage": "1 a 300",
        # Libelle BRUT du preset lu par _scm_cedant_situation_maritale_display.
        "selas_situation_maritale": "Marié(e) sous le régime de la séparation de biens",
    }
    monkeypatch.setattr(shell, "st", _StScmStub(session_state))

    praticien = {
        "civilite": "Monsieur",
        "prenom": "Jean",
        "nom": "Dupont",
        "situation_maritale": "marie",
        "genre": Gender.MASCULIN,
        "nationalite": "française",
    }
    societe = {"denomination": "SELAS EXEMPLE", "capital_social": "1 000", "ville_rcs": "Lyon"}

    scm_ctx = shell._render_scm_cession_form(
        True,
        praticien=praticien,
        societe=societe,
        profession_label="médecin",
        ordre={"departement_ordre": "Rhône"},
        prefix="selas",
    )

    assert scm_ctx is not None
    # Le maillon UI a injecte le libelle COMPLET ACCENTUE (regime separation, accorde
    # au masculin) — pas la valeur collapsee « marie ».
    assert scm_ctx.cedant.situation_maritale == (
        "marié sous le régime de séparation de biens"
    )
    # La denomination de la SEL cessionnaire suit bien la fiche societe (SELAS EXEMPLE).
    assert scm_ctx.cessionnaire.denomination == "SELAS EXEMPLE"


def test_selas_cession_normalise_variante_structure_scm() -> None:
    # TACHE A : build_generation_context realigne scm_cession.variante_structure sur
    # la structure du dossier ('selas'), sinon l'orchestrateur n'emet pas la cession
    # SCM (service._scm_cession_enabled) et le generateur SCM leve.
    from sydel_doc_engine.front_app import selas_multi_slice as sms
    from sydel_doc_engine.scenarios.selarl import scm_cession_fixture

    scm = scm_cession_fixture()
    assert scm.variante_structure == "selarl"
    payload = {**_selas_payload(), "scm_cession_context": scm}
    ctx = sms.build_generation_context(payload)
    assert ctx.scm_cession is not None
    assert ctx.scm_cession.variante_structure == "selas"


def test_selas_cession_vendeur_selectionnable(tmp_path: Path, monkeypatch) -> None:
    # #11 (onglet 24) : en SELAS, un menu permet de choisir le vendeur du cabinet
    # parmi les associes ; le wording « associe unique » disparait.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-vendeur")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)

    selectbox_keys = {str(s.key) for s in app.selectbox}
    assert "selas_cession_vendeur_index" in selectbox_keys
    # #10 (onglet 24) : le type de cabinet est derive de la profession -> pas de menu.
    assert "selas_cession_meta_type_cabinet" not in selectbox_keys
    # R16 (Rafael 2026-06-23) : plus de choix d'etape (acte/compromis generes ensemble).
    assert "selas_cession_meta_etape" not in selectbox_keys
    checkbox_labels = " ".join(str(c.label).lower() for c in app.checkbox)
    assert "associe unique" not in checkbox_labels
    assert "associé unique" not in checkbox_labels


def test_selas_profession_field_removed(tmp_path: Path, monkeypatch) -> None:
    # #9 (onglet 24) : le champ « profession » de l'associe est retire du formulaire
    # (redondant avec la qualification ; titre « Docteur » derive cote moteur ->
    # comparution inchangee). cf. docs/review/QUESTIONS_RAFAEL.md #9.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-prof")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert not any(k.endswith("_profession") for k in keys)


def test_selas_adresses_sur_une_ligne(tmp_path: Path, monkeypatch) -> None:
    # O24-03 (onglet 24) : toutes les adresses sur UNE ligne, pas de champ separe.
    # Le siege et l'adresse personnelle sont des champs texte uniques ; les anciennes
    # cles structurees (siege_num/voie/cp/ville) et la case « siege = lieu d'exercice »
    # n'existent plus (la case migre sur le CABINET en cession, O24-12).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-adresses")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    # siege sur une ligne, valeur pre-remplie complete
    assert "selas_siege_adresse" in keys
    siege = next(w for w in app.text_input if str(w.key) == "selas_siege_adresse")
    assert siege.value == "5 place du Centre, 69000 Lyon"
    # adresse perso de l'associe 0 sur une ligne
    assert "selas_associe_0_adresse" in keys
    # plus aucun champ structuré (siège), ni double champ siège, ni case « siège = lieu d'exercice »
    assert "selas_siege_voie" not in keys
    assert "selas_siege_num" not in keys
    assert "selas_siege" not in keys  # plus de champ « Siege (adresse affichee) » libre (doublon)
    assert "selas_siege_same_as_lieu_exercice" not in {str(c.key) for c in app.checkbox}
    # O24-03 : l'adresse de l'ORDRE aussi sur une ligne (plus de CP/Ville ordre séparés)
    assert "selas_ordre_adresse" in keys
    assert "selas_ordre_cp" not in keys
    assert "selas_ordre_ville" not in keys


@pytest.mark.parametrize(
    "label, prefix, oneline_key",
    [
        # O24-03 (re-Akainu T4) : adresse de l'ORDRE sur UNE ligne pour les types restants
        # (SELAS multi deja couvert par test_selas_adresses_sur_une_ligne). Le SELARL suit
        # sa convention de suffixe « _ligne » (comme siege_ligne / adresse_ligne).
        ("SELARL", "selarl", "selarl_ordre_adresse_ligne"),
        ("SCM", "scm", "scm_ordre_adresse"),
        ("SPFPL dentistes — cession", "spfpl_cession", "spfpl_cession_ordre_adresse"),
    ],
)
def test_ordre_adresse_sur_une_ligne(
    label, prefix, oneline_key, tmp_path: Path, monkeypatch
) -> None:
    # L'adresse de l'ordre est un champ texte UNIQUE ; les anciens champs separes
    # (« *_ordre_cp » / « *_ordre_ville » / « *_ordre_adresse_ligne_1 ») n'existent plus
    # comme widgets.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / f"ui-ordre-{prefix}")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(label)
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert oneline_key in keys
    assert f"{prefix}_ordre_cp" not in keys
    assert f"{prefix}_ordre_ville" not in keys
    assert f"{prefix}_ordre_adresse_ligne_1" not in keys


def test_selas_uni_medecin_ordre_adresse_sur_une_ligne(tmp_path: Path, monkeypatch) -> None:
    # O24-03 (re-Akainu T4) : SELAS unipersonnelle medecin — adresse de l'ordre sur UNE
    # ligne (pas de bouton de donnees de test pour ce type ; le widget existe des le
    # rendu du formulaire). Plus de champs separes CP/Ville ordre.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-ordre-selas-uni")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS unipersonnelle médecin")
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert "selas_uni_medecin_ordre_adresse" in keys
    assert "selas_uni_medecin_ordre_cp" not in keys
    assert "selas_uni_medecin_ordre_ville" not in keys
    assert "selas_uni_medecin_ordre_adresse_ligne_1" not in keys


def test_su1_selas_uni_situation_un_seul_champ_menu(tmp_path: Path, monkeypatch) -> None:
    # SU1 (Albane 2026-06-25) : la SELAS unipersonnelle aligne son deroule matrimonial
    # sur la SELARL — UN SEUL champ « Situation matrimoniale » (menu preset). Le regime
    # matrimonial ET le regime de communaute (DOC-005/006) sont DERIVES du libelle ;
    # plus de double champ texte (situation + regime) ni de case a cocher dediee.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-su1-selas-uni")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS unipersonnelle médecin")
    app = app.run(timeout=180)

    selectbox_keys = {str(w.key) for w in app.selectbox}
    text_keys = {str(w.key) for w in app.text_input}
    # UN SEUL champ : le menu « Situation matrimoniale ».
    assert "selas_uni_medecin_situation" in selectbox_keys
    # Plus AUCUN champ texte situation/regime separe, ni case a cocher dediee.
    assert "selas_uni_medecin_situation_maritale" not in text_keys
    assert "selas_uni_medecin_regime_matrimonial" not in text_keys
    assert "selas_uni_medecin_regime_communautaire" not in text_keys
    assert "selas_uni_medecin_regime_communautaire" not in selectbox_keys
    # SU2 (M2 Akainu) : le connecteur grammatical existe sur le formulaire SELAS uni et
    # offre « des » (sinon « conseil departemental des Hauts de Seine » est inatteignable).
    connecteur_box = next(
        w for w in app.selectbox if str(w.key) == "selas_uni_medecin_ordre_connecteur"
    )
    assert "des" in [str(o) for o in connecteur_box.options]


def test_r5_selas_uni_destinataire_forme_longue_connecteur_des(tmp_path: Path) -> None:
    # R5 (Albane) SUPERSEDE SU2 (25/06) : le destinataire revient a la FORME LONGUE
    # « Conseil départemental de l’Ordre des <profession_pluriel> <connecteur> <departement> »
    # (profession PUIS departement, forme du modele d'Albane, confirme 2026-07-01).
    # Cas connecteur « des » (departement « Hauts de Seine » -> « des Hauts de Seine ») : le
    # rendu porte alors deux « des » (« ... de l’Ordre des médecins des Hauts de Seine »), ce
    # qui est l'accord correct du departement « (les) Hauts-de-Seine ».
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    payload = _selas_uni_medecin_payload()
    payload["departement_ordre"] = "Hauts de Seine"
    payload["connecteur_departement"] = "des"
    generated = uni.generate_dossier(payload, tmp_path / "selas-uni-r5")
    text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "demande_inscription_ordre.docx")
    )
    assert "Conseil départemental de l’Ordre des médecins des Hauts de Seine" in text


@pytest.mark.parametrize(
    "label, prefix, has_siege, perso_keys",
    [
        # O24-03 (propagation Q4) : siege + adresse perso sur UNE ligne pour TOUS les
        # types restants (civils via repeater partage, SAS double-saisie, SPFPL).
        ("SCI", "sci", True, ("sci_associe_0_adresse",)),
        ("SCM", "scm", True, ("scm_associe_0_adresse",)),
        ("SCS", "scs", True, ("scs_associe_0_adresse",)),
        ("SCI IRIS", "sci_iris", True, ("sci_iris_associe_1_adresse",)),
        ("SPFPL médecins (forme SAS)", "sas", True, ("sas_adresse",)),
        ("SPFPL dentistes — apport", "spfpl_apport", True, ("spfpl_apport_adresse",)),
        (
            "SPFPL dentistes — cession",
            "spfpl_cession",
            True,
            ("spfpl_cession_adresse",),
        ),
    ],
)
def test_o24_03_adresses_une_ligne_par_type(
    tmp_path: Path, monkeypatch, label: str, prefix: str, has_siege: bool, perso_keys
) -> None:
    """O24-03 : plus AUCUN champ separe No/Voie/CP/Ville (siege ni perso) ; un champ
    une-ligne a la place. Verifie les cles de widget apres « donnees de test », puis
    que la generation reste propre (DOCX sans token residuel)."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / f"ui-{prefix}")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value(label)
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    # siege sur UNE ligne (civils : « {prefix}_siege_adresse » ; SAS/SPFPL : « {prefix}_siege »),
    # aucune cle de composant separee
    if has_siege:
        assert f"{prefix}_siege_adresse" in keys or f"{prefix}_siege" in keys
        for comp in ("siege_num", "siege_voie", "siege_cp", "siege_ville"):
            assert f"{prefix}_{comp}" not in keys, f"{prefix}_{comp} doit avoir disparu (O24-03)"
    # adresse personnelle sur UNE ligne, aucune cle de composant separee
    for pk in perso_keys:
        assert pk in keys, f"{pk} (champ une-ligne) doit exister"
        base = pk[: -len("_adresse")]  # ex. « sas » ou « sci_associe_0 »
        for comp in ("adresse_num", "adresse_voie", "adresse_cp", "adresse_ville"):
            stray = f"{base}_{comp}"
            assert stray not in keys, f"{stray} doit avoir disparu (O24-03)"

    # La generation doit rester possible et propre (bouton actif, pas de blocage).
    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)


def test_o24_03_spfpl_cession_cible_siege_une_ligne(tmp_path: Path, monkeypatch) -> None:
    """O24-03 : le siege de la societe cible (cession SPFPL) est sur UNE ligne — plus de
    grille No/Voie/CP/Ville ni double-saisie avec le champ « Siege cible (affiche) ».
    Retour Rafael 2026-07-07 : le champ « Siege cible (affiche) » (`cible_siege`) et la
    saisie « Valeur globale apportee » (`apport_valeur_globale`, fusionnee dans
    « Montant de l'apport ») ont DISPARU du formulaire."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-spfpl-cession")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SPFPL dentistes — cession")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert "spfpl_cession_cible_siege_cession" in keys
    for comp in ("cible_siege_num", "cible_siege_voie", "cible_siege_cp", "cible_siege_ville"):
        assert f"spfpl_cession_{comp}" not in keys, f"{comp} doit avoir disparu (O24-03)"
    # Rafael 2026-07-07 : plus de champ « Siege cible (affiche) » ni « Valeur globale ».
    assert "spfpl_cession_cible_siege" not in keys
    assert "spfpl_cession_apport_valeur_globale" not in keys
    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)


def test_spfpl_apport_cible_siege_une_ligne_et_valeur_globale_derivee(
    tmp_path: Path, monkeypatch
) -> None:
    """Retour Rafael 2026-07-07 (APPORT) : le bloc « Societe cible » porte la ligne
    unique + case « Meme adresse que le siege de la SPFPL » (patron cession reutilise)
    a la place du champ libre « Siege cible (affiche) » ; « Valeur globale apportee »
    n'est plus saisie (derivee de « Montant de l'apport »). La generation APPORT
    complete (bundle creation + DOC-041/042/043) reste possible et aboutit."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-spfpl-apport")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SPFPL dentistes — apport")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert "spfpl_apport_cible_siege_cession" in keys  # ligne unique partagee
    assert "spfpl_apport_cible_siege" not in keys  # champ libre supprime
    assert "spfpl_apport_apport_valeur_globale" not in keys  # fusion : plus de saisie
    checkbox_keys = {str(w.key) for w in app.checkbox}
    assert "spfpl_apport_cible_siege_meme_spfpl" in checkbox_keys  # case meme-adresse

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)
    # Generation apport COMPLETE de bout en bout (verifie la derivation cible_siege +
    # apport_valeur_globale jusqu'aux DOCX : le ZIP + les telechargements apparaissent).
    generate_button.click()
    app = app.run(timeout=300)
    labels = [item.label for item in app.get("download_button")]
    assert any("ZIP" in label for label in labels)
    assert any("contrat_apport" in label for label in labels)


def test_o24_03_selarl_adresses_une_ligne(tmp_path: Path, monkeypatch) -> None:
    """O24-03 : SELARL (type de reference, gold) — adresse perso ET siege sur UNE ligne.
    Plus aucun champ separe Numero-et-voie / CP / Ville. La conversion alimente les memes
    cles que l'ancienne grille -> generateur et DOCX gold inchanges (verifie par le test
    line-by-line dedie)."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selarl-adresses")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    # SELARL est le type par defaut ; bouton « Generer des donnees de test » dedie.
    next(b for b in app.button if str(b.key) == "clean_generate_test_data").click()
    app = app.run(timeout=180)

    keys = {str(w.key) for w in app.text_input}
    assert "selarl_adresse_ligne" in keys
    assert "selarl_siege_ligne" in keys
    for stray in (
        "selarl_adresse_voie",
        "selarl_adresse_cp",
        "selarl_adresse_ville",
        "selarl_siege_voie",
        "selarl_siege_cp",
        "selarl_siege_ville",
    ):
        assert stray not in keys, f"{stray} doit avoir disparu (O24-03)"


def test_parse_address_full_tolere_formes_usuelles() -> None:
    # O24-03 (onglet 24) : la saisie sur UNE ligne doit accepter les formes françaises
    # courantes (virgule après le numéro, sans virgule), pas seulement « num voie, cp ville ».
    from sydel_doc_engine.front_app.selas_multi_slice import _parse_address_full

    attendu = ("12", "rue de la Paix", "75001", "Paris")
    for forme in (
        "12 rue de la Paix, 75001 Paris",
        "12, rue de la Paix, 75001 Paris",
        "12 rue de la Paix 75001 Paris",
    ):
        a = _parse_address_full(forme)
        assert a is not None, forme
        assert (a.num_voie, a.voie, a.cp, a.ville) == attendu, forme

    # re-Akainu 2026-06-23 (MAJEUR O24-03) : une voie contenant une ANNÉE ne doit jamais voir
    # l'année prise pour un code postal (« 8 Mai 1945 » : 1945 = 4 chiffres ; le CP français =
    # DERNIER groupe de 5 chiffres). Sinon cp/ville ET l'affichage sont corrompus silencieusement.
    a = _parse_address_full("10 avenue du 8 Mai 1945, 33700 Mérignac")
    assert a is not None
    assert (a.num_voie, a.voie, a.cp, a.ville) == (
        "10",
        "avenue du 8 Mai 1945",
        "33700",
        "Mérignac",
    )
    assert a.adresse_affichee == "10 avenue du 8 Mai 1945, 33700 Mérignac"
    a = _parse_address_full("5 rue du 11 Novembre 1918, 69100 Villeurbanne")
    assert a is not None and a.cp == "69100" and a.ville == "Villeurbanne"

    # Bug Albane 2026-07-08 : une adresse avec COMPLEMENT en tete (nom de residence /
    # batiment / lieu-dit AVANT le numero) etait rejetee (None) -> generation bloquee
    # a tort (« adresse/CP/ville de l'ordre requis ») alors que le champ etait rempli.
    # Le complement est colle au numero (num_voie), ordre preserve, zero perte.
    a = _parse_address_full("Maison Blanche 14 boulevard Carabacel, 06000 NICE")
    assert a is not None
    assert (a.num_voie, a.voie, a.cp, a.ville) == (
        "Maison Blanche 14",
        "boulevard Carabacel",
        "06000",
        "NICE",
    )
    assert a.adresse_affichee == "Maison Blanche 14 boulevard Carabacel, 06000 NICE"
    a = _parse_address_full("Résidence Les Tilleuls, 12 rue de la Paix, 75001 Paris")
    assert a is not None and a.cp == "75001" and a.num_voie == "Résidence Les Tilleuls, 12"
    # Non-regression : lieu-dit SANS aucun numero reste None (arbitrage Albane en attente).
    assert _parse_address_full("Place de l'Église, 75001 Paris") is None

    # re-Akainu tour 2 (NITPICK O24-03) : suffixe bis/ter/quater détaché du numéro, en
    # MAJUSCULES comme en minuscules.
    a = _parse_address_full("12 BIS rue de la Paix 75001 Paris")
    assert a is not None and a.num_voie == "12 BIS" and a.voie == "rue de la Paix"

    # re-Akainu tour 2 (MAJEUR O24-03) : le numéro de tête est REQUIS (cohérence avec
    # _validate_common_docs qui exige siege_num / signataire_adresse_num). Une adresse sans
    # numéro (lieu-dit, place) renvoie None tant qu'Albane n'a pas tranché ce périmètre
    # (docs/review/QUESTIONS_RAFAEL.md) — pas d'extrapolation.
    assert _parse_address_full("Lieu-dit Le Bourg, 12340 Bozouls") is None

    assert _parse_address_full("") is None
    assert _parse_address_full("pas une adresse") is None


def test_selarl_valeur_nominale_affichee_dans_un_champ(tmp_path: Path, monkeypatch) -> None:
    # O24-05 (onglet 24) : valeur nominale « calculée automatiquement ET affichée dans le
    # champ concerné » — y compris SELARL (était un st.caption gris, pas un champ).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selarl-vn")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    labels = [str(w.label) for w in app.text_input]
    assert any("Valeur nominale d'une part (calculee)" in s for s in labels), labels
    # re-Akainu 2026-06-23 (MINEUR O24-05) : le verbatim exige « calculée automatiquement ET
    # affichée ». On prouve la moitié « calculée » : capital 1000 / 100 parts -> value == '10'
    # injectée DANS le champ (désactivé), pas seulement la présence du libellé.
    app.number_input(key="selarl_capital_social").set_value(1000)
    app.number_input(key="selarl_nb_parts_total").set_value(100)
    app = app.run(timeout=180)
    vn = next(
        w
        for w in app.text_input
        if "Valeur nominale d'une part (calculee)" in str(w.label)
    )
    assert vn.value == "10"
    assert vn.disabled is True


def test_selarl_scm_cedee_valeur_nominale_calculee(tmp_path: Path, monkeypatch) -> None:
    # re-Akainu tour 2 (MINEUR O24-05) : « valeur nominale calculée ET affichée » s'applique
    # AUSSI à la SCM cédée (sous-formulaire de cession), pas seulement aux 6 types principaux.
    # Le champ devient désactivé + auto-calculé (capital SCM / nb parts SCM), plus de saisie libre.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selarl-scm-vn")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELARL")
    app = app.run(timeout=180)
    app.checkbox(key="selarl_scm").set_value(True)
    app = app.run(timeout=180)
    # F1 (Albane 2026-07-10) : le flux SCM ne préremplit plus les champs -> on saisit
    # capital 3 000 / 300 parts pour vérifier la dérivation de la valeur nominale (= 10).
    _set_text_widget(app, "selarl_cession_scm_cedee_capital_social", "3 000")
    _set_text_widget(app, "selarl_cession_scm_cedee_nb_parts_total", "300")
    app = app.run(timeout=180)
    # le champ valeur nominale de la SCM cédée : désactivé + auto-calculé (label distinct de
    # la valeur nominale de la société pour éviter un DuplicateWidgetID).
    scm_vn = next(
        w for w in app.text_input
        if "Valeur nominale d'une part de SCM (calculee)" in str(w.label)
    )
    assert scm_vn.disabled is True
    # Capital 3 000 / 300 parts saisis -> valeur nominale 10 dérivée et affichée.
    assert scm_vn.value == "10"
    # plus aucun champ « Valeur nominale d'une part » LIBRE (éditable) ne subsiste.
    assert not any(
        str(w.label).strip().startswith("Valeur nominale d'une part") and not w.disabled
        for w in app.text_input
    )


def test_selas_cession_vendeur_regime_complet(tmp_path: Path, monkeypatch) -> None:
    # O24-11 (onglet 24) : « toutes ses informations soient retranscrites » -> la situation
    # matrimoniale du vendeur est reprise en LIBELLE COMPLET (preset), pas aplatie en « marié »,
    # pour que la cession dérive le bon régime (séparation, etc.).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    preset_sep = "Marié(e) sous le régime de la séparation de biens"
    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-vendeur-regime")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    app.selectbox(key="selas_associe_0_situation").set_value(preset_sep)
    app = app.run(timeout=180)
    # re-Akainu 2026-06-23 (MAJEUR O24-11) : le conjoint doit être capté pour TOUT régime marié
    # (séparation incluse), pas seulement la communauté légale -> les champs conjoint apparaissent.
    conj_keys = {str(w.key) for w in app.text_input}
    assert "selas_associe_0_conjoint_prenom" in conj_keys
    assert "selas_associe_0_conjoint_nom" in conj_keys
    for w in app.text_input:
        if str(w.key) == "selas_associe_0_conjoint_prenom":
            w.set_value("Alex")
        elif str(w.key) == "selas_associe_0_conjoint_nom":
            w.set_value("Separe")
    app = app.run(timeout=180)
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)
    # le libellé complet (préset) est propagé POUR LA DÉRIVATION DU RÉGIME, pas le mot aplati
    assert app.session_state["selas_situation_maritale"] == preset_sep


def test_selas_vendeur_situation_dissociee_du_regime() -> None:
    # re-Akainu 2026-06-23 (MAJEUR O24-11) : le libellé BRUT du preset sert à DÉRIVER le RÉGIME,
    # mais l'AFFICHAGE de la situation doit être COLLAPSÉ + accentué (« marié »/« mariée »), jamais
    # le libellé brut « Marie(e) sous le régime ... » (sinon l'acte rend non accentué + régime
    # doublé + artefact « (e) »). On verrouille la dissociation au niveau du contrat de fonctions.
    from sydel_doc_engine.domain.enums import Gender
    from sydel_doc_engine.front_app.field_derivations import matrimonial_status_value
    from sydel_doc_engine.front_app.shell import _situation_display, _vendeur_regime_label

    preset = "Marié(e) sous le régime de la communauté universelle"
    # régime dérivé du libellé BRUT
    assert _vendeur_regime_label(preset) == "communauté universelle"
    # affichage = valeur COLLAPSÉE accentuée (ce qui part dans situation_maritale du praticien)
    assert _situation_display(matrimonial_status_value(preset), Gender.FEMININ) == "mariée"
    assert _situation_display(matrimonial_status_value(preset), Gender.MASCULIN) == "marié"
    # le libellé brut (avec « (e) ») ne doit JAMAIS être ce qui s'affiche
    affiche = _situation_display(matrimonial_status_value(preset), Gender.FEMININ)
    assert "(e)" not in affiche and "regime" not in affiche.lower()


def test_selas_pacs_affiche_partenaire(tmp_path: Path, monkeypatch) -> None:
    # Albane 6.3/7.3 (RATIFIE 2026-07-06) : SUPERSEDE l'exclusion PACS (O24-11). Le partenaire
    # d'un associe PACSE est desormais capte et REPRIS dans les docs (« pacsé avec {partenaire} »)
    # -> choisir « Pacsé(e) » DOIT afficher les champs prenom/nom du partenaire (comme le mariage).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-pacs")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    app.selectbox(key="selas_associe_0_situation").set_value("Pacsé(e)")
    app = app.run(timeout=180)
    keys = {str(w.key) for w in app.text_input}
    assert "selas_associe_0_conjoint_prenom" in keys
    assert "selas_associe_0_conjoint_nom" in keys


def test_selas_cession_cabinet_meme_adresse_lieu_exercice(tmp_path: Path, monkeypatch) -> None:
    # O24-12 (onglet 24) : « adresse du cabinet -> ajouter une case "meme adresse que le
    # lieu d'exercice" et reporter les donnees si cochee ». La case est sur le CABINET
    # (cession), pas sur le siege.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-cabinet")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    # rendre le LIEU D'EXERCICE distinct du siège -> test discriminant (sinon on ne sait pas
    # si le cabinet reprend le lieu d'exercice ou le simple pré-remplissage du siège).
    next(
        w for w in app.text_input if str(w.key) == "selas_adresse_lieu_exercice"
    ).set_value("99 avenue Distincte, 75001 Paris")
    app = app.run(timeout=180)
    # activer la cession de cabinet liberal
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)
    assert "selas_cabinet_meme_lieu_exercice" in {str(c.key) for c in app.checkbox}
    # décochée : le cabinet = le SIÈGE (pré-rempli), PAS le lieu d'exercice distinct
    cab = next(w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse")
    assert cab.value == "5 place du Centre, 69000 Lyon"
    # cochée -> le cabinet reprend le LIEU D'EXERCICE distinct (preuve du report O24-12)
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(True)
    app = app.run(timeout=180)
    cab = next(w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse")
    assert cab.value == "99 avenue Distincte, 75001 Paris"
    # re-Akainu 2026-06-23 (MAJEUR O24-12) : DÉCOCHER doit DÉFAIRE le report -> retour au SIÈGE
    # (sinon l'adresse du cabinet reste polluée par le lieu d'exercice après un décochage).
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(False)
    app = app.run(timeout=180)
    cab = next(w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse")
    assert cab.value == "5 place du Centre, 69000 Lyon"
    # re-Akainu tour 2 (MINEUR O24-12) : une adresse saisie MANUELLEMENT ne doit pas être perdue
    # à la décoche -> saisie manuelle, coche (report), décoche : on RETROUVE la saisie manuelle.
    next(
        w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse"
    ).set_value("7 rue Manuelle, 13000 Marseille")
    app = app.run(timeout=180)
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(True)
    app = app.run(timeout=180)
    cab = next(w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse")
    assert cab.value == "99 avenue Distincte, 75001 Paris"  # report écrase l'affichage
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(False)
    app = app.run(timeout=180)
    cab = next(w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse")
    assert cab.value == "7 rue Manuelle, 13000 Marseille"  # saisie manuelle restaurée


def test_selas_cession_cabinet_meme_adresse_etat_disabled_aux_3_phases(
    tmp_path: Path, monkeypatch
) -> None:
    # O24-12 (MINEUR 3, re-Akainu T5) : scénario « lieu renseigné » — on asserte EXPLICITEMENT
    # l'état `disabled` du champ adresse cabinet aux 3 phases (décoché -> coché -> décoché),
    # pas seulement sa valeur. Coché + lieu renseigné = champ DÉSACTIVÉ (miroir non éditable) ;
    # décoché = ÉDITABLE.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-disabled")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    next(
        w for w in app.text_input if str(w.key) == "selas_adresse_lieu_exercice"
    ).set_value("99 avenue Distincte, 75001 Paris")
    app = app.run(timeout=180)
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)

    def _cab():
        return next(
            w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse"
        )

    # Phase 1 — décoché : champ ÉDITABLE.
    assert _cab().disabled is False
    # Phase 2 — coché (lieu renseigné) : champ DÉSACTIVÉ (miroir du lieu d'exercice).
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(True)
    app = app.run(timeout=180)
    assert _cab().disabled is True
    assert _cab().value == "99 avenue Distincte, 75001 Paris"
    # Phase 3 — décoché : champ RE-ÉDITABLE.
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(False)
    app = app.run(timeout=180)
    assert _cab().disabled is False


def test_selas_cession_cabinet_meme_adresse_lieu_vide_reste_editable(
    tmp_path: Path, monkeypatch
) -> None:
    # O24-12 (MINEUR 3, re-Akainu T5) : scénario « lieu d'exercice VIDE ». La case existe
    # (verbatim « ajouter une case », inconditionnel) mais cocher ne reporte RIEN (rien à
    # recopier) -> le champ reste ÉDITABLE même coché. Une saisie manuelle survit au cycle
    # coche -> décoche (jamais écrasée par un report fantôme).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-lieu-vide")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    # Lieu d'exercice VIDE -> rien à reporter dans l'adresse cabinet.
    next(
        w for w in app.text_input if str(w.key) == "selas_adresse_lieu_exercice"
    ).set_value("")
    app = app.run(timeout=180)
    app.checkbox(key="selas_cession_on").set_value(True)
    app = app.run(timeout=180)

    def _cab():
        return next(
            w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse"
        )

    # Phase 1 — décoché : champ ÉDITABLE.
    assert _cab().disabled is False
    # Phase 2 — coché mais lieu VIDE : champ reste ÉDITABLE (pas de miroir possible).
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(True)
    app = app.run(timeout=180)
    assert _cab().disabled is False
    # Saisie MANUELLE alors que la case est cochée (champ éditable car lieu vide).
    next(
        w for w in app.text_input if str(w.key) == "selas_cession_cabinet_adresse"
    ).set_value("12 rue Saisie, 31000 Toulouse")
    app = app.run(timeout=180)
    # Phase 3 — décoché : la saisie manuelle est RETROUVÉE (jamais écrasée).
    app.checkbox(key="selas_cabinet_meme_lieu_exercice").set_value(False)
    app = app.run(timeout=180)
    assert _cab().disabled is False
    assert _cab().value == "12 rue Saisie, 31000 Toulouse"


def test_cession_type_et_label_derives_de_la_profession() -> None:
    # O24-10 (onglet 24) : le menu « type de cabinet » est supprimé, le type + la nature
    # du fonds sont DÉRIVÉS de la profession. Bug Akainu : « chirurgien-dentiste » (tiret,
    # forme du menu SELAS) ne matchait pas PROFESSION_DENTISTE (« chirurgien_dentiste »,
    # underscore) -> un cabinet DENTAIRE générait des docs MÉDICAUX. La dérivation doit
    # être tolérante à la forme.
    from sydel_doc_engine.front_app.shell import _cession_default_type, _profession_label

    for dentiste in ("chirurgien-dentiste", "chirurgien_dentiste", "Chirurgien-dentiste"):
        assert _cession_default_type(dentiste) == "dentaire"
        assert _profession_label(dentiste) == "chirurgien-dentiste"
    assert _cession_default_type("médecin") == "medical"
    assert _profession_label("médecin") == "médecin"


def test_front_selas_dentiste_pluri_uses_dentiste_corpus(
    tmp_path: Path, monkeypatch
) -> None:
    # Cas unique « SELAS pluripersonnelle » (retours Rafael 2026-06-23) : la profession
    # se choisit DANS le formulaire. En selectionnant « chirurgien-dentiste », la
    # generation doit produire les STATUTS DENTISTE (corpus dentiste), pas medecin.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-dentiste")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)

    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Choisir la profession dentiste dans le formulaire -> bascule le moteur sur le
    # corpus dentiste (plus de pre-reglage par cle de type).
    app.selectbox(key="selas_profession_choice").set_value("chirurgien-dentiste")
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)
    generate_button.click()
    app = app.run(timeout=180)

    statuts_files = list(
        (tmp_path / "ui-selas-dentiste").rglob(_SELAS_STATUTS_NAME)
    )
    assert statuts_files, "Statuts SELAS dentiste non generes."
    text = _docx_text(statuts_files[0])
    # Marqueurs PROPRES au corpus dentiste (absents du corpus medecin).
    assert "l’exercice en commun de la profession de chirurgien-dentiste" in text
    assert "Conseil Départemental de l’Ordre des Chirurgiens-dentistes" in text


def test_front_selas_change_dirigeant_generates(tmp_path: Path, monkeypatch) -> None:
    # Reunion 2026-06-09, feature #1 : on peut designer un autre associe que le
    # premier comme dirigeant ; ses champs DNC se saisissent SOUS lui, et la
    # generation reussit. Prouve le cablage UI checkbox -> president -> DNC.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-dir")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)

    # Prefill (2 associes, dirigeant = associe 0 par defaut).
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Designer l'associe 1 comme dirigeant a la place de l'associe 0.
    app.checkbox(key="selas_associe_0_is_dirigeant").set_value(False)
    app.checkbox(key="selas_associe_1_is_dirigeant").set_value(True)
    app = app.run(timeout=180)

    def set_text(key: str, value: str) -> None:
        for widget in app.text_input:
            if str(widget.key) == key:
                widget.set_value(value)
                return
        raise KeyError(key)

    # DNC du NOUVEAU dirigeant (associe 1) : #8 / B4 — l'adresse et la date sont
    # reprises de l'associe 1 (deja prefill), SEULE la filiation se saisit sous sa
    # case « Dirigeant » (plus de re-saisie d'adresse / date).
    set_text("selas_associe_1_sig_nom_pere", "Paul Martin")
    set_text("selas_associe_1_sig_nom_mere", "Marie Martin")
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    assert not any("Blocage" in item.value for item in app.caption)
    generate_button.click()
    app = app.run(timeout=180)

    download_labels = [item.label for item in app.get("download_button")]
    assert f"Telecharger {_SELAS_STATUTS_NAME}" in download_labels
    assert "Telecharger le dossier ZIP" in download_labels


def test_selas_parse_associe_birthdate_handles_french_long_form() -> None:
    # #8 (onglet 24) : la date de naissance de la DNC est DERIVEE de l'unique champ
    # texte de l'associe (plus de double saisie via un picker dedie). Le parseur
    # accepte la forme longue francaise ET « JJ/MM/AAAA ». Promu en helper PARTAGE
    # (field_derivations) — un seul parseur pour selas_multi, sas, sasu_holding.
    from datetime import date

    from sydel_doc_engine.front_app.field_derivations import parse_associe_birthdate

    assert parse_associe_birthdate("1 janvier 1980") == date(1980, 1, 1)
    assert parse_associe_birthdate("1er janvier 1980") == date(1980, 1, 1)
    assert parse_associe_birthdate("2 février 1982") == date(1982, 2, 2)
    assert parse_associe_birthdate("02/02/1982") == date(1982, 2, 2)
    assert parse_associe_birthdate("") is None
    assert parse_associe_birthdate("pas une date") is None


def test_selas_deux_directeurs_generaux_avertit(tmp_path: Path, monkeypatch) -> None:
    # #7 (onglet 24) : fonctions de direction non cumulatives -> un seul Directeur General admis.
    # KAN-2 (Rafael, rejeté 2×) : ce conflit de rôle NE BLOQUE PLUS la génération ; l'avertissement
    # « Un seul Directeur Général est admis » reste affiché, mais le bouton reste actif.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-dg")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Associe 1 devient dirigeant, puis les deux dirigeants prennent le role DG.
    app.checkbox(key="selas_associe_1_is_dirigeant").set_value(True)
    app = app.run(timeout=180)
    app.selectbox(key="selas_associe_0_role_dirigeant").set_value("Directeur Général")
    app.selectbox(key="selas_associe_1_role_dirigeant").set_value("Directeur Général")
    app = app.run(timeout=180)

    # KAN-2 : le message de conflit de rôle passe de BLOCAGE (caption « À compléter ») à
    # AVERTISSEMENT (st.info), le bouton restant actif.
    infos = " ".join(item.value for item in app.info)
    assert "Un seul Directeur Général est admis" in infos
    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False


def test_selas_role_directeur_general_associe(tmp_path: Path, monkeypatch) -> None:
    # R3 (Rafael 2026-06-23) : le role « Directeur Général Associé » MANQUAIT ; il doit
    # etre proposable. Un Président + un Directeur Général Associé ne bloque pas (les
    # DG Associes peuvent etre plusieurs, contrairement au President et au DG uniques).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-dga")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Associe 1 devient dirigeant ; le role « Directeur Général Associé » est propose.
    app.checkbox(key="selas_associe_1_is_dirigeant").set_value(True)
    app = app.run(timeout=180)
    role_sb = next(
        s for s in app.selectbox if str(s.key) == "selas_associe_1_role_dirigeant"
    )
    assert "Directeur Général Associé" in list(role_sb.options)
    role_sb.set_value("Directeur Général Associé")
    app = app.run(timeout=180)

    # R7 : ce dirigeant a sa propre DNC -> sa filiation est requise pour generer.
    for w in app.text_input:
        if str(w.key) == "selas_associe_1_sig_nom_pere":
            w.set_value("Paul Martin")
        elif str(w.key) == "selas_associe_1_sig_nom_mere":
            w.set_value("Marie Martin")
    app = app.run(timeout=180)

    # Président (associe 0) + DG Associé (associe 1, filiation remplie) -> pas de blocage.
    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False


def test_selas_zero_president_ne_bloque_plus(tmp_path: Path, monkeypatch) -> None:
    # O24-07 (onglet 24) : un Président est attendu (« soit président (un seul) »).
    # KAN-2 (Rafael, rejeté 2×) : « Tous les documents doivent pouvoir être générés. » L'absence
    # de Président désigné NE BLOQUE PLUS la génération (le moteur retombe sur le 1er associé
    # physique). ⚠️ FLAG produit : ce repli silencieux peut requalifier un DG en Président —
    # à confirmer avec le sachant métier (tension KAN-2 « tout générer » vs O24-07).
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-no-pres")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)
    # L'associé 0 (dirigeant par défaut « Président ») est requalifié « Directeur Général ».
    next(
        s for s in app.selectbox if str(s.key) == "selas_associe_0_role_dirigeant"
    ).set_value("Directeur Général")
    app = app.run(timeout=180)
    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False  # KAN-2 : non bloquant (repli 1er associé physique)


def _fake_st_session(monkeypatch, state: dict) -> None:
    """Remplace le `st` du module SELAS par un faux portant session_state = state.

    Permet de tester _collect_dirigeants_nomines_indices / _validate_roles_dirigeants sur
    le chemin PROGRAMMATIQUE (sans run AppTest), là où la requalification silencieuse
    O24-07 se produisait (re-Akainu 2026-06-23)."""
    from types import SimpleNamespace

    from sydel_doc_engine.front_app import selas_multi_slice as sms

    monkeypatch.setattr(sms, "st", SimpleNamespace(session_state=state))


def _phys(n: int) -> list:
    from types import SimpleNamespace

    return [SimpleNamespace(type_personne="personne_physique") for _ in range(n)]


def test_selas_collect_dirigeants_respecte_role_explicite(monkeypatch) -> None:
    # Racine O24-07 (re-Akainu 2026-06-23, MAJEUR) : un dirigeant explicitement « Directeur
    # Général » ne doit JAMAIS être requalifié « Président » par la collecte des dirigeants,
    # même via le chemin programmatique (président_index résolu par fallback, sans Président).
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    _fake_st_session(
        monkeypatch,
        {
            "selas_associe_0_is_dirigeant": True,
            "selas_associe_0_role_dirigeant": "Directeur Général",
            "selas_associe_1_is_dirigeant": True,
            "selas_associe_1_role_dirigeant": "Directeur Général Associé",
        },
    )
    # président_index = 0 (premier dirigeant coché ; AUCUN « Président » explicite)
    dirigeants = dict(sms._collect_dirigeants_nomines_indices(_phys(2), 0))
    assert dirigeants[0] == "Directeur Général"  # PAS « Président » : plus de requalification
    assert dirigeants[1] == "Directeur Général Associé"


def test_selas_deux_presidents_bloque(monkeypatch) -> None:
    # O24-07 (« soit président (UN SEUL) ») : 2 associés « Président » -> blocage.
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    _fake_st_session(
        monkeypatch,
        {
            "selas_associe_0_is_dirigeant": True,
            "selas_associe_0_role_dirigeant": "Président",
            "selas_associe_1_is_dirigeant": True,
            "selas_associe_1_role_dirigeant": "Président",
        },
    )
    blockers = sms._validate_roles_dirigeants(_phys(2))
    assert any("Un seul Président" in b for b in blockers)


def test_selas_president_plusieurs_dga_ok(monkeypatch) -> None:
    # O24-07 (« directeur général associé (PLUSIEURS) ») : Président + ≥2 DGA -> aucun blocage.
    from sydel_doc_engine.front_app import selas_multi_slice as sms

    _fake_st_session(
        monkeypatch,
        {
            "selas_associe_0_is_dirigeant": True,
            "selas_associe_0_role_dirigeant": "Président",
            "selas_associe_1_is_dirigeant": True,
            "selas_associe_1_role_dirigeant": "Directeur Général Associé",
            "selas_associe_2_is_dirigeant": True,
            "selas_associe_2_role_dirigeant": "Directeur Général Associé",
        },
    )
    assert sms._validate_roles_dirigeants(_phys(3)) == []


def test_selas_dnc_une_par_dirigeant(tmp_path: Path, monkeypatch) -> None:
    # R7 (Rafael 2026-06-23) : une declaration de non-condamnation PAR dirigeant
    # (President + chaque DG / DG Associe), chacune nommee par son nom.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-dnc-multi")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # Associe 1 devient un 2e dirigeant (DG Associe) + sa filiation (requise pour sa DNC).
    app.checkbox(key="selas_associe_1_is_dirigeant").set_value(True)
    app = app.run(timeout=180)
    next(
        s for s in app.selectbox if str(s.key) == "selas_associe_1_role_dirigeant"
    ).set_value("Directeur Général Associé")
    for w in app.text_input:
        if str(w.key) == "selas_associe_1_sig_nom_pere":
            w.set_value("Paul Martin")
        elif str(w.key) == "selas_associe_1_sig_nom_mere":
            w.set_value("Marie Martin")
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    generate_button.click()
    app = app.run(timeout=180)

    names = {
        p.name
        for p in (tmp_path / "ui-selas-dnc-multi").rglob("declaration_non_condamnation*.docx")
    }
    assert "declaration_non_condamnation_Durand.docx" in names  # president
    assert "declaration_non_condamnation_Martin.docx" in names  # DG Associe


def test_selas_situation_communaute_genere_docs_conjoint(tmp_path: Path, monkeypatch) -> None:
    # R10/R11 (Rafael 2026-06-23) : la situation matrimoniale est un MENU ; choisir un
    # regime de COMMUNAUTE fait apparaitre les champs conjoint et genere la renonciation
    # (DOC-005) + l'avertissement au conjoint (DOC-006). Plus de case a cocher dediee.
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell
    from sydel_doc_engine.front_app.field_derivations import (
        MATRIMONIAL_STATUS_MARRIED_COMMUNAUTE,
    )

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-selas-commu")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SELAS pluripersonnelle")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    # LIVE-02 : AUCUNE case à cocher « régime communautaire » dans le formulaire SELAS
    # (le verbatim dit « Supprimer » ; remplacée par le menu situation matrimoniale).
    assert not any("communautaire" in str(c.label).lower() for c in app.checkbox)

    # Associe 0 : choisir le regime de communaute -> les champs conjoint apparaissent.
    app.selectbox(key="selas_associe_0_situation").set_value(
        MATRIMONIAL_STATUS_MARRIED_COMMUNAUTE
    )
    app = app.run(timeout=180)
    for w in app.text_input:
        if str(w.key) == "selas_associe_0_conjoint_prenom":
            w.set_value("Alex")
        elif str(w.key) == "selas_associe_0_conjoint_nom":
            w.set_value("Durand")
    app = app.run(timeout=180)

    generate_button = next(
        b for b in app.button if str(b.key) == "clean_typed_generate_dossier"
    )
    assert generate_button.disabled is False
    generate_button.click()
    app = app.run(timeout=180)

    names = {p.name for p in (tmp_path / "ui-selas-commu").rglob("*.docx")}
    assert any("renonciation" in n for n in names), names  # DOC-005
    assert any("avertissement" in n for n in names), names  # DOC-006


def test_front_today_button_fills_date_non_selarl() -> None:
    # Retour Rafael 2026-06-09 : bouton « Aujourd'hui » sur les dates de TOUS les
    # types (existait deja en SELARL). Preuve sur SCI (non-SELARL) : presence du
    # bouton + remplissage effectif de la date du jour.
    from datetime import date as _date_cls

    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app.field_derivations import format_french_date

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=120)

    today_buttons = [b for b in app.button if str(b.key).endswith("_today")]
    assert today_buttons  # bouton « Aujourd'hui » present sur les dates SCI

    # Vider la date de signature puis cliquer « Aujourd'hui » -> date du jour.
    sig_key = "sci_signature_date"
    next(w for w in app.text_input if str(w.key) == sig_key).set_value("")
    app = app.run(timeout=120)
    next(b for b in app.button if str(b.key) == f"{sig_key}_today").click()
    app = app.run(timeout=120)
    assert next(w for w in app.text_input if str(w.key) == sig_key).value == (
        format_french_date(_date_cls.today())
    )


# --- Retours Albane 2026-06-17 : formulaires civils ---------------------------


def test_civil_forme_sociale_derived_from_structure() -> None:
    # §18.1 : la forme sociale (libelle) n'est plus saisie, elle est derivee.
    assert css.civil_forme_sociale("SCM") == "société civile de moyens"
    assert css.civil_forme_sociale("SCI") == "société civile"
    assert css.civil_forme_sociale("SCI IRIS") == "société civile"
    assert css.civil_forme_sociale("SCS") == "société civile"


def test_civil_build_derives_form_fields() -> None:
    # SCREEN-2 / §18.2 : valeur nominale = capital / nb parts (auto).
    # §18.1 : forme sociale derivee. §18.3 : duree figee 99. §18.4 : lieu = ville siege.
    payload = _civil_base(
        "SCI",
        "sci",
        [_pp("Jean", "Durand", 40, 1, 40, 400), _pp("Alice", "Martin", 60, 41, 100, 600)],
    )
    # On vide les champs derives du payload : le build doit les recalculer.
    payload["forme_sociale"] = ""
    payload["valeur_nominale_part"] = ""
    payload["signature_lieu"] = ""
    payload["duree_societe"] = ""
    ctx = css.build_generation_context(payload)
    assert ctx.statuts_civils.valeur_nominale_part == "10"  # 1000 / 100
    assert ctx.statuts_civils.forme_sociale == "société civile"
    assert ctx.statuts_civils.duree_societe == "99"
    assert ctx.signature.lieu == payload["siege_ville"]


def test_civil_scm_keeps_societe_civile_de_moyens() -> None:
    # §18.1 : la SCM porte « societe civile de moyens » meme si le payload est vide.
    payload = _civil_base(
        "SCM",
        "scm",
        [_pp("Jean", "Durand", 70, 1, 70, 700), _pp("Alice", "Martin", 30, 71, 100, 300)],
    )
    payload["forme_sociale"] = ""
    ctx = css.build_generation_context(payload)
    assert ctx.statuts_civils.forme_sociale == "société civile de moyens"


def test_repeater_assigns_cumulative_part_ranges() -> None:
    # §18.5 : les plages parts debut/fin ne sont plus saisies, elles se derivent de
    # l'ordre + du nombre de parts (contigues 1..N, sans trou ni chevauchement).
    from sydel_doc_engine.domain.models import (
        StatutsCivilsAssocie,
        StatutsCivilsParts,
    )
    from sydel_doc_engine.front_app.associe_repeater import (
        _assign_cumulative_part_ranges,
    )

    def _a(nb: int) -> StatutsCivilsAssocie:
        return StatutsCivilsAssocie(
            type_personne="personne_physique",
            parts=StatutsCivilsParts(nb=nb),
        )

    associes = [_a(40), _a(35), _a(25)]
    _assign_cumulative_part_ranges(associes)
    assert [(a.parts.debut, a.parts.fin) for a in associes] == [(1, 40), (41, 75), (76, 100)]
    assert associes[1].parts.plage_affichee == "41 à 75"  # N4 : « à » accentue


def test_a26_pv5_dirigeant_identite_phrase_via_flux_reel() -> None:
    # A26-PV5 (+ Akainu B1) : la phrase d'identite du dirigeant dans le PV SELAS reprend le format
    # EXACT du modele Albane et porte la profession REGLEMENTEE (« chirurgien-dentiste »), JAMAIS le
    # titre derive « Docteur ». Test sur le FLUX REEL : StatutsCivilsAssocie (profession="Docteur",
    # qualification_principale="chirurgien-dentiste") -> _build_dirigeants_nomines_payload ->
    # _build_dirigeants_nomines -> identite_phrase. (Un test qui hardcoderait la profession
    # masquerait le piege « Docteur » — regle 65.)
    from datetime import date as _date

    from sydel_doc_engine.domain.enums import Gender as _G
    from sydel_doc_engine.domain.models import Address as _A
    from sydel_doc_engine.domain.models import StatutsCivilsAssocie
    from sydel_doc_engine.front_app.selas_multi_slice import (
        _build_dirigeants_nomines,
        _build_dirigeants_nomines_payload,
        _selas_dirigeant_identite_phrase,
    )

    adresse = _A(adresse_affichee="31B Boulevard de Sévigné, 35700 RENNES")
    associe = StatutsCivilsAssocie(
        type_personne="personne_physique", genre=_G.MASCULIN, civilite_affichage="Monsieur",
        prenom="Jean-Guillaume", prenoms="Jean-Guillaume", nom="FUCHS",
        date_naissance=_date(1994, 2, 20), ville_naissance="RENNES", departement_naissance="35",
        nationalite="française",
        profession="Docteur",  # TITRE derive (le piege Akainu B1)
        qualification_principale="chirurgien-dentiste",  # profession REGLEMENTEE (source correcte)
        situation_maritale=(
            "marié sous le régime de la séparation des biens avec société d’acquêts, "
            "avec Madame Eva ROUAULT"
        ),
        adresse_personnelle=adresse,
    )
    payload_list = _build_dirigeants_nomines_payload([associe], 0)
    dirigeants = _build_dirigeants_nomines({"dirigeants_nomines": payload_list}, 0, adresse)
    assert len(dirigeants) == 1
    phrase = dirigeants[0].identite_phrase
    assert phrase == (
        "Monsieur Jean-Guillaume FUCHS, chirurgien-dentiste, de nationalité française, "
        "né le 20 février 1994 à RENNES (35), marié sous le régime de la séparation des biens "
        # Rafael/Albane 2026-07-09 : « demeurant [adresse] » -> « demeurant au [adresse] ».
        "avec société d’acquêts, avec Madame Eva ROUAULT, demeurant au 31B Boulevard de Sévigné, "
        "35700 RENNES"
    )
    assert "Docteur" not in phrase  # Akainu B1 : jamais le titre, la profession reglementee

    # Sans profession -> None (repli byte-identique sur la reconstruction par champs du PV).
    entry = payload_list[0]
    assert _selas_dirigeant_identite_phrase({**entry, "profession": ""}, adresse) is None
    # Genre feminin -> « née ».
    assert "née le" in _selas_dirigeant_identite_phrase({**entry, "genre": _G.FEMININ}, adresse)


def test_repeater_nationalite_dropdown_lowercased() -> None:
    # §SCREEN-1 : nationalite en deroulant (NATIONALITY_PRESETS), sortie lowercased.
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=120)

    # Le selectbox de nationalite de l'associe 0 existe (plus de text_input libre).
    choice = next(
        s for s in app.selectbox if str(s.key) == "sci_associe_0_nationalite_choice"
    )
    assert list(choice.options) == [
        "Française",
        "Belge",
        "Portugaise",
        "Suisse",
        "Luxembourgeoise",
        "Autre",
    ]
    assert not any(str(w.key) == "sci_associe_0_nationalite" for w in app.text_input)


def test_civil_non_scm_does_not_collect_profession() -> None:
    # §18.6 : profession demandee uniquement pour la SCM. SCI : pas de champ.
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=120)
    assert not any(str(w.key) == "sci_associe_0_profession" for w in app.text_input)

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=120)
    app.selectbox(key="clean_dossier_type").set_value("SCM")
    app = app.run(timeout=120)
    assert any(str(w.key) == "scm_associe_0_profession" for w in app.text_input)


# --- §17.1 : SELAS UNIPERSONNELLE medecin (DOC-018, generateur rebranche) ------


def test_selas_uni_medecin_present_in_dropdown() -> None:
    # §17.1 : le cas etait orphelin (generateur DOC-018 sans entree de menu). Il
    # doit desormais figurer dans la deroulante, par son nom, generation activee.
    option = dossier_type_by_label("SELAS unipersonnelle médecin")
    assert option.key == "selas_uni_medecin_v1"
    assert option.structure == "SELAS uni medecin"
    assert option.generation_enabled is True
    assert option.slice_module == "sydel_doc_engine.front_app.selas_uni_medecin_slice"
    # La SELARL reste en premiere position (defaut historique).
    assert dossier_type_labels()[0] == "SELARL"


def _selas_uni_medecin_payload():
    return {
        "denomination": "SELAS MARTIN",
        "capital_social": "1000",
        "nb_actions_total": 1000,
        "duree": "99 ans",
        "ville_rcs": "Paris",
        "lieu_exercice_adresse": "12 avenue de la Republique, 75011 Paris",
        "siege_num": "10",
        "siege_voie": "rue de la Paix",
        "siege_cp": "75002",
        "siege_ville": "Paris",
        "banque_nom": "BANQUE EXEMPLE",
        "banque_adresse": "1 boulevard Haussmann, 75009 Paris",
        "exercice_debut": "1er janvier",
        "exercice_fin": "31 decembre",
        "exercice_cloture": "31 decembre 2026",
        "civilite": "Monsieur",
        "prenom": "Camille",
        "nom": "Martin",
        "date_naissance": date(1980, 1, 2),
        "ville_naissance": "Paris",
        "departement_naissance": "75",
        "nationalite": "francaise",
        "titre_affichage": "Docteur",
        "adresse_num": "5",
        "adresse_voie": "5 rue Royale",
        "adresse_cp": "75008",
        "adresse_ville": "Paris",
        "nom_pere": "Pierre Martin",
        "nom_mere": "Anne Martin",
        "situation_maritale": "marie",
        "regime_matrimonial": "communaute legale",
        "conjoint_civilite": "Madame",
        "conjoint_prenom": "Alice",
        "conjoint_nom": "Martin",
        "ordre_conseil": "Ordre des medecins",
        "departement_ordre": "Paris",
        "numero_ordre": "12345",
        "numero_rpps": "10000000001",
        "ordre_ville": "Paris",
        "ordre_cp": "75008",
        "ordre_adresse_ligne_1": "1 rue de l'Ordre",
        "signature_lieu": "Paris",
        "signature_date": date(2026, 5, 14),
    }


def test_selas_uni_medecin_generates_doc018_bundle(tmp_path: Path) -> None:
    # §17.1 : le bundle de creation expose DOC-018 (statuts SELAS medecin) + tronc
    # commun. Genere proprement (aucun token residuel) et porte la forme SELAS.
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    payload = _selas_uni_medecin_payload()
    plan = uni.build_selas_uni_medecin_plan(payload)
    assert plan.can_generate is True
    # ANO-045 : un unipersonnel est toujours attestable -> DOC-045 (attestation
    # souscripteurs) fait partie du bundle de creation SELAS.
    assert plan.document_codes == (
        "DOC-018",
        "DOC-001",
        "DOC-002",
        "DOC-003",
        "DOC-004",
        "DOC-034",
        "DOC-045",
    )

    generated = uni.generate_dossier(payload, tmp_path / "selas-uni-medecin")
    _assert_bundle_clean(
        generated,
        {
            "Statuts SELAS MARTIN.docx",
            "declaration_non_condamnation.docx",
            "autorisation_domiciliation.docx",
            "procuration.docx",
            "pv_nomination_dirigeant.docx",
            "demande_inscription_ordre.docx",
            "attestation_capital_souscripteurs_selas.docx",
        },
    )
    statuts_text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "Statuts SELAS MARTIN.docx")
    )
    assert "SELAS MARTIN" in statuts_text


def test_selas_uni_medecin_empty_conjoint_warns() -> None:
    # Retour Rafael 2026-06-25 (#2) : un associe MARIE sans conjoint (vraie donnee manquante).
    # KAN-2 (Rafael, rejeté 2×) : ce manque NE BLOQUE PLUS -> avertissement, génération possible
    # (la zone conjoint sort en « (À COMPLÉTER : … ) »). Le payload de base est marie.
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    payload = _selas_uni_medecin_payload()
    payload["conjoint_prenom"] = ""
    payload["conjoint_nom"] = ""
    plan = uni.build_selas_uni_medecin_plan(payload)
    assert plan.can_generate is True
    assert plan.blockers == ()
    assert any("conjoint" in w.lower() for w in plan.warnings)


def test_selas_uni_medecin_celibataire_can_generate(tmp_path: Path) -> None:
    # Retour Rafael 2026-06-25 (#2 — Akainu B1) : un associe CELIBATAIRE (champs conjoint vides,
    # comme le front les pose desormais) doit pouvoir GENERER (plus de blockers conjoint
    # inconditionnels). Symetrique du test marie-sans-conjoint ci-dessus.
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    payload = _selas_uni_medecin_payload()
    payload["situation_maritale"] = "célibataire"
    payload["conjoint_civilite"] = ""
    payload["conjoint_prenom"] = ""
    payload["conjoint_nom"] = ""
    plan = uni.build_selas_uni_medecin_plan(payload)
    assert plan.can_generate is True
    assert not any("conjoint" in b.lower() for b in plan.blockers)

    generated = uni.generate_dossier(payload, tmp_path / "selas-uni-celibataire")
    statuts_text = _docx_text(
        next(p for p in generated.docx_paths if p.name == "Statuts SELAS MARTIN.docx")
    )
    assert "célibataire" in statuts_text
    assert "sous le régime de" not in statuts_text
    assert "avec Madame" not in statuts_text


def test_selas_uni_medecin_pv_uses_actions_not_parts(tmp_path: Path) -> None:
    # Coherence 2026-06-22 : une SELAS = societe par actions ; le PV de nomination doit dire
    # « actions », pas « parts » (il ignorait capital.type_titre='actions').
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    generated = uni.generate_dossier(_selas_uni_medecin_payload(), tmp_path / "selas-pv")
    pv = next(p for p in generated.docx_paths if p.name == "pv_nomination_dirigeant.docx")
    text = _docx_text(pv)
    assert "propriétaire de toutes les actions" in text
    assert "propriétaire de toutes les parts" not in text


def test_selas_uni_medecin_regime_communautaire_generates_doc005_006(tmp_path: Path) -> None:
    # RAF-001 (regression) : SELAS uni medecin + regime communautaire bloquait la
    # generation (CODE-RC-001 : date_courrier_avertissement absente de l'adaptateur
    # SELARL). Aucun test n'exercait ce chemin -> le bug est passe. On verrouille :
    # regime actif -> DOC-005/006 generes proprement, comme tous les autres types.
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    payload = _selas_uni_medecin_payload()
    payload["regime_communautaire"] = True
    plan = uni.build_selas_uni_medecin_plan(payload)
    assert plan.can_generate is True
    assert "DOC-005" in plan.document_codes
    assert "DOC-006" in plan.document_codes

    generated = uni.generate_dossier(payload, tmp_path / "selas-uni-medecin-regime")
    names = {p.name for p in generated.docx_paths}
    assert "lettre_renonciation_associe.docx" in names
    assert "lettre_avertissement_conjoint.docx" in names
    for path in generated.docx_paths:
        _assert_clean(_docx_text(path))


def test_selas_uni_medecin_context_is_selas_actions() -> None:
    # §17.1 : le contexte derive du parcours SELARL uni est bien transforme en
    # SELAS medecin (structure SELAS, overlay selas_medecin, titres = actions,
    # dirigeant President). Garde-fou anti-detournement du chemin SELARL.
    from sydel_doc_engine.front_app import selas_uni_medecin_slice as uni

    ctx = uni.build_generation_context(_selas_uni_medecin_payload())
    assert ctx.structure == "SELAS"
    assert ctx.statuts_sel.overlay == "selas_medecin"
    assert ctx.societe.forme_sociale_abregee == "SELAS"
    assert ctx.capital.type_titre == "actions"
    assert ctx.dirigeant_nomine.fonction_affichage == "président"
    assert ctx.dirigeant_nomine.duree_mandat == "illimitée"


def test_su3_scs2_force_signature_siege_et_decision_signature_divergent(tmp_path: Path) -> None:
    # SU3/SCS2/B1 (Albane 2026-06-25, lock Akainu) : avec une ville de signature DIVERGENTE du
    # siege ET une date de decision DIFFERENTE de la signature, le doc genere DOIT forcer
    # signature = siege et decision/reunion = date de signature. Valeurs divergentes = vrai test.
    from docx import Document

    from sydel_doc_engine.front_app import civil_statuts_slice as css
    from sydel_doc_engine.front_app import (
        sas_slice,
        selas_multi_slice,
        selas_uni_medecin_slice,
        spfpl_slice,
    )

    siege, divergent = "Lyon", "VilleSignatureDivergente"
    # ANNEE divergente aussi (Akainu re-gate m1) : 2024 != 2026 -> attrape annee_lettres du PV.
    sig_date, decision_div = date(2026, 5, 26), date(2024, 3, 7)

    def _full_text(p: Path) -> str:
        d = Document(p)
        parts = [pa.text for pa in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    parts += [pa.text for pa in c.paragraphs]
        return "\n".join(parts)

    def _ov(payload: dict) -> dict:
        p = dict(payload)
        p.update({
            "siege_ville": siege, "signature_lieu": divergent,
            "signature_date": sig_date, "decision_date": decision_div,
        })
        return p

    cases = [
        ("SCS", lambda d: css.generate_dossier(_ov(_civil_base(
            "SCS", "scs",
            [_pp("Jean", "Durand", 60, 1, 60, 600, role="commandite"),
             _pp("Alice", "Martin", 40, 61, 100, 400, role="commanditaire")])), d)),
        ("SAS", lambda d: sas_slice.generate_dossier(_ov(_sas_payload()), d)),
        ("SELAS-multi", lambda d: selas_multi_slice.generate_dossier(_ov(_selas_payload()), d)),
        ("SPFPL", lambda d: spfpl_slice.generate_dossier(_ov(_spfpl_payload("SPFPL cession")), d)),
        ("SELAS-uni", lambda d: selas_uni_medecin_slice.generate_dossier(
            _ov(_selas_uni_medecin_payload()), d)),
    ]
    for label, gen in cases:
        generated = gen(tmp_path / label)
        for path in generated.docx_paths:
            text = _full_text(path)
            # SU3 : la ville de signature divergente ne doit JAMAIS apparaitre (forcee au siege).
            assert divergent not in text, f"{label} ville signature (SU3)"
            # SCS2/B1 : la date de decision divergente (7 mars 2024) ne doit pas apparaitre (PV).
            assert "sept mars" not in text, f"{label} date decision (SCS2/B1)"
            assert "07/03/2024" not in text, f"{label} date decision chiffres"
            # B1 annee : l'annee divergente en lettres ne doit pas apparaitre (annee_lettres PV).
            assert "deux mille vingt-quatre" not in text, f"{label} annee decision (B1)"


def test_apporteur_apport_M1_absent_donne_marqueur_pas_capital() -> None:
    # M1 (Akainu 2026-06-26) : apport individuel ABSENT -> marqueur visible
    # « (À COMPLÉTER : apport individuel) », JAMAIS un repli sur le capital (qui
    # reintroduirait le bug A26-40 : capital affiche au lieu de l'apport reel).
    from types import SimpleNamespace

    sans_apport = SimpleNamespace(apport=None)
    res = selas_multi_slice._apporteur_apport(sans_apport)
    assert res is not None
    assert "À COMPLÉTER" in res.montant
    assert "À COMPLÉTER" in res.montant_lettres
    assert "1 000" not in res.montant  # surtout PAS le capital

    # Apport individuel present -> repris tel quel (chiffre + lettres).
    avec_apport = SimpleNamespace(
        apport=SimpleNamespace(montant="500", montant_lettres="cinq cents")
    )
    res2 = selas_multi_slice._apporteur_apport(avec_apport)
    assert res2.montant == "500"
    assert res2.montant_lettres == "cinq cents"


# ---------------------------------------------------------------------------
# R29-06 (Rafael) : selecteur de date (calendrier) + bouton « Aujourd'hui »
# PARTOUT. Tests ADVERSARIAUX : on assert que, pour un champ date route au helper
# `date_input_with_today`, l'app rendue expose BIEN les trois widgets — le champ
# texte `{key}`, le calendrier `{key}_cal` ET le bouton `{key}_today`. Sans ces
# assertions, une regression (champ retombe en text_input nu) passerait verte.
# ---------------------------------------------------------------------------


def _widget_keys(app) -> set[str]:
    """Toutes les cles de widgets rendus (text_input + date_input + button)."""
    keys: set[str] = set()
    for collection in (app.text_input, app.date_input, app.button):
        for widget in collection:
            if widget.key is not None:
                keys.add(str(widget.key))
    return keys


def _assert_date_picker_trio(app, base_key: str) -> None:
    """Le champ date `base_key` expose UN champ texte editable + le bouton « Aujourd'hui ».

    Retour Rafael 2026-07-02 : le calendrier `st.date_input` (`{base_key}_cal`) a ete RETIRE —
    il s'affichait comme un champ date VIDE en double (« tout seul ») au-dessus du champ
    labellise. Il ne doit donc PLUS exister ; seuls le champ texte + le bouton subsistent."""
    keys = _widget_keys(app)
    assert base_key in keys, f"champ texte editable absent : {base_key}"
    assert f"{base_key}_cal" not in keys, (
        f"calendrier `{base_key}_cal` encore present -> le champ date en double devait etre "
        "retire (retour Rafael 2026-07-02, « champs dates tout seul au-dessus »)"
    )
    assert f"{base_key}_today" in keys, (
        f"bouton « Aujourd'hui » absent pour {base_key}"
    )


def test_r29_06_repeater_naissance_expose_calendrier_et_aujourdhui() -> None:
    """Rafael 2026-07-02 : la date de naissance d'un associe du REPEATER (SCI) expose UN champ
    texte editable + le bouton `{key}_today`, SANS calendrier `{key}_cal` (retire, cf. helper)."""
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="clean_dossier_type").set_value("SCI")
    app = app.run(timeout=180)
    next(b for b in app.button if "test_data" in str(b.key)).click()
    app = app.run(timeout=180)

    _assert_date_picker_trio(app, "sci_associe_0_date_naissance")

    # NON-REGRESSION : la saisie verbatim francaise reste possible sur le champ texte
    # (meme cle) -> le helper n'a pas casse la saisie LIVE-03.
    _set_text_widget(app, "sci_associe_0_date_naissance", "1er aout 1980")
    app = app.run(timeout=180)
    assert (
        app.session_state["sci_associe_0_date_naissance"] == "1er aout 1980"
    ), "la saisie verbatim doit etre conservee telle quelle dans la cle texte"


def test_r29_06_cession_dates_exposent_calendrier_et_aujourdhui(monkeypatch, tmp_path) -> None:
    """Rafael 2026-07-02 : les dates du sous-formulaire CESSION (date du bail, date d'acquisition
    du cabinet) — via `_cession_date` sur le helper — exposent UN champ texte + le bouton
    `{key}_today`, SANS calendrier `{key}_cal` (retire : plus de champ date « tout seul »)."""
    from streamlit.testing.v1 import AppTest

    from sydel_doc_engine.front_app import shell

    monkeypatch.setattr(shell, "ARTIFACTS_DIR", tmp_path / "ui-r2906-cession")
    app = AppTest.from_file("src/sydel_doc_engine/front_app/app.py").run(timeout=180)
    app.selectbox(key="selarl_profession").set_value("Medecin")
    app = app.run(timeout=180)
    app.button(key="clean_generate_test_data").click()
    app = app.run(timeout=180)
    assert app.checkbox(key="selarl_cession").value is True

    _assert_date_picker_trio(app, "selarl_cession_bail_date_bail")
    _assert_date_picker_trio(app, "selarl_cession_cabinet_origine_date")

    # NON-REGRESSION : la saisie texte cession reste fonctionnelle (meme cle).
    app.text_input(key="selarl_cession_bail_date_bail").set_value("1er aout 2021")
    app = app.run(timeout=180)
    assert app.session_state["selarl_cession_bail_date_bail"] == "1er aout 2021"
